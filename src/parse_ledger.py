# -*- coding: utf-8 -*-
"""监测表 → 规整数据（长表 + 宽表 + 元数据 + 断言）

唯一输入：input/<期次>/监测表.xlsx（馆员手工维护的 22 工作簿台账）
唯一产出：out/<期次>/ 下的 ledger_long.parquet / wide/*.parquet / meta.json

原则：只读单元格原值，不做任何“合理修正”；脏数据（%%、备注行）靠
白名单与清洗规则处理，处理痕迹记录进 meta.json 的 clean_log。
"""
import json
import re
import sys
import time
from pathlib import Path

import openpyxl
import pandas as pd

DISCIPLINES_EN = [
    "Chemistry", "Plant & Animal Science", "Microbiology", "Engineering",
    "Materials Science", "Mathematics", "Environment/Ecology",
    "Biology & Biochemistry", "Pharmacology & Toxicology", "Geosciences",
    "Agricultural Sciences", "Computer Science", "Clinical Medicine",
    "Molecular Biology & Genetics", "Physics", "Social Sciences, general",
    "Multidisciplinary", "Neuroscience & Behavior", "Space Science",
    "Immunology", "Psychiatry/Psychology", "Economics & Business",
]

DISCIPLINES_CN = {
    "Chemistry": "化学",
    "Plant & Animal Science": "植物与动物学",
    "Microbiology": "微生物学",
    "Engineering": "工程学",
    "Materials Science": "材料科学",
    "Mathematics": "数学",
    "Environment/Ecology": "环境生态学",
    "Biology & Biochemistry": "生物学与生物化学",
    "Pharmacology & Toxicology": "药理学与毒理学",
    "Geosciences": "地球科学",
    "Agricultural Sciences": "农业科学",
    "Computer Science": "计算机科学",
    "Clinical Medicine": "临床医学",
    "Molecular Biology & Genetics": "分子生物学与遗传学",
    "Physics": "物理学",
    "Social Sciences, general": "社会科学",
    "Multidisciplinary": "多学科",
    "Neuroscience & Behavior": "神经系统科学与行为学",
    "Space Science": "空间科学",
    "Immunology": "免疫学",
    "Psychiatry/Psychology": "精神病学与心理学",
    "Economics & Business": "经济与商业",
}

WIDE_SHEETS = ["总被引", "论文数", "CNCI", "接近度", "ESI被引阈值", "TOP论文"]
WIDE_METRIC = {
    "总被引": "total_cites",
    "论文数": "papers",
    "CNCI": "cnci",
    "接近度": "proximity_1pct",
    "ESI被引阈值": "threshold_1pct",
    "TOP论文": "top_papers",
}


def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("％", "").replace("%%", "")
    s = s.replace("%", "")
    if s in ("", "—", "-", "√"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


_PERIOD_RE = re.compile(r"(20\d{2})\D*(\d{1,2})\D*$")


def _norm_period(s):
    """脏期次 → YYYYMM。2023.3 / 2023-03 / 2023年3月 / 2023/3 / 20233 都收。

    只在自愈 overlay 打开时启用（见 `load_ledger`）——默认口径不做任何"合理修正"，
    这是本模块开篇的原则。归一是纯数字重排，幂等、可回退。
    """
    t = "" if s is None else str(s).strip()
    m = _PERIOD_RE.fullmatch(t)
    return f"{m.group(1)}{int(m.group(2)):02d}" if m else t


def _canon(name, alias):
    """学科名归一：命中别名表的换成规范名，未命中原样返回。

    `alias` 默认空 = 完全按现状走（认不出就跳过）。别名表由馆员口径提供，
    模块不臆测相似度 —— 猜错会把数据挂到别的学科上，且全程不报错。
    """
    return (alias or {}).get(name, name)


def _parse_scope(overlay, sheet):
    """从 overlay 里取出作用于某张工作表的读表参数。"""
    ov = overlay or {}
    return {
        "header_row": (ov.get("header_offset") or {}).get(sheet, 0),
        "group_size": (ov.get("group_size") or {}).get(sheet),
        "normalize": bool(ov.get("normalize_period")),
        "alias": ov.get("discipline_alias") or {},
    }


def parse_wide(ws, clean_log, header_row=0, normalize=False, alias=None, diag=None):
    """学科×期次 宽表：首行期次，首列学科英文名。

    `header_row` / `normalize` / `alias` 为自愈 overlay 入口，默认值＝既有口径。
    `diag` 是诊断出口，记下**首列原样出现的名字**——宽表 reindex 到 22 个规范名后，
    光看 `df.index` 已经分不清"少了谁"，断言与《诊断报告》都要靠这份原始名单。
    """
    rows = list(ws.iter_rows(values_only=True))
    header = rows[header_row] if header_row < len(rows) else []
    periods = []
    for h in header[1:]:
        s = "" if h is None else str(h).strip()
        if s:
            periods.append(_norm_period(s) if normalize else s)
    data, unknown = {}, []
    for r in rows[header_row + 1:]:
        raw = "" if r[0] is None else str(r[0]).strip()
        name = _canon(raw, alias)
        if name not in DISCIPLINES_EN:
            if raw:
                unknown.append(raw)      # 静默跳过的名字要留证，否则断言只会说"少了一个"
            continue
        if name in data:
            continue  # 只取首个块（CNCI 表上下两块，报告用第一块）
        data[name] = [_num(v) for v in r[1:1 + len(periods)]]
    df = pd.DataFrame(data, index=periods).T
    df.index.name = "discipline"
    df = df.reindex(DISCIPLINES_EN)
    # ⚠ reindex 之后 index 恒为 22 个规范名 —— 断言里**不能**用 df.index 判"名单一致"，
    #   那两条断言会恒真。把真正被认出来的名字走 diag 通道送出去。
    #   （不用 df.attrs：pandas 3 会把 attrs 写进 parquet 元数据，害得产物字节不再可复现。）
    if diag is not None:
        diag.setdefault("known_names", {})[ws.title] = sorted(data)
        diag.setdefault("unknown_names", {})[ws.title] = sorted(set(unknown))
    clean_log.append(f"{ws.title}: 宽表 {len(periods)} 期 × {len(df)} 学科"
                     + (f"（首列未识别：{sorted(set(unknown))}）" if unknown else ""))
    return df


def parse_yun_da_rank(ws, clean_log):
    """云大ESI排名：两块 4 列并排（期次/国际/国内/国内高校）。"""
    rows = list(ws.iter_rows(values_only=True))
    recs = []
    for r in rows[1:]:
        for base in (0, 6):
            period = r[base]
            if period is None or not str(period).strip():
                continue
            recs.append({
                "period": str(period).strip(),
                "rank_intl": _num(r[base + 1]),
                "rank_cn": _num(r[base + 2]),
                "rank_cn_uni": _num(r[base + 3]),
            })
    df = pd.DataFrame(recs).drop_duplicates("period").set_index("period").sort_index()
    clean_log.append(f"{ws.title}: {len(df)} 期机构排名")
    return df


def parse_grouped(ws, group_size, period_row, metric_names, clean_log, alias=None):
    """每期一组（组宽 group_size）的学科宽表，组名在 period_row 行。"""
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[period_row]
    groups = []
    j = 1
    while j < len(hdr):
        s = "" if hdr[j] is None else str(hdr[j]).strip()
        if s:
            groups.append((j, s))
        j += 1
    out, unknown = {}, []
    for r in rows[period_row + 1:]:
        raw = "" if r[0] is None else str(r[0]).strip()
        name = _canon(raw, alias)
        if name not in DISCIPLINES_EN:
            if raw:
                unknown.append(raw)
            continue
        out[name] = {}
        for start, period in groups:
            vals = [_num(v) for v in r[start:start + group_size]]
            out[name][period] = vals
    per = [p for _, p in groups]
    df = pd.DataFrame(
        [[d, p, *out[d][p]] for d in out for p in per],
        columns=["discipline", "period"] + metric_names,
    )
    clean_log.append(f"{ws.title}: {len(groups)} 期 × {len(out)} 学科 × {group_size} 指标"
                     + (f"（首列未识别：{sorted(set(unknown))}）" if unknown else ""))
    return df


def parse_transposed(ws, clean_log, header_row=0, alias=None):
    """国际/国内排名：学科 3 列组（组名第 1 行），期次为行。"""
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[header_row] if header_row < len(rows) else []
    groups = []
    j = 1
    while j < len(hdr):
        s = "" if hdr[j] is None else str(hdr[j]).strip()
        if s:
            groups.append((j, s))
        j += 1
    recs = []
    for r in rows[header_row + 2:]:
        period = "" if r[0] is None else str(r[0]).strip()
        if not period:
            continue
        for start, disc in groups:
            recs.append({
                "discipline": _canon(disc, alias), "period": period,
                "n_inst_1pct": _num(r[start]),
                "rank": _num(r[start + 1]),
                "rank_pct": _num(r[start + 2]),
            })
    df = pd.DataFrame(recs)
    dropped = []
    if not df.empty:
        dropped = sorted(set(df.discipline) - set(DISCIPLINES_EN))
        df = df[df.discipline.isin(DISCIPLINES_EN)]
    clean_log.append(f"{ws.title}: 转置表 {len(groups)} 学科 × {df.period.nunique()} 期"
                     + (f"（表头未识别：{dropped}）" if dropped else ""))
    return df


def parse_threshold_1permille(ws, clean_log, header_row=0, normalize=False, alias=None):
    rows = list(ws.iter_rows(values_only=True))
    head = rows[header_row] if header_row < len(rows) else []
    periods = [(_norm_period(h) if normalize else str(h).strip())
               for h in head[1:] if h and str(h).strip()]
    data, unknown = {}, []
    for r in rows[header_row + 1:]:
        raw = "" if r[0] is None else str(r[0]).strip()
        name = _canon(raw, alias)
        if name not in DISCIPLINES_EN:
            if raw:
                unknown.append(raw)
            continue
        data[name] = [_num(v) for v in r[1:1 + len(periods)]]
    df = pd.DataFrame(data, index=periods).T
    df.index.name = "discipline"
    clean_log.append(f"{ws.title}: {periods}"
                     + (f"（首列未识别：{sorted(set(unknown))}）" if unknown else ""))
    return df


def parse_avg_cites(ws, clean_log):
    """篇均被引：每期 3 列组（云大/中国大陆/国际）。"""
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[1]
    groups = []
    j = 1
    while j < len(hdr):
        s = "" if hdr[j] is None else str(hdr[j]).strip()
        if s:
            groups.append((j, s))
        j += 1
    recs = []
    for r in rows[3:]:
        name = "" if r[0] is None else str(r[0]).strip()
        if name not in DISCIPLINES_EN:
            continue
        for start, period in groups:
            recs.append({
                "discipline": name, "period": period,
                "avg_cites_ynd": _num(r[start]),
                "avg_cites_cn": _num(r[start + 1]),
                "avg_cites_global": _num(r[start + 2]),
            })
    df = pd.DataFrame(recs)
    clean_log.append(f"{ws.title}: {len(groups)} 期 × {df.discipline.nunique()} 学科")
    return df


def parse_top_summary(ws, clean_log):
    """TOP论文：宽表 + 合计行 + 年份行。（表头期次有脏值 2023.3/2023.5，统一归一）"""
    rows = list(ws.iter_rows(values_only=True))
    hdr = ["" if h is None else str(h).strip() for h in rows[0][1:]]
    periods = [re.sub(r"\.(\d)$", r".0\1", h).replace(".", "") for h in hdr if h]
    data, year_rows, total = {}, {}, None
    for r in rows[1:]:
        name = "" if r[0] is None else str(r[0]).strip()
        if name in DISCIPLINES_EN:
            data[name] = [_num(v) for v in r[1:1 + len(periods)]]
        elif re.fullmatch(r"20\d\d", name):
            year_rows[name] = [_num(v) for v in r[1:1 + len(periods)]]
        elif name.startswith("高被引论文合计"):
            total = {p: _num(v) for p, v in zip(periods, r[1:1 + len(periods)])}
    df = pd.DataFrame(data, index=periods).T.reindex(DISCIPLINES_EN)
    df.index.name = "discipline"
    yr = pd.DataFrame(year_rows, index=periods).T
    clean_log.append(f"{ws.title}: 宽表 {len(periods)} 期 + 年份行 {list(year_rows)}")
    return df, yr, total


def parse_q1q4(ws, clean_log):
    """Q1-Q4：季度 × 实体（云大/中国大陆/全球）× Q1-Q4 占比。"""
    rows = list(ws.iter_rows(values_only=True))
    hdr_q, hdr_e = rows[0], rows[1]
    groups, cur_q = [], None
    for j in range(1, len(hdr_q)):
        q = hdr_q[j]
        if q is not None and str(q).strip():
            cur_q = str(q).strip()
        e = hdr_e[j]
        if e is not None and str(e).strip():
            groups.append((j, cur_q, str(e).strip()))
    recs = []
    for r in rows[3:]:
        name = "" if r[0] is None else str(r[0]).strip()
        if name not in DISCIPLINES_EN:
            continue
        for start, q, ent in groups:
            vals = [_num(v) for v in r[start:start + 4]]
            recs.append({"discipline": name, "quarter": q, "entity": ent,
                         "q1": vals[0], "q2": vals[1], "q3": vals[2], "q4": vals[3]})
    df = pd.DataFrame(recs)
    clean_log.append(f"{ws.title}: {df.quarter.nunique()} 季度 × {df.entity.nunique()} 实体 × {df.discipline.nunique()} 学科")
    return df


def parse_province(ws, clean_log):
    """高校学科发展：9 校指标行 + 前1%学科概况块。"""
    rows = list(ws.iter_rows(values_only=True))
    unis = ["" if v is None else str(v).strip() for v in rows[0][1:10]]
    metrics = {}
    for r in rows[2:9]:
        label = "" if r[0] is None else str(r[0]).strip()
        if not label:
            continue
        metrics[label] = {u: _num(r[i + 1]) for i, u in enumerate(unis)}
    df = pd.DataFrame(metrics).T
    entry_block = []
    for r in rows[24:]:
        name = "" if r[0] is None else str(r[0]).strip()
        v = _num(r[1])
        if name and name != "7所高校入围ESI前1%学科概况":
            entry_block.append((name, v))
    clean_log.append(f"{ws.title}: {list(metrics)} × 9 校；前1%学科概况 {len(entry_block)} 行")
    return df, entry_block


def parse_forecast_gap(ws, clean_log):
    """实际与预测差值：期次范围 × 学科 × 指标 × 来源(incite/esi/差值)。"""
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    groups = []
    j = 2
    while j < len(hdr):
        s = "" if hdr[j] is None else str(hdr[j]).strip()
        if s and j % 4 == 2:
            groups.append((j, s))
        j += 1
    CN_TO_EN = {v: k for k, v in DISCIPLINES_CN.items()}
    recs, cur_range = [], None
    for r in rows[2:]:
        rng = "" if r[0] is None else str(r[0]).strip()
        src = "" if r[1] is None else str(r[1]).strip()
        if rng:
            cur_range = rng
        if src in ("incite", "esi", "差值"):
            for start, disc in groups:
                disc_en = CN_TO_EN.get(disc)
                if disc_en is None:
                    continue
                recs.append({
                    "win": cur_range, "source": src, "discipline": disc_en,
                    "papers": _num(r[start]), "cites": _num(r[start + 1]),
                    "rank_intl": _num(r[start + 2]), "rank_cn": _num(r[start + 3]),
                })
    df = pd.DataFrame(recs)
    clean_log.append(f"{ws.title}: {len(groups)} 学科 × {df.win.nunique()} 期范围 × 3 来源")
    return df


#: 中文学科名 → 规范英文名的反查表（正向表见 `DISCIPLINES_CN`）。
CN2EN = {v: k for k, v in DISCIPLINES_CN.items()}

#: 「学科进入时间」表里出现过的**手写变体** → 规范中文名。
#: 🔴 「生物与生物化学」必须显式列出：台账规范名是「**生物学**与生物化学」，
#:    差一个「学」字 → 精确匹配落空 → 会被「化学」子串截胡（2026-09-19 实测踩到）。
ENTRY_CN_ALIAS = {
    "微生物": "微生物学",
    "生物与生物化学": "生物学与生物化学",
    "生物与生化": "生物学与生物化学",
}

#: 退出事件写在学科列的**注释文本**里，形如「2025年5月数学退出」。
ENTRY_OUT_RE = re.compile(r"^(\d{4})年(\d{1,2})月(.+?)退出$")


def _cn2en(name, alias=None):
    """中文学科名 → 规范英文名。认不出时原样返回（由调用方记诊断，不臆测）。

    匹配顺序是硬要求：**精确 → 按别名长度降序做包含**。任意顺序的宽松包含会让
    「生物与生物化学」命中「化学」（后者是前者的子串）。
    """
    s = str(name or "").strip()
    s = ENTRY_CN_ALIAS.get((alias or {}).get(s, s), (alias or {}).get(s, s))
    if s in CN2EN:
        return CN2EN[s]
    for k in sorted(CN2EN, key=len, reverse=True):
        if k and k in s:
            return CN2EN[k]
    return s


def parse_entry_events(ws, clean_log, alias=None):
    """解析「学科进入时间」工作表 → `[(期次, 学科英文名, "in"/"out")]`。

    这张表是**馆员手填的流水账**，不是结构化状态矩阵。三条必须遵守（2026-09-19 实测定）：

    1. **年月靠向下填充**：年份只在每年首行写一次，后续行留空（Excel 合并单元格的表现）。
       逐行取 `row[0]`/`row[1]` 会漏掉 201307、201701、202503 等事件。
    2. **退出写在学科列的注释文本里**（如「2025年5月数学退出」），只能靠正则提取；
       学科行的「目前ESI前1%学科」√ 列**只描述当前状态，不是每期状态**。
    3. **学科名是中文且有手写变体**，见 `ENTRY_CN_ALIAS`。

    返回 (events, unparsed)，`unparsed` 是需要人工看一眼的行（无法定位年月或学科名）。
    """
    rows = list(ws.iter_rows(values_only=True))
    events, unparsed = [], []
    year = month = None
    for r in rows[1:]:
        y, m, name, _cur = (list(r) + [None] * 4)[:4]
        if y is not None:
            year = int(y)
        if m:
            month = m
        s = str(name or "").strip()
        if not s:
            continue
        mo = ENTRY_OUT_RE.match(s)
        if mo:
            events.append((f"{int(mo.group(1))}{int(mo.group(2)):02d}",
                           _cn2en(mo.group(3), alias), "out"))
            continue
        if "退出" in s:                       # 含「退出」但年月格式认不出
            unparsed.append(s)
            continue
        if not (year and month):
            unparsed.append(s)
            continue
        ym = f"{year}{str(month).replace('月', '').zfill(2)}"
        events.append((ym, _cn2en(s, alias), "in"))

    known = {d for _, d, _ in events if d in DISCIPLINES_EN}
    unknown = sorted({d for _, d, _ in events if d not in DISCIPLINES_EN})
    clean_log.append(f"学科进入时间: 解析 {len(events)} 条事件"
                     f"（进入 {sum(1 for e in events if e[2] == 'in')} / "
                     f"退出 {sum(1 for e in events if e[2] == 'out')}），"
                     f"覆盖 {len(known)} 个学科")
    if unknown:
        clean_log.append(f"[告警] 学科进入时间: 学科名未识别 {unknown} —— 这些事件未参与计算")
    if unparsed:
        clean_log.append(f"[告警] 学科进入时间: {len(unparsed)} 行无法定位年月，需人工核对：{unparsed}")
    return events, unparsed


def build_entered_1pct(events, periods, disciplines):
    """事件表 → 学科 × 期次的 0/1 矩阵（1 = 该期已在 ESI 前 1% 内）。

    🔴 必须**按时序交替回放** in/out，不能「先全部 in 再全部 out」：
    微生物学 202505 退出、202507 再进入，先加后删会把 202507 之后的它误删掉。
    同一期次既有 in 又有 out 时，**out 最后执行**（＝该期判为未进入，保守）。
    """
    merged = sorted([e for e in events if e[1] in disciplines],
                    key=lambda x: (x[0], 0 if x[2] == "in" else 1))
    periods = [str(p) for p in periods]
    df = pd.DataFrame(0.0, index=pd.Index(list(disciplines), name="discipline"),
                      columns=periods)
    for p in periods:
        cur = set()
        for ym, d, typ in merged:
            if ym <= p:
                if typ == "in":
                    cur.add(d)
                else:
                    cur.discard(d)
        for d in cur:
            df.loc[d, p] = 1.0
    return df


def load_ledger(path, overlay=None):
    """读台账 → 规整数据。`overlay` 是自愈注入的**读表口径**（默认 None ＝ 既有口径）。

    overlay 只改变"怎么读"，不改变"读了什么"：台账文件全程只读，
    任何 overlay 都可以通过不传它来完全回退。字段见
    `src/agent/rules/repair_rules.yaml`，由 `src/agent/heal.py` 生成。

        header_offset      {工作表名: 表头行号(0 起)}
        group_size         {工作表名: 每期列数}
        normalize_period   bool  脏期次归一为 YYYYMM
        discipline_alias   {别名: 规范名}
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    clean_log: list[str] = []
    diag: dict = {"overlay": overlay or {},
                  "overlay_applied": sorted(overlay) if overlay else []}
    sc = {n: _parse_scope(overlay, n) for n in
          ("总被引", "论文数", "CNCI", "接近度", "ESI被引阈值", "TOP论文",
           "千分之一被引阈值", "新国际排名", "新国内排名", "国际排名", "国内排名")}
    # 宽表 6 张用同一组读表参数（表头行 / 脏期次归一 / 学科别名）
    wide_kw = {"header_row": sc["总被引"]["header_row"],
               "normalize": sc["总被引"]["normalize"],
               "alias": sc["总被引"]["alias"]}
    wide = {WIDE_METRIC[n]: parse_wide(wb[n], clean_log, diag=diag, **wide_kw)
            for n in WIDE_SHEETS}
    ynd_rank = parse_yun_da_rank(wb["云大ESI排名"], clean_log)
    new_intl = parse_grouped(wb["新国际排名"], sc["新国际排名"]["group_size"] or 3,
                             sc["新国际排名"]["header_row"],
                             ["n_inst_1pct", "rank", "rank_pct"], clean_log,
                             alias=sc["新国际排名"]["alias"])
    new_cn = parse_grouped(wb["新国内排名"], sc["新国内排名"]["group_size"] or 3,
                           sc["新国内排名"]["header_row"],
                           ["n_inst_1pct", "rank", "rank_pct"], clean_log,
                           alias=sc["新国内排名"]["alias"])
    intl = parse_transposed(wb["国际排名"], clean_log, header_row=sc["国际排名"]["header_row"],
                            alias=sc["国际排名"]["alias"])
    cn = parse_transposed(wb["国内排名"], clean_log, header_row=sc["国内排名"]["header_row"],
                          alias=sc["国内排名"]["alias"])
    thr_1p = parse_threshold_1permille(
        wb["千分之一被引阈值"], clean_log, header_row=sc["千分之一被引阈值"]["header_row"],
        normalize=sc["千分之一被引阈值"]["normalize"],
        alias=sc["千分之一被引阈值"]["alias"])
    avg = parse_avg_cites(wb["篇均被引"], clean_log)
    q1q4 = parse_q1q4(wb["Q1-Q4"], clean_log)
    top, top_years, top_total = parse_top_summary(wb["TOP论文"], clean_log)
    prov, prov_entry = parse_province(wb["高校学科发展"], clean_log)
    gap = parse_forecast_gap(wb["实际与预测差值"], clean_log)
    # 「某期是否已进入 ESI 前 1%」的**唯一权威来源**是这张事件流水表。
    # 🔴 不能用接近度反推：未进入学科的接近度来自 InCites，与 ESI 阈值不是同一个量，
    #    `接近度 >= 1` 会多算（实测第 1 期 202601 多算数学 1 个：14 vs 原件 13）。
    tc = wide["total_cites"]
    entry_events, entry_unparsed = parse_entry_events(
        wb["学科进入时间"], clean_log, alias=(overlay or {}).get("discipline_alias"))
    entered = build_entered_1pct(entry_events, tc.columns, tc.index)
    # 诊断信息：供自愈规则匹配与《诊断报告》引用（只读，不含台账明细值）
    diag["periods_total_cites"] = [str(c) for c in wide["total_cites"].columns]
    diag["entry_unparsed"] = entry_unparsed
    return {
        "wide": wide,
        "entered_1pct": entered,
        "entry_events": entry_events,
        "ynd_rank": ynd_rank,
        "new_intl": new_intl,
        "new_cn": new_cn,
        "rank_intl": intl,
        "rank_cn": cn,
        "threshold_1permille": thr_1p,
        "avg_cites": avg,
        "q1q4": q1q4,
        "top_papers": top,
        "top_by_year": top_years,
        "top_total": top_total,
        "province": prov,
        "province_entry": prov_entry,
        "forecast_gap": gap,
        "clean_log": clean_log,
        "diagnostics": diag,
    }


def _latest_periods(ledger):
    """当前期 = 总被引宽表最后一个有数据的期次；上期 = 其前一个。"""
    tc = ledger["wide"]["total_cites"]
    filled = [c for c in tc.columns if tc[c].notna().any()]
    if len(filled) < 2:
        raise ValueError("台账被引表有效期次不足 2 期，无法定位当期")
    return filled[-1], filled[-2]


def _known_names(ledger, metric, df):
    """该宽表**真正被认出来**的学科名。

    没有诊断记录时（例如由旧版产物/外部构造的 ledger 进来）退回 `df.index`，
    但因为已经 reindex，那等于"全都认出来了"——所以 `run_assertions` 的调用方
    必须走 `load_ledger`，别自己拼 ledger 字典。
    """
    title = next((k for k, v in WIDE_METRIC.items() if v == metric), None)
    nm = (ledger.get("diagnostics") or {}).get("known_names") or {}
    return set(nm.get(title, df.index))


def run_assertions(ledger):
    """机器断言闸：结构级不变量。返回 (ok, report)。"""
    report, ok = [], True

    def check(cond, msg):
        nonlocal ok
        report.append(("PASS" if cond else "FAIL", msg))
        ok = ok and cond

    P, _Q = _latest_periods(ledger)
    TOP_BLANK_OK = {"Multidisciplinary", "Neuroscience & Behavior", "Immunology"}
    for m, df in ledger["wide"].items():
        # ⚠ 宽表在 parse_wide 里已 reindex 到 22 个规范名，df.index 恒等 —— 必须用
        #   diagnostics 里「真正认出来的名字」判名单，否则这两条断言永远为真
        #   （学科名拼错也放行，最后只在"期无缺失"上朦胧地报一句）。
        known = _known_names(ledger, m, df)
        missing = sorted(set(DISCIPLINES_EN) - known)
        check(len(known) == 22, f"宽表 {m} 学科数=22 (实际 {len(known)}"
                                + (f"，缺 {missing}" if missing else "") + ")")
        check(not missing, f"宽表 {m} 学科名单一致"
                           + (f"（缺 {missing}）" if missing else ""))
        check(P in df.columns, f"宽表 {m} 含期次 {P}")
        if m == "top_papers":
            na_set = set(df[P][df[P].isna()].index)
            check(na_set <= TOP_BLANK_OK, f"宽表 {m} {P} 期空白仅限 {TOP_BLANK_OK} (实际 {sorted(na_set)})")
        else:
            check(df[P].notna().sum() == 22, f"宽表 {m} {P} 期无缺失")
        future = [c for c in df.columns if c > P]
        for fc in future:
            if m == "cnci":
                continue  # InCites 月更先行，允许已贴
            check(df[fc].isna().all(), f"宽表 {m} {fc} 列为空（下期未贴数）")
        bad = df.map(lambda v: isinstance(v, str) and "%" in v).any().any() if df.map(lambda v: isinstance(v, str)).any().any() else False
        check(not bad, f"宽表 {m} 无残留 %% 脏值")

    y = ledger["ynd_rank"]
    check(P in y.index and pd.notna(y.loc[P, "rank_intl"]), f"云大ESI排名 {P} 已贴数")
    if P == "202607":
        check(y.loc[P, "rank_intl"] == 995, "金样(202607)：国际排名=995")
    for df, nm in [(ledger["new_intl"], "新国际排名"), (ledger["new_cn"], "新国内排名")]:
        check(df.period.nunique() >= 5, f"{nm} 期组数≥5")
        check(df.discipline.nunique() == 22, f"{nm} 学科数=22")
        check((df.period == P).sum() == 22, f"{nm} 含 {P} 全 22 行")
    for df, nm in [(ledger["rank_intl"], "国际排名"), (ledger["rank_cn"], "国内排名")]:
        check(df.period.nunique() >= 40, f"{nm} 期数≥40")
        check((df.period == P).sum() == 22, f"{nm} 含 {P} 全 22 行")
    check(P in ledger["threshold_1permille"].columns, f"千分之一阈值含 {P}")
    ent = ledger["entered_1pct"]
    check(ent.shape[0] == 22, f"进入前1%矩阵 学科数=22 (实际 {ent.shape[0]})")
    check(P in ent.columns, f"进入前1%矩阵 含期次 {P}")
    check(ent[P].notna().sum() == 22, f"进入前1%矩阵 {P} 期无缺失")
    if "202601" in ent.columns:
        # 🔴 金样锚点：第 1 期人工原件写「ESI前1%学科13个」。用接近度>=1 会算成 14（多算数学）。
        #    这条断言锁死新口径，防止哪天被改回接近度口径而无人察觉。
        check(int(ent["202601"].sum()) == 13,
              f"金样(202601)：前1%学科数=13（人工原件口径，实际 {int(ent['202601'].sum())}）")
    check(ledger["top_total"] is not None and ledger["top_total"].get(P) is not None, f"顶尖论文合计 {P} 已贴数")
    if P == "202607":
        check(ledger["top_total"].get(P) == 274, "金样(202607)：顶尖论文合计=274")
    check(ledger["province"].shape == (6, 9), f"省高校表 6 指标 × 9 校 (实际 {ledger['province'].shape})")
    return ok, report


def _jsonable(o):
    if isinstance(o, dict):
        return {k: _jsonable(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [_jsonable(v) for v in o]
    if hasattr(o, "item") and not isinstance(o, (str,)):
        try:
            return o.item()
        except Exception:
            pass
    if isinstance(o, float):
        return None if o != o else o
    if o is None or isinstance(o, (str, int, float, bool)):
        return o
    return str(o)


#: 一次运行的**成品** —— 只有它们会被当报告看。中间产物（宽表 / 大表 / 图）每一步都
#: 直接覆盖同名文件，留旧的无害；成品则必须先撤下：否则一次在 charts / docx 步失败的
#: 运行会留下「本次的新宽表 + 上一次的旧报告」，看板照旧展示旧报告、对勘本拿旧槽位配
#: 新格子，而页面上没有任何警报（2026-09-16 取证复现）。
FINAL_PRODUCTS = ("报告初稿_*.docx", "红色件清单.md", "red_slots.json")

#: 撤下的成品去哪。放期次目录**内部**而不是 `out/` 下：让"这一期的历史版本"跟着这一期走，
#: 人工翻档不用两头找；也避开 `out/` 根目录的期次发现逻辑（`server.jobs.PERIOD_DIR_RE`
#: 只认六位数字开头的目录名，`_superseded` 不会被当成一期）。
SUPERSEDED_DIR = "_superseded"


def quarantine_final_products(out_dir, run_dir=None):
    """把期内已有的成品挪进 `_superseded/<时间戳>/`，返回挪走的文件名。

    时机是这件事的全部：必须在**本次确定要改写这个期次目录之后、写成品之前**做。

    - 挪早了（比如解析之前）：一次注定失败的运行会把上一版好报告也藏起来，
      而它其实一个字节都没被覆盖过。
    - 挪晚了（写完成品之后）：等于没挪。

    两条调用路径（`agent/flow.py` 的编排、`run_pipeline.py` 的直跑）都在"当期算出来了"
    的那一刻调它，口径因此一致：**成品只在它所在目录开始被本次运行改写之后才撤下**。
    若解析这一步就没过（目录压根没被碰过），上一版继续有效。
    """
    out = Path(out_dir)
    if not out.is_dir():
        return []
    todo = [p for pat in FINAL_PRODUCTS for p in out.glob(pat) if p.is_file()]
    if not todo:
        return []
    dst = out / SUPERSEDED_DIR / time.strftime("%Y%m%dT%H%M%S")
    try:
        dst.mkdir(parents=True, exist_ok=True)
        moved = []
        for p in todo:
            p.replace(dst / p.name)
            moved.append(p.name)
        (dst / "_来源.txt").write_text(
            f"{time.strftime('%Y-%m-%d %H:%M:%S')}：{out.name} 期开始一次新运行，"
            f"旧成品先撤到这里。\n"
            f"发起方运行目录：{run_dir or '（未提供）'}\n"
            f"若那次运行失败，本期期内将**没有**成品（＝未出稿），"
            f"而不是让看板继续展示这一版 —— 因为期内其余文件已经是新数据了。\n",
            encoding="utf-8")
        return sorted(moved)
    except OSError:
        return []          # 搬迁/留痕失败不阻断生产（与 agent/trace.py 同口径）


def write_outputs(ledger, path, out_dir):
    """把已解析台账落盘为规整数据 + meta，返回 (ok, meta)。"""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    ok, report = run_assertions(ledger)
    P, Q = _latest_periods(ledger)
    next_cols = [c for c in ledger["wide"]["total_cites"].columns if c > P]

    long_frames = []
    for m, df in ledger["wide"].items():
        long_frames.append(df.reset_index().melt(id_vars="discipline", var_name="period", value_name="value").assign(metric=m))
    for key, cols in [("new_intl", ["n_inst_1pct", "rank", "rank_pct"]),
                      ("new_cn", ["n_inst_1pct", "rank", "rank_pct"])]:
        long_frames.append(ledger[key].melt(id_vars=["discipline", "period"], value_vars=cols,
                                            var_name="metric", value_name="value").assign(source=key))
    for key, cols in [("rank_intl", ["n_inst_1pct", "rank", "rank_pct"]),
                      ("rank_cn", ["n_inst_1pct", "rank", "rank_pct"])]:
        long_frames.append(ledger[key].melt(id_vars=["discipline", "period"], value_vars=cols,
                                            var_name="metric", value_name="value").assign(source=key))
    long_df = pd.concat(long_frames, ignore_index=True)
    long_df.to_parquet(out_dir / "ledger_long.parquet", index=False)

    wide_dir = out_dir / "wide"
    wide_dir.mkdir(exist_ok=True)
    for m, df in ledger["wide"].items():
        df.to_parquet(wide_dir / f"{m}.parquet")
    ledger["ynd_rank"].reset_index().to_parquet(wide_dir / "ynd_rank.parquet")
    ledger["top_papers"].to_parquet(wide_dir / "top_papers.parquet")
    pd.Series(ledger["top_total"], name="total").rename_axis("period").reset_index().to_parquet(wide_dir / "top_total.parquet")
    ledger["threshold_1permille"].to_parquet(wide_dir / "threshold_1permille.parquet")
    ledger["entered_1pct"].to_parquet(wide_dir / "entered_1pct.parquet")
    ledger["q1q4"].to_parquet(wide_dir / "q1q4.parquet")
    ledger["province"].to_parquet(wide_dir / "province.parquet")
    ledger["forecast_gap"].to_parquet(wide_dir / "forecast_gap.parquet")

    meta = {
        "input": str(path),
        "assertions_ok": ok,
        "assertions": report,
        "clean_log": ledger["clean_log"],
        "provinces": ledger["province"].columns.tolist(),
        "province_entry_block": ledger["province_entry"],
        "top_total_by_period": ledger["top_total"],
        "period_current": P,
        "period_prev": Q,
        "period_next_empty": next_cols[0] if next_cols else None,
        # 自愈相关：本次读表用的 overlay 与未识别学科名（只读诊断，供《诊断报告》引用）
        "diagnostics": ledger.get("diagnostics", {}),
    }
    (out_dir / "meta.json").write_text(json.dumps(_jsonable(meta), ensure_ascii=False, indent=1), encoding="utf-8")
    return ok, meta


def main(path, out_dir=None):
    ledger = load_ledger(path)
    if out_dir is None:
        P, _ = _latest_periods(ledger)
        out_dir = Path("out") / P
    ok, meta = write_outputs(ledger, path, out_dir)
    print(json.dumps(_jsonable(meta), ensure_ascii=False, indent=1))
    return ok


if __name__ == "__main__":
    src = Path(sys.argv[1] if len(sys.argv) > 1 else r"input\202607\监测表.xlsx")
    dst = sys.argv[2] if len(sys.argv) > 2 else None
    ok = main(src, dst)
    sys.exit(0 if ok else 1)
