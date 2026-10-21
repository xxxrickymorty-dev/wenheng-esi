# -*- coding: utf-8 -*-
"""跨校对比取数（灵活报告第四节）。

只读 ESI 侧归档，不碰台账、不调模型。两条数据路径：

1. `数据与复现/ESI数据/_归档统一命名/<期次>/<期次>ESI总表模板.xlsx`
   —— 带中文校名与顶尖/高被引/热点列；目前只有 202601–202607 四期。
2. `数据与复现/ESI数据/<期次>/<学科英文名>.xlsx`
   —— 每期都有；列为 位次 / 机构 / 国家 / 发文 / 被引 / 篇均 / 顶尖。

校名解析：中文名优先查归档总表的「中国机构信息」（945 条），
查不到再到当期学科表里按英文名兜底。用户每次指定，不设默认名单。
"""
from __future__ import annotations

import functools
import pathlib
import re

import openpyxl

from .datasource import cn_name

REPO = pathlib.Path(__file__).resolve().parents[2]
WORKSPACE = REPO.parent
ESI_ROOT = WORKSPACE / "数据与复现" / "ESI数据"
TEMPLATE_DIR = "_归档统一命名"


def _display_path(period: str, filename: str) -> str:
    """报告里写的来源：仓库相对路径，不把本机绝对路径带进成品。"""
    full = ESI_ROOT / period / filename
    try:
        return full.relative_to(WORKSPACE).as_posix()
    except ValueError:                                          # 服务器数据根不在仓库旁
        return f"数据与复现/ESI数据/{period}/{filename}"


class PeerError(Exception):
    """对标高校解析/取数失败（人话信息，直接可回给用户）。"""


def _norm(s) -> str:
    s = str(s or "").upper()
    s = s.replace("&", " AND ")
    s = re.sub(r"[-‐–—,，.。()（）/]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    s = re.sub(r"\b(CHINA|MAINLAND)\b", "", s)
    return re.sub(r"\s+", " ", s).strip()


def _clean_en(s) -> str:
    """去掉 ESI 表里的地区后缀：`SOUTHWEST FORESTRY UNIVERSITY - CHINA`。"""
    return re.sub(r"\s*[-–—]\s*CHINA\s*$", "", str(s or "").strip(), flags=re.I)


def _split_queries(queries) -> list[str]:
    if queries is None:
        return []
    if isinstance(queries, str):
        raw = [queries]
    else:
        raw = [str(x) for x in queries]
    out: list[str] = []
    for chunk in raw:
        for part in re.split(r"[,，、;；\n]+|\s+和\s+", chunk):
            p = part.strip()
            if p and p not in out:
                out.append(p)
    return out


def _latest_template() -> pathlib.Path | None:
    base = ESI_ROOT / TEMPLATE_DIR
    if not base.is_dir():
        return None
    files = sorted(base.glob("*/*ESI总表模板.xlsx"))
    return files[-1] if files else None


def _clean_cn(s) -> str:
    s = str(s or "").strip()
    s = re.sub(r"\s*[-–—]\s*(中国|CHINA)\s*$", "", s, flags=re.I)
    return s


@functools.lru_cache(maxsize=4)
def _name_map(template: str) -> dict[str, tuple[str, str]]:
    """英文/中文规范名 → (英文名, 中文名)。

    除「中国机构信息」外，还把 22 张学科表的（英文单位, 单位）一并收进来——
    总表缺的校名（如西南林业大学）在学科表里往往有。
    """
    p = pathlib.Path(template)
    if not p.exists():
        return {}
    wb = openpyxl.load_workbook(p, data_only=True)
    out: dict[str, tuple[str, str]] = {}

    def add(en, cn):
        en = _clean_en(en)
        cn = _clean_cn(cn)
        if not en:
            return
        if cn and cn not in ("#N/A", "None", "N/A"):
            out.setdefault(_norm(cn), (en, cn))
        out.setdefault(_norm(en), (en, cn or en))

    if "中国机构信息" in wb.sheetnames:
        for row in wb["中国机构信息"].iter_rows(min_row=2, values_only=True):
            add(row[0] if len(row) > 0 else "", row[1] if len(row) > 1 else "")
    for sheet in wb.sheetnames:
        if sheet in ("中国机构信息", "学校总表", "上榜学科", "大陆上榜学科明细"):
            continue
        for row in wb[sheet].iter_rows(min_row=2, values_only=True):
            if len(row) > 2:
                add(row[1], row[2])
    return out


def _discipline_sheet(disc_en: str) -> str:
    return cn_name(disc_en) or disc_en


def _template_rows(period: str, disc_en: str):
    p = ESI_ROOT / TEMPLATE_DIR / period / f"{period}ESI总表模板.xlsx"
    if not p.exists():
        return None
    wb = openpyxl.load_workbook(p, data_only=True)
    sheet = _discipline_sheet(disc_en)
    if sheet not in wb.sheetnames:
        return None
    rows = []
    for r in wb[sheet].iter_rows(min_row=2, values_only=True):
        if not r or not r[1]:
            continue
        def num(v):
            try:
                return float(str(v).replace(",", ""))
            except (TypeError, ValueError):
                return None
        def pick(i):
            return r[i] if len(r) > i else None
        rows.append({
            "rank": num(pick(0)), "en": _clean_en(pick(1)),
            "cn": str(pick(2) or "").strip(),
            "country": str(pick(3) or "").strip(),
            "papers": num(pick(4)), "cites": num(pick(5)),
            "cites_per_paper": num(pick(6)), "top_papers": num(pick(7)),
            "hcp": num(pick(8)), "hot": num(pick(9)),
        })
    return rows, p.name, sheet


def _raw_rows(period: str, disc_en: str):
    p = ESI_ROOT / period / f"{disc_en}.xlsx"
    if not p.exists():
        return None
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header_at = None
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if header_at is None:
            if r and any(str(x or "").strip().lower() == "institutions" for x in r):
                header_at = i
            continue
        if not r or not r[1]:
            continue
        def num(v):
            try:
                return float(str(v).replace(",", ""))
            except (TypeError, ValueError):
                return None
        rows.append({
            "rank": num(r[0]), "en": _clean_en(r[1]), "cn": "",
            "country": str(r[2] or "").strip(),
            "papers": num(r[3]), "cites": num(r[4]),
            "cites_per_paper": num(r[5]), "top_papers": num(r[6]),
            "hcp": None, "hot": None,
        })
    return rows, p.name, ws.title


def _period_dirs() -> list[str]:
    if not ESI_ROOT.is_dir():
        return []
    return sorted(p.name for p in ESI_ROOT.iterdir()
                  if p.is_dir() and re.fullmatch(r"\d{6}", p.name))


def fetch(period: str, disc_en: str) -> dict:
    """读某期某学科的机构表。返回 `{rows, source, period, discipline}`。

    优先原表（每期都有、列口径稳定）；当期缺失或为空时，**只向前**回退到
    最近一个有数据的期次，并在返回值里标明实际期次（不改任何数据）。
    """
    period = str(period)
    candidates = [period] + [p for p in reversed(_period_dirs()) if p < period]
    chosen = None
    for p in candidates:
        got = _raw_rows(p, disc_en)
        if got is not None and got[0]:
            chosen = (p, *got)
            break
        if got is None:
            tpl = _template_rows(p, disc_en)
            if tpl is not None and tpl[0]:
                chosen = (p, *tpl)
                break
    if chosen is None:
        raise PeerError(f"{period} 及之前都没有「{cn_name(disc_en) or disc_en}」的机构表，"
                        f"无法做跨校对比（未改动任何数据）。")
    actual, rows, filename, sheet = chosen
    tpl = _latest_template()
    name_map = _name_map(str(tpl)) if tpl else {}
    for r in rows:
        if not r.get("cn"):
            hit = name_map.get(_norm(r["en"]))
            if hit:
                r["cn"] = hit[1]
    return {"rows": rows, "source": filename, "sheet": sheet,
            "period": actual, "requested_period": period, "discipline": disc_en}


def _match(rows, name_map, query) -> tuple[str, str]:
    """查询名 → (英文名, 中文名)。先查总表，再查当期学科表。"""
    nq = _norm(query)
    hit = name_map.get(nq)
    if hit:
        return hit
    # 中文去掉“大学/学院”后缀再试一次
    short = re.sub(r"(大学|学院|研究院|研究所)$", "", str(query).strip())
    if short and short != query:
        hit = name_map.get(_norm(short))
        if hit:
            return hit
    for r in rows:
        if _norm(r["en"]) == nq or _norm(r["cn"]) == nq:
            return r["en"], r["cn"] or r["en"]
    for r in rows:
        if nq and (nq in _norm(r["en"]) or (r["cn"] and nq in _norm(r["cn"]))):
            return r["en"], r["cn"] or r["en"]
    raise PeerError(f"没找到高校「{query}」。请用中文校名（如“昆明理工大学”）"
                    f"或 ESI 英文名；可先确认该校在该学科是否上榜。")


def _ynu(rows) -> dict | None:
    for r in rows:
        if _norm(r["en"]) == "YUNNAN UNIVERSITY" or r["cn"] == "云南大学":
            return r
    return None


def _domestic_rank(rows, target) -> int | None:
    """大陆机构按全球位次计国内序位（只数 CHINA MAINLAND）。"""
    if not target or target.get("rank") is None:
        return None
    cn = [r for r in rows if "CHINA MAINLAND" in str(r.get("country", "")).upper()
          and r.get("rank") is not None]
    cn.sort(key=lambda r: r["rank"])
    for i, r in enumerate(cn, 1):
        if r is target:
            return i
    return None


_METRICS = (("rank", "全球位次"), ("papers", "发文量"), ("cites", "总被引"),
            ("cites_per_paper", "篇均被引"), ("top_papers", "顶尖论文"))


def compare(period: str, disc_en: str, queries) -> dict:
    """云南大学 vs 用户指定高校（同一学科）。"""
    data = fetch(period, disc_en)
    rows = data["rows"]
    tpl = _latest_template()
    name_map = _name_map(str(tpl)) if tpl else {}
    ynu = _ynu(rows)
    if ynu is None:
        raise PeerError(f"{period}「{cn_name(disc_en) or disc_en}」机构表里没有云南大学。")

    wanted = []
    for q in _split_queries(queries):
        en, cn = _match(rows, name_map, q)
        if _norm(en) == _norm(ynu["en"]):
            continue
        if all(_norm(en) != _norm(w["en"]) for w in wanted):
            wanted.append({"en": en, "cn": cn})

    picked = [ynu] + [next((r for r in rows if _norm(r["en"]) == _norm(w["en"])), None)
                      for w in wanted]
    out_rows = []
    for r, w in zip(picked, [None] + wanted):
        if r is None:
            assert w is not None
            out_rows.append({"en": w["en"], "cn": w["cn"], "present": False,
                             "values": {k: None for k, _ in _METRICS},
                             "rank_cn": None})
            continue
        out_rows.append({
            "en": r["en"], "cn": r["cn"] or (w["cn"] if w else r["en"]),
            "present": True,
            "values": {k: r.get(k) for k, _ in _METRICS},
            "rank_cn": _domestic_rank(rows, r),
        })

    actual = data["period"]
    return {
        "period": actual, "requested_period": period, "discipline": disc_en,
        "discipline_cn": cn_name(disc_en) or disc_en,
        "ynu": {"en": ynu["en"], "cn": ynu["cn"] or "云南大学"},
        "rows": out_rows, "requested": [w["cn"] for w in wanted],
        "n_global": len(rows),
        "source": {"dataset": _display_path(actual, data["source"]),
                   "sheet": data["sheet"], "period": actual},
        "metrics": [{"key": k, "cn": v} for k, v in _METRICS],
    }


def rows_for_prompt(period: str, disc_en: str, queries) -> list[dict]:
    """给报告正文用的平表（避免正文层再碰 openpyxl）。"""
    return compare(period, disc_en, queries)["rows"]
