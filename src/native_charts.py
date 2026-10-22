# -*- coding: utf-8 -*-
"""Word **原生图表**（OOXML `word/charts/chartN.xml`）构建与插入。

背景与动机（2026-10-22 rick 定调「所有图表都要 Word 可编辑」）：
  现有路径是 matplotlib 出 PNG 再插进 docx —— 观感 1:1，但**评审方无法改数**。
  本模块把图改成 Word 原生 chart part：右键「编辑数据」可开内嵌工作簿改数，
  数字改动直接反映到图上。

## 一个 chart 由三件东西组成（缺一件就不是「可编辑」）

| 部件 | 作用 | 缺了会怎样 |
|---|---|---|
| `word/charts/chartN.xml` | `c:chartSpace`（图形 + `c:numCache` 数值缓存） | 没有图 |
| `word/embeddings/*.xlsx` + `<c:externalData>` | 「编辑数据」打开的**数据源** | 右键「编辑数据」报错，改不了数（实测 HRESULT 0xB0D7019E） |
| 段落里的 `<c:chart r:id>` | 图形框引用 | 图不显示 |

🔴 **实测（2026-10-22，Word 16.0 + `probe_word.ps1`）**：
  · 有内嵌工作簿时，**工作簿是唯一数据源**——只改 `numCache` 不改工作簿，图**不变**；
    改工作簿（+0.15）后图变（chart 区域 6722 px / 131 行变化）；
  · 摘掉工作簿后，`numCache` 变成唯一数据源，改 cache 图也变（9436 px）。
  所以本模块**同时写工作簿与 numCache**，两者必须同源同值。

## 与 `charts.py` 的关系

`charts.py`（matplotlib PNG 路径）**一个字节都不动**；本模块只是并行的第二条渲染路径，
由开关控制（见 `fill_docx.NATIVE_CHARTS_ENV`）。纵轴的「逐学科自适应 + 整齐步长」口径
**复用 `charts._nice_step`**，避免两条路径算出不同的轴。

## 做不到的部分（原生图表能力边界，详见 `docs/native-charts-第一段验证报告.md`）

1. **数值标签不自动避让**：Word/WPS 没有「躲开别的字」的机制。统一用 `dLblPos="t"`，
   边缘点的标签会压到纵轴刻度字（matplotlib 路径靠 `_labels` 的贪心避让解决）。
2. **图例只能选固定位**（`t/b/l/r/tr`），不能「坐标区外右上、距轴顶 1 行字高」。
3. **图上自由文字**（如 1‰ 图的「ESI1‰ 基线 = 100%」）没有对应 API。
4. 字体走 `a:latin/a:ea`，Word 缺字时用**它自己的**回退，不能像 matplotlib 那样钉住字体列表。
"""
from __future__ import annotations

import io
import pathlib
import re
import sys

sys.path.insert(0, str(pathlib.Path(__file__).parent))

from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from docx.opc.packuri import PackURI  # noqa: E402
from docx.opc.part import Part  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402

import charts  # noqa: E402  复用 _nice_step / _slug / DETAIL_TABLES（口径单一来源）

CHART_CT = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET = "Sheet1"                       # 必须与 c:f 里的表名一致

C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

# 主题色（与 charts.py 同值：系列1 蓝 / 系列2 绿 / 系列3 黄）
C_BLUE, C_GREEN, C_YELLOW = "4472C4", "70AD47", "FFC000"
C_ORANGE, C_GRAY = "ED7D31", "A5A5A5"
AXIS_LINE, GRID_LINE, INK = "D1D5DB", "E8EAED", "1F2937"

EMU_PER_CM = 360000
SLOT_W_CM, SLOT_H_CM = 8.26, 5.50      # 母版槽位硬约束（与 fill_docx 一致）
EA_FONT, LATIN_FONT = "宋体", "Times New Roman"

_BLIP = "\x00"                         # 占位：不需要


# ---------------------------------------------------------------- 小工具
def _esc(s):
    """XML 转义（文本节点与属性值通用）。

    🔴 属性值里的 `"` 必须转义：数值格式码 `0.##"%"` 自带引号，直接插进
    `formatCode="…"` 会把 XML 属性截断 → **Word 拒绝打开整份文档**（2026-10-22 实测）。
    """
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace('"', "&quot;").replace("'", "&apos;"))


def _col(n):
    """1-based 列号 → 字母（本报告最多 8 列，够用）。"""
    return chr(ord("A") + n - 1)


def _txpr(size_pt=8, *, ea=EA_FONT, latin=LATIN_FONT, bold=False, color=None):
    """图表内文字样式（`sz` 单位＝百分之一磅）。"""
    c = f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>' if color else ""
    b = ' b="1"' if bold else ""
    return ('<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr'
            f' sz="{int(size_pt * 100)}"{b}>'
            f'<a:latin typeface="{latin}"/><a:ea typeface="{ea}"/>'
            f'<a:cs typeface="{latin}"/>{c}</a:defRPr></a:pPr>'
            '<a:endParaRPr lang="zh-CN"/></a:p></c:txPr>')


def _str_cache(vals):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{_esc(v)}</c:v></c:pt>' for i, v in enumerate(vals))
    return f'<c:ptCount val="{len(vals)}"/>{pts}'


def _num_cache(vals, fmt="General"):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{v!r}</c:v></c:pt>' for i, v in enumerate(vals))
    return f'<c:formatCode>{fmt}</c:formatCode><c:ptCount val="{len(vals)}"/>{pts}'


def _hex(c):
    return str(c).lstrip("#")


# ---------------------------------------------------------------- 数据工作簿
def build_data_xlsx(cats, series, *, sheet=SHEET):
    """内嵌数据工作簿：A 列＝系列名（首格写「期次」），第 1 行＝类目，其余＝数值。"""
    from openpyxl import Workbook  # noqa: PLC0415

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws["A1"] = "期次"
    for j, c in enumerate(cats):
        ws.cell(row=1, column=2 + j, value=str(c))
    for i, s in enumerate(series):
        ws.cell(row=2 + i, column=1, value=str(s["name"]))
        for j, v in enumerate(s["values"]):
            ws.cell(row=2 + i, column=2 + j, value=None if v is None else float(v))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- 系列
def _dlabels(fmt, *, show=True, size=7, hide=(), pos="t"):
    """系列数值标签。`hide` 里的点索引单独 `delete=1`（如只标末点的图）。

    🔴 元素顺序＝CT_DLbls：`dLbl*` → `delete` → `numFmt`/`spPr`/`txPr`/`dLblPos`/`show*`。
    顺序错了 Word 会宽容，WPS 与 schema 校验器不宽容（照 Word 自己生成的 chart 抄的序）。

    🔴 `pos` 必须按图型选：**折线/散点** 可用 `t`（上）、`b`、`l`、`r`、`ctr`；
    **柱状/条形** 只认 `outEnd`（外侧末端）/`inEnd`/`inBase`/`ctr`/`bestFit`。
    给柱状图写 `dLblPos="t"` 会让 **Word 直接拒绝打开整份文档**（2026-10-22 实测，
    报「Word 在读取该文件时遇到错误」；隔离到一张柱状图上二分出来的）。
    """
    if not show:
        return '<c:dLbls><c:delete val="1"/></c:dLbls>'
    per_point = "".join(
        f'<c:dLbl><c:idx val="{k}"/><c:delete val="1"/></c:dLbl>' for k in sorted(hide))
    return (
        "<c:dLbls>" + per_point +
        f'<c:numFmt formatCode="{_esc(fmt)}" sourceLinked="0"/>'
        '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
        + _txpr(size) +
        f'<c:dLblPos val="{pos}"/>'
        '<c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/>'
        '<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/>'
        "</c:dLbls>")


def _ser_line(idx, name, vals, cats, *, color, label_fmt, show_labels=True,
              dash=None, marker="circle", marker_size=5, marker_fill="FFFFFF",
              lw=12700, hide_labels=(), smooth=False, val_fmt=None):
    ncat = len(cats)
    ln = f'<a:ln w="{lw}" cap="flat"><a:solidFill><a:srgbClr val="{_hex(color)}"/></a:solidFill>'
    ln += f'<a:prstDash val="{dash}"/>' if dash else "<a:round/>"
    ln += "</a:ln>"
    m = ""
    if marker and marker != "none":
        m = (f'<c:marker><c:symbol val="{marker}"/><c:size val="{marker_size}"/>'
             '<c:spPr><a:solidFill>'
             f'<a:srgbClr val="{_hex(marker_fill)}"/></a:solidFill>'
             f'<a:ln w="9525"><a:solidFill><a:srgbClr val="{_hex(color)}"/>'
             "</a:solidFill></a:ln></c:spPr></c:marker>")
    else:
        m = '<c:marker><c:symbol val="none"/></c:marker>'
    return (
        "<c:ser>"
        f'<c:idx val="{idx}"/><c:order val="{idx}"/>'
        f'<c:tx><c:strRef><c:f>{SHEET}!$A${idx + 2}</c:f><c:strCache>'
        f'<c:ptCount val="1"/><c:pt idx="0"><c:v>{_esc(name)}</c:v></c:pt>'
        "</c:strCache></c:strRef></c:tx>"
        f"<c:spPr>{ln}</c:spPr>" + m + _dlabels(label_fmt, show=show_labels, hide=hide_labels) +
        f"<c:cat><c:strRef><c:f>{SHEET}!$B$1:${_col(ncat + 1)}$1</c:f><c:strCache>"
        f"{_str_cache(cats)}</c:strCache></c:strRef></c:cat>"
        f'<c:val><c:numRef><c:f>{SHEET}!$B${idx + 2}:${_col(ncat + 1)}${idx + 2}</c:f>'
        f'<c:numCache>{_num_cache([float(v) for v in vals], val_fmt or "General")}'
        "</c:numCache></c:numRef></c:val>"
        f'<c:smooth val="{1 if smooth else 0}"/>'
        "</c:ser>")


def _ser_bar(idx, name, vals, cats, *, color, label_fmt, show_labels=True, hide_labels=(),
             dlbl_pos="outEnd"):
    """柱系列。🔴 顺序＝CT_BarSer：`spPr` → `invertIfNegative` → `dLbls` → `cat` → `val`。

    🔴 `dlbl_pos` 按分组方式定：**clustered** 用 `outEnd`（柱外侧）；
    **stacked / percentStacked** 只认 `ctr`/`inBase`/`inEnd` —— 给堆叠柱写 `outEnd`
    会让 **Word 拒绝打开整份文档**（2026-10-22 实测）。
    """
    ncat = len(cats)
    return (
        "<c:ser>"
        f'<c:idx val="{idx}"/><c:order val="{idx}"/>'
        f'<c:tx><c:strRef><c:f>{SHEET}!$A${idx + 2}</c:f><c:strCache>'
        f'<c:ptCount val="1"/><c:pt idx="0"><c:v>{_esc(name)}</c:v></c:pt>'
        "</c:strCache></c:strRef></c:tx>"
        '<c:spPr><a:solidFill>'
        f'<a:srgbClr val="{_hex(color)}"/></a:solidFill><a:ln><a:noFill/></a:ln></c:spPr>'
        '<c:invertIfNegative val="0"/>'
        + _dlabels(label_fmt, show=show_labels, hide=hide_labels, pos=dlbl_pos) +
        f"<c:cat><c:strRef><c:f>{SHEET}!$B$1:${_col(ncat + 1)}$1</c:f><c:strCache>"
        f"{_str_cache(cats)}</c:strCache></c:strRef></c:cat>"
        f'<c:val><c:numRef><c:f>{SHEET}!$B${idx + 2}:${_col(ncat + 1)}${idx + 2}</c:f>'
        f'<c:numCache>{_num_cache([float(v) for v in vals])}</c:numCache></c:numRef></c:val>'
        "</c:ser>")


# ---------------------------------------------------------------- 坐标轴
def _scaling(*, reverse=False, minimum=None, maximum=None, log=False):
    s = f'<c:scaling><c:orientation val="{"maxMin" if reverse else "minMax"}"/>'
    if log:
        s += '<c:logBase val="10"/>'
    if maximum is not None:
        s += f'<c:max val="{maximum}"/>'
    if minimum is not None:
        s += f'<c:min val="{minimum}"/>'
    return s + "</c:scaling>"


def _val_axis(ax_id, cross_ax, *, pos="l", reverse=False, numfmt="General", major_unit=None,
              minimum=None, maximum=None, gridlines=True, tick_label=True, fontsize=8):
    grid = ""
    if gridlines:
        grid = ('<c:majorGridlines><c:spPr><a:ln w="9525"><a:solidFill>'
                f'<a:srgbClr val="{GRID_LINE}"/></a:solidFill></a:ln></c:spPr>'
                "</c:majorGridlines>")
    return (
        f'<c:valAx><c:axId val="{ax_id}"/>'
        + _scaling(reverse=reverse, minimum=minimum, maximum=maximum) +
        '<c:delete val="0"/>' + f'<c:axPos val="{pos}"/>' + grid +
        f'<c:numFmt formatCode="{_esc(numfmt)}" sourceLinked="0"/>'
        '<c:majorTickMark val="none"/><c:minorTickMark val="none"/>'
        f'<c:tickLblPos val="{"nextTo" if tick_label else "none"}"/>'
        '<c:spPr><a:ln w="9525"><a:solidFill>'
        f'<a:srgbClr val="{AXIS_LINE}"/></a:solidFill></a:ln></c:spPr>'
        + _txpr(fontsize) +
        f'<c:crossAx val="{cross_ax}"/><c:crosses val="autoZero"/>'
        '<c:crossBetween val="between"/>'
        + (f'<c:majorUnit val="{major_unit}"/>' if major_unit else "") +
        "</c:valAx>")


def _cat_axis(ax_id, cross_ax, *, pos="b", reverse=False, crosses=None, fontsize=8,
              rotation=None, delete=False):
    tx = _txpr(fontsize)
    if rotation:
        tx = tx.replace('<a:bodyPr/>', f'<a:bodyPr rot="{rotation}"/>')
    # 🔴 反向纵轴时 `crosses=autoZero` 会把**类目轴顶到上方**（0 落在轴外，Word 贴最近端）
    cross = crosses or "autoZero"
    return (
        f'<c:catAx><c:axId val="{ax_id}"/>'
        f'<c:scaling><c:orientation val="{"maxMin" if reverse else "minMax"}"/></c:scaling>'
        f'<c:delete val="{1 if delete else 0}"/>' + f'<c:axPos val="{pos}"/>'
        '<c:majorTickMark val="none"/><c:minorTickMark val="none"/>'
        f'<c:tickLblPos val="{"none" if delete else "nextTo"}"/>'
        '<c:spPr><a:ln w="9525"><a:solidFill>'
        f'<a:srgbClr val="{AXIS_LINE}"/></a:solidFill></a:ln></c:spPr>'
        + tx + f'<c:crossAx val="{cross_ax}"/><c:crosses val="{cross}"/>'
        '<c:auto val="1"/><c:lblAlgn val="ctr"/><c:lblOffset val="100"/>'
        '<c:noMultiLvlLbl val="0"/></c:catAx>')


# ---------------------------------------------------------------- chartSpace 组装
def _title_xml(text, size=11, bold=True):
    if not text:
        return ""
    b = ' b="1"' if bold else ""
    return (
        "<c:title><c:tx><c:rich><a:bodyPr/><a:lstStyle/><a:p><a:pPr>"
        f'<a:defRPr sz="{int(size * 100)}"{b}/></a:pPr>'
        f'<a:r><a:rPr lang="zh-CN" sz="{int(size * 100)}"{b}/><a:t>{_esc(text)}</a:t></a:r>'
        "</a:p></c:rich></c:tx><c:layout/><c:overlay val=\"0\"/></c:title>")


def _legend_xml(pos, size=8):
    if not pos:
        return '<c:legend><c:delete val="1"/></c:legend>'
    return (
        f'<c:legend><c:legendPos val="{pos}"/><c:overlay val="0"/>'
        '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>' + _txpr(size) +
        "</c:legend>")


def chart_space_xml(plot_area, *, title=None, legend_pos="t", legend_size=8,
                    external=True, title_size=11):
    """组装完整 `c:chartSpace`。子元素顺序＝CT_ChartSpace 的 schema 顺序，错序 Word 报损坏。

    ⚠ `c:dTable`（原生数据表）属于 **plotArea** 的子元素，不能挂在这里 —— 由
    `line_chart_xml` / `combo_chart_xml` 拼进 `plot_area`（放错位置 Word 虽能容忍，但不合规）。
    """
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<c:chartSpace xmlns:c="{C_NS}" xmlns:a="{A_NS}" xmlns:r="{R_NS}">'
        '<c:date1904 val="0"/><c:lang val="zh-CN"/><c:roundedCorners val="0"/>'
        "<c:chart>" + _title_xml(title, title_size) +
        f'<c:autoTitleDeleted val="{0 if title else 1}"/>'
        + plot_area + _legend_xml(legend_pos, legend_size) +
        '<c:plotVisOnly val="1"/><c:dispBlanksAs val="gap"/></c:chart>'
        # 图表区无填充无边框（原件观感；Word 默认会给图表区描边）
        '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
        + ('<c:externalData r:id="rId1"><c:autoUpdate val="0"/></c:externalData>'
           if external else "") +
        "</c:chartSpace>")


# ---------------------------------------------------------------- 图型
def line_chart_xml(series, cats, *, reverse=False, numfmt="General", label_fmt=None,
                   major_unit=None, minimum=None, maximum=None, title=None,
                   legend_pos="t", cat_fontsize=8, val_fontsize=8, gridlines=True,
                   dtable=None):
    """折线图（多系列）。series: [{name, values, color, dash, marker, show_labels, hide_labels}]"""
    ser_xml = []
    for i, s in enumerate(series):
        ser_xml.append(_ser_line(
            i, s["name"], s["values"], cats, color=s.get("color", C_BLUE),
            label_fmt=s.get("label_fmt", label_fmt or numfmt),
            show_labels=s.get("show_labels", True), dash=s.get("dash"),
            marker=s.get("marker", "circle"), marker_size=s.get("marker_size", 5),
            marker_fill=s.get("marker_fill", "FFFFFF"), lw=s.get("lw", 12700),
            hide_labels=s.get("hide_labels", ()), smooth=s.get("smooth", False),
            val_fmt=s.get("val_fmt")))
    plot = (
        '<c:plotArea><c:layout/><c:lineChart>'
        '<c:grouping val="standard"/><c:varyColors val="0"/>'
        + "".join(ser_xml) +
        '<c:axId val="111111111"/><c:axId val="222222222"/>'
        "</c:lineChart>"
        + _cat_axis("111111111", "222222222",
                    crosses=("max" if reverse else "autoZero"), fontsize=cat_fontsize)
        + _val_axis("222222222", "111111111", reverse=reverse, numfmt=numfmt,
                    major_unit=major_unit, minimum=minimum, maximum=maximum,
                    gridlines=gridlines, fontsize=val_fontsize)
        + (dtable or "") +
        '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr></c:plotArea>')
    return chart_space_xml(plot, title=title, legend_pos=legend_pos)


def bar_chart_xml(series, cats, *, bar_dir="col", grouping="clustered", gap_width=150,
                  overlap=-27, numfmt="General", label_fmt=None, reverse=False,
                  major_unit=None, minimum=None, maximum=None, title=None,
                  legend_pos="r", cat_fontsize=8, val_fontsize=8, gridlines=True,
                  cat_rotation=None):
    """柱状图：`bar_dir` col（竖）/ bar（横）；`grouping` clustered / stacked / percentStacked。

    🔴 `gap_width` 的 schema 上限是 **500**（ST_GapAmount，0–500）。matplotlib 的
    `width=0.16` 换算成 gapWidth 是 (1−0.16)/0.16×100 = **525 > 500** —— 超范围会让
    **Word 拒绝打开整份文档**（2026-10-22 实测二分定位）。这里夹到 500（柱宽 1/6≈0.167，
    与 0.16 的观感差肉眼不可辨）；`overlap` 同理夹到 −100..100。
    """
    gap_width = max(0, min(500, int(round(gap_width))))
    overlap = max(-100, min(100, int(round(overlap))))
    ser_xml = []
    for i, s in enumerate(series):
        ser_xml.append(_ser_bar(
            i, s["name"], s["values"], cats, color=s.get("color", C_BLUE),
            label_fmt=s.get("label_fmt", label_fmt or numfmt),
            show_labels=s.get("show_labels", True),
            hide_labels=s.get("hide_labels", ())))
    plot = (
        '<c:plotArea><c:layout/><c:barChart>'
        f'<c:barDir val="{bar_dir}"/><c:grouping val="{grouping}"/>'
        '<c:varyColors val="0"/>'
        + "".join(ser_xml) +
        f'<c:gapWidth val="{gap_width}"/><c:overlap val="{overlap}"/>'
        '<c:axId val="111111111"/><c:axId val="222222222"/>'
        "</c:barChart>"
        + _cat_axis("111111111", "222222222", pos=("l" if bar_dir == "bar" else "b"),
                    reverse=reverse, crosses=("max" if reverse else "autoZero"),
                    fontsize=cat_fontsize, rotation=cat_rotation)
        + _val_axis("222222222", "111111111",
                    pos=("b" if bar_dir == "bar" else "l"),
                    reverse=reverse, numfmt=numfmt, major_unit=major_unit,
                    minimum=minimum, maximum=maximum, gridlines=gridlines,
                    fontsize=val_fontsize)
        + '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr></c:plotArea>')
    return chart_space_xml(plot, title=title, legend_pos=legend_pos)


def combo_chart_xml(groups, axes, cats, *, title=None, legend_pos="b", cat_fontsize=8,
                    val_fontsize=8, dtable=None):
    """组合图：多个 chart group 共用一套坐标轴（如「柱 + 红线基线」「柱 + 折线双轴」）。

    `groups`：[{"kind": "bar"|"line", "series": [...], "cat_ax": id, "val_ax": id, ...}]
       bar 额外认 `bar_dir` / `grouping` / `gap_width` / `overlap`；
    `axes`：  [{"id":…, "kind": "cat"|"val", "cross": 对侧轴 id, "pos":…, "reverse":…, …}]
       val 额外认 `numfmt` / `major_unit` / `minimum` / `maximum` / `gridlines`。

    坐标轴声明顺序＝XML 里 plotArea 子元素的顺序（schema 要求所有 chart group 在轴之前）。

    🔴 `id` 只是本函数内的**逻辑名**，写进 XML 前会换成数字：
    `c:axId/@val` 的类型是 `ST_AxId = xsd:unsignedInt`，写 `val="c1"` 会让
    **Word 拒绝打开整份文档**（2026-10-22 实测，最小组合图即复现）。
    """
    axnum = {a["id"]: str(111111111 * (i + 1)) for i, a in enumerate(axes)}
    body = []
    # 🔴 `c:idx`/`c:order` 在**整张图**里必须唯一（Word 自己生成的组合图是 bar 0,1 + line 2）。
    #    两个 group 各自从 0 开始计数 → Word 拒绝打开整份文档（2026-10-22 实测）。
    #    同时它也是内嵌工作簿的行号依据，所以必须与 `build_data_xlsx` 的系列顺序一致。
    gidx = 0
    for g in groups:
        ser = []
        for s in g["series"]:
            if g["kind"] == "bar":
                stacked = g.get("grouping", "clustered") in ("stacked", "percentStacked")
                ser.append(_ser_bar(gidx, s["name"], s["values"], cats,
                                    color=s.get("color", C_BLUE),
                                    label_fmt=s.get("label_fmt", "General"),
                                    show_labels=s.get("show_labels", False),
                                    hide_labels=s.get("hide_labels", ()),
                                    dlbl_pos=s.get("dlbl_pos",
                                                  "inBase" if stacked else "outEnd")))
            else:
                ser.append(_ser_line(gidx, s["name"], s["values"], cats,
                                     color=s.get("color", C_BLUE),
                                     label_fmt=s.get("label_fmt", "General"),
                                     show_labels=s.get("show_labels", False),
                                     dash=s.get("dash"), marker=s.get("marker", "none"),
                                     marker_size=s.get("marker_size", 5),
                                     marker_fill=s.get("marker_fill", "FFFFFF"),
                                     lw=s.get("lw", 12700),
                                     hide_labels=s.get("hide_labels", ()),
                                     val_fmt=s.get("val_fmt")))
            gidx += 1
        ax_ids = (f'<c:axId val="{axnum[g["cat_ax"]]}"/>'
                  f'<c:axId val="{axnum[g["val_ax"]]}"/>')
        if g["kind"] == "bar":
            gap = max(0, min(500, int(round(g.get("gap_width", 150)))))
            ovl = max(-100, min(100, int(round(g.get("overlap", -27)))))
            body.append('<c:barChart>'
                        f'<c:barDir val="{g.get("bar_dir", "col")}"/>'
                        f'<c:grouping val="{g.get("grouping", "clustered")}"/>'
                        '<c:varyColors val="0"/>' + "".join(ser) +
                        f'<c:gapWidth val="{gap}"/><c:overlap val="{ovl}"/>' + ax_ids +
                        "</c:barChart>")
        else:
            body.append('<c:lineChart><c:grouping val="standard"/><c:varyColors val="0"/>'
                        + "".join(ser) + ax_ids + "</c:lineChart>")

    ax_xml = []
    for a in axes:
        if a["kind"] == "cat":
            ax_xml.append(_cat_axis(axnum[a["id"]], axnum[a["cross"]], pos=a.get("pos", "b"),
                                    reverse=a.get("reverse", False),
                                    crosses=a.get("crosses"),
                                    fontsize=a.get("fontsize", cat_fontsize),
                                    rotation=a.get("rotation"),
                                    delete=a.get("delete", False)))
        else:
            ax_xml.append(_val_axis(axnum[a["id"]], axnum[a["cross"]], pos=a.get("pos", "l"),
                                    reverse=a.get("reverse", False),
                                    numfmt=a.get("numfmt", "General"),
                                    major_unit=a.get("major_unit"),
                                    minimum=a.get("minimum"), maximum=a.get("maximum"),
                                    gridlines=a.get("gridlines", True),
                                    tick_label=a.get("tick_label", True),
                                    fontsize=a.get("fontsize", val_fontsize)))
    plot = ('<c:plotArea><c:layout/>' + "".join(body) + "".join(ax_xml) +
            (dtable or "") +
            '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr></c:plotArea>')
    return chart_space_xml(plot, title=title, legend_pos=legend_pos)


    return chart_space_xml(plot, title=title, legend_pos=legend_pos)


def data_table_xml(horz=True, vert=True, outline=True, keys=False, fmt="General"):
    """图表下方的**原生数据表**（`c:dTable`，Excel「数据表」同款，原件趋势图用的就是这个）。"""
    return ('<c:dTable><c:showHorzBorder val="%d"/><c:showVertBorder val="%d"/>'
            '<c:showOutline val="%d"/><c:showKeys val="%d"/>'
            f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="700">'
            '<a:latin typeface="Times New Roman"/><a:ea typeface="宋体"/></a:defRPr>'
            '</a:pPr><a:endParaRPr lang="zh-CN"/></a:p></c:txPr>'
            "</c:dTable>") % (int(horz), int(vert), int(outline), int(keys))


def pie_chart_xml(series, cats, *, doughnut=False, hole=42, title=None, legend_pos="r",                  label_fmt="0.0%", show_legend_labels=True):
    """饼图 / 环形图（单系列）。"""
    s = series[0]
    ncat = len(cats)
    ser = (
        "<c:ser>"
        f'<c:idx val="0"/><c:order val="0"/>'
        f'<c:tx><c:strRef><c:f>{SHEET}!$A$2</c:f><c:strCache>'
        f'<c:ptCount val="1"/><c:pt idx="0"><c:v>{_esc(s["name"])}</c:v></c:pt>'
        "</c:strCache></c:strRef></c:tx>"
        '<c:spPr><a:ln w="9525"><a:solidFill><a:srgbClr val="FFFFFF"/>'
        "</a:solidFill></a:ln></c:spPr>"
        f'<c:dLbls><c:numFmt formatCode="{_esc(label_fmt)}" sourceLinked="0"/>' + _txpr(8) +
        '<c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/>'
        '<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/>'
        "</c:dLbls>"
        f"<c:cat><c:strRef><c:f>{SHEET}!$B$1:${_col(ncat + 1)}$1</c:f><c:strCache>"
        f"{_str_cache(cats)}</c:strCache></c:strRef></c:cat>"
        f'<c:val><c:numRef><c:f>{SHEET}!$B$2:${_col(ncat + 1)}$2</c:f>'
        f'<c:numCache>{_num_cache([float(v) for v in s["values"]])}</c:numCache>'
        "</c:numRef></c:val></c:ser>")
    tag = "doughnutChart" if doughnut else "pieChart"
    inner = f'<c:holeSize val="{hole}"/>' if doughnut else ""
    plot = (f'<c:plotArea><c:layout/><c:{tag}><c:varyColors val="1"/>' + ser + inner +
            "</c:" + tag + ">"
            '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr></c:plotArea>')
    return chart_space_xml(plot, title=title, legend_pos=legend_pos)


# ---------------------------------------------------------------- 纵轴自适应（复用 charts 口径）
def auto_axis(series, *, include=None, intervals=5, pad=0.08, lo_pad=None, hi_pad=None,
              floor_zero=False, zero_base=False, integer=False, clamp_zero=True,
              invert=False):
    """返回 `(lo, hi, step, ticks)` —— 与 `charts._auto_ylim` **同一套口径**。

    `charts._auto_ylim` 要一个 matplotlib `ax`（为了算 reserve_top 的像素），这里只需要
    数值轴范围，所以复刻同一段数学、并复用同一个 `charts._nice_step`。
    `tests/test_native_charts.py` 用真数据逐学科比对两者的刻度列表，防止两条路径漂移。
    """
    vals = []
    for s in series:
        vals.extend(float(v) for v in s if v is not None and v == v)
    if include is not None:
        vals.extend([float(include)] if not isinstance(include, (list, tuple)) else
                    [float(v) for v in include])
    if not vals:
        return None
    lo, hi = min(vals), max(vals)
    if zero_base:
        lo, hi = min(lo, 0.0), max(hi, 0.0)
    if floor_zero:
        lo = 0.0
    span = hi - lo
    if span <= 0:
        span = max(abs(hi), abs(lo), 1.0) * 0.2
        lo, hi = lo - span / 2, hi + span / 2
        span = hi - lo
    q_lo = lo if (floor_zero and lo_pad is None) else lo - span * (pad if lo_pad is None else lo_pad)
    q_hi = hi + span * (pad if hi_pad is None else hi_pad)
    if clamp_zero and lo >= 0:
        q_lo = max(q_lo, 0.0)
    step = charts._nice_step(q_hi - q_lo, intervals, integer)
    t_lo = round(step * int((q_lo / step) // 1), 10)
    t_hi = round(step * int(-((-q_hi / step) // 1)), 10)
    n = int(round((t_hi - t_lo) / step))
    ticks = [round(t_lo + i * step, 10) for i in range(n + 1)]
    return t_lo, t_hi, step, ticks


# ---------------------------------------------------------------- 插入 docx
def _next_partname(package, folder, base, ext):
    used = {p.partname for p in package.iter_parts()}
    i = 1
    while PackURI(f"/word/{folder}/{base}{i}.{ext}") in used:
        i += 1
    return PackURI(f"/word/{folder}/{base}{i}.{ext}")


def insert_chart_in_paragraph(p, xml_text, doc, *, w_cm=SLOT_W_CM, h_cm=SLOT_H_CM,
                              doc_pr_id=9001, name="原生图表", xlsx_blob=None):
    """把 chart 部件（+ 内嵌工作簿）挂到 doc 上，并在段落 p 末尾插 `wp:inline` 图形框。

    返回 `(rId, partname)`。
    """
    package = doc.part.package
    partname = _next_partname(package, "charts", "chart", "xml")
    validate_chart_xml(xml_text)
    part = Part(partname, CHART_CT, xml_text.encode("utf-8"), package)
    if xlsx_blob is not None:
        emb_name = _next_partname(package, "embeddings", "Microsoft_Excel_Worksheet", "xlsx")
        emb = Part(emb_name, XLSX_CT, xlsx_blob, package)
        part.relate_to(emb, RT.PACKAGE)          # → chartN.xml.rels 的 rId1
    rId = doc.part.relate_to(part, RT.CHART)
    cx, cy = int(w_cm * EMU_PER_CM), int(h_cm * EMU_PER_CM)
    drawing = (
        "<w:drawing "
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        f'xmlns:a="{A_NS}" xmlns:c="{C_NS}" xmlns:r="{R_NS}">'
        '<wp:inline distT="0" distB="0" distL="0" distR="0">'
        f'<wp:extent cx="{cx}" cy="{cy}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:docPr id="{doc_pr_id}" name="{_esc(name)}"/>'
        "<wp:cNvGraphicFramePr/>"
        "<a:graphic>"
        '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
        f'<c:chart r:id="{rId}"/>'
        "</a:graphicData></a:graphic></wp:inline></w:drawing>")
    p._p.append(parse_xml(drawing))
    return rId, partname


_MAX_GAP_WIDTH = 500           # ST_GapAmount 上限（超了 Word 拒开整份文档，2026-10-22 实测）
_LINE_ONLY_DLBLPOS = ("t", "b", "l", "r")   # 折线/散点专属；柱状图用会让 Word 拒开


def validate_chart_xml(xml):
    """插入前拦住三类会让 **Word 拒绝打开整份文档** 的写法。

    2026-10-22 实测（逐项二分，报告见 `docs/native-charts-第一段验证报告.md`）：
      1. 柱状/条形图的 `c:dLblPos` 用了折线专属位（t/b/l/r）——Word 直接报
         「Word 在读取该文件时遇到错误」，整份报告打不开，不是"图不显示"；
      2. `c:gapWidth` > 500（matplotlib `width=0.16` 换算出来是 525）；
      3. `formatCode="0.##"%"` 引号没转义 → 属性被截断。
    前两条这里查；第三条由 `_esc` 保证，顺手也查一遍"属性引号是否成对"。
    校验器很便宜（lxml 解析 + 正则），但能挡住"报告打不开"这种最贵的失败。
    """
    from lxml import etree  # noqa: PLC0415

    try:
        etree.fromstring(xml.encode("utf-8"))
    except etree.XMLSyntaxError as e:
        raise ValueError(f"图表 XML 不是合法 XML：{e}") from e

    m = re.search(r'<c:gapWidth val="(-?\d+)"/>', xml)
    if m and not (0 <= int(m.group(1)) <= _MAX_GAP_WIDTH):
        raise ValueError(f"gapWidth={m.group(1)} 超出 0–{_MAX_GAP_WIDTH}（Word 会拒绝打开文档）")

    if "<c:barChart>" in xml or "<c:bar3DChart>" in xml:
        for pos in re.findall(r'<c:dLblPos val="([^"]*)"/>', xml):
            if pos in _LINE_ONLY_DLBLPOS:
                raise ValueError(f"柱状图用了折线专属的 dLblPos=\"{pos}\""
                                 "（Word 会拒绝打开文档；柱状图请用 outEnd/inEnd/inBase/ctr）")
        if re.search(r'<c:grouping val="(?:stacked|percentStacked)"/>', xml):
            for pos in re.findall(r'<c:dLblPos val="([^"]*)"/>', xml):
                if pos == "outEnd":
                    raise ValueError("堆叠柱用了 dLblPos=\"outEnd\""
                                     "（Word 会拒绝打开文档；堆叠柱只认 ctr/inBase/inEnd）")

    for m in re.finditer(r'<c:numFmt[^>]*>', xml):
        if m.group(0).count('"') % 2:
            raise ValueError(f"numFmt 属性引号不成对（格式码里的 \" 没转义）：{m.group(0)}")
    return True


def count_chart_parts(doc):
    """文档里 chart 图形框的数量（含母版残留的原生图表）。"""
    from docx.oxml.ns import qn  # noqa: PLC0415
    n = 0
    for p in doc.paragraphs:
        n += len(p._element.findall(".//" + qn("c:chart")))
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    n += len(p._element.findall(".//" + qn("c:chart")))
    return n
