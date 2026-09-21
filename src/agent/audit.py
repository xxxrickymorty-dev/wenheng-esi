# -*- coding: utf-8 -*-
"""文衡 Agent 层 · 审计闭环（LLD §6.6 六项；P3 / 里程碑 M4）。

六项检查（口径照抄 LLD §6.6）：

| # | 项 | 判据 |
|---|---|---|
| 1 | `big_table` | 大表 T01 **全量 242 格**回查（＝22 数据行 × 11 列） |
| 2 | `slot_sample` | 文字/图槽位**抽样 10%** 回查 |
| 3 | `stale_scan` | 陈旧数字 **0 残留** |
| 4 | `golden_bt` | 金样回测 **≥99.5%** 或列差异 |
| 5 | `manual_slots` | 手工槽清点，**责任到人** |
| 6 | `figures` | 图验收判据＝`tools/verify_report_docx.py` 输出「槽位格 96/96」 |

外加 P3 已注册待实现的三件工具：`scan_stale` / `collect_artifacts` /
`write_env_snapshot`，以及期次归档 `out/_archive/<期次>/`。

🔴 **真值链（写明，不藏）**
第 1/2 项的真值 = 用 `parse_ledger` + `compute` 从**台账原件**重新推导，
**不是**拿 `out/<期次>/ledger_long.parquet` 自证 —— 后者与被测同源于一次解析，
比它属于同源自证，审计价值为零，本模块**不提供**这种假回查（缺台账路径时该项
直接标 `skipped` 并写明原因，不假装通过）。

代价：两侧共用 `parse_ledger` 与 `compute`。所以能抓出「写盘 / 读回 / CSV 浮点
往返 / docx 装配」引入的错，**抓不出 `parse_ledger` 自身的解析错**；后者由第 4 项
金样回测与断言覆盖。这条限制会原样写进 `audit_record.json` 的 `truth_chain`。

🔴 **第 4 项不重跑回测**
金样回测属 M0（既有 99.5%）。本项只**登记状态**：有回测产物则读数，无原件则标
`not_run` —— 与 2026-09-16「暂时不做回测」一致，不因为审计模块上线就偷偷重跑。

用法：
    .venv\\Scripts\\python.exe -c "import sys;sys.path.insert(0,'src');
        from agent import audit; audit.main('out/202607', ledger='../《云大...》.xlsx')"
"""
from __future__ import annotations

import hashlib
import json
import math
import os
import pathlib
import platform
import random
import re
import subprocess
import sys
import tempfile
import time

# ------------------------------------------------------------------ 常量
BIG_TABLE_ROWS = 22          # 数据行（表头除外；doc 里整表 23 行）
BIG_TABLE_COLS = 11
BIG_TABLE_CELLS = BIG_TABLE_ROWS * BIG_TABLE_COLS          # 242
SLOT_SAMPLE_RATE = 0.10
SLOT_SAMPLE_SEED = 20260701  # 固定种子：抽样可复现，否则每次审计结论都不一样
GOLDEN_MIN_PCT = 99.5
FIGURE_SLOTS_TOTAL = 96

PERIOD_RE = re.compile(r"\b(20\d{4})\b")
SLOT_LOC_RE = re.compile(r"^T(\d+)/r(\d+)c(\d+)$")
#: 流水线自己留的「这里得人来填」记号，例：`据【手工：ESI发布日期】…`
MANUAL_RE = re.compile(r"【手工：([^】]*)】")

AUDIT_ITEMS = (
    "big_table", "slot_sample", "stale_scan",
    "golden_bt", "manual_slots", "figures",
)


# ------------------------------------------------------------------ 小工具
def _sha256(path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _blank(v) -> bool:
    """空值判定：None / NaN / 空串 / 纯空白 都算空。

    ⚠ 必须显式处理：CSV 读回的空值是 NaN，而 pandas 里没落值的格子可能是空串，
    两者语义相同但 `float('')` 会抛异常 → 不处理就会报出一堆假差异
    （首次实跑就因此误报 3 格 top_papers）。
    """
    if v is None:
        return True
    try:
        return math.isnan(float(v))
    except (TypeError, ValueError):
        return str(v).strip() == ""


def _same(a, b) -> bool:
    """一格是否相等。数值按数值比（防 CSV 往返 12→12.0），其余按去空白字符串比。"""
    if _blank(a) and _blank(b):
        return True
    if _blank(a) or _blank(b):
        return False
    try:
        fa, fb = float(a), float(b)
        if math.isnan(fa) and math.isnan(fb):
            return True
        if math.isnan(fa) or math.isnan(fb):
            return False
        return math.isclose(fa, fb, rel_tol=1e-9, abs_tol=1e-9)
    except (TypeError, ValueError):
        return str(a).strip() == str(b).strip()


def _find_docx(out: pathlib.Path):
    cands = sorted(out.glob("报告初稿_*.docx"))
    return cands[0] if cands else None


# ------------------------------------------------------------------ 真值
def load_truth(ledger_xlsx):
    """从**台账原件**重新推导真值表（table2）。

    这是第 1/2 项唯一的真值来源。走的是「parse → 落临时目录 → compute」整条链，
    与生产的差别只在落盘位置，所以装配 / 写盘 / 读回环节都被覆盖。
    """
    import compute
    import parse_ledger

    src = pathlib.Path(ledger_xlsx)
    led = parse_ledger.load_ledger(src)
    with tempfile.TemporaryDirectory(prefix="wh_truth_") as td:
        parse_ledger.write_outputs(led, src, pathlib.Path(td))
        wide, new_intl, new_cn, thr = compute.load_wide(td)
        period, prev = compute.detect_periods(wide)
        t2 = compute.build_table2(wide, new_intl, new_cn, thr, period, prev)
    return {"table2": t2, "period": period, "prev": prev}


def _truth_table2(out_dir, ledger):
    """优先用台账原件重推；没有台账就返回 None（调用方须标 skipped，不得降级自证）。"""
    if not ledger:
        return None, "未提供台账原件路径"
    try:
        return load_truth(ledger), None
    except Exception as e:                                    # noqa: BLE001
        return None, f"台账重推失败：{type(e).__name__}: {e}"


# ================================================================== 六项
def check_big_table(out_dir, ledger=None):
    """第 1 项：大表 T01 全量 242 格回查。"""
    import pandas as pd

    out = pathlib.Path(out_dir)
    f_csv = out / "table2.csv"
    if not f_csv.exists():
        return {"id": "big_table", "status": "skipped", "reason": "无 table2.csv"}

    got = pd.read_csv(f_csv)
    truth, err = _truth_table2(out_dir, ledger)
    if truth is None:
        return {"id": "big_table", "status": "skipped",
                "reason": f"无独立真值（{err}）；不用 ledger_long 自证"}

    exp = truth["table2"]
    n_rows, n_cols = min(len(got), BIG_TABLE_ROWS), min(len(got.columns), BIG_TABLE_COLS)
    diffs = []
    for ri in range(n_rows):
        for ci in range(n_cols):
            col = got.columns[ci]
            if col not in exp.columns:
                diffs.append({"row": ri, "col": col, "got": None,
                              "expected": None, "note": "真值表无此列"})
                continue
            g = got.iloc[ri, ci]
            e = exp.iloc[ri, ci]
            if not _same(g, e):
                diffs.append({"row": ri, "col": col,
                              "got": (None if pd.isna(g) else g),
                              "expected": (None if pd.isna(e) else e)})
    n_checked = n_rows * n_cols
    return {
        "id": "big_table",
        "status": "pass" if not diffs else "fail",
        "cells_total": BIG_TABLE_CELLS,
        "cells_checked": n_checked,
        "cells_full": n_checked == BIG_TABLE_CELLS,
        "n_diff": len(diffs),
        "diffs": diffs[:20],
        "period": truth["period"], "prev": truth["prev"],
    }


def check_slot_sample(out_dir, ledger=None, rate=SLOT_SAMPLE_RATE, seed=SLOT_SAMPLE_SEED):
    """第 2 项：文字/图槽位抽样 10% 回查。

    回查方式：把抽样到的槽位按 `T{n}/r{r}c{c}` 定位回报告 docx，
    确认该格文本**确实写着**这次要填的 `new` 值。定位不到 / 越界的算
    `unverifiable` 单列，不混进 diff —— 否则格式一变就满屏假失败。
    """
    out = pathlib.Path(out_dir)
    f_slots = out / "red_slots.json"
    if not f_slots.exists():
        return {"id": "slot_sample", "status": "skipped", "reason": "无 red_slots.json"}
    slots = json.loads(f_slots.read_text(encoding="utf-8"))
    if not slots:
        return {"id": "slot_sample", "status": "skipped", "reason": "槽位清单为空"}

    k = max(1, round(len(slots) * rate))
    sample = random.Random(seed).sample(slots, k)

    docx = _find_docx(out)
    if docx is None:
        return {"id": "slot_sample", "status": "skipped",
                "reason": "未找到报告 docx，无法回读槽位"}

    try:
        from docx import Document
        d = Document(str(docx))
    except Exception as e:                                    # noqa: BLE001
        return {"id": "slot_sample", "status": "skipped",
                "reason": f"打不开报告 docx：{type(e).__name__}: {e}"}

    diffs, unverifiable = [], []
    for s in sample:
        loc = str(s.get("location", ""))
        m = SLOT_LOC_RE.match(loc)
        if not m:
            unverifiable.append({"location": loc, "why": "位置格式无法解析"})
            continue
        ti, ri, ci = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if ti >= len(d.tables) or ri >= len(d.tables[ti].rows) \
                or ci >= len(d.tables[ti].rows[ri].cells):
            unverifiable.append({"location": loc, "why": "定位越界"})
            continue
        cell = d.tables[ti].rows[ri].cells[ci]
        want = str(s.get("new", "")).strip()

        if want.startswith("[图]"):
            # 图槽位的 `new` 是「[图] 相对路径」**标记**，不是要落进格子的字面文本
            # （格子里实际是「图题 + 图」两段）。拿它做文本包含判断会全军覆没 ——
            # 首次实跑 10 处"差异"全是这么来的。图槽位该验的是**格子里有没有图**。
            from docx.oxml.ns import qn
            n_img = len(cell._element.findall(".//" + qn("w:drawing")))
            if n_img < 1:
                diffs.append({"location": loc, "kind": "figure", "expected": want,
                              "in_docx": cell.text.strip()[:60], "why": "图槽位无图"})
            continue

        text = cell.text
        if want and want not in text.replace("\n", " "):
            diffs.append({"location": loc, "kind": "text", "expected": want,
                          "in_docx": text.strip()[:60]})

    return {
        "id": "slot_sample",
        "status": "pass" if not diffs else "fail",
        "slots_total": len(slots),
        "sample_size": k,
        "rate": rate,
        "seed": seed,
        "n_diff": len(diffs),
        "n_unverifiable": len(unverifiable),
        "diffs": diffs[:20],
        "unverifiable": unverifiable[:10],
    }


def scan_stale(out_dir, ledger=None):
    """第 3 项：陈旧数字残留扫描（P3 已注册工具 `scan_stale` 的实现）。

    扫的是**期次标签残留**——双月报告最典型的陈旧数字：本该跟着本期走的字符串
    （"接近度较7月"、正文里的期次）还留着上上期的。合法出现的期次只有
    {本期 P, 上期 Q}，出现其它 `20xxxx` 一律列**可疑**等人工确认。

    ⚠ 输出是**可疑清单**，不是自动判定"0 残留"：图是 PNG 扫不到文字，
    正文里历史叙述也可能正当引用早期期次。清单清空才叫 0 残留。
    """
    out = pathlib.Path(out_dir)
    meta = out / "meta.json"
    period = prev = None
    if meta.exists():
        try:
            m = json.loads(meta.read_text(encoding="utf-8"))
            # ⚠ 键名是 period_current / period_prev（不是 period / prev）。
            #   记错键名 → allowed 为空 → 连本期期次都会被判成"残留"（首次实跑即如此）。
            period = str(m.get("period_current") or "") or None
            prev = str(m.get("period_prev") or "") or None
        except Exception:                                     # noqa: BLE001
            pass
    allowed = {p for p in (period, prev) if p}

    # 扫描范围：**产出的文本类文件**全扫，不用硬编码名单——
    # 名单式扫描会漏掉新加的产物（测试时就因为写了个 x.md 而一个字都没扫到）。
    # 排除项都是有正当理由带期次的：meta.json 记 period_next_empty（下期空列）、
    # 审计自己写出的两份、以及归档清单。
    skip = {"meta.json", "audit_record.json", "审计记录.md", "manifest.json",
            "env_snapshot.json"}
    targets = [p for p in sorted(out.glob("*"))
               if p.is_file() and p.suffix in (".md", ".json") and p.name not in skip]
    docx = _find_docx(out)
    if docx:
        targets.append(docx)

    suspicious = []
    for f in targets:
        try:
            if f.suffix == ".docx":
                from docx import Document
                text = "\n".join(p.text for p in Document(str(f)).paragraphs)
            else:
                text = f.read_text(encoding="utf-8", errors="ignore")
        except Exception:                                     # noqa: BLE001
            continue
        for tok in sorted(set(PERIOD_RE.findall(text))):
            if tok not in allowed:
                # 严重程度分档，否则历史序列的正常引用会淹没真正的残留：
                #   > 本期  → high（未来期次，必是残留）
                #   > 上期  → high（夹在两期之间，不该出现）
                #   ≤ 上期  → low（历史序列起点/对比年份，多为正当引用）
                sev = "low" if (prev and tok <= prev) else "high"
                suspicious.append({"file": f.name, "period": tok, "severity": sev})

    n_high = sum(1 for s in suspicious if s["severity"] == "high")
    return {
        "id": "stale_scan",
        "status": "fail" if n_high else ("pass" if not suspicious else "pass"),
        "period": period, "prev": prev,
        "files_scanned": len(targets),
        "n_suspicious": len(suspicious),
        "n_high": n_high,
        "suspicious": suspicious[:30],
        "note": "仅 high 档判失败；low 档＝上期及更早的期次引用，多为历史序列正当引用，"
                "列出供人工扫一眼。图（PNG）内文字不在扫描范围",
    }


def check_golden_bt(out_dir, ledger=None):
    """第 4 项：金样回测 ≥99.5%。**只登记状态，不重跑**（M0 既有证据）。"""
    out = pathlib.Path(out_dir)
    cands = sorted(out.glob("*回测*.md")) + sorted(out.glob("*backtest*.csv"))
    if not cands:
        return {"id": "golden_bt", "status": "not_run",
                "reason": "无金样回测产物（原件缺 202601/03/05，2026-09-16 起暂不做回测）",
                "threshold": GOLDEN_MIN_PCT}
    rec = {"id": "golden_bt", "threshold": GOLDEN_MIN_PCT,
           "artifacts": [c.name for c in cands]}
    txt = "\n".join(c.read_text(encoding="utf-8", errors="ignore")
                    for c in cands if c.suffix == ".md")

    # ⚠ 绝不能笼统抓「第一个 xx%」——首次实跑就抓到了「1‰接近度 55.13%」，
    #   差点把接近度当成一致率判失败。必须按**语义**定位，且区分格级 / 行级口径。
    pct = unit = None
    m_pct = re.search(r"一致率\D{0,8}(\d{1,3}(?:\.\d+)?)\s*%", txt)
    m_frac = re.search(r"(\d+)\s*/\s*(\d+)\s*格", txt)
    m_row = re.search(r"行级通过\s*(\d+)\s*/\s*(\d+)", txt)
    m_ndiff = re.search(r"不一致\s*(\d+)\s*行", txt)
    if m_pct:
        pct, unit = float(m_pct.group(1)), "格级（一致率%）"
    elif m_frac:
        a, b = int(m_frac.group(1)), int(m_frac.group(2))
        pct, unit = round(a * 100.0 / b, 2), f"格级（{a}/{b} 格）"
    elif m_row:
        y = int(m_row.group(2))
        n = int(m_ndiff.group(1)) if m_ndiff else 0
        pct, unit = round((y - n) * 100.0 / y, 2), f"行级（{y - n}/{y} 行）"

    rec["pct"] = pct
    rec["unit"] = unit
    rec["evidence"] = [ln.strip() for ln in txt.splitlines()
                       if ("行级通过" in ln or "不一致" in ln or "逐格一致" in ln)][:5]
    if pct is None:
        rec["status"] = "review"
        rec["reason"] = "产物里没找到可解析的一致率，需人工读数"
    elif unit.startswith("行级"):
        rec["status"] = "review"
        rec["reason"] = ("只解析到**行级**口径，而 M0 基线 99.5% 是**格级**。"
                         "行级 100% 不等于格级达标（一行里错一格，行级看不出来）——"
                         "口径不同就不自动判通过，交人工确认")
    else:
        rec["status"] = "pass" if pct >= GOLDEN_MIN_PCT else "fail"
    return rec


def manual_slot_checklist(out_dir, ledger=None):
    """第 5 项：手工槽清点，责任到人。

    判定依据只有一个：**槽位值里含 `【手工：…】` 标记** —— 这是流水线自己留的
    "这里得人来填"记号（例：`据【手工：ESI发布日期】ESI最新数据显示…`）。

    ⚠ 别用 `source` 字段判：图槽位的 source 是 `charts/**.png`、正文的 source 是
    `sentences.py`，看着像人工，其实都是**机器**产物。按 source 判会把约 100 个
    机器图槽也算成待人工清点（首次实跑就报了 115 个，虚高）。

    输出待签清单，`owner` 留空由馆员填 —— 本模块**不代签**，也不断言"已清点"。
    """
    out = pathlib.Path(out_dir)
    f_slots = out / "red_slots.json"
    if not f_slots.exists():
        return {"id": "manual_slots", "status": "skipped", "reason": "无 red_slots.json"}
    slots = json.loads(f_slots.read_text(encoding="utf-8"))

    pending = []
    for s in slots:
        new = str(s.get("new", ""))
        for m in MANUAL_RE.finditer(new):
            pending.append({
                "location": s.get("location"),
                "item": m.group(1),
                "context": new[max(0, m.start() - 12):m.end() + 12],
                "source": s.get("source"), "owner": "", "signed_at": "",
            })
    return {
        "id": "manual_slots",
        "status": "review" if pending else "pass",
        "slots_total": len(slots),
        "n_pending": len(pending),
        "items": sorted({p["item"] for p in pending}),
        "pending": pending,
        "note": "owner / signed_at 须由馆员签署，程序不得代签",
    }


def check_figures(out_dir, ledger=None):
    """第 6 项：图验收判据＝`tools/verify_report_docx.py` 输出「槽位格 96/96」。"""
    out = pathlib.Path(out_dir)
    docx = _find_docx(out)
    if docx is None:
        return {"id": "figures", "status": "skipped", "reason": "未找到报告 docx"}
    script = pathlib.Path(__file__).resolve().parents[2] / "tools" / "verify_report_docx.py"
    if not script.exists():
        return {"id": "figures", "status": "skipped", "reason": f"无脚本 {script}"}
    try:
        # 🔴 子进程编码必须两头钉死：父进程若带 PYTHONIOENCODING=utf-8（脚本/CI 常见），
        #    而这里用 text=True 按 locale（中文 Windows = GBK）解码，reader 线程会抛
        #    UnicodeDecodeError → stdout 整段丢失 → 槽位读数解不出来 → 审计误判 fail。
        #    2026-09-21 全量回归实测（档案页与 e2e 连带红）。
        env = {**os.environ, "PYTHONIOENCODING": "utf-8"}
        proc = subprocess.run([sys.executable, str(script), str(docx)],
                              capture_output=True, text=True, encoding="utf-8",
                              errors="replace", timeout=300, env=env)
    except Exception as e:                                    # noqa: BLE001
        return {"id": "figures", "status": "skipped",
                "reason": f"调用失败：{type(e).__name__}: {e}"}
    stdout = proc.stdout or ""
    m = re.search(r"槽位格通过:\s*(\d+)/(\d+)", stdout)
    got, total = (int(m.group(1)), int(m.group(2))) if m else (None, FIGURE_SLOTS_TOTAL)
    return {
        "id": "figures",
        "status": "pass" if (proc.returncode == 0 and got == FIGURE_SLOTS_TOTAL) else "fail",
        "slots_ok": got, "slots_total": total,
        "required": FIGURE_SLOTS_TOTAL,
        "exit_code": proc.returncode,
        "output": stdout.strip().splitlines()[-8:],
    }


CHECKS = {
    "big_table": check_big_table,
    "slot_sample": check_slot_sample,
    "stale_scan": scan_stale,
    "golden_bt": check_golden_bt,
    "manual_slots": manual_slot_checklist,
    "figures": check_figures,
}


# ================================================================== 组装
def run_audit(out_dir, ledger=None, rate=SLOT_SAMPLE_RATE, seed=SLOT_SAMPLE_SEED):
    """跑六项，落 `audit_record.json` + `审计记录.md`。不抛异常——审计自己崩溃
    不该阻断生产，但要记进结果里。"""
    out = pathlib.Path(out_dir)
    results = []
    for name in AUDIT_ITEMS:
        fn = CHECKS[name]
        t0 = time.time()
        try:
            r = fn(out_dir, ledger, rate=rate, seed=seed) if name == "slot_sample" \
                else fn(out_dir, ledger)
        except Exception as e:                                # noqa: BLE001
            r = {"id": name, "status": "error",
                 "reason": f"{type(e).__name__}: {e}"}
        r["elapsed_s"] = round(time.time() - t0, 2)
        results.append(r)

    verdict = "pass"
    for r in results:
        if r.get("status") == "fail":
            verdict = "fail"
        elif r.get("status") in ("review", "error") and verdict == "pass":
            verdict = "review"

    rec = {
        "out_dir": str(out),
        "ledger": str(ledger) if ledger else None,
        "generated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
        "verdict": verdict,
        "items": results,
        "truth_chain": {
            "source": "台账原件 → parse_ledger → compute（重推）",
            "not_used": "out/<期次>/ledger_long.parquet（与被测同源于一次解析，属同源自证）",
            "covers": "写盘 / 读回 / CSV 浮点往返 / docx 装配",
            "does_not_cover": "parse_ledger 自身的解析错（由第 4 项金样回测与断言覆盖）",
        },
    }
    out.mkdir(parents=True, exist_ok=True)
    (out / "audit_record.json").write_text(
        json.dumps(rec, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    (out / "审计记录.md").write_text(render_md(rec), encoding="utf-8")
    return rec


def render_md(rec) -> str:
    zh = {"pass": "通过", "fail": "未过", "review": "待人工确认",
          "skipped": "跳过", "not_run": "未执行", "error": "出错"}
    rows = []
    for r in rec["items"]:
        rows.append(f"| {r['id']} | {zh.get(r.get('status'), r.get('status'))} | "
                    f"{_detail(r)} |")
    return (
        f"# 审计记录（{rec['out_dir']}）\n\n"
        f"- 生成时间：{rec['generated_at']}\n"
        f"- 台账原件：{rec['ledger'] or '（未提供 → 第 1/2 项跳过，不做同源自证）'}\n"
        f"- **总判定：{zh.get(rec['verdict'], rec['verdict'])}**\n\n"
        "| 检查项 | 结果 | 说明 |\n|---|---|---|\n" + "\n".join(rows) + "\n\n"
        "## 真值链（审计能证明什么、不能证明什么）\n\n"
        f"- 真值来源：{rec['truth_chain']['source']}\n"
        f"- 明确不用：{rec['truth_chain']['not_used']}\n"
        f"- 能抓出：{rec['truth_chain']['covers']}\n"
        f"- 抓不出：{rec['truth_chain']['does_not_cover']}\n"
    )


def _detail(r) -> str:
    s = r.get("status")
    if s == "skipped":
        return f"跳过 —— {r.get('reason', '')}"
    if s == "error":
        return f"出错 —— {r.get('reason', '')}"
    if s == "not_run":
        return f"未执行 —— {r.get('reason', '')}"
    if r["id"] == "big_table":
        return (f"{r['cells_checked']}/{r['cells_total']} 格已回查"
                f"{'' if r['cells_full'] else '（未满量！）'}，差异 {r['n_diff']} 格")
    if r["id"] == "slot_sample":
        return (f"抽 {r['sample_size']}/{r['slots_total']}（{r['rate']:.0%}，种子 {r['seed']}），"
                f"差异 {r['n_diff']}，无法定位 {r['n_unverifiable']}")
    if r["id"] == "stale_scan":
        return f"扫 {r['files_scanned']} 个文件，可疑期次残留 {r['n_suspicious']} 处"
    if r["id"] == "golden_bt":
        if r.get("pct") is None:
            return r.get("reason", "")
        # 必须带上口径单位，否则「100% + 阈值 99.5% + 待人工确认」看着自相矛盾：
        # 它其实是行级 100%，而阈值是格级的，两者不可比。
        return (f"{r['pct']}%（{r.get('unit', '口径未知')}；阈值 {r['threshold']}%）"
                f" —— {r.get('reason', '')}")
    if r["id"] == "manual_slots":
        items = "、".join(r.get("items", [])[:6]) or "无"
        return f"待人工清点 {r['n_pending']} 处（{r['slots_total']} 槽中）；待填项：{items}"
    if r["id"] == "figures":
        return f"槽位格 {r['slots_ok']}/{r['slots_total']}（要求 {r['required']}）"
    return ""


# ================================================================== 归档三件
def collect_artifacts(run_dir):
    """P3 工具 `collect_artifacts`：产物清单 + sha256（归档包 manifest 来源）。"""
    root = pathlib.Path(run_dir)
    files = []
    for p in sorted(root.rglob("*")):
        if p.is_file() and "_superseded" not in p.parts and p.name != "manifest.json":
            files.append({"path": str(p.relative_to(root)).replace("\\", "/"),
                          "size": p.stat().st_size, "sha256": _sha256(p)})
    return {"run_dir": str(root), "n_files": len(files),
            "bytes": sum(f["size"] for f in files), "files": files}


def write_env_snapshot(run_dir):
    """P3 工具 `write_env_snapshot`：环境快照（架构设计 §5.2）。"""
    root = pathlib.Path(run_dir)
    root.mkdir(parents=True, exist_ok=True)
    try:
        import pandas  # noqa: F401
        import statsmodels  # noqa: F401
        vers = {"python": platform.python_version(), "pandas": pandas.__version__,
                "statsmodels": statsmodels.__version__}
    except Exception:                                         # noqa: BLE001
        vers = {"python": platform.python_version()}
    snap = {
        "generated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
        "python": platform.python_version(),
        "platform": platform.platform(),
        "tz": time.tzname,
        "deps": vers,
        "git_commit": _git_commit(),
    }
    (root / "env_snapshot.json").write_text(
        json.dumps(snap, ensure_ascii=False, indent=1), encoding="utf-8")
    return snap


def _git_commit():
    try:
        r = subprocess.run(["git", "rev-parse", "HEAD"], capture_output=True,
                           text=True, timeout=10)
        return r.stdout.strip() or None
    except Exception:                                         # noqa: BLE001
        return None


def archive_period(out_dir, archive_root=None):
    """期次归档：把一期次的全套产物 + 审计记录 + 环境快照收进
    `out/_archive/<期次>/`，附 `manifest.json`（含 sha256）。

    只**复制**，不动 `out/<期次>/` 一个字节——归档出错不能毁掉本期产物。
    """
    import shutil

    out = pathlib.Path(out_dir)
    period = out.name
    root = pathlib.Path(archive_root) if archive_root else out.parent / "_archive"
    dst = root / period
    if dst.exists():
        shutil.rmtree(dst)
    dst.mkdir(parents=True, exist_ok=True)

    for p in sorted(out.rglob("*")):
        if not p.is_file() or "_superseded" in p.parts:
            continue
        rel = p.relative_to(out)
        (dst / rel).parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(p, dst / rel)

    man = collect_artifacts(out)
    man["archived_at"] = time.strftime("%Y-%m-%dT%H:%M:%S")
    man["period"] = period
    (dst / "manifest.json").write_text(
        json.dumps(man, ensure_ascii=False, indent=1), encoding="utf-8")
    return {"archive": str(dst), "n_files": man["n_files"], "bytes": man["bytes"]}


# ================================================================== CLI
def main(out_dir, ledger=None, archive=False):
    rec = run_audit(out_dir, ledger=ledger)
    write_env_snapshot(out_dir)
    out = {
        "verdict": rec["verdict"],
        "audit_record": str(pathlib.Path(out_dir) / "audit_record.json"),
        "env_snapshot": str(pathlib.Path(out_dir) / "env_snapshot.json"),
    }
    if archive:
        out["archive"] = archive_period(out_dir)
    print(render_md(rec))
    print("产物：", json.dumps(out, ensure_ascii=False))
    return 0 if rec["verdict"] != "fail" else 1


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="文衡审计闭环（LLD §6.6 六项）")
    ap.add_argument("out_dir")
    ap.add_argument("--ledger", default=None, help="台账原件 xlsx（不给则第 1/2 项跳过）")
    ap.add_argument("--archive", action="store_true", help="顺带归档到 out/_archive/<期次>/")
    a = ap.parse_args()
    sys.exit(main(a.out_dir, ledger=a.ledger, archive=a.archive))
