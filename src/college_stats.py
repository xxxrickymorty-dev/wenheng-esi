# -*- coding: utf-8 -*-
"""顶尖论文「学院统计」：读馆员的《<期次>论文清单.xlsx》→ 每学院发文量/被引频次。

清单由来（馆员每期提供）：列 = 序号 | WOS收录号 | 被引频次 | ESI学科 | 出版年 | 云大作者 |
热点论文/研究前沿；「云大作者」形如 `Zhu, Lei//生命科学学院` 或 `Chen, Wen/通讯作者/地球科学学院`
（单双斜杠混用）→ 取最后一段为学院/机构，再按**映射表**归一。

产物：
  `out/<期次>/college_stats.json`        报告取数用（charts 与 fill_docx 共用）
  `out/<期次>/institution_review.md/.json` 签认清单（未归一 / 可疑写法；只给人看，不进数据路径）

学院名归一（2026-10-22 起）
────────────────────────────────────────────────────────────────
`原写法 → 规范名` 不再写死在代码里，而是 `config/college_map.yaml`：
每条带 `by`（确认人）/ `at`（确认时间）/ `basis`（依据），可审、可签认、可追认。
流程：**机器提可疑 → 人签认 → 写回映射表 → 重跑**（同 DISCCA「机构匹配」页的分工，
不是它的那套页面）。机器只做「可解释的规则」（去括号 / 全半角 / 编辑距离 / 包含），
不许用模型猜——红线：**运行期 LLM 不进数据路径**，模型给的建议只能落在
`institution_suggestions.md` 里，人签认后才生效（见本文件「AI 建议」一节）。

改口径＝改本文件的 `_PARSER_VERSION` **加**映射表的 `parser_version`，两处一起 +1，
旧缓存才会失效重解析（2026-09-24 踩过：加了别名没动版本，报告里天文所仍是两行）。
"""
from __future__ import annotations

import argparse
import difflib
import json
import pathlib
import re
import unicodedata
from collections import Counter

import openpyxl
import yaml

CAND_DIRS = ("input/{period}", "../数据与复现", "out/{period}", "{ledger_dir}")

REPO_ROOT = pathlib.Path(__file__).resolve().parents[1]

#: 映射表（数据文件，入库）：`原写法 → 规范名` + 确认人/时间/依据。
MAP_PATH = REPO_ROOT / "config" / "college_map.yaml"

#: 映射表缺席时的**兜底**（新 clone 少了 config/ 也不至于把天文所拆成两行）。
#: 表在时以表为准；两者必须一致，由 `tests/test_institution_map.py` 钉住。
UNIT_ALIAS = {
    "天文研究所": "中国（西南）天文研究所",
}

#: 解析器版本：改动 `parse()` 的口径（别名表、取段规则…）就 +1 ——
#: 缓存 `college_stats.json` 版本不匹配会**重新解析**，避免"改了规则但旧缓存照用"
#: （2026-09-24 实测踩过：加了别名，报告里天文所仍是两行）。
#: ⚠ 改映射表同样要 +1，且要与表里的 `parser_version` 相等（`load_map` 会核对）。
_PARSER_VERSION = 2

REVIEW_MD = "institution_review.md"
REVIEW_JSON = "institution_review.json"
SUGGEST_MD = "institution_suggestions.md"

#: 缺省确认人/时间（映射表可覆盖）：见 config/college_map.yaml 的 default_by/default_at。
_PENDING_MARK = "待"


# ──────────────────────────────────────────────────────────── 映射表（数据）
def load_map(path=None) -> dict:
    """读 `config/college_map.yaml` → 归一后的映射表（含每条的有效确认人/时间/依据）。

    返回：{"path","source","version","updated","owner","notes",
           "aliases": {原写法: 规范名}, "canonical": [规范名…], "rows": [每条明细]}

    · 表缺席 → 退回内置 `UNIT_ALIAS`（`source="builtin"`），不抛异常；
    · 表在但 `parser_version` 与 `_PARSER_VERSION` 不等 → 抛 ValueError（口径没对齐，
      宁可不跑也不静默用旧缓存）；表在但缺 `by`/`basis` → 抛 ValueError（没签认不许进表）。
    """
    p = pathlib.Path(path) if path is not None else MAP_PATH
    if not p.exists():
        rows = [{"college": v, "by": "内置兜底（UNIT_ALIAS）", "at": "", "basis": "内置兜底",
                 "aliases": [{"form": k, "by": "内置兜底（UNIT_ALIAS）", "at": "",
                              "basis": "内置兜底"}]} for k, v in UNIT_ALIAS.items()]
        return {"path": None, "source": "builtin", "version": _PARSER_VERSION, "updated": "",
                "owner": "", "notes": [], "aliases": dict(UNIT_ALIAS),
                "canonical": sorted(set(UNIT_ALIAS.values())), "rows": rows}

    doc = yaml.safe_load(p.read_text(encoding="utf-8")) or {}
    ver = doc.get("parser_version")
    if ver != _PARSER_VERSION:
        raise ValueError(
            f"映射表 {p} 声明 parser_version={ver!r}，代码 _PARSER_VERSION={_PARSER_VERSION}；"
            f"改映射表要两处一起 +1（旧缓存才会失效重解析）")
    d_by, d_at = doc.get("default_by") or "", doc.get("default_at") or ""
    aliases: dict[str, str] = {}
    canonical: list[str] = []
    rows: list[dict] = []
    for e in doc.get("colleges") or []:
        name = str(e.get("college") or "").strip()
        if not name:
            raise ValueError(f"映射表 {p}：有一条缺 college 名")
        canonical.append(name)
        row = {"college": name, "by": str(e.get("by") or d_by).strip(),
               "at": str(e.get("at") or d_at).strip(),
               "basis": str(e.get("basis") or "").strip(), "aliases": []}
        if not row["basis"] or not row["by"]:
            raise ValueError(f"映射表 {p}：{name} 缺 by/basis（没签认的依据不许进表）")
        for a in e.get("aliases") or []:
            form = str(a.get("form") or "").strip()
            if not form:
                raise ValueError(f"映射表 {p}：{name} 下有一条别名缺 form")
            if form == name:
                raise ValueError(f"映射表 {p}：{name} 下别名与规范名相同（多余的一条）")
            if form in aliases and aliases[form] != name:
                raise ValueError(
                    f"映射表 {p}：写法「{form}」同时映射到 {aliases[form]} 与 {name}（冲突）")
            al = {"form": form, "by": str(a.get("by") or d_by).strip(),
                  "at": str(a.get("at") or d_at).strip(),
                  "basis": str(a.get("basis") or "").strip()}
            if not al["basis"] or not al["by"]:
                raise ValueError(f"映射表 {p}：别名「{form}」缺 by/basis")
            aliases[form] = name
            row["aliases"].append(al)
        rows.append(row)
    return {"path": _rel(p), "source": "file", "version": ver,
            "updated": str(doc.get("updated") or ""), "owner": str(doc.get("owner") or ""),
            "notes": [str(n) for n in (doc.get("notes") or [])],
            "aliases": aliases, "canonical": canonical, "rows": rows}


def is_pending(by: str) -> bool:
    """确认人里带「待」＝还没签认（沿用 repair_rules.yaml「未授权不执行」的口径）。"""
    return _PENDING_MARK in (by or "")


def _rel(p: pathlib.Path) -> str:
    """落进产物里的路径写法：仓库内走相对（`config/college_map.yaml`），
    仓库外照实给（换台机器不成立的路径不硬编成相对，口径见 `src/display_path.py`）。"""
    try:
        return pathlib.Path(p).resolve().relative_to(REPO_ROOT).as_posix()
    except (ValueError, OSError):
        return pathlib.Path(p).as_posix()


def normalize_unit(unit: str, m: dict) -> str:
    """一个写法 → 规范名（不在表里就原样返回）。"""
    return (m.get("aliases") or {}).get(unit, unit)


# ──────────────────────────────────────────────────────────── 解析
def find_list(ledger_dir: pathlib.Path, period: str) -> pathlib.Path | None:
    """按优先级找清单：input/<期次>/ → ../数据与复现/ → out/<期次>/。"""
    for d in (pathlib.Path("input") / period, pathlib.Path("..") / "数据与复现",
              pathlib.Path(ledger_dir), pathlib.Path(ledger_dir).parent / period):
        if not d.is_dir():
            continue
        hits = sorted(q for q in d.glob("*论文清单*.xlsx") if not q.name.startswith("~$"))
        if hits:
            return hits[0]
    return None


def _col_index(hdr: list[str], *keys: str) -> int | None:
    for i, h in enumerate(hdr):
        if any(k in h for k in keys):
            return i
    return None


def unit_of(cell: object) -> str:
    """「云大作者」单元格 → 末段写法（未归一）。单双斜杠混用都认。"""
    s = re.sub(r"\s+", "", str(cell or ""))
    parts = [x for x in re.split(r"/+", s) if x]
    return parts[-1] if parts else ""


def _cites_of(cell: object) -> int:
    try:
        return int(float(cell))                                  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return 0


def _read_rows(path: pathlib.Path) -> tuple[list[str], list[list]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    hdr = [str(c).strip() if c is not None else "" for c in rows[0]]
    return hdr, [r for r in rows[1:] if r is not None and any(c is not None for c in r)]


def tally(rows: list[list], hdr: list[str]) -> dict:
    """清单行 → 写法统计（**未归一**的原写法）：parse 与签认清单共用这一处口径。

    ⚠ 与旧口径逐字对齐：末段为空的那些行**照样**计进 `papers[""]`（因此 `n_papers`
    含它们），只是不进 `units`——改动这里会让 `college_stats.json` 与现产对不上。
    """
    ia = _col_index(hdr, "作者")
    ic = _col_index(hdr, "被引频次")
    ie = _col_index(hdr, "学科")
    if ia is None or ic is None:
        raise ValueError(f"清单表头认不出「云大作者」/「被引频次」列：{hdr}")
    papers, cites, samples = Counter(), Counter(), {}
    no_unit: list[str] = []
    for r in rows:
        raw = unit_of(r[ia] if ia < len(r) else "")
        if not raw:
            no_unit.append(str(r[ia])[:40] if ia < len(r) else "")
        papers[raw] += 1
        cites[raw] += _cites_of(r[ic] if ic < len(r) else 0)
        if raw:
            samples.setdefault(raw, str(r[ia])[:60] if ia < len(r) else "")
    return {"ia": ia, "ic": ic, "ie": ie, "papers": papers, "cites": cites,
            "samples": samples, "no_unit": no_unit}


def units_of(papers: Counter, cites: Counter, m: dict) -> list[dict]:
    """写法统计 → 归一后的单元列表（与 `college_stats.json` 的 `units` 同口径）。

    同一机构的多种写法在这里并成一行；排序固定为（发文量降序, 被引降序）。
    """
    merged_p, merged_c = Counter(), Counter()
    for raw in papers:
        u = normalize_unit(raw, m)
        merged_p[u] += papers[raw]
        merged_c[u] += cites[raw]
    return sorted(({"unit": u, "papers": merged_p[u], "cites": merged_c[u]}
                   for u in merged_p if u), key=lambda x: (-x["papers"], -x["cites"]))


def parse(path: pathlib.Path, *, map_path=None) -> dict:
    """清单 → 学院统计（口径见模块注释）。返回结构与 `college_stats.json` 逐字段一致。"""
    m = load_map(map_path)
    hdr, rows = _read_rows(pathlib.Path(path))
    t = tally(rows, hdr)
    units = units_of(t["papers"], t["cites"], m)
    rows_out = [[("" if c is None else str(c)) for c in r[:7]] for r in rows]
    n_subjects = 0
    if t["ie"] is not None:
        n_subjects = len({re.sub(r"\s+", "", str(r[t["ie"]]))
                          for r in rows if r[t["ie"]] not in (None, "")})
    return {"parser": _PARSER_VERSION, "source": str(path), "units": units,
            "rows": rows_out, "header": hdr[:7],
            "n_papers": sum(t["papers"].values()), "n_cites": sum(t["cites"].values()),
            "n_units": len(units), "n_subjects": n_subjects, "no_unit": t["no_unit"][:10]}


# ──────────────────────────────────────────────────────────── 签认清单
#: 可疑判定规则（**可解释**，不许用模型猜）。每条给出「为什么怀疑」。
RULES = {
    "R1": "去括号后同名",
    "R2": "括号内的写法同名",
    "R3": "只差全半角/空白/标点",
    "R4": "疑似错拼（编辑距离）",
    "R5": "疑似截断（一方是另一方的子串）",
}

_FUZZY_MIN = 0.80        # R4 相似度阈值
_SHORT_MIN = 4           # R4/R5 里短的那一方至少 4 字（挡住「医学院 vs 药学院」这类）


def _norm(s: str) -> str:
    """归一化键：全角→半角、去掉空白与标点（只用于比对，不改数据）。"""
    return re.sub(r"[\s\W_]+", "", unicodedata.normalize("NFKC", s or ""))


def _core(s: str) -> str:
    """去掉括号及其内容后的核心名（「软件学院（人工智能学院）」→「软件学院」）。"""
    return _norm(re.sub(r"[（(][^（()）]*[)）]", "", s or ""))


def _parens(s: str) -> list[str]:
    return [_norm(x) for x in re.findall(r"[（(]([^（()）]*)[)）]", s or "")]


def _rule_of(a: str, b: str) -> str | None:
    """两个写法之间的可疑规则（按精确→模糊排序，返回第一条命中的）。"""
    na, nb = _norm(a), _norm(b)
    if not na or not nb:
        return None
    if na == nb:
        return "R3"
    if _core(a) == _core(b):
        return "R1"
    pa, pb = _parens(a), _parens(b)
    if nb in pa or na in pb:
        return "R2"
    short, long_ = (na, nb) if len(na) <= len(nb) else (nb, na)
    if len(short) >= _SHORT_MIN and short in long_:
        return "R5"
    if (min(len(na), len(nb)) >= _SHORT_MIN
            and abs(len(na) - len(nb)) <= 3
            and difflib.SequenceMatcher(None, na, nb).ratio() >= _FUZZY_MIN):
        return "R4"
    return None


def suspects(papers: Counter, m: dict) -> list[dict]:
    """写法两两比对 → 可疑对（含规则名、篇数、是否已由映射表签认）。"""
    forms = sorted(f for f in papers if f)          # 空写法（无末段）不参与两两比对
    canonical = set(m.get("canonical") or [])
    out = []
    for i, a in enumerate(forms):
        for b in forms[i + 1:]:
            if a in canonical and b in canonical:
                continue                      # 两条都是规范名 → 不提示合并（如 云南大学/附属医院）
            rule = _rule_of(a, b)
            if rule is None:
                continue
            signed = normalize_unit(a, m) == normalize_unit(b, m)
            out.append({"rule": rule, "why": RULES[rule], "a": a, "b": b,
                        "a_papers": papers[a], "b_papers": papers[b],
                        "signed": signed,
                        "target": normalize_unit(a, m) if signed else ""})
    out.sort(key=lambda d: (d["signed"], -min(d["a_papers"], d["b_papers"]),
                            d["rule"], d["a"], d["b"]))
    return out


def review_of(state: dict, *, period: str = "", map_path=None) -> dict:
    """现产统计 + 映射表 → 签认清单数据（未归一 / 可疑 / 已签认命中 / 待追认）。

    只读：不改 `state`，也不参与报告取数。
    """
    m = load_map(map_path)
    hdr = [str(h) for h in (state.get("header") or [])]
    rows = [list(r) for r in (state.get("rows") or [])]
    try:
        t = tally(rows, hdr)
    except ValueError:                                # 表头认不出 → 空清单，照实报
        t = {"papers": Counter(), "cites": Counter(), "samples": {}, "no_unit": []}
    papers, cites = t["papers"], t["cites"]
    canonical = set(m.get("canonical") or [])
    items = []
    for raw in sorted((f for f in papers if f), key=lambda x: (-papers[x], x)):
        if raw in m.get("aliases", {}):
            status = "alias"
        elif raw in canonical:
            status = "canonical"
        else:
            status = "unmapped"
        items.append({"form": raw, "papers": papers[raw], "cites": cites[raw],
                      "status": status, "target": normalize_unit(raw, m),
                      "sample": t["samples"].get(raw, "")})
    recomputed = [(u["unit"], u["papers"], u["cites"]) for u in units_of(papers, cites, m)]
    now = [(u.get("unit"), u.get("papers"), u.get("cites"))
           for u in (state.get("units") or [])]
    hit = [{"form": it["form"], "target": it["target"], "papers": it["papers"],
            "cites": it["cites"], **{k: v for k, v in _alias_row(it["form"], m).items()
                                     if k in ("by", "at", "basis")}}
           for it in items if it["status"] == "alias"]
    pending = [{"kind": "条目", "form": r["college"], "by": r["by"], "at": r["at"],
                "basis": r["basis"]} for r in m["rows"] if is_pending(r["by"])]
    for r in m["rows"]:
        for a in r["aliases"]:
            if is_pending(a["by"]):
                pending.append({"kind": "别名", "form": a["form"], "by": a["by"],
                                "at": a["at"], "basis": a["basis"]})
    return {
        "period": period, "source": str(state.get("source") or ""),
        "parser": state.get("parser"), "n_papers": state.get("n_papers"),
        "n_cites": state.get("n_cites"), "n_units": state.get("n_units"),
        "map": {"path": m.get("path"), "source": m.get("source"), "version": m.get("version"),
                "updated": m.get("updated"), "owner": m.get("owner"), "notes": m["notes"],
                "n_canonical": len(canonical), "n_aliases": len(m.get("aliases") or {})},
        "items": items, "suspects": suspects(papers, m), "hits": hit, "pending": pending,
        "no_unit": t["no_unit"][:10],
        "unseen_aliases": sorted(k for k in (m.get("aliases") or {}) if k not in papers),
        "consistent": recomputed == now,
    }


def _alias_row(form: str, m: dict) -> dict:
    for r in m["rows"]:
        for a in r["aliases"]:
            if a["form"] == form:
                return a
    return {}


def _disp(p) -> str:
    """给人看的路径写法：仓库内的走相对（`config/college_map.yaml`），
    仓库外/别的盘照实给绝对路径（口径见 `src/display_path.py`）。"""
    try:
        import display_path                                          # noqa: PLC0415

        s = str(p or "")
        if s and display_path.is_absolute(s):
            return display_path.to_repo_relative(s)
        return s.replace("\\", "/")
    except Exception:                                                # noqa: BLE001
        return str(p or "")


def render_review_md(rev: dict) -> str:
    """签认清单 → Markdown（给馆员看；行序固定，同一输入必得同一份）。"""
    mp = rev["map"]
    unmapped = [it for it in rev["items"] if it["status"] == "unmapped"]
    susp = rev["suspects"]
    L = []
    L.append(f"# {rev['period']} 学院/机构名签认清单")
    L.append("")
    L.append("> 自动生成、**可重算**（同一「论文清单 + 映射表」必得同一份）；本文件只给人看，"
             "不参与报告取数。")
    L.append(f"> 清单：`{_disp(rev['source'])}`　映射表：`{_disp(mp['path'])}`"
             f"（version {mp['version']} · updated {mp['updated']}）")
    L.append(f"> 规范名 {mp['n_canonical']} 个 · 本期写法 {len(rev['items'])} 个 · "
             f"待签认 {len(unmapped)} 条 · 可疑 {len(susp)} 对"
             f"（已签认 {sum(1 for s in susp if s['signed'])}）")
    L.append("")
    if not rev["consistent"]:
        L.append("## ⚠ 缓存与映射表不一致")
        L.append("")
        L.append("按当前映射表重算的单元与 `college_stats.json` 现产**不一致** —— 缓存来自旧口径。"
                 "按映射表表头的次序做：改 `parser_version` 与 `_PARSER_VERSION` 两处一起 +1 → 重跑。")
        L.append("")
    L.append("## 一、待签认（未命中映射表的写法）")
    L.append("")
    if unmapped:
        L.append("| # | 写法 | 论文量 | 被引 | 机器建议的候选 | 样例（云大作者原文） |")
        L.append("|---|---|---:|---:|---|---|")
        for i, it in enumerate(unmapped, 1):
            cands = [s for s in susp if it["form"] in (s["a"], s["b"])]
            hint = "、".join(sorted({s["b"] if s["a"] == it["form"] else s["a"]
                                     for s in cands})) or "—"
            L.append(f"| {i} | {it['form']} | {it['papers']} | {it['cites']} | {hint} "
                     f"| `{it['sample']}` |")
    else:
        L.append("本期没有未归一的写法 ✅（映射表覆盖了清单里出现的全部写法）")
    L.append("")
    L.append("## 二、可疑写法（机器只提可解释的规则，签不签认由人定）")
    L.append("")
    if susp:
        L.append("| # | 规则 | 写法 A | 写法 B | 机器判断 | 状态 |")
        L.append("|---|---|---|---|---|---|")
        for i, s in enumerate(susp, 1):
            if s["signed"]:
                state = f"已签认 → {s['target']}"
            else:
                state = "待签认"
            L.append(f"| {i} | {s['rule']} {s['why']} | {s['a']}（{s['a_papers']} 篇） "
                     f"| {s['b']}（{s['b_papers']} 篇） | {'同一机构' if s['signed'] else '疑似同一机构'} "
                     f"| {state} |")
    else:
        L.append("没有可疑写法。")
    L.append("")
    L.append("## 三、已签认映射的本期命中")
    L.append("")
    if rev["hits"]:
        L.append("| 原写法 | → 规范名 | 论文量 | 被引 | 确认人 | 确认时间 |")
        L.append("|---|---|---:|---:|---|---|")
        for h in rev["hits"]:
            L.append(f"| {h['form']} | {h['target']} | {h['papers']} | {h['cites']} "
                     f"| {h['by']} | {h['at']} |")
    else:
        L.append("本期没有命中别名的写法。")
    L.append("")
    if rev["unseen_aliases"]:
        L.append(f"（映射表里本期未出现的别名 {len(rev['unseen_aliases'])} 条："
                 f"{'、'.join(rev['unseen_aliases'])}）")
        L.append("")
    if rev["no_unit"]:
        L.append(f"（另有 {len(rev['no_unit'])} 行「云大作者」列取不出机构末段，未计入单元："
                 f"{'、'.join(f'`{x}`' for x in rev['no_unit'][:3])}…）")
        L.append("")
    L.append("## 四、待签认人追认的既有映射")
    L.append("")
    if rev["pending"]:
        n_all = mp["n_canonical"] + mp["n_aliases"]
        if len(rev["pending"]) >= n_all:
            L.append(f"**整表 {n_all} 行都还没签认**（确认人里带「待」）。两种签法："
                     f"整表一次追认（把 `config/college_map.yaml` 的 `default_by` 改成签认人、"
                     f"`default_at` 改成当天）；或逐条改每行的 `by`/`at`。")
            L.append("")
        L.append("| 类别 | 写法/规范名 | 确认人 | 确认时间 | 依据（全文见映射表） |")
        L.append("|---|---|---|---|---|")
        for p in rev["pending"]:
            basis = p["basis"] if len(p["basis"]) <= 48 else p["basis"][:47] + "…"
            L.append(f"| {p['kind']} | {p['form']} | {p['by']} | {p['at']} | {basis} |")
        L.append("")
        L.append("追认＝把「确认人」改成签认人姓名、把时间改成签认当天，"
                 "并同步 `parser_version` 与 `_PARSER_VERSION`（两处一起 +1）。")
    else:
        L.append("全部映射都已签认 ✅")
    L.append("")
    if mp["notes"]:
        L.append("## 五、口径备注（映射表 notes）")
        L.append("")
        L.extend(f"- {n}" for n in mp["notes"])
        L.append("")
    L.append("## 六、怎么签认")
    L.append("")
    L.append("1. 打开 `config/college_map.yaml`，在对应 `college` 下补一条 `aliases`：")
    L.append("   `form`（清单里的原写法）+ `by`（确认人）+ `at`（确认时间）+ `basis`（依据）；")
    L.append("2. 把文件里的 `parser_version` 与 `src/college_stats.py` 的 `_PARSER_VERSION` **一起 +1**；")
    L.append("3. 重跑该期（`uv run python src\\run_pipeline.py --period <期次>` 或看板一键生成）；")
    L.append("4. 本文件里该写法会从「一、待签认」移到「三、已签认命中」，"
             "`college_stats.json` 里两行并成一行。")
    L.append("")
    return "\n".join(L)


def write_review(ledger_dir, state: dict, *, period: str = "", map_path=None,
                 rev: dict | None = None) -> pathlib.Path:
    """写 `institution_review.md` + `.json`（确定性：同一输入逐字节一致）。"""
    d = pathlib.Path(ledger_dir)
    d.mkdir(parents=True, exist_ok=True)
    rev = rev if rev is not None else review_of(state, period=period, map_path=map_path)
    (d / REVIEW_MD).write_text(render_review_md(rev), encoding="utf-8")
    (d / REVIEW_JSON).write_text(
        json.dumps(rev, ensure_ascii=False, indent=1, sort_keys=False), encoding="utf-8")
    return d / REVIEW_MD


# ──────────────────────────────────────────────────────────── AI 建议（默认关；整段可摘除）
# 🔴 红线：运行期 LLM **不进数据路径**。
#   本节的产物只是一份**建议文件**（`institution_suggestions.md`），
#   它不参与 `college_stats.json` 的任何数字；建议要由人签认写回
#   `config/college_map.yaml` 之后才生效（流程与「机器提可疑 → 人签认」同一条）。
#   默认关：`SUGGESTIONS_ENABLED = False`（或 CLI 不带 `--suggest`）时，
#   整段代码不会被调用，全流程不依赖它。
#   要**整段摘除**：删本节 → 删 `_write_review_and_maybe_suggest` 末尾那两行
#   → 删 CLI 的 `--suggest` 参数与两处透传（都带 `# AI 建议` 标记）。
#   将来接模型时，唯一进出点是 `src/ai/_attempt.py`（见 CONTEXT.md 红线 1），
#   调用方把 `suggest_fn` 传进来即可，本节不直接调模型。
SUGGESTIONS_ENABLED = False


def suggestions_of(rev: dict, *, suggest_fn=None) -> list[dict]:
    """对「待签认」的写法给候选规范名（默认只用可解释的规则，不调模型）。"""
    canonical = [it["target"] for it in rev["items"] if it["status"] == "canonical"]
    canonical += [h["target"] for h in rev["hits"]]
    canonical = sorted(set(c for c in canonical if c))
    out = []
    for it in rev["items"]:
        if it["status"] != "unmapped":
            continue
        cands = []
        for s in rev["suspects"]:
            if it["form"] == s["a"]:
                cands.append((s["b"], f"{s['rule']} {s['why']}"))
            elif it["form"] == s["b"]:
                cands.append((s["a"], f"{s['rule']} {s['why']}"))
        close = difflib.get_close_matches(_norm(it["form"]), [_norm(c) for c in canonical],
                                          n=3, cutoff=0.6)
        by_norm = {_norm(c): c for c in canonical}
        for c in close:
            cands.append((by_norm[c], "R6 与规范名相近（编辑距离）"))
        out.append({"form": it["form"], "papers": it["papers"], "cites": it["cites"],
                    "candidates": sorted(set(cands)), "source": "规则"})
    if suggest_fn is not None:                      # 模型建议（经 src/ai/ 边界注入）
        for s in suggest_fn(out) or []:
            for row in out:
                if row["form"] == s.get("form"):
                    row["candidates"] = sorted(set(row["candidates"])
                                               | {(str(c), "模型建议") for c in s.get("candidates") or []})
                    row["source"] = "规则+模型"
    return out


def render_suggestions_md(rev: dict, sugg: list[dict]) -> str:
    L = [f"# {rev['period']} 学院/机构名归一 · 建议（仅供参考）", "",
         "> 🔴 **本文件不进数据路径**：`college_stats.json` 的数字与它无关。",
         "> 建议要由人签认写回 `config/college_map.yaml`（补 `by`/`at`/`basis`）后才生效。",
         "> 默认关：`--suggest` 或 `SUGGESTIONS_ENABLED = True` 才产出本文件。", ""]
    if not sugg:
        L.append("本期没有待签认的写法，无需建议 ✅")
        return "\n".join(L) + "\n"
    L += ["| 写法 | 论文量 | 被引 | 候选规范名 | 来源 |", "|---|---:|---:|---|---|"]
    for s in sugg:
        cands = "；".join(f"{c}（{w}）" for c, w in s["candidates"]) or "—"
        L.append(f"| {s['form']} | {s['papers']} | {s['cites']} | {cands} | {s['source']} |")
    return "\n".join(L) + "\n"


def write_suggestions(ledger_dir, rev: dict, *, suggest_fn=None) -> pathlib.Path:
    """写 `institution_suggestions.md`（只有开关打开时才被调用）。"""
    d = pathlib.Path(ledger_dir)
    d.mkdir(parents=True, exist_ok=True)
    p = d / SUGGEST_MD
    p.write_text(render_suggestions_md(rev, suggestions_of(rev, suggest_fn=suggest_fn)),
                 encoding="utf-8")
    return p
# ────────────────────────────────────────────────────── AI 建议结束（以上整段可摘除）


# ──────────────────────────────────────────────────────────── 读缓存 / 落盘
def load(ledger_dir: pathlib.Path, period: str, *, save: bool = False,
         map_path=None, suggest: bool = False) -> dict | None:
    """取该期学院统计：缓存命中直接用，口径不匹配则重解析。

    `save=True`（真实运行）时顺带刷新 `institution_review.md/.json` 签认清单。
    """
    ld = pathlib.Path(ledger_dir)
    j = ld / "college_stats.json"
    if j.exists():
        try:
            st = json.loads(j.read_text(encoding="utf-8"))
            if st.get("parser") == _PARSER_VERSION:
                if save:
                    _write_review_and_maybe_suggest(ld, st, period, map_path, suggest)
                return st
        except Exception:                                          # noqa: BLE001
            pass
    p = find_list(ld, period)
    if p is None:
        return None
    st = parse(p, map_path=map_path)
    if save:
        ld.mkdir(parents=True, exist_ok=True)
        j.write_text(json.dumps(st, ensure_ascii=False, indent=1), encoding="utf-8")
        _write_review_and_maybe_suggest(ld, st, period, map_path, suggest)
    return st


def _write_review_and_maybe_suggest(ld: pathlib.Path, st: dict, period: str,
                                    map_path, suggest: bool) -> None:
    rev = review_of(st, period=period, map_path=map_path)
    write_review(ld, st, period=period, map_path=map_path, rev=rev)
    if suggest or SUGGESTIONS_ENABLED:            # AI 建议：默认关，关掉时全流程不依赖它
        write_suggestions(ld, rev)


# ──────────────────────────────────────────────────────────── 命令行
def _check_against(state: dict, other: dict) -> list[str]:
    """两个统计逐字段比对 → 差异行（字段名：现产 vs 重建）。"""
    diffs = []
    for k in sorted(set(state) | set(other)):
        a, b = state.get(k), other.get(k)
        if a != b:
            diffs.append(f"{k}: 现产={_brief(a)} != 重建={_brief(b)}")
    return diffs


def _brief(v) -> str:
    s = json.dumps(v, ensure_ascii=False)
    return s if len(s) <= 120 else s[:117] + "…"


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="学院统计：解析 / 重建 / 比对 / 签认清单")
    ap.add_argument("ledger_dir", nargs="?", default="out/202609", help="期次产物目录")
    ap.add_argument("--source", help="指定论文清单（默认按 find_list 找）")
    ap.add_argument("--out", help="统计写到别的 json（默认写回 ledger_dir/college_stats.json）")
    ap.add_argument("--check", help="只解析该清单并与 ledger_dir 现产逐字段比对，不落盘")
    ap.add_argument("--ignore-source", action="store_true",
                    help="比对时忽略 source 字段（同一份清单换个路径时用）")
    ap.add_argument("--review-only", action="store_true", help="只用现产缓存重写签认清单")
    ap.add_argument("--map", help="换一张映射表（演练用；默认 config/college_map.yaml）")
    ap.add_argument("--suggest", action="store_true", help="额外产建议文件（默认关）")
    a = ap.parse_args(argv)
    ld = pathlib.Path(a.ledger_dir)
    per = re.search(r"(\d{6})", ld.name)
    per = per.group(1) if per else ""

    if a.check:
        now = json.loads((ld / "college_stats.json").read_text(encoding="utf-8"))
        other = parse(pathlib.Path(a.check), map_path=a.map)
        if a.ignore_source:
            now = {k: v for k, v in now.items() if k != "source"}
            other = {k: v for k, v in other.items() if k != "source"}
        diffs = _check_against(now, other)
        print(f"现产：{now.get('n_papers')} 篇 / {now.get('n_cites')} 引 / "
              f"{now.get('n_units')} 个单元　重建：{other['n_papers']} 篇 / "
              f"{other['n_cites']} 引 / {other['n_units']} 个单元")
        if a.ignore_source:
            print("（按 --ignore-source 跳过 source 字段：同一份清单换个路径时用）")
        if diffs:
            print(f"差异 {len(diffs)} 处：")
            for d in diffs:
                print("  " + d)
            return 1
        print("逐字段一致：0 差异 ✅")
        return 0

    if a.review_only:
        st = json.loads((ld / "college_stats.json").read_text(encoding="utf-8"))
        rev = review_of(st, period=per, map_path=a.map)
        p = write_review(ld, st, period=per, map_path=a.map, rev=rev)
        if a.suggest or SUGGESTIONS_ENABLED:
            write_suggestions(ld, rev)
        print(f"签认清单 → {p}")
        return 0

    if a.source or a.out:
        src = pathlib.Path(a.source) if a.source else find_list(ld, per)
        if src is None:
            raise SystemExit(f"没找到 {per or ld} 的论文清单（可用 --source 指定）")
        st = parse(src, map_path=a.map)
        dst = pathlib.Path(a.out) if a.out else (ld / "college_stats.json")
        dst.parent.mkdir(parents=True, exist_ok=True)
        dst.write_text(json.dumps(st, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"清单 {st['source']}: {st['n_papers']} 篇 / {st['n_subjects']} 学科 / "
              f"{st['n_units']} 个学院机构 / 被引 {st['n_cites']} → {dst}")
        return 0

    st = load(ld, per, save=True, map_path=a.map, suggest=a.suggest)
    if st is None:
        raise SystemExit("没找到论文清单")
    print(f"清单 {st['source']}: {st['n_papers']} 篇 / {st['n_subjects']} 学科 / "
          f"{st['n_units']} 个学院机构 / 被引 {st['n_cites']}")
    for u in st["units"][:8]:
        print(f"   {u['unit']:24s} {u['papers']:>3} 篇  {u['cites']:>6} 次")
    return 0


if __name__ == "__main__":
    import sys

    sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
