# -*- coding: utf-8 -*-
"""监测版 Word 导出：消费 report._report_data() 的同一份数据装配，出简洁公文版式。

数字与 HTML 同源（共用 _report_data，无第二套换算）；导出不调模型（路由侧
insight=None，缺省即无解读段）。v1 取舍：图表（svg）不放 docx，趋势章保留章首注
与进位时间线文字。落盘 monitor_data/<校名>/<期次>.docx，已存在直接返回（幂等）。
"""
from __future__ import annotations

import pathlib

import docx
from docx.shared import Pt

from monitor import report as mr


def _small(doc, text: str) -> None:
    """注脚式小字段落（meta／note／页脚）。"""
    p = doc.add_paragraph(text)
    for r in p.runs:
        r.font.size = Pt(9.5)


def _table(doc, t: dict) -> None:
    """table 规格 → Table Grid 表（表头行加粗；空表落 placeholder 合并行）。"""
    tab = doc.add_table(rows=1, cols=len(t["head"]))
    tab.style = "Table Grid"
    for i, head in enumerate(t["head"]):
        tab.rows[0].cells[i].paragraphs[0].add_run(head).bold = True
    if not t["rows"] and t.get("placeholder"):
        row = tab.add_row()
        merged = row.cells[0].merge(row.cells[-1])
        merged.paragraphs[0].add_run(t["placeholder"])
        return
    for vals in t["rows"]:
        cells = tab.add_row().cells
        for i, v in enumerate(vals):
            cells[i].text = str(v)


def export(cn: str, period: str, cache: dict, insight: dict | None = None) -> pathlib.Path:
    """导出一所学校的 Word 版监测报告，返回文件路径（已存在直接返回，幂等）。"""
    d = mr._report_data(cn, period, cache, insight)
    out = mr.MONITOR_ROOT / cn / f"{period}.docx"
    if out.exists():
        return out
    doc = docx.Document()
    doc.add_heading(f'{d["cn"]} · 学科监测报告 · {d["period"]}', level=0)
    _small(doc, d["meta"])
    for n in d["nums"]:
        p = doc.add_paragraph()
        p.add_run(f'{n["label"]}：').bold = True
        p.add_run(n["value"] + (n.get("suffix") or ""))
        if n.get("note"):
            p.add_run(f'（{n["note"]}）')
    if d.get("absent_note"):
        _small(doc, d["absent_note"])
    for sec in d["sections"]:
        doc.add_heading(f'{sec["num"]}、{sec["title"]}', level=1)
        if sec.get("table"):
            _table(doc, sec["table"])
        for note in sec.get("notes", []):
            _small(doc, note)
        # 图表（svg）v1 不放 docx：趋势章的章首注随 notes、进位时间线随 timeline 保留
        tl = sec.get("timeline")
        if tl:
            p = doc.add_paragraph()
            p.add_run(tl["cap"]).bold = True
            for item in tl["items"]:
                doc.add_paragraph(item)
    ins = d.get("insight")
    if ins and (ins.get("text") or "").strip():
        doc.add_heading("六、解读", level=1)
        for pp in [x for x in ins["text"].split("\n\n") if x.strip()]:
            doc.add_paragraph(pp)
        _small(doc, mr.insight_how(ins.get("mode")) + "。")
    _small(doc, "派生算法：" + " · ".join(d["algorithms"]) + "。")
    _small(doc, d["footer"]["source"])
    _small(doc, d["footer"]["engine"])
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
