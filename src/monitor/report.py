# -*- coding: utf-8 -*-
"""监测版报告引擎（最简版）：只吃 ESI 全局件，产出单文件 HTML 站位报告。

设计对应《监测版包契约-ESIonly（2026-10-31）》：全局件对任何学校一字不差，
所以**按期次建一次缓存**（机构总表＋22 榜本校行＋阈值＋±3 邻近窗），渲染毫秒级。
缓存与产物都落在 `monitor_data/`（不在任何只读数据红线区内，源数据只读）。
本模块零 LLM、零网络：数字全部可回溯到包内原始表。
"""
from __future__ import annotations

import datetime
import html
import json
import pathlib
import re

import openpyxl

from monitor import schools as sch

WENHENG_ROOT = pathlib.Path(__file__).resolve().parents[2]
ESI_ROOT = WENHENG_ROOT.parent / "数据与复现" / "ESI数据"
MONITOR_ROOT = WENHENG_ROOT / "monitor_data"

STEMS = ["Agricultural Sciences", "Biology & Biochemistry", "Chemistry", "Clinical Medicine",
         "Computer Science", "Economics & Business", "Engineering", "Environment Ecology",
         "Geosciences", "Immunology", "Materials Science", "Mathematics", "Microbiology",
         "Molecular Biology & Genetics", "Multidisciplinary", "Neuroscience & Behavior",
         "Pharmacology & Toxicology", "Physics", "Plant & Animal Science",
         "Psychiatry Psychology", "Social Sciences, General", "Space Science"]
CN = {"Agricultural Sciences": "农业科学", "Biology & Biochemistry": "生物学与生物化学",
      "Chemistry": "化学", "Clinical Medicine": "临床医学", "Computer Science": "计算机科学",
      "Economics & Business": "经济与商业", "Engineering": "工程学",
      "Environment Ecology": "环境生态学", "Geosciences": "地球科学", "Immunology": "免疫学",
      "Materials Science": "材料科学", "Mathematics": "数学", "Microbiology": "微生物学",
      "Molecular Biology & Genetics": "分子生物学与遗传学", "Multidisciplinary": "多学科",
      "Neuroscience & Behavior": "神经系统科学与行为学",
      "Pharmacology & Toxicology": "药理学与毒理学", "Physics": "物理学",
      "Plant & Animal Science": "植物与动物学", "Psychiatry Psychology": "精神病学与心理学",
      "Social Sciences, General": "社会科学", "Space Science": "空间科学"}

_norm = lambda s: " ".join(str(s).upper().replace("/", " ").split())
_f = lambda x: float(x)


def latest_period() -> str | None:
    """最新一个带机构总表的 ESI 期次目录；没有返回 None。"""
    if not ESI_ROOT.is_dir():
        return None
    ds = sorted((d.name for d in ESI_ROOT.iterdir()
                 if re.fullmatch(r"20\d{4}", d.name) and (d / "Institutions.xlsx").exists()),
                reverse=True)
    return ds[0] if ds else None


def cache_path(period: str) -> pathlib.Path:
    return MONITOR_ROOT / f"_cache_{period}.json"


def build_cache(period: str) -> dict:
    """按期次建缓存：一次读盘，全名单受益。产物写入 monitor_data/。"""
    base = ESI_ROOT / period
    en_names = [s["en"] for s in sch.SCHOOLS]
    en_set = set(en_names)

    def load_ranking(path: pathlib.Path):
        wb = openpyxl.load_workbook(path)
        rows = [r for r in wb.active.iter_rows(values_only=True)
                if any(c is not None and str(c).strip() for c in r)]
        hdr = next(i for i, r in enumerate(rows) if str(r[1]).strip() == "Institutions")
        data = sorted((r for r in rows[hdr + 1:] if r[1]), key=lambda r: int(_f(r[0])))
        wb.close()
        return data

    # ---- 机构总表：全量（±3 邻近需要全球行）＋ CHINA MAINLAND 子集（省内关键词） ----
    inst = load_ranking(base / "Institutions.xlsx")
    neighbors: dict[str, list] = {n: [] for n in en_names}
    for i, r in enumerate(inst):
        name = str(r[1]).strip()
        if name in en_set:
            k = int(_f(r[0]))
            rows = []
            for x in inst[max(0, k - 4):k + 3]:
                rows.append([str(x[1]).strip(), str(x[2]).strip(), int(_f(x[0])),
                             _f(x[3]), _f(x[4]), _f(x[5]), _f(x[9])])
            neighbors[name] = rows
    cn_inst = [[str(r[1]).strip(), int(_f(r[0])), _f(r[3]), _f(r[4]), _f(r[5]), _f(r[9])]
               for r in inst
               if str(r[2]).strip() in ("CHINA MAINLAND", "N/A")]  # N/A＝属地未标注（如部属医院），按名关键词进省内

    # ---- 阈值 + 22 榜本校行（国内排名在建缓存时一次算好） ----
    wbt = openpyxl.load_workbook(base / "ThresholdESI.xlsx")
    trows = [r for r in wbt.active.iter_rows(values_only=True)
             if any(c is not None and str(c).strip() for c in r)]
    wbt.close()
    thr1 = {norm_row[0]: _f(norm_row[3]) for norm_row in trows
            if str(norm_row[0]).strip() not in ("THRESHOLD-ESI", "RESEARCH FIELDS")
            and norm_row[3] is not None and str(norm_row[3]).strip()}
    thr1 = {_norm(k): v for k, v in thr1.items()}

    fields: dict[str, dict] = {}
    for stem in STEMS:
        data = load_ranking(base / f"{stem}.xlsx")
        m = len(data)
        line = (m + 1) // 10
        line_cites = _f(data[line - 1][4]) if 0 < line <= m else None
        rows: dict[str, list] = {}
        for i, r in enumerate(data):
            name = str(r[1]).strip()
            if name not in en_set:
                continue
            rank = int(_f(r[0]))
            cn_rank = 1 + sum(1 for x in data[:i] if str(x[2]).strip() == "CHINA MAINLAND")
            rows[name] = [rank, cn_rank, _f(r[3]), _f(r[4]), _f(r[5]), _f(r[9]), _f(r[12])]
        fields[stem] = {"m": m, "line": line, "line_cites": line_cites, "schools": rows}

    meta_path = base / "_下载信息.json"
    download = ""
    if meta_path.exists():
        try:
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
            download = f"{meta.get('首次下载', '')}/{meta.get('最近下载', '')}"
        except (json.JSONDecodeError, OSError):
            download = ""
    cache = {"period": period, "built": datetime.date.today().isoformat(),
             "download": download, "n_inst": len(inst),
             "fields": fields, "thr1": thr1,
             "cn_inst": cn_inst, "neighbors": neighbors}
    MONITOR_ROOT.mkdir(parents=True, exist_ok=True)
    cache_path(period).write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")
    return cache


def load_cache(period: str) -> dict | None:
    p = cache_path(period)
    if not p.exists():
        return None
    return json.loads(p.read_text(encoding="utf-8"))


# ---------------------------------------------------------------- 渲染
_e = html.escape
_fmt = lambda v: f"{v:,.0f}"


def _school_total(cache: dict, en: str) -> list | None:
    for row in cache["cn_inst"]:
        if row[0] == en:
            return row
    return None


_CSS = """
:root{--bg:#faf9f6;--card:#fff;--ink:#17181a;--ink-2:#3a3c40;--muted:#6f7278;--line:#e4e1da;
--line-2:#efede7;--ok:#2f6f5e;--ok-bg:#e8f1ed;--warn:#9a6b14;--warn-bg:#f6efdf;--radius:5px}
*{margin:0;padding:0;box-sizing:border-box}
body{background:var(--bg);color:var(--ink);font-family:"PingFang SC","Microsoft YaHei",system-ui,sans-serif;
font-size:15.5px;line-height:1.6;-webkit-font-smoothing:antialiased;font-variant-numeric:tabular-nums}
.wrap{max-width:1000px;margin:0 auto;padding:0 clamp(18px,4vw,40px)}
header.top{padding:24px 0 0;display:flex;justify-content:space-between;gap:14px;font-size:12.5px;color:var(--muted);flex-wrap:wrap}
h1{font-size:clamp(1.6rem,1.3rem+1.6vw,2.2rem);font-weight:600;letter-spacing:-.02em;margin-top:clamp(34px,5vh,54px)}
.lead{margin-top:12px;max-width:66ch;color:var(--ink-2)}.lead b{background:#f0ece1;padding:0 .18em;font-weight:600}
.meta{margin-top:14px;font-size:12.5px;color:var(--muted)}
.nums{margin-top:22px;display:grid;grid-template-columns:repeat(4,1fr);gap:10px}
.nm{background:var(--card);border:1px solid var(--line);border-radius:var(--radius);padding:13px 15px}
.nm .nl{font-size:11.5px;color:var(--muted)}.nm .nv{font-size:clamp(1.2rem,1.1rem+.6vw,1.5rem);font-weight:600;margin-top:2px}
.nm .nx{font-size:11.5px;color:var(--muted)}
section{margin-top:clamp(40px,6vh,60px)}
h2{font-size:1.28rem;font-weight:600;display:flex;gap:10px;align-items:baseline}
h2 .num{font-size:12.5px;color:var(--muted);font-family:ui-monospace,Consolas,monospace}
table{width:100%;border-collapse:collapse;font-size:13.5px;margin-top:16px}
th{font-weight:600;text-align:left;padding:8px 10px 8px 0;border-bottom:1px solid var(--line);font-size:12.5px;color:var(--ink-2)}
th.n,td.n{text-align:right;padding-right:12px}
td{padding:9px 10px 9px 0;border-bottom:1px solid var(--line-2);vertical-align:top;color:var(--ink-2)}
td b{font-weight:600;color:var(--ink)}
td .en{display:block;font-size:11px;color:var(--muted);font-family:ui-monospace,Consolas,monospace}
td.n{font-variant-numeric:tabular-nums;white-space:nowrap}
.pill{font-size:11px;padding:2px 8px;border-radius:99px;white-space:nowrap}
.pill.ok{background:var(--ok-bg);color:var(--ok)}
.pill.pm{background:var(--warn-bg);color:var(--warn);font-weight:600}
.note{font-size:12.5px;color:var(--muted);margin-top:10px}
footer{margin-top:clamp(48px,8vh,76px);padding:22px 0 46px;border-top:1px solid var(--line);font-size:12.5px;color:var(--muted);line-height:1.8}
@media (max-width:760px){.nums{grid-template-columns:repeat(2,1fr)}}
"""


def render(cn: str, period: str, cache: dict, insight: dict | None = None) -> str:
    """最简版站位报告：纯全局件，无论文清单、无图表。insight＝AI/骨架解读（可缺省）。"""
    s = sch.find(cn)
    if s is None:
        raise KeyError(cn)
    en = s["en"]
    tag = "985 工程" if s["is985"] else "211 工程"

    rows = {stem: d["schools"].get(en) for stem, d in cache["fields"].items()}
    present = {stem: r for stem, r in rows.items() if r}
    total = _school_total(cache, en)
    entered, pm, near = [], [], []
    for stem in STEMS:
        r = rows.get(stem)
        if not r:
            continue
        t1 = cache["thr1"].get(_norm(stem))
        prox = (r[3] / t1) if t1 else None
        d = {"stem": stem, "cn": CN[stem], "r": r, "prox": prox}
        (entered if prox and prox >= 1 else near).append(d)
        if r[0] <= cache["fields"][stem]["line"]:
            pm.append(d)
    entered.sort(key=lambda d: -d["r"][3])
    pm_set = set()
    # 站位表按被引降序；1‰ 集合同步算
    pairs = sorted(((stem, r) for stem, r in present.items()), key=lambda kv: -kv[1][3])
    for stem, r in pairs:
        if r[0] <= cache["fields"][stem]["line"]:
            pm_set.add(stem)

    def field_row(kv):
        stem, r = kv
        line = cache["fields"][stem]["line"]
        badge = ('<span class="pill pm">前 1‰</span>' if r[0] <= line
                 else '<span class="pill ok">前 1%</span>')
        return (f'<tr><td><b>{_e(CN[stem])}</b><span class="en">{_e(stem)}</span></td>'
                f'<td class="n">{r[0]} / {cache["fields"][stem]["m"]}</td>'
                f'<td class="n">{r[1]}</td><td class="n">{_fmt(r[2])}</td>'
                f'<td class="n">{_fmt(r[3])}</td><td class="n">{r[4]:.2f}</td>'
                f'<td class="n">{r[5]:.0f}</td><td class="n">{r[6]:.0f}</td><td>{badge}</td></tr>')

    def inst_row(r):
        return (f'<tr><td><b>{_e(str(r[0]).title())}</b></td><td class="n">{r[1]}</td>'
                f'<td class="n">{_fmt(r[2])}</td><td class="n">{_fmt(r[3])}</td>'
                f'<td class="n">{r[4]:.2f}</td><td class="n">{r[5]:.0f}</td></tr>')

    pm_rows = "".join(
        f'<tr><td><b>{_e(CN[stem])}</b></td>'
        f'<td class="n">{r[0]} / {cache["fields"][stem]["m"]}</td>'
        f'<td class="n">第 {cache["fields"][stem]["line"]} 位</td>'
        f'<td class="n">{_fmt(cache["fields"][stem]["line_cites"])}</td></tr>'
        for stem, r in pairs if r[0] <= cache["fields"][stem]["line"])
    near_rows = "".join(
        f'<tr><td><b>{_e(d["cn"])}</b></td><td class="n">{_fmt(d["r"][3])}</td>'
        f'<td class="n">{_fmt(cache["thr1"].get(_norm(d["stem"]), 0))}</td>'
        f'<td class="n">{d["prox"]:.2f}</td></tr>'
        for d in sorted(near, key=lambda x: -(x["prox"] or 0)))
    near_block = (
        f'<table><thead><tr><th>学科</th><th class="n">本校被引</th>'
        f'<th class="n">前 1% 阈值</th><th class="n">接近度</th></tr></thead>'
        f'<tbody>{near_rows}</tbody></table>' if near_rows else
        (f'<p class="note">上榜学科全部已在前 1% 内（最小接近度 '
         f'{min(d["prox"] for d in entered):.1f}×阈值），无临阈学科。</p>' if entered
         else '<p class="note">本期全球机构总表未收录该校，无法计算接近度。</p>'))
    miss_list = [CN[stem] for stem in STEMS if stem not in present]

    if total:
        nums = (f'<div class="nm"><div class="nl">入围 ESI 前 1% 学科</div>'
                f'<div class="nv">{len(entered)}<small> / 22</small></div>'
                f'<div class="nx">被引 ≥ 该学科机构阈值</div></div>'
                f'<div class="nm"><div class="nl">全球机构总排名</div>'
                f'<div class="nv">{total[1]:,}<small> / {cache["n_inst"]:,}</small></div>'
                f'<div class="nx">按总被引</div></div>'
                f'<div class="nm"><div class="nl">总被引</div><div class="nv">{_fmt(total[3])}</div>'
                f'<div class="nx">论文 {_fmt(total[2])} · 篇均 {total[4]:.2f}</div></div>'
                f'<div class="nm"><div class="nl">进入全球前 1‰ 的学科</div>'
                f'<div class="nv">{len(pm_set)}</div>'
                f'<div class="nx">{"、".join(_e(CN[stem]) for stem, _rr in pairs if stem in pm_set) or "—"}</div></div>')
        absent = ""
    else:
        nums = ""
        absent = ('<p class="note" style="margin-top:18px">本期全球机构总表未收录该校'
                  f'（ESI 名 <code>{_e(en)}</code>）——可能低于收录线或名称待人工核对。'
                  '下方表格只呈现 22 学科榜里搜得到的学科行。</p>')
    table_body = "".join(field_row(kv) for kv in pairs) or (
        '<tr><td colspan="9" class="note">22 学科榜内未搜到该校——低于收录线或未入册，'
        '本期不出学科行。</td></tr>')
    nb = cache["neighbors"].get(en) or []

    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{_e(cn)} · 学科监测报告 · {period}</title>
<style>{_CSS}</style>
</head>
<body>
<header class="top wrap"><span>监测版 · 包契约验证件 · {_e(tag)}</span><span>生成 {_e(cache["built"])}</span></header>
<div class="wrap">
<h1>{_e(cn)} · 学科监测报告</h1>
<p class="lead">期 1 · 站位报告。数据面＝<b>{_e(period)} 期 ESI 全局件</b>（机构总表、22 学科全球榜、阈值表）——
未取论文清单，论文层分析不在本版。所有数字可回溯到包内原始表。</p>
<p class="meta">数据源：Clarivate ESI {_e(period)} 期（下载 {_e(cache.get("download") or "—")}）·
{_e(cn)}上榜 {len(present)} / 22 学科 · ESI 名 <code>{_e(en)}</code></p>
<div class="nums">{nums}</div>
{absent}

<section>
<h2><span class="num">一</span>学科站位（上榜 {len(present)} 学科，按被引降序）</h2>
<table>
<thead><tr><th>学科</th><th class="n">全球排名</th><th class="n">国内排名*</th><th class="n">论文</th><th class="n">被引</th><th class="n">篇均</th><th class="n">高被引</th><th class="n">热点</th><th>档位</th></tr></thead>
<tbody>{table_body}</tbody>
</table>
<p class="note">* 国内排名＝榜内 CHINA MAINLAND 计数推导，非 ESI 原生指标。排名写法「全球第 N / 榜内 M 所」。</p>
</section>

<section>
<h2><span class="num">二</span>前 1‰ 俱乐部（{len(pm)} 学科）</h2>
{f'<table><thead><tr><th>学科</th><th class="n">全球排名</th><th class="n">1‰ 线</th><th class="n">线上被引</th></tr></thead><tbody>{pm_rows}</tbody></table>' if pm_rows else '<p class="note">本期无学科进入前 1‰。</p>'}
<p class="note">1‰ 线＝榜内第 ⌊(m+1)/10⌋ 位机构的被引（m＝榜内机构数），与 <code>extract_1permille</code> 同法。</p>
</section>

<section>
<h2><span class="num">三</span>进位预警</h2>
{near_block}
{f'<p class="note">未上榜 {len(miss_list)} 学科（{"、".join(_e(m) for m in miss_list)}）低于榜内收录线或未导出，无被引数据，接近度不可算，不外推。</p>' if miss_list else ''}
</section>

<section>
<h2><span class="num">四</span>对标 · 总排名邻近（±3 名）</h2>
<table>
<thead><tr><th>机构</th><th class="n">总排名</th><th class="n">论文</th><th class="n">被引</th><th class="n">篇均</th><th class="n">高被引</th></tr></thead>
<tbody>{"".join(f'<tr><td><b>{_e(str(r[0]).title())}</b></td><td class="n">{r[2]}</td><td class="n">{_fmt(r[3])}</td><td class="n">{_fmt(r[4])}</td><td class="n">{r[5]:.2f}</td><td class="n">{r[6]:.0f}</td></tr>' for r in nb)}</tbody>
</table>
<p class="note">自动邻近窗；跨省自由对比用插件页右上角「＋」进队列，一键出对比件（不限数量）。</p>
</section>

{_insight_section(insight)}

<footer>
数据来源：Clarivate ESI（Essential Science Indicators）{_e(period)} 期全局件，云南大学图书馆订阅账号下载
（{_e(cache.get("download") or "—")}）；本报告是监测版包契约的验证件，不是对外报告，数字未经该校核对。<br>
期 1 无历史台账，趋势章与预测章随期次自建后解锁（期 3+／期 6+）。
派生算法：国内排名＝榜内 CHINA MAINLAND 计数 · 1‰ 线＝第 ⌊(m+1)/10⌋ 位被引 · 接近度＝被引 ÷ 前 1% 机构阈值。<br>
生成：文衡监测版引擎 <code>src/monitor/</code>（{_e(cache["built"])}）。
</footer>
</div>
</body>
</html>
"""


def _insight_section(insight: dict | None) -> str:
    """篇尾解读段：AI 模型文本或句骨架（mode 如实标注；缺省＝无此段）。"""
    if not insight or not (insight.get("text") or "").strip():
        return ""
    how = "模型解读（数字已过盾校验）" if insight.get("mode") == "model" else "规则拼装（写作位未开或模型不可用）"
    body = _e(insight["text"]).replace("；", "；\n")
    return f"""
<section>
<h2><span class="num">五</span>解读</h2>
<p style="margin-top:14px;max-width:70ch;color:var(--ink-2)">{body}</p>
<p class="note">{_e(how)}。</p>
</section>"""


def generate(cn: str, period: str | None = None) -> pathlib.Path:
    """生成（或复用）一所学校的站位报告，返回文件路径。"""
    period = period or latest_period()
    if not period:
        raise FileNotFoundError("数据与复现/ESI数据/ 下没有可用的 ESI 期次目录")
    cache = load_cache(period) or build_cache(period)
    ins: dict | None = None
    try:
        from monitor import insight as _ins                              # noqa: PLC0415
        ins = _ins.insight(cn, period, cache)
    except Exception:                                                   # noqa: BLE001
        ins = None
    out = MONITOR_ROOT / cn / f"{period}.html"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(render(cn, period, cache, insight=ins), encoding="utf-8")
    return out
