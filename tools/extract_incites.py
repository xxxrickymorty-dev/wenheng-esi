# -*- coding: utf-8 -*-
"""从 InCites 快照（+ 同期 ESI 原始包）抽取台账当期列：国际排名 / 国内排名 / CNCI / 总被引 / 论文数 / 篇均被引；
并维护台账里一张小 sheet「InCites快照」做快照登记。

口径（2026-09-23 以 202607 期 22 学科全量回放确定，逐格 22/22 命中）
--------------------------------------------------------------------
每学科一组（新国际排名/新国内排名 里 3 列 = 分母 | 分子 | 百分位）：

    · 国际排名：分母 = 该学科 ESI 机构榜行数 m；分子 = 云大在 ESI 榜的「序号」
      （不在 ESI 前 1% 榜里 → 分子取 **InCites 机构榜**「排名」列）
    · 国内排名：分母 = ESI 榜里的中国大陆条目数；分子 = 云大在其中的序号
      （不在榜 → InCites「中国大陆机构按被引降序」的云大序号）
    · 百分位 = 分子 ÷ 分母（未进入的学科可 > 1，台账原样口径，不裁剪）
    · CNCI（云大块）= InCites 机构榜云大行「学科规范化的引文影响力」
    · 总被引 / 论文数 / 篇均被引：在 ESI 榜的学科由 ESI 侧填（本工具不碰）；
      **不在榜的学科**由本工具从 InCites 云大行写（cites / papers / cites÷papers）

写回纪律：只重写目标 sheet 的 XML（zip 级），其余成员逐字节复制；写前 XML 良构自检；
与台账现有值冲突时**拒绝写入**（口径变化须人工确认）；--write 自动备份到 归档/台账备份/。

用法（wenheng/ 下，用 .venv 的 python）
----------------------------------------
    python tools\\extract_incites.py 202607                     # 预览 + 与台账逐格比对（只读）
    python tools\\extract_incites.py 202607 --out tmp\\led.xlsx  # 写副本
    python tools\\extract_incites.py 202607 --write              # 写台账（自动备份）
    python tools\\extract_incites.py preview [--snapshot 目录]    # 只算 InCites 侧（无 ESI 包时）
    python tools\\extract_incites.py registry [--write]           # 维护「InCites快照」登记 sheet
"""
from __future__ import annotations

import argparse
import datetime as dt
import pathlib
import re
import shutil
import sys
import xml.etree.ElementTree as ET
import zipfile

import openpyxl

REPO = pathlib.Path(__file__).resolve().parent.parent            # wenheng/
ROOT = REPO.parent                                              # 工作区根
LEDGER = ROOT / "数据与复现" / "云大全学科排名、阈值动态检测表.xlsx"
INCITES_ROOT = ROOT / "数据与复现" / "incites"
ESI_ROOT = ROOT / "数据与复现" / "ESI数据"
SCRATCH = REPO / "tmp"
REGISTRY_SHEET = "InCites快照"

sys.path.insert(0, str(REPO / "src"))
import ledger_patch as LP                                              # noqa: E402
import parse_ledger as PL                                              # noqa: E402
from extract_matrix import SNAP_MAP                                    # noqa: E402

SUBJECTS = list(PL.DISCIPLINES_EN)
CELL_RE = re.compile(r'<c r="([A-Z]+)(\d+)"([^>]*?)(?:/>|>(.*?)</c>)', re.S)
ROW_RE = re.compile(r'<row r="(\d+)"([^>]*?)>(.*?)</row>', re.S)


# ------------------------------------------------------------------ 基础
def _num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.strip().replace(",", ""))
        except ValueError:
            return None
    return None


def _fmt(v) -> str:
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return repr(v) if isinstance(v, float) else str(v)


def _eq(a, b, tol=1e-9) -> bool:
    x, y = _num(a), _num(b)
    if x is None or y is None:
        return (a in (None, "")) == (b in (None, ""))
    return abs(x - y) <= tol * max(1.0, abs(x), abs(y))


def col_letter(idx: int) -> str:
    s = ""
    while idx:
        idx, r = divmod(idx - 1, 26)
        s = chr(65 + r) + s
    return s


def col_index(s: str) -> int:
    n = 0
    for ch in s:
        n = n * 26 + ord(ch.upper()) - 64
    return n


def _resolve_file(d: pathlib.Path, subject: str) -> pathlib.Path | None:
    """学科名 → 目录文件（兼容 `/`→空格 与 04.30 期 `xxxlsx.xlsx` 笔误）。"""
    cands = [subject + ".xlsx", subject.replace("/", " ") + ".xlsx", subject + "xlsx.xlsx"]
    for n in cands:
        if (d / n).exists():
            return d / n
    low = {q.name.lower(): q for q in d.iterdir() if q.is_file()}
    for n in cands:
        if n.lower() in low:
            return low[n.lower()]
    return None


# ------------------------------------------------------------------ 读源
def read_esi_list(path: pathlib.Path) -> dict:
    """ESI 机构榜 → {m, ypos, cn, cn_pos}（位次优先取「序号」列，缺序号按行序顺延）。"""
    wb = openpyxl.load_workbook(path, data_only=True)  # ESI 新导出 dimension 常坏，普通模式才读得全
    try:
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    hi = next((i for i, r in enumerate(rows[:12])
               if r and any(str(c).strip() == "Institutions" for c in r if c is not None)), None)
    if hi is None:
        raise ValueError(f"未识别 ESI 机构榜表头：{path.name}")
    hdr = [str(c).strip() if c is not None else "" for c in rows[hi]]
    ni = hdr.index("Institutions")
    gi = next((j for j, h in enumerate(hdr) if "Countries" in h), None)
    di = next((j for j, h in enumerate(hdr) if "Documents" in h), None)
    ci = next((j for j, h in enumerate(hdr) if h.strip() == "Cites"), None)
    data = []                                   # [(序号, 大写机构名, 大写国家, docs, cites)]
    seq = 0
    for r in rows[hi + 1:]:
        if not r:
            continue
        if any(c and "copyright" in str(c).lower() for c in r[:3]):
            break
        inst = str(r[ni]).strip() if ni < len(r) and r[ni] is not None else ""
        if not inst:
            continue
        s = _num(r[0]) if len(r) else None
        seq = int(s) if s and s > 0 else seq + 1
        ctry = str(r[gi]).strip().upper() if gi is not None and gi < len(r) and r[gi] else ""
        docs = _num(r[di]) if di is not None and di < len(r) else None
        cites = _num(r[ci]) if ci is not None and ci < len(r) else None
        data.append((seq, inst.upper(), ctry, docs, cites))
    yrow = next((d for d in data if "YUNNAN UNIVERSITY" in d[1]), None)
    ypos = yrow[0] if yrow else None
    cn = sum(1 for _s, _i, c, _d, _x in data if "CHINA" in c)
    cn_pos = sum(1 for s, _i, c, _d, _x in data if ypos is not None and s <= ypos and "CHINA" in c) or None
    return {"m": len(data), "ypos": ypos, "cn": cn, "cn_pos": cn_pos,
            "y_docs": int(yrow[3]) if yrow and yrow[3] is not None else None,
            "y_cites": int(yrow[4]) if yrow and yrow[4] is not None else None}


def read_esi(esi_dir: pathlib.Path) -> dict:
    out = {}
    for s in SUBJECTS:
        p = _resolve_file(esi_dir, s)
        if p is None:
            raise SystemExit(f"[缺件] ESI 包 {esi_dir.name} 里没有 {s}")
        out[s] = read_esi_list(p)
    return out


def read_incites_snapshot(snap_dir: pathlib.Path) -> dict:
    out = {}
    for s in SUBJECTS:
        p = _resolve_file(snap_dir, s)
        if p is None:
            raise SystemExit(f"[缺件] InCites 快照 {snap_dir.name} 里没有 {s}")
        pp = p
        wb = openpyxl.load_workbook(str(pp), read_only=True, data_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            it = ws.iter_rows(values_only=True)
            hdr = [str(c).strip() if c is not None else "" for c in next(it)]

            def find(*keys, _h=hdr):
                for i, h in enumerate(_h):
                    if all(k in h for k in keys):
                        return i
                raise ValueError(f"{pp.name} 表头缺列 {keys}")

            ri, ci, pi = find("排名"), find("被引频次"), find("论文数")
            cni, li = find("学科规范化"), find("国家或地区")
            rows = list(it)
        finally:
            wb.close()
        y = next((r for r in rows if str(r[0] or "").strip().lower() == "yunnan university"), None)
        if y is None:
            raise SystemExit(f"[缺行] {snap_dir.name}/{p.name} 里没有 Yunnan University")
        cn = [r for r in rows if str(r[li] or "").strip().upper() == "CHINA MAINLAND"]
        cn_sorted = sorted(cn, key=lambda r: -(_num(r[ci]) or 0))
        cn_rank = next((k for k, r in enumerate(cn_sorted, 1)
                        if str(r[0]).strip().lower() == "yunnan university"), None)
        out[s] = {"rank": int(_num(y[ri]) or 0), "cn_rank": cn_rank, "cn_n": len(cn),
                  "cnci": float(_num(y[cni]) or 0), "cites": int(_num(y[ci]) or 0),
                  "papers": int(_num(y[pi]) or 0)}
    return out


def compute(snap_dir: pathlib.Path, esi_dir: pathlib.Path | None) -> dict:
    ic = read_incites_snapshot(snap_dir)
    esi = read_esi(esi_dir) if esi_dir is not None else None
    table = {}
    for s in SUBJECTS:
        i = ic[s]
        e = esi.get(s) if esi else None
        if e is None:
            den_i = den_c = None
            num_i, num_c = i["rank"], i["cn_rank"]
        else:
            den_i, den_c = e["m"], e["cn"]
            num_i = e["ypos"] or i["rank"]
            num_c = e["cn_pos"] or i["cn_rank"]
        table[s] = {
            "in_esi": bool(e and e["ypos"]),
            "intl": (den_i, num_i, (num_i / den_i) if den_i else None),
            "cn": (den_c, num_c, (num_c / den_c) if den_c else None),
            "cnci": i["cnci"], "cites": i["cites"], "papers": i["papers"],
            "avg": (i["cites"] / i["papers"]) if i["papers"] else None,
        }
    return table


# ------------------------------------------------------------------ 台账读
def read_grid(led: pathlib.Path, sheet: str):
    return LP.read_grid(led, sheet)


def gtext(g, shared, rn, cn):
    cell = g.get(rn, {}).get(cn)
    return LP.cell_text(cell[0], cell[1], shared) if cell else None


def subject_rows(g, shared, names: set) -> dict:
    out = {}
    for rn in sorted(g):
        t = str(gtext(g, shared, rn, "A") or "").strip()
        if t in names and t not in out:
            out[t] = rn
    return out


def ss_index(shared, text: str):
    for i, s in enumerate(shared):
        if str(s).strip() == text:
            return i
    return None


def text_op(shared, text: str):
    """文本格内容 → ('t="s"' 或 't="inlineStr"', inner)。"""
    i = ss_index(shared, text)
    if i is not None:
        return ('t="s"', f"<v>{i}</v>")
    return ('t="inlineStr"', f'<is><t xml:space="preserve">{text}</t></is>')


def _drop_t(attrs: str) -> str:
    return re.sub(r'\s+t="[^"]*"', "", attrs or "")


def _A(attrs: str) -> str:
    """单元格属性串规范化：非空时补一个前导空格（拼 <c> 用）。"""
    s = (attrs or "").strip()
    return " " + s if s else ""


# ops: {(行, 列): ("set", 值) | ("new", style, t_attr, inner)}
def edit_cells(xml: str, ops: dict) -> tuple[str, int]:
    changed = 0
    todo = set(ops)

    def fix(m):
        nonlocal changed
        rn = int(m.group(1))
        if not any(r == rn for r, _c in todo):
            return m.group(0)
        cells = {cn: (a, i) for cn, _r, a, i in CELL_RE.findall(m.group(3))}
        for (er, ec), op in ops.items():
            if er != rn:
                continue
            if op[0] == "set" and ec in cells:
                v = op[1]
                cells[ec] = (_drop_t(cells[ec][0]),
                             f"<v>{_fmt(v)}</v>" if v is not None else None)
                changed += 1
            elif ec not in cells:
                if op[0] == "set":
                    style = _style_fallback(cells)
                    v = op[1]
                    cells[ec] = (_drop_t(style), f"<v>{_fmt(v)}</v>" if v is not None else None)
                else:
                    _k, style, t_attr, inner = op
                    cells[ec] = (f'{_drop_t(style)} {t_attr}'.strip(), inner)
                changed += 1
        body = "".join(f'<c r="{cn}{rn}"{_A(a)}>{i}</c>' if i is not None
                       else f'<c r="{cn}{rn}"{_A(a)}/>'
                       for cn, (a, i) in sorted(cells.items(), key=lambda kv: col_index(kv[0])))
        spans = max((col_index(cn) for cn in cells), default=1)
        row_attrs = re.sub(r'spans="1:\d+"', f'spans="1:{spans}"', m.group(2))
        return f'<row r="{rn}"{row_attrs}>{body}</row>'

    return ROW_RE.sub(fix, xml), changed


def _style_fallback(cells: dict) -> str:
    for cn in ("B", "C"):
        if cn in cells:
            return cells[cn][0]
    return ""


def append_row(xml: str, rn: int, ops: dict, spans_max: int, shared) -> str:
    cells = {}
    for (_er, cn), op in ops.items():
        if op[0] == "set":
            cells[cn] = (op[1].get("style", ""), f"<v>{_fmt(op[1]['value'])}</v>") \
                if isinstance(op[1], dict) else ("", f"<v>{_fmt(op[1])}</v>")
        else:
            _k, style, t_attr, inner = op
            cells[cn] = (f'{_drop_t(style)} {t_attr}'.strip(), inner)
    body = "".join(f'<c r="{cn}{rn}"{_A(a)}>{i}</c>' if i is not None
                   else f'<c r="{cn}{rn}"{_A(a)}/>'
                   for cn, (a, i) in sorted(cells.items(), key=lambda kv: col_index(kv[0])))
    return xml.replace("</sheetData>", f'<row r="{rn}" spans="1:{spans_max}">{body}</row></sheetData>')


def insert_row(xml: str, rn: int, ops: dict, spans_max: int) -> str:
    """按行号**插入**新行（保持物理排序；两块并排的表必须用这个，不能用 append_row）。"""
    cells = {}
    for (_er, cn), op in ops.items():
        if op[0] == "set":
            cells[cn] = ("", f"<v>{_fmt(op[1])}</v>")
        else:
            _k, style, t_attr, inner = op
            cells[cn] = (f'{_drop_t(style)} {t_attr}'.strip(), inner)
    body = "".join(f'<c r="{cn}{rn}"{_A(a)}>{i}</c>' if i is not None
                   else f'<c r="{cn}{rn}"{_A(a)}/>'
                   for cn, (a, i) in sorted(cells.items(), key=lambda kv: col_index(kv[0])))
    row = f'<row r="{rn}" spans="1:{spans_max}">{body}</row>'
    for m in ROW_RE.finditer(xml):
        if int(m.group(1)) > rn:
            return xml[:m.start()] + row + xml[m.start():]
    return xml.replace("</sheetData>", row + "</sheetData>")


def fix_dimension(xml: str, max_col: int, max_row: int) -> str:
    return re.sub(r'<dimension ref="A1:[A-Z]+\d+"/>',
                  f'<dimension ref="A1:{col_letter(max_col)}{max_row}"/>', xml)


def add_merge(xml: str, ref: str) -> str:
    m = re.search(r'<mergeCells count="(\d+)">', xml)
    if not m:
        return xml
    return xml.replace("</mergeCells>", f'<mergeCell ref="{ref}"/></mergeCells>').replace(
        m.group(0), f'<mergeCells count="{int(m.group(1)) + 1}">')


def max_row_xml(xml: str) -> int:
    return max((int(m.group(1)) for m in ROW_RE.finditer(xml)), default=1)


# ------------------------------------------------------------------ 事务
class LedgerTx:
    """一次读入台账 zip，累加改写，最后原子写出（非目标成员逐字节不变）。"""

    def __init__(self, path: pathlib.Path):
        self.path = pathlib.Path(path)
        with zipfile.ZipFile(self.path) as zf:
            self.infos = zf.infolist()
            self.members = {i.filename: zf.read(i.filename) for i in self.infos}
        self.replace: dict[str, bytes] = {}
        self.added: dict[str, bytes] = {}
        self.meta_changed: set[str] = set()

    def member_of(self, sheet: str) -> str:
        wbxml = self.members["xl/workbook.xml"].decode("utf-8")
        pairs = dict(re.findall(r'<sheet[^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbxml))
        if sheet not in pairs:
            raise KeyError(f"台账里没有 sheet「{sheet}」")
        rels = self.members["xl/_rels/workbook.xml.rels"].decode("utf-8")
        rmap = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
        t = rmap[pairs[sheet]].lstrip("/")
        return t if t.startswith("xl/") else "xl/" + t

    def sheet_xml(self, sheet: str) -> str:
        return self.members[self.member_of(sheet)].decode("utf-8")

    def shared(self) -> list:
        return ["".join(re.findall(r"<t[^>]*>(.*?)</t>", si, re.S))
                for si in re.findall(r"<si>(.*?)</si>",
                                     self.members["xl/sharedStrings.xml"].decode("utf-8"), re.S)]

    def put_sheet(self, sheet: str, xml: str) -> None:
        ET.fromstring(xml)
        self.replace[self.member_of(sheet)] = xml.encode("utf-8")

    def add_sheet(self, name: str, xml: str) -> None:
        ET.fromstring(xml)
        used = set()
        for i in self.infos:
            m = re.match(r"xl/worksheets/sheet(\d+)\.xml$", i.filename)
            if m:
                used.add(int(m.group(1)))
        n = max(used, default=0) + 1
        part = f"xl/worksheets/sheet{n}.xml"
        self.added[part] = xml.encode("utf-8")
        wb = self.members["xl/workbook.xml"].decode("utf-8")
        sid = max((int(x) for x in re.findall(r'<sheet[^>]*sheetId="(\d+)"', wb)), default=0) + 1
        rels_txt = self.members["xl/_rels/workbook.xml.rels"].decode("utf-8")
        rid = max((int(x) for x in re.findall(r'Id="rId(\d+)"', rels_txt)), default=0) + 1
        wb = wb.replace("</sheets>", f'<sheet name="{name}" sheetId="{sid}" r:id="rId{rid}"/></sheets>')
        rels = rels_txt.replace(
            "</Relationships>",
            f'<Relationship Id="rId{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            f'relationships/worksheet" Target="worksheets/sheet{n}.xml"/></Relationships>')
        ct = self.members["[Content_Types].xml"].decode("utf-8").replace(
            "</Types>",
            f'<Override PartName="/{part}" ContentType="application/vnd.openxmlformats-officedocument.'
            f'spreadsheetml.worksheet+xml"/></Types>')
        self.members["xl/workbook.xml"] = wb.encode("utf-8")
        self.members["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")
        self.members["[Content_Types].xml"] = ct.encode("utf-8")
        self.meta_changed |= {"xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml"}

    def commit(self) -> dict:
        tmp = self.path.with_suffix(".patching.xlsx")
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
            for info in self.infos:
                zo.writestr(info, self.replace.get(info.filename, self.members[info.filename]))
            for name, data in self.added.items():
                zo.writestr(name, data)
        shutil.move(str(tmp), str(self.path))
        return {"changed": sorted(set(self.replace) | self.meta_changed),
                "added": sorted(self.added)}


# ------------------------------------------------------------------ 写计划
def plan_wide(tx: LedgerTx, sheet: str, period: str, values: dict) -> tuple[str, dict]:
    """宽表（行=学科、列=期次）：CNCI / 总被引 / 论文数。values: {学科: 值或 None}。"""
    xml = tx.sheet_xml(sheet)
    g, shared = read_grid(tx.path, sheet)
    periods = [(col_index(cn), cn) for cn, (a, i) in g[1].items()
               if re.fullmatch(r"20\d{4}", str(LP.cell_text(a, i, shared) or "").strip())]
    last_i, last_cn = max(periods)
    period_col = next((cn for _i, cn in periods
                       if str(gtext(g, shared, 1, cn) or "").strip() == period), None)
    append = period_col is None
    if append:
        period_col = col_letter(last_i + 1)
    prev_col = col_letter(col_index(period_col) - 1)
    row_of = subject_rows(g, shared, set(SUBJECTS))
    ops = {}
    for s, v in values.items():
        if v is None or s not in row_of:
            continue
        rn = row_of[s]
        if period_col in g.get(rn, {}):
            ops[(rn, period_col)] = ("set", v)
        else:
            style = g.get(rn, {}).get(prev_col, ("", None))[0]
            ops[(rn, period_col)] = ("new", style, "", f"<v>{_fmt(v)}</v>")
    if append:
        t_attr, inner = text_op(shared, period)
        ops[(1, period_col)] = ("new", g.get(1, {}).get(last_cn, ("", None))[0], t_attr, inner)
    xml, n = edit_cells(xml, ops)
    xml = fix_dimension(xml, col_index(period_col), max_row_xml(xml))
    return xml, {"cells": n}


def plan_group3(tx: LedgerTx, sheet: str, period: str, table: dict, kind: str) -> tuple[str, dict]:
    """新国际排名 / 新国内排名：每期 3 列（分母|分子|百分位），行=22 学科（第 3 行起）。"""
    labels = ("ESI1%学科数", "排名", "排名百分位")
    xml = tx.sheet_xml(sheet)
    g, shared = read_grid(tx.path, sheet)
    starts = sorted(col_index(cn) for cn, (a, i) in g[1].items()
                    if re.fullmatch(r"20\d{4}", str(LP.cell_text(a, i, shared) or "").strip()))
    period_col = next((cn for cn, (a, i) in g[1].items()
                       if str(LP.cell_text(a, i, shared) or "").strip() == period), None)
    row_of = subject_rows(g, shared, set(SUBJECTS))
    ops = {}
    if period_col is None:
        base, prev_base = starts[-1] + 3, starts[-1]
        t_attr, inner = text_op(shared, period)
        ops[(1, col_letter(base))] = ("new", g[1][col_letter(prev_base)][0], t_attr, inner)
        for k, lab in enumerate(labels):
            t_attr, inner = text_op(shared, lab)
            style = g[2][col_letter(prev_base + k)][0] if col_letter(prev_base + k) in g.get(2, {}) else ""
            ops[(2, col_letter(base + k))] = ("new", style, t_attr, inner)
        if "<mergeCells" in xml:
            xml = add_merge(xml, f"{col_letter(base)}1:{col_letter(base + 2)}1")
    else:
        base, prev_base = col_index(period_col), col_index(period_col) - 3
    for s, rec in table.items():
        rn = row_of.get(s)
        if rn is None:
            continue
        for k, v in enumerate(rec["intl"] if kind == "intl" else rec["cn"]):
            cn = col_letter(base + k)
            if cn in g.get(rn, {}):
                ops[(rn, cn)] = ("set", v)
            else:
                prev_cn = col_letter(prev_base + k)
                style = g.get(rn, {}).get(prev_cn, ("", None))[0]
                ops[(rn, cn)] = ("new", style, "", f"<v>{_fmt(v)}</v>")
    xml, n = edit_cells(xml, ops)
    xml = fix_dimension(xml, base + 2, max_row_xml(xml))
    return xml, {"cells": n}


def plan_transposed(tx: LedgerTx, sheet: str, period: str, table: dict, kind: str) -> tuple[str, dict]:
    """国际排名 / 国内排名（转置）：行=期次、列=22 学科×3（分母|分子|百分位）。"""
    xml = tx.sheet_xml(sheet)
    g, shared = read_grid(tx.path, sheet)
    starts = sorted((col_index(cn), cn) for cn, (a, i) in g[1].items()
                    if cn != "A" and str(LP.cell_text(a, i, shared) or "").strip())
    if len(starts) != len(SUBJECTS):
        raise SystemExit(f"[警告] {sheet} 第 1 行学科组 {len(starts)} 个（应 22），先人工核对版式")
    name_of = {cn: str(LP.cell_text(g[1][cn][0], g[1][cn][1], shared) or "").strip() for _i, cn in starts}
    span_max = starts[-1][0] + 2
    prow = next((rn for rn in sorted(g)
                 if str(gtext(g, shared, rn, "A") or "").strip() == period), None)
    last_period_row = max((rn for rn in g if str(gtext(g, shared, rn, "A") or "").strip()), default=2)
    ops = {}
    if prow is None:
        cand = max(g)
        if str(gtext(g, shared, cand, "A") or "").strip() or any(
                i is not None and str(i).strip() for _a, i in g[cand].values()):
            prow = cand + 1
        else:
            prow = cand
        prev_row = last_period_row
    else:
        prev_row = max((rn for rn in g if rn < prow and str(gtext(g, shared, rn, "A") or "").strip()),
                       default=None)
    for ci, cn in starts:
        rec = table.get(name_of[cn])
        if rec is None:
            continue
        for k, v in enumerate(rec["intl"] if kind == "intl" else rec["cn"]):
            tcn = col_letter(ci + k)
            if tcn in g.get(prow, {}):
                ops[(prow, tcn)] = ("set", v)
            else:
                style = g.get(prev_row, {}).get(tcn, ("", None))[0] if prev_row else ""
                ops[(prow, tcn)] = ("new", style, "", f"<v>{_fmt(v)}</v>")
    if prow not in g or "A" not in g[prow]:
        t_attr, inner = text_op(shared, period)
        style = g.get(prev_row, {}).get("A", ("", None))[0] if prev_row else ""
        ops[(prow, "A")] = ("new", style, t_attr, inner)
    if prow in g:
        xml, n = edit_cells(xml, ops)
    else:
        xml = append_row(xml, prow, ops, span_max, shared)
        n = len(ops)
    xml = fix_dimension(xml, span_max, max(max_row_xml(xml), prow))
    return xml, {"cells": n}


def plan_avg(tx: LedgerTx, sheet: str, period: str, table: dict) -> tuple[str, dict]:
    """篇均被引：每期 3 列（云大|中国大陆|国际）；只写**非 ESI 学科**的云大格。"""
    labels = ("云大", "中国大陆", "国际")
    xml = tx.sheet_xml(sheet)
    g, shared = read_grid(tx.path, sheet)
    starts = sorted((col_index(cn), cn) for cn, (a, i) in g[2].items()
                    if re.fullmatch(r"20\d{4}", str(LP.cell_text(a, i, shared) or "").strip()))
    period_col = next((cn for cn, (a, i) in g[2].items()
                       if str(LP.cell_text(a, i, shared) or "").strip() == period), None)
    row_of = subject_rows(g, shared, set(SUBJECTS))
    ops = {}
    if period_col is None:
        base, prev_base = starts[-1][0] + 3, starts[-1][0]
        t_attr, inner = text_op(shared, period)
        ops[(2, col_letter(base))] = ("new", g[2][col_letter(prev_base)][0] if 2 in g else "",
                                      t_attr, inner)
        for k, lab in enumerate(labels):
            t_attr, inner = text_op(shared, lab)
            style = g[3][col_letter(prev_base + k)][0] if 3 in g and col_letter(prev_base + k) in g[3] else ""
            ops[(3, col_letter(base + k))] = ("new", style, t_attr, inner)
        if "<mergeCells" in xml:
            xml = add_merge(xml, f"{col_letter(base)}2:{col_letter(base + 2)}2")
        vcol, prev_cn = col_letter(base), col_letter(prev_base)
    else:
        vcol, prev_cn = period_col, col_letter(col_index(period_col) - 3)
    for s, rec in table.items():
        if rec["in_esi"] or rec["avg"] is None:
            continue
        rn = row_of.get(s)
        if rn is None:
            continue
        if vcol in g.get(rn, {}):
            ops[(rn, vcol)] = ("set", rec["avg"])
        else:
            style = g.get(rn, {}).get(prev_cn, ("", None))[0]
            ops[(rn, vcol)] = ("new", style, "", f"<v>{_fmt(rec['avg'])}</v>")
    xml, n = edit_cells(xml, ops)
    xml = fix_dimension(xml, col_index(vcol) + 2, max_row_xml(xml))
    return xml, {"cells": n}


# ------------------------------------------------------------------ Q1–Q4（InCites 研究领域）
_Q1Q4_ENTITIES = (
    ("云南大学", ("云南大学研究领域.xlsx", "云大 研究领域.xlsx", "云大研究领域.xlsx")),
    ("中国大陆", ("中国研究领域.xlsx",)),
    ("全球", ("全球 研究领域.xlsx", "全球 研究领域 .xlsx", "全球研究领域.xlsx")),
)
_Q1Q4_LABELS = ("Q1 %", "Q2%", "Q3 %", "Q4 %")


def read_q1q4(snap_dir: pathlib.Path) -> tuple[dict, list]:
    """{学科: {实体: (Q1%, Q2%, Q3%, Q4%)}} + 缺件说明。来源=快照的三张研究领域文件。"""
    out, notes = {}, []
    for ent, cands in _Q1Q4_ENTITIES:
        p = next((snap_dir / n for n in cands if (snap_dir / n).exists()), None)
        if p is None:
            notes.append(f"缺 {ent}「研究领域」文件（{'、'.join(cands)}）")
            continue
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            it = ws.iter_rows(values_only=True)
            hdr = [str(c).strip() if c is not None else "" for c in next(it)]
            qi = []
            for n in (1, 2, 3, 4):
                idx = next((i for i, h in enumerate(hdr)
                            if h.startswith(f"Q{n}") and "百分比" in h), None)
                if idx is not None:
                    qi.append(idx)
            if len(qi) != 4:
                notes.append(f"{p.name} 缺 Q1–Q4 百分比列")
                continue
            for r in it:
                name = str(r[0] or "").strip()
                if name in SUBJECTS:
                    out.setdefault(name, {})[ent] = tuple(float(_num(r[i]) or 0) for i in qi)
        finally:
            wb.close()
    return out, notes


def plan_q1q4(tx: LedgerTx, period: str, table: dict) -> tuple[str, dict]:
    """Q1-Q4：每期 12 列（3 实体 × Q1..Q4），行=22 学科（第 4 行起）。"""
    sheet = "Q1-Q4"
    xml = tx.sheet_xml(sheet)
    g, shared = read_grid(tx.path, sheet)
    label = f"{period[:4]}.{period[4:]}"
    starts = sorted(col_index(cn) for cn, (a, i) in g[1].items()
                    if re.fullmatch(r"20\d{2}\.\d{2}", str(LP.cell_text(a, i, shared) or "").strip()))
    period_col = next((cn for cn, (a, i) in g[1].items()
                       if str(LP.cell_text(a, i, shared) or "").strip() == label), None)
    row_of = subject_rows(g, shared, set(SUBJECTS))
    ops = {}
    if period_col is None:
        base, prev_base = starts[-1] + 12, starts[-1]
        t_attr, inner = text_op(shared, label)
        ops[(1, col_letter(base))] = ("new", g[1][col_letter(prev_base)][0], t_attr, inner)
        for e, (ent, _c) in enumerate(_Q1Q4_ENTITIES):
            t_attr, inner = text_op(shared, ent)
            style = g[2][col_letter(prev_base + e * 4)][0] if 2 in g else ""
            ops[(2, col_letter(base + e * 4))] = ("new", style, t_attr, inner)
            for k, lab in enumerate(_Q1Q4_LABELS):
                t_attr, inner = text_op(shared, lab)
                style = g[3][col_letter(prev_base + k)][0] if 3 in g else ""
                ops[(3, col_letter(base + e * 4 + k))] = ("new", style, t_attr, inner)
        if "<mergeCells" in xml:
            xml = add_merge(xml, f"{col_letter(base)}1:{col_letter(base + 11)}1")
            for e in range(3):
                xml = add_merge(xml, f"{col_letter(base + e * 4)}2:{col_letter(base + e * 4 + 3)}2")
    else:
        base, prev_base = col_index(period_col), col_index(period_col) - 12
    for s, per_ent in table.items():
        rn = row_of.get(s)
        if rn is None:
            continue
        for e, (ent, _c) in enumerate(_Q1Q4_ENTITIES):
            vals = per_ent.get(ent)
            if vals is None:
                continue
            for k, v in enumerate(vals):
                cn = col_letter(base + e * 4 + k)
                if cn in g.get(rn, {}):
                    ops[(rn, cn)] = ("set", v)
                else:
                    style = g.get(rn, {}).get(col_letter(prev_base + e * 4 + k), ("", None))[0]
                    ops[(rn, cn)] = ("new", style, "", f"<v>{_fmt(v)}</v>")
    xml, n = edit_cells(xml, ops)
    xml = fix_dimension(xml, max(base + 11, max((col_index(c) for c in (period_col,) if c), default=0)),
                        max_row_xml(xml))
    return xml, {"cells": n}


# ------------------------------------------------------------------ 冲突门
def existing_values(led: pathlib.Path, period: str) -> dict:
    """台账该期次现有值 → {(sheet, 学科, 字段): 值}（字段：'v' / 'avg' / '0'~'2'）"""
    out = {}
    for sheet in ("CNCI", "总被引", "论文数", "新国际排名", "新国内排名", "篇均被引"):
        g, shared = read_grid(led, sheet)
        if sheet in ("新国际排名", "新国内排名"):
            col = next((cn for cn, (a, i) in g[1].items()
                        if str(LP.cell_text(a, i, shared) or "").strip() == period), None)
            if col is None:
                continue
            base, first_row = col_index(col), 3
        elif sheet == "篇均被引":
            col = next((cn for cn, (a, i) in g[2].items()
                        if str(LP.cell_text(a, i, shared) or "").strip() == period), None)
            if col is None:
                continue
            base, first_row = col_index(col), 4
        else:
            col = next((cn for cn, (a, i) in g[1].items()
                        if str(LP.cell_text(a, i, shared) or "").strip() == period), None)
            if col is None:
                continue
            base, first_row = col_index(col), 2
        row_of = subject_rows(g, shared, set(SUBJECTS))
        for s, rn in row_of.items():
            if rn < first_row:
                continue
            if sheet in ("新国际排名", "新国内排名"):
                for k in range(3):
                    v = gtext(g, shared, rn, col_letter(base + k))
                    if v not in (None, ""):
                        out[(sheet, s, str(k))] = v
            elif sheet == "篇均被引":
                v = gtext(g, shared, rn, col_letter(base))
                if v not in (None, ""):
                    out[(sheet, s, "avg")] = v
            else:
                v = gtext(g, shared, rn, col)
                if v not in (None, ""):
                    out[(sheet, s, "v")] = v
    for sheet in ("国际排名", "国内排名"):
        g, shared = read_grid(led, sheet)
        prow = next((rn for rn in sorted(g)
                     if str(gtext(g, shared, rn, "A") or "").strip() == period), None)
        if prow is None:
            continue
        starts = sorted((col_index(cn), cn) for cn, (a, i) in g[1].items()
                        if cn != "A" and str(LP.cell_text(a, i, shared) or "").strip())
        for ci, cn in starts:
            s = str(LP.cell_text(g[1][cn][0], g[1][cn][1], shared) or "").strip()
            for k in range(3):
                v = gtext(g, shared, prow, col_letter(ci + k))
                if v not in (None, ""):
                    out[(sheet, s, str(k))] = v
    # Q1-Q4（12 列/期：3 实体 × Q1..Q4）
    g, shared = read_grid(led, "Q1-Q4")
    label = f"{period[:4]}.{period[4:]}"
    qcol = next((cn for cn, (a, i) in g[1].items()
                 if str(LP.cell_text(a, i, shared) or "").strip() == label), None)
    if qcol is not None:
        base = col_index(qcol)
        for s, rn in subject_rows(g, shared, set(SUBJECTS)).items():
            for e, (ent, _c) in enumerate(_Q1Q4_ENTITIES):
                for k in range(4):
                    v = gtext(g, shared, rn, col_letter(base + e * 4 + k))
                    if v not in (None, ""):
                        out[("Q1-Q4", s, f"{ent}.{k}")] = v
    return out


def conflicts(table: dict, existing: dict) -> list:
    bad = []
    for s, rec in table.items():
        checks = [("CNCI", "v", rec["cnci"])]
        for sheet, kind in (("国际排名", "intl"), ("新国际排名", "intl"),
                            ("国内排名", "cn"), ("新国内排名", "cn")):
            for k in range(3):
                checks.append((sheet, str(k), rec[kind][k]))
        if not rec["in_esi"]:
            checks += [("总被引", "v", rec["cites"]), ("论文数", "v", rec["papers"]),
                       ("篇均被引", "avg", rec["avg"])]
        for sheet, field, new in checks:
            if new is None:
                continue
            old = existing.get((sheet, s, field))
            if old is not None and not _eq(old, new):
                bad.append((sheet, s, field, old, new))
    return bad


# ------------------------------------------------------------------ 展示
def show(table: dict, esi_missing: bool):
    print(f"{'学科':30s} {'国际(分母/分子/百分位)':>26s} {'国内(分母/分子/百分位)':>26s} "
          f"{'CNCI':>8s} {'InCites 被引/论文/篇均':>22s}")
    for s in SUBJECTS:
        r = table[s]

        def fmt3(t):
            return "/".join("-" if v is None else (f"{v:.4f}" if isinstance(v, float) else str(v))
                            for v in t)

        print(f"{s:30s} {fmt3(r['intl']):>26s} {fmt3(r['cn']):>26s} {r['cnci']:>8.4f} "
              f"{r['cites']:>8d}/{r['papers']:<5d}/{r['avg']:>7.3f}"
              f"{'  [ESI 在榜]' if r['in_esi'] else ''}")
    if esi_missing:
        print("\n[i] 无 ESI 原始包：分母/位次列为空；其余为 InCites 侧值。")


# ------------------------------------------------------------------ 登记
def snapshot_meta(d: pathlib.Path) -> dict:
    have = [s for s in SUBJECTS if _resolve_file(d, s)]
    total = ynu = 0
    for s in have:
        p = _resolve_file(d, s)
        assert p is not None
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        try:
            rows = wb[wb.sheetnames[0]].max_row
        finally:
            wb.close()
        if rows and rows > 3:
            total += rows - 3
        # 云大在榜：导出要么 inlineStr（新，在 sheet1.xml）、要么 sharedStrings（旧），两处都扫
        with zipfile.ZipFile(p) as zf:
            hit = False
            for member in ("xl/sharedStrings.xml", "xl/worksheets/sheet1.xml"):
                try:
                    if b"yunnan university" in zf.read(member).lower():
                        hit = True
                        break
                except KeyError:
                    continue
            if hit:
                ynu += 1
    deploy = ""
    for probe in ("云南大学研究领域.xlsx", "中国研究领域.xlsx", "Chemistry.xlsx"):
        p = d / probe
        if not p.exists():
            continue
        try:
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            try:
                ws = wb[wb.sheetnames[-1]]
                txt = "\n".join(str(c) for r in ws.iter_rows(values_only=True) for c in r if c)
            finally:
                wb.close()
            mm = re.search(r"数据集更新日期\s*([\d-]+)", txt)
            if mm:
                deploy = mm.group(1)
                break
        except Exception:                                       # noqa: BLE001
            continue
    mm = re.search(r"(\d{4})\.(\d{2})\.(\d{2})$", d.name)
    cutoff = f"{mm.group(1)}-{mm.group(2)}-{mm.group(3)}" if mm else ""
    period = next((p for p, dn in SNAP_MAP.items() if dn == d.name), "")
    notes = []
    if len(have) == 0:
        notes.append("老命名体系（无 22 英文学科文件）")
    elif len(have) < 22:
        notes.append(f"缺 {22 - len(have)} 学科文件")
    if (d / "测试").is_dir():
        notes.append("含 测试/ 子目录")
    return {"dir": d.name, "cutoff": cutoff, "deploy": deploy, "n": len(have),
            "total": total, "ynu": ynu, "period": period, "notes": "; ".join(notes)}


def registry_rows() -> list:
    rows = [["快照目录", "WoS内容截止", "数据集更新", "学科文件", "机构条目合计", "云大在榜学科",
             "期次", "登记时间", "备注"]]
    today = dt.date.today().isoformat()
    for d in sorted(INCITES_ROOT.glob("数据覆盖*")):
        if not d.is_dir():
            continue
        m = snapshot_meta(d)
        rows.append([m["dir"], m["cutoff"], m["deploy"], f"{m['n']}/22", str(m["total"]),
                     f"{m['ynu']}/22", m["period"] or "待定", today, m["notes"]])
    return rows


def registry_xml(rows: list) -> str:
    body = []
    for rn, row in enumerate(rows, 1):
        cells = "".join(
            f'<c r="{col_letter(ci)}{rn}" t="inlineStr"><is><t xml:space="preserve">{v}</t></is></c>'
            for ci, v in enumerate(row, 1))
        body.append(f'<row r="{rn}">' + cells + "</row>")
    widths = "".join(f'<col min="{i}" max="{i}" width="{w}" customWidth="1"/>'
                     for i, w in enumerate([30, 13, 12, 10, 14, 12, 9, 12, 26], 1))
    dim = f"A1:{col_letter(len(rows[0]))}{len(rows)}"
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            f'<dimension ref="{dim}"/>'
            '<sheetViews><sheetView workbookViewId="0">'
            '<pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>'
            '</sheetView></sheetViews>'
            f'<cols>{widths}</cols><sheetData>{"".join(body)}</sheetData></worksheet>')


# ------------------------------------------------------------------ 命令
def _do_write(tx: LedgerTx, ledger: pathlib.Path, args, tag: str, preview_log: dict | None = None):
    if args.write:
        bak = LP.backup(ledger, tag=tag)
        print(f"[备份] {bak}")
    if args.out:
        out = pathlib.Path(args.out)
        out.parent.mkdir(parents=True, exist_ok=True)
        if out.resolve() != ledger.resolve():          # 同路径（就地写副本）不复制
            shutil.copy2(ledger, out)
        tx.path = out
    result = tx.commit()
    print(f"[写] {tx.path}")
    if preview_log:
        for k, v in preview_log.items():
            print(f"    {k}: {v['cells']} 格")
    print(f"    改动成员 {result['changed']}；新增成员 {result['added']}")


def cmd_period(args) -> int:
    period = args.target
    ledger = pathlib.Path(args.ledger)
    snap_name = SNAP_MAP.get(period) or args.snapshot
    if not snap_name:
        raise SystemExit(f"[拒绝] {period} 未登记 InCites 快照映射（SNAP_MAP：{sorted(SNAP_MAP)}）；"
                         f"可显式 --snapshot")
    snap_dir = INCITES_ROOT / snap_name
    if not snap_dir.is_dir():
        raise SystemExit(f"[拒绝] 快照目录不存在：{snap_dir}")
    esi_dir = pathlib.Path(args.esi_root) / period
    have_esi = esi_dir.is_dir()
    if not have_esi:
        print(f"[!] ESI 原始包未到位（{esi_dir}）：分母/位次算不了，只预览 InCites 侧。")
    table = compute(snap_dir, esi_dir if have_esi else None)
    q1_table, q1_notes = read_q1q4(snap_dir)
    for n in q1_notes:
        print(f"[!] {n}")
    show(table, esi_missing=not have_esi)

    if have_esi:
        ex = existing_values(ledger, period)
        bad = conflicts(table, ex)
        for s, per_ent in q1_table.items():
            for ent, vals in per_ent.items():
                for k, v in enumerate(vals):
                    old = ex.get(("Q1-Q4", s, f"{ent}.{k}"))
                    if old is not None and not _eq(old, v):
                        bad.append(("Q1-Q4", s, f"{ent}.{k}", old, v))
        if bad:
            print(f"\n[!] 与台账现有值冲突 {len(bad)} 处（口径变化须人工确认）：")
            for sheet, s, f, o, n in bad[:15]:
                print(f"    {sheet} · {s} · {f}: 台账 {o} ≠ 新算 {n}")
            if not args.force:
                print("    → 拒绝写入；确认无误后方可 --force")
                return 1
        elif ex:
            print(f"\n[✓] 与台账现有值逐格一致（抽查 {len(ex)} 个存量格）")
    if not (args.write or args.out):
        print("\n（预览模式。--out 写副本 / --write 写台账）")
        return 0
    if not have_esi:
        print("[拒绝] ESI 包未到位时不允许写台账（会让分母/位次为空）。")
        return 1

    tx = LedgerTx(ledger)
    log = {}

    def apply(sheet: str, result: tuple[str, dict]):
        xml, lg = result
        tx.put_sheet(sheet, xml)
        log[sheet] = lg

    apply("CNCI", plan_wide(tx, "CNCI", period, {s: r["cnci"] for s, r in table.items()}))
    for sheet, kind in (("新国际排名", "intl"), ("新国内排名", "cn")):
        apply(sheet, plan_group3(tx, sheet, period, table, kind))
    for sheet, kind in (("国际排名", "intl"), ("国内排名", "cn")):
        apply(sheet, plan_transposed(tx, sheet, period, table, kind))
    apply("总被引", plan_wide(tx, "总被引", period,
                              {s: (None if r["in_esi"] else r["cites"]) for s, r in table.items()}))
    apply("论文数", plan_wide(tx, "论文数", period,
                              {s: (None if r["in_esi"] else r["papers"]) for s, r in table.items()}))
    apply("篇均被引", plan_avg(tx, "篇均被引", period, table))
    if q1_table:
        apply("Q1-Q4", plan_q1q4(tx, period, q1_table))
    _do_write(tx, ledger, args, tag=f"回填{period}前", preview_log=log)
    return 0


def cmd_registry(args) -> int:
    ledger = pathlib.Path(args.ledger)
    rows = registry_rows()
    print(f"{'目录':34s} {'WoS截止':11s} {'更新':11s} {'文件':6s} {'机构':>8s} {'云大':5s} {'期次':7s}")
    for r in rows[1:]:
        print(f"{r[0]:34s} {r[1]:11s} {r[2]:11s} {r[3]:6s} {r[4]:>8s} {r[5]:5s} {r[6]:7s} {r[8]}")
    if not (args.write or args.out):
        print("\n（预览模式。--out 写副本 / --write 写台账）")
        return 0
    tx = LedgerTx(ledger)
    xml = registry_xml(rows)
    if REGISTRY_SHEET in set(re.findall(r'<sheet[^>]*name="([^"]+)"',
                                        tx.members["xl/workbook.xml"].decode("utf-8"))):
        tx.put_sheet(REGISTRY_SHEET, xml)
    else:
        tx.add_sheet(REGISTRY_SHEET, xml)
    if args.write:
        print(f"[备份] {LP.backup(ledger, tag='写InCites快照登记前')}")
    if args.out:
        out = pathlib.Path(args.out)
        out.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(ledger, out)
        tx.path = out
    result = tx.commit()
    print(f"[写] {tx.path}\n    改动 {result['changed']}；新增 {result['added']}")
    return 0


def cmd_preview(args) -> int:
    p = pathlib.Path(args.snapshot)
    snap_dir = p if p.is_absolute() else INCITES_ROOT / p
    if not snap_dir.is_dir():
        raise SystemExit(f"[拒绝] 快照目录不存在：{snap_dir}")
    table = compute(snap_dir, None)
    show(table, esi_missing=True)
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description="InCites 快照 → 台账当期列（预览/写回）+ 快照登记")
    ap.add_argument("target", help="期次（如 202607）或 registry / preview")
    ap.add_argument("--snapshot", default=None, help="显式指定快照目录名或路径（默认走 SNAP_MAP 登记）")
    ap.add_argument("--esi-root", default=str(ESI_ROOT))
    ap.add_argument("--ledger", default=str(LEDGER))
    ap.add_argument("--out", default=None, help="写副本（不动台账）")
    ap.add_argument("--write", action="store_true", help="写真台账（自动备份）")
    ap.add_argument("--force", action="store_true", help="允许覆盖与台账不一致的现有值")
    args = ap.parse_args()
    if args.target == "registry":
        return cmd_registry(args)
    if args.target == "preview":
        if not args.snapshot:
            dirs = sorted(q for q in INCITES_ROOT.glob("数据覆盖*") if q.is_dir())
            args.snapshot = str(dirs[-1].name)
        return cmd_preview(args)
    return cmd_period(args)


if __name__ == "__main__":
    sys.exit(main())
