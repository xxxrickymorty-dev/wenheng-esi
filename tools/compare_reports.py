# -*- coding: utf-8 -*-
"""原件 vs 生成稿：结构对照（骨架 / 图表 / 表格 / 文本 / 占位）。

与 src/backtest.py 互补：backtest 只比数字（大表 + 概述 + 省表云大行），
本脚本比"报告形态"——正文图表去向、槽位图数量、省域表是否为空、
顶尖论文清单差异、【手工：…】占位与母版残留标记。

用法（workdir=wenheng，python 用仓库 venv）：
  .venv\\Scripts\\python.exe <本技能目录>\\scripts\\compare_reports.py ^
      "out\\202607\\报告初稿_202607（试排）.docx" ^
      "..\\云南大学学科信息参考2026年第4期.docx" ^
      "out\\202607\\对照差异报告_202607结构.md"

输出：stdout 打印 + 可选第三个参数写 Markdown。
基线（202607 实测）见 SKILL.md「原件 vs 生成稿差异基线」。
返回码：0=正常（差异如实列出）；2=参数/文件错误。
"""
import hashlib
import pathlib
import sys
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from docx.oxml.ns import qn

PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

TEXT_ANCHORS = [
    "ESI最新数据显示", "与ESI前1‰接近度较高", "国内排名继续呈现较大降幅",
    "与上期数据相比", "前1%学科继续保持", "从国内排名上看",
    "接近度达到", "从预测情况看", "云南高校ESI前1%学科总数", "下图显示的是",
]

# 关键表形状：形状 → (标签, 说明)
KEY_TABLES = [
    ((13, 10), "T省表", "云南省高校对比（13 校 × 10 列）"),
    ((23, 12), "T矩阵", "学科×高校矩阵（22 学科 × 11 高校）"),
    (None, "T清单", "顶尖论文清单（275 行，按最长表识别）"),
]


def load_rels(path):
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/_rels/document.xml.rels"))
    return {r.get("Id"): r.get("Target") for r in root.findall("{%s}Relationship" % PKG_REL_NS)}


def load_media(path):
    with zipfile.ZipFile(path) as z:
        return {n: hashlib.md5(z.read(n)).hexdigest()
                for n in z.namelist() if n.startswith("word/media/")}


def body_graphics(doc, relmap):
    out = []
    for i, p in enumerate(doc.paragraphs):
        for b in p._element.findall(".//" + qn("a:blip")):
            out.append(("img", i, p.text.strip()[:45], relmap.get(b.get(qn("r:embed")), "?")))
        for c in p._element.findall(".//" + qn("c:chart")):
            out.append(("chart", i, p.text.strip()[:45], relmap.get(c.get("{%s}id" % R_NS), "?")))
    return out


def table_info(doc):
    info = []
    for ti, tb in enumerate(doc.tables):
        n_img = len(tb._tbl.findall(".//" + qn("a:blip")))
        n_ch = len(tb._tbl.findall(".//" + qn("c:chart")))
        first = tb.rows[0].cells[0].text.strip()[:26] if tb.rows else ""
        if ti == 0:
            first = " | ".join(c.text.strip()[:8] for c in tb.rows[0].cells[:6])
        info.append((ti, len(tb.rows), len(tb.columns), n_img, n_ch, first))
    return info


def find_by_shape(doc, shape):
    for tb in doc.tables:
        if (len(tb.rows), len(tb.columns)) == shape:
            return tb
    return None


def longest_table(doc):
    return max(doc.tables, key=lambda t: len(t.rows))


def cell_diff(tn, to):
    diffs, n = [], 0
    for r in range(min(len(tn.rows), len(to.rows))):
        for c in range(min(len(tn.columns), len(to.columns))):
            a = tn.rows[r].cells[c].text.strip().replace("\n", "")
            b = to.rows[r].cells[c].text.strip().replace("\n", "")
            n += 1
            if a != b:
                diffs.append((r, c, a[:40], b[:40]))
    return n, diffs


def get_paras(doc, anchors):
    out = {}
    for a in anchors:
        for p in doc.paragraphs:
            t = p.text.strip()
            if a in t:
                out[a] = t
                break
    return out


def markers(doc):
    import re
    ph, mk = [], []
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        ph += re.findall(r"【手工：[^】]+】", t)
        if ("（图）" in t) or ("更新图" in t):
            mk.append((i, t[:70]))
    return ph, mk


def main(new_path, old_path, out_md=None):
    dn, do = Document(new_path), Document(old_path)
    rn, ro = load_rels(new_path), load_rels(old_path)
    mn, mo = load_media(new_path), load_media(old_path)
    gn, go = body_graphics(dn, rn), body_graphics(do, ro)

    L = []
    add = L.append
    add("# 原件 vs 生成稿 · 结构对照")
    add("")
    add(f"- 生成稿：`{new_path}`")
    add(f"- 原件：`{old_path}`")
    add("")
    add("## 1. 骨架")
    add("")
    add(f"- 生成稿：段落 {len(dn.paragraphs)} · 表格 {len(dn.tables)} · 正文图形 {len(gn)} · 媒体 {len(mn)}")
    add(f"- 原件：段落 {len(do.paragraphs)} · 表格 {len(do.tables)} · 正文图形 {len(go)} · 媒体 {len(mo)}")
    same = sorted(set(mn) & set(mo))
    eq = [n for n in same if mn[n] == mo[n]]
    add(f"- 媒体同名同内容：{len(eq)} 张" + (f"（{', '.join(eq)}）" if eq else ""))
    rev = {}
    for n, h in mo.items():
        rev.setdefault(h, []).append(n)
    cross = [f"{n} == {rev[mn[n]][0]}" for n in sorted(mn)
             if mn[n] in rev and n not in same]
    if cross:
        add(f"- 跨名同内容（图片复用）：{len(cross)} 张——" + "；".join(cross))
    add("")
    add("## 2. 正文图形（段落级）")
    add("")
    add("| 文稿 | 类型 | 段落 | 位置文本 | 目标 |")
    add("|---|---|---|---|---|")
    for tag, g in (("生成稿", gn), ("原件", go)):
        for kind, i, t, tgt in g:
            add(f"| {tag} | {kind} | P{i:04d} | {t} | {tgt} |")
    add("")
    add("## 3. 表格清单（行×列 / 槽位图数 / 原生图数）")
    add("")
    add("| 表 | 生成稿 | 原件 |")
    add("|---|---|---|")
    for (ti, r1, c1, i1, h1, f1), (_, r2, c2, i2, h2, f2) in zip(table_info(dn), table_info(do)):
        add(f"| T{ti:02d} | {r1}×{c1} imgs={i1} charts={h1} | {r2}×{c2} imgs={i2} charts={h2} |")
    add("")
    add("## 4. 关键表逐格差异")
    add("")
    for shape, tag, desc in KEY_TABLES:
        tn = find_by_shape(dn, shape) if shape else longest_table(dn)
        to = find_by_shape(do, shape) if shape else longest_table(do)
        if tn is None or to is None:
            add(f"- {tag}（{desc}）：未找到")
            continue
        n, diffs = cell_diff(tn, to)
        add(f"- **{tag}**（{desc}）：{len(tn.rows)}×{len(tn.columns)}，比对 {n} 格，差异 **{len(diffs)}** 格")
        for r, c, a, b in diffs[:6]:
            add(f"    - r{r}c{c}: 生成={a!r} 原件={b!r}")
    add("")
    add("## 5. 关键正文段对照（各截 110 字）")
    add("")
    pn, po = get_paras(dn, TEXT_ANCHORS), get_paras(do, TEXT_ANCHORS)
    for a in TEXT_ANCHORS:
        add(f"### {a}")
        add(f"- 生成稿: {pn.get(a, '（无）')[:110]}")
        add(f"- 原件  : {po.get(a, '（无）')[:110]}")
    add("")
    add("## 6. 占位与残留标记")
    add("")
    pn_, mn_ = markers(dn)
    po_, mo_ = markers(do)
    add(f"- 生成稿：【手工】{len(pn_)} 处；残留标记（（图）/更新图）{len(mn_)} 处")
    for i, t in mn_:
        add(f"    - P{i:04d} {t}")
    add(f"- 原件：【手工】{len(po_)} 处；残留标记 {len(mo_)} 处")
    for i, t in mo_:
        add(f"    - P{i:04d} {t}")
    add("")
    add("---")
    add("*基线（202607 实测）：大表 21/22（多学科口径）、概述全对；"
        "重绘 3 图（排名/1%接近度/1‰潜力值）；缺 1‰趋势图 + 顶尖论文分布 2 图；"
        "T19/T20 全空；T21 274/275；【手工】16 处；残留标记 3 处。*")

    text = "\n".join(L)
    print(text)
    if out_md:
        pathlib.Path(out_md).write_text(text, encoding="utf-8")
        print(f"\n[已写入 {out_md}]")
    return 0


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(2)
    new, old = sys.argv[1], sys.argv[2]
    for p in (new, old):
        if not pathlib.Path(p).exists():
            print(f"文件不存在: {p}")
            sys.exit(2)
    sys.exit(main(new, old, sys.argv[3] if len(sys.argv) > 3 else None))
