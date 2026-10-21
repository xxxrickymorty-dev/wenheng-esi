# -*- coding: utf-8 -*-
"""换馆沙盘（决赛推广证据）：把「换馆=换配置包」做成可核验的契约检查。

用法（在 wenheng/ 下）：
    uv run python tools\\migration_check.py

它只做一件事：拿一份**台账副本**，按引擎的读表契约逐项检查——
表在不在、22 个学科名认不认、最新期次列有没有、母版在不在。
不跑流水线、不改原件；结论写 tmp/_migration_sandbox/readiness.json。

契约事实（母本名 / 表清单 / 学科名单 / 母版 / 数据源目录）全部来自
`ledger_contract`——本工具不再各写一份，换馆时只改契约模块与配置。

诚实边界：副本的数据仍是云大样本，这次验证的是「引擎不依赖校名、只依赖台账契约」；
真实换馆还要换学科清单、数据源与母版封面（页面里列了这四件事）。
"""
from __future__ import annotations

import datetime
import json
import pathlib
import sys

REPO = pathlib.Path(__file__).resolve().parents[1]
for p in (str(REPO), str(REPO / "src")):
    if p not in sys.path:
        sys.path.insert(0, p)

import openpyxl                                                  # noqa: E402

import ledger_contract as contract                               # noqa: E402

WORK = REPO / "tmp" / "_migration_sandbox"


def latest_period(ws) -> str:
    """在前 6 行里找 6 位期次，返回最大者；找不到返回空串。"""
    periods = []
    for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
        for v in row:
            s = str(v or "").strip()
            if len(s) == 6 and s.isdigit() and s.startswith("20"):
                periods.append(s)
    return max(periods) if periods else ""


def main() -> int:
    src = contract.require(REPO)
    WORK.mkdir(parents=True, exist_ok=True)
    copy = contract.sandbox_copy(src, WORK / "兄弟馆_监测表.xlsx")

    # 不用 read_only=True：仓里有守卫（test_guard_no_readonly）专门钉这件事；
    # 这里只读副本、不存盘，正常加载即可。
    wb = openpyxl.load_workbook(copy, data_only=True)
    try:
        checks = []

        missing_sheets = [s for s in contract.READINESS_SHEETS if s not in wb.sheetnames]
        checks.append({
            "id": "sheets",
            "label": f"必需的 {len(contract.READINESS_SHEETS)} 张工作表",
            "ok": not missing_sheets,
            "detail": "齐全" if not missing_sheets else "缺：" + "、".join(missing_sheets),
        })

        names = []
        if "总被引" in wb.sheetnames:
            for row in wb["总被引"].iter_rows(min_col=1, max_col=1, values_only=True):
                s = str(row[0] or "").strip()
                if s:
                    names.append(s)
        name_set = set(names)
        missing_disc = [f"{en}/{cn}" for en, cn in contract.DISCIPLINES_CN.items()
                        if en not in name_set and cn not in name_set]
        checks.append({
            "id": "disciplines", "label": "22 个规范学科名",
            "ok": not missing_disc,
            "detail": "22 / 22 认得出（中/英任一）" if not missing_disc else "缺：" + "、".join(missing_disc),
        })

        period = latest_period(wb["总被引"]) if "总被引" in wb.sheetnames else ""
        checks.append({
            "id": "period", "label": "最新期次列",
            "ok": bool(period),
            "detail": f"最新 {period}" if period else "没找到 6 位期次列",
        })

        master = contract.master_path(REPO)
        checks.append({
            "id": "master", "label": "报告结构母版",
            "ok": master.exists(),
            "detail": "在" if master.exists() else f"缺 master/{contract.MASTER_NAME}",
        })

        sources = [p for p in (contract.esi_root(REPO), contract.incites_root(REPO))
                   if p.exists()]
        checks.append({
            "id": "sources", "label": "数据源目录",
            "ok": len(sources) == 2,
            "detail": "ESI / InCites 都在" if len(sources) == 2 else "缺 ESI数据 或 incites 目录",
        })
    finally:
        wb.close()

    ready = all(c["ok"] for c in checks)
    story = {
        "generated_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "school": "某兄弟馆（契约演练）",
        "ledger": copy.name,
        "checks": checks,
        "ready": ready,
        "note": "数据仍是云大样本；这次验证的是引擎不依赖校名、只依赖台账契约。",
    }
    (WORK / "readiness.json").write_text(
        json.dumps(story, ensure_ascii=False, indent=2), encoding="utf-8")

    print("=" * 64)
    print("换馆沙盘 · 契约检查（只读副本，不跑流水线）")
    print("=" * 64)
    for c in checks:
        print(("PASS " if c["ok"] else "FAIL ") + c["label"] + " — " + c["detail"])
    print("-" * 64)
    print("结论：" + ("契约齐全，换馆只差配置。" if ready else "还有缺项，见上。"))
    print(f"故事卡：{(WORK / 'readiness.json').relative_to(REPO)}")
    return 0 if ready else 1


if __name__ == "__main__":
    raise SystemExit(main())
