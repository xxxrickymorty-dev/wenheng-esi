# -*- coding: utf-8 -*-
"""InCites 导出表头乱码：诊断 + 还原（本项目 2026-09-19 反推确立）。

## 乱码是怎么产生的（实测确认，两层叠加）

**第一层：UTF-8 文本被按 GB18030 解码**（下载的 CSV 被 Excel 用中文码页打开，
或有人另存时选错了编码）。特征：中文变成 `鍚嶇О`／`鏈熷垔涓` 这类字，
并可能夹带**私用区码点**（U+E000–U+F8FF，如 `\ue160`）。

**第二层：部分字节被替换成 `?`(0x3F)**。这是导出/另存过程中真实丢字节，
**该字节永久损失**。特征：`\xe6\x95\x3f` 这种「高位字节后跟 0x3F」。

## 为什么之前的修法失败

  `s.encode("gbk")`  → **失败**（`'gbk' codec can't encode character '\ue160'`）
  `s.encode("gb18030")` → **成功**   ← 必须用 GB18030

Python 的 `gbk`/`cp936` 表**不含** GBK 用户自定义区 → 私用区码点编不回去；
`gb18030` 含完整映射，所以只有它能完成往返。这不是"选哪个编码读一次"的问题，
而是**必须用 gb18030 反编码、再用 utf-8 重解**。

## 能修到什么程度

  · 第一层（纯 mojibake）        → **完全还原**（`Q1 鏈熷垔涓\ue160殑璁烘枃` → `Q1 期刊中的论文`）
  · 第一层 + 第二层（含 `?`）    → **只能还原到洞为止**；洞位解成 U+FFFD，
                                    尾段因对齐正确仍可解（这一点很关键：`?` 占 1 字节，
                                    所以**没有打乱 2 字节配对**）
  · 洞位本身                     → 不可逆。要么按**表头词表前缀匹配**补全，要么重新导出

## 所以「将来要按列名精确匹配」的正确做法，优先级从高到低

  1. **别依赖表头文字**——用语义判据（本项目 `backfill_q1q4.q_cols` 的做法）。
  2. 新下载的文件：读 CSV 一律 `encoding="utf-8-sig"`，失败退 `gbk`；**别用 Excel 双击**。
  3. 已有存档：`unmojibake()` 批量还原；含 `?` 的用 `repair_with_vocab()` 按词表补。
  4. 要长期干净的源：请馆员按 `数据与复现/incites/Incites 导出标题行文字乱码问题的解决方法*.docx`
     的流程**重新导出**（记事本删 `"索引号"` 引号 → 另存 TXT 选 UTF-8 → 拖进空白 Excel）。

用法：
    python tools/fix_incites_encoding.py --scan <目录>        # 只诊断
    python tools/fix_incites_encoding.py --selftest            # 跑自检
"""
import argparse
import csv
import pathlib
import re
import sys
import unicodedata

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

PUA = lambda ch: 0xE000 <= ord(ch) <= 0xF8FF


def _mojibake_score(s):
    """越大越像乱码：私用区 + 高频乱码字 + 常见正确字缺席。"""
    if not isinstance(s, str) or not s:
        return 0
    score = sum(1 for ch in s if PUA(ch)) * 3
    score += sum(s.count(w) for w in ("鍚", "鎺", "璁", "枃", "鏂", "鐨", "涓", "鏈", "熷", "鐧", "姣", "寮", "垔"))
    if any("\u4e00" <= ch <= "\u9fff" for ch in s) and not score:
        return 0
    return score


def unmojibake(s, strict=False):
    """乱码 → 正确文本。失败返回 None。

    🔴 必须用 **gb18030**：`gbk`/`cp936` 不含用户自定义区，遇到 `\\ue160` 直接抛错。
    🔴 `?`(0x3F) 是丢字节留下的洞，占 1 字节、**不会打乱 2 字节配对**，
       所以用 `errors="replace"` 解能保住洞之后的尾段；洞位得到 U+FFFD。
    """
    if not isinstance(s, str) or not s:
        return s
    dec = "utf-8" if strict else "utf-8"
    try:
        raw = s.encode("gb18030", errors="strict")
    except UnicodeEncodeError:
        return None
    try:
        return raw.decode(dec, errors="strict")
    except UnicodeDecodeError:
        if strict:
            return None
        return raw.decode(dec, errors="replace")


def repair_with_vocab(s, vocab):
    """含 U+FFFD / `?` 的表头，按**前缀**在词表里找唯一候选。

    洞通常吃掉最后一个字（如 `Q1 期刊中论文的百分?`），所以用「去掉洞后续」做前缀匹配。
    """
    cand = unmojibake(s)
    if cand is None or ("\ufffd" not in cand and "?" not in cand):
        return cand
    stem = re.split(r"[\ufffd?]", cand)[0]
    stem = stem.rstrip()
    if len(stem) < 4:
        return cand
    hits = [v for v in vocab if v.startswith(stem)]
    return hits[0] if len(hits) == 1 else cand


def collect_vocab(paths):
    """从**干净的新期导出**里收集规范表头，当词表用。"""
    import openpyxl
    vocab = set()
    for p in paths:
        try:
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            for r in ws.iter_rows(min_row=1, max_row=1, values_only=True):
                for c in r:
                    if isinstance(c, str) and c.strip() and not _mojibake_score(c):
                        vocab.add(c.strip())
            wb.close()
        except Exception:
            continue
    return sorted(vocab)


def analyze(path, vocab=None):
    """→ [(列序, 原文, 还原值, 状态)]；状态 ∈ OK / FIXED / PARTIAL / FAIL"""
    import openpyxl
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "gbk"):
            try:
                rows = list(csv.reader(raw.decode(enc).splitlines()))
                break
            except Exception:
                continue
        else:
            rows = list(csv.reader(raw.decode("utf-8", errors="replace").splitlines()))
        head = rows[0] if rows else []
    else:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        head = list(next(wb[wb.sheetnames[0]].iter_rows(min_row=1, max_row=1, values_only=True)))
        wb.close()

    out = []
    for i, c in enumerate(head):
        if not isinstance(c, str) or not c.strip():
            continue
        if not _mojibake_score(c):
            out.append((i, c, c, "OK"))
            continue
        fixed = unmojibake(c)
        if fixed is None:
            out.append((i, c, None, "FAIL"))
            continue
        st = "PARTIAL" if ("\ufffd" in fixed or "?" in fixed) else "FIXED"
        if st == "PARTIAL" and vocab:
            alt = repair_with_vocab(c, vocab)
            if alt and "\ufffd" not in alt and "?" not in alt:
                fixed, st = alt, "VOCAB"
        out.append((i, c, fixed, st))
    return out


def selftest():
    cases = [
        ("Q1 鏈熷垔涓\ue160殑璁烘枃", "Q1 期刊中的论文"),
        ("鏈熷垔涓\ue160殑璁烘枃", "期刊中的论文"),
        ("鍚嶇О", "名称"),
        ("鎺掑悕", "排名"),
        ("鍥介檯鍚堜綔璁烘枃", "国际合作论文"),
        ("% DOAJ 閲戣壊鏂囨。", "% DOAJ 金色文档"),
    ]
    bad = 0
    for raw, want in cases:
        got = unmojibake(raw)
        flag = "OK " if got == want else "FAIL"
        if got != want:
            bad += 1
        print(f"  [{flag}] {raw!r}\n         → {got!r}（期望 {want!r}）")
    # 干表头不该被"还原"
    for clean in ("名称", "Q1 期刊中论文的百分比", "Web of Science 论文数"):
        assert _mojibake_score(clean) == 0, f"{clean!r} 被误判为乱码"
    print(f"\n  干表头误判检查：通过")
    print(f"  自检：{'全部通过 ✅' if not bad else f'{bad} 项失败 ❌'}")
    return 1 if bad else 0


def main():
    ap = argparse.ArgumentParser(description="InCites 表头乱码诊断/还原（GB18030 往返）")
    ap.add_argument("--scan", metavar="DIR", help="扫描目录下所有 csv/xlsx 的表头")
    ap.add_argument("--file", help="只查一个文件")
    ap.add_argument("--vocab-from", metavar="DIR",
                    help="从该目录（放**干净的新期导出**）收集规范表头，用于补 `?` 洞")
    ap.add_argument("--limit", type=int, default=0, help="最多列几个文件（0=全部）")
    ap.add_argument("--show", type=int, default=6, help="每个文件最多展示几条问题表头")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()

    if a.selftest:
        return selftest()
    if not a.scan and not a.file:
        ap.print_help()
        return 0

    vocab = None
    if a.vocab_from:
        vp = [p for p in pathlib.Path(a.vocab_from).rglob("*")
              if p.suffix.lower() in (".csv", ".xlsx")]
        vocab = collect_vocab(vp)
        print(f"词表：{len(vocab)} 条（来自 {len(vp)} 个干净文件）\n")

    files = [pathlib.Path(a.file)] if a.file else sorted(
        p for p in pathlib.Path(a.scan).rglob("*") if p.suffix.lower() in (".csv", ".xlsx"))
    if a.limit:
        files = files[:a.limit]

    tot = {"OK": 0, "FIXED": 0, "VOCAB": 0, "PARTIAL": 0, "FAIL": 0}
    per_snap = {}                      # 顶层子目录（= 一次快照）→ 汇总
    n_files = 0
    for p in files:
        # 🔴 43 个快照里文件名是重复的（每个都有 Agricultural Sciences.xlsx），
        #    只 print(p.name) 无法定位到哪一期 → 必须输出相对路径并按快照汇总。
        rel = p.relative_to(pathlib.Path(a.scan)) if a.scan else pathlib.Path(p.name)
        snap = rel.parts[0] if len(rel.parts) > 1 else "."
        try:
            rows = analyze(p, vocab)
        except Exception as e:
            print(f"[跳过] {rel.as_posix()}: {type(e).__name__} {e}")
            continue
        stat = {k: sum(1 for r in rows if r[3] == k) for k in tot}
        for k in tot:
            tot[k] += stat[k]
        d = per_snap.setdefault(snap, {k: 0 for k in tot} | {"_files": 0, "_badfiles": 0})
        for k in tot:
            d[k] += stat[k]
        d["_files"] += 1
        if stat["FIXED"] or stat["VOCAB"] or stat["PARTIAL"] or stat["FAIL"]:
            n_files += 1
            d["_badfiles"] += 1
            print(f"\n{rel.as_posix()}")
            print(f"   OK {stat['OK']} / 可还原 {stat['FIXED']} / 词表补全 {stat['VOCAB']} / "
                  f"部分还原 {stat['PARTIAL']} / 无法还原 {stat['FAIL']}")
            shown = 0
            for i, raw, fx, st in rows:
                if st == "OK" or shown >= a.show:
                    continue
                shown += 1
                print(f"   [{st}] 列{i}: {raw[:60]!r}")
                print(f"          → {str(fx)[:60]!r}")
    print(f"\n合计（{len(files)} 个文件，有问题的 {n_files} 个）：")
    print(f"   干净 {tot['OK']} / 可完全还原 {tot['FIXED']} / 词表补全 {tot['VOCAB']} / "
          f"只能部分还原 {tot['PARTIAL']} / 无法还原 {tot['FAIL']}")
    if per_snap:
        bad = {s: d for s, d in per_snap.items() if d["_badfiles"]}
        print(f"\n\n── 按快照目录汇总（共 {len(per_snap)} 个，其中 {len(bad)} 个有脏表头）──")
        print(f"{'快照目录':<34}{'文件':>5}{'脏文件':>6}{'可还原':>7}"
              f"{'词表补全':>8}{'部分还原':>8}{'无法还原':>8}")
        for s in sorted(bad, key=lambda x: -bad[x]["_badfiles"]):
            d = bad[s]
            print(f"{s:<34}{d['_files']:>5}{d['_badfiles']:>6}{d['FIXED']:>7}"
                  f"{d['VOCAB']:>8}{d['PARTIAL']:>8}{d['FAIL']:>8}")
        clean = len(per_snap) - len(bad)
        if clean:
            print(f"（其余 {clean} 个快照表头全干净）")
    if tot["PARTIAL"] or tot["FAIL"]:
        print("\n   ⚠ 「部分/无法还原」的表头里有被 `?` 顶掉的字节，**不可逆**——"
              "加 --vocab-from 按词表补，或请馆员重新导出。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
