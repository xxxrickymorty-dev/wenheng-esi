# -*- coding: utf-8 -*-
"""任务5 副本验证：逐项核对 Q1-Q4 插入结果，任何一项不过就报错。"""
import pathlib
import re
import sys
import warnings
import xml.etree.ElementTree as ET
import zipfile

import openpyxl

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = pathlib.Path(__file__).resolve().parents[2]    # 工作区根（本文件在 wenheng/tools/ 下）
BK = ROOT / "归档" / "台账备份"
# 对「回填 Q1-Q4 前」备份与「当前真台账」比对，验证只有 sheet15 变化。
ORI = BK / "云大全学科排名、阈值动态检测表.回填Q1Q4前.bak.xlsx"
NEW = ROOT / "数据与复现" / "云大全学科排名、阈值动态检测表.xlsx"

NEWP = [f"{y}{m:02d}" for y in range(2019, 2025) for m in (1, 3, 5, 7, 9, 11)
        if not (y == 2019 and m < 3) and not (y == 2024 and m > 11)]
SHIFT = len(NEWP) * 12
DISC = ["Chemistry", "Plant & Animal Science", "Microbiology", "Engineering",
        "Materials Science", "Mathematics", "Environment/Ecology",
        "Biology & Biochemistry", "Pharmacology & Toxicology", "Geosciences",
        "Agricultural Sciences", "Computer Science", "Clinical Medicine",
        "Molecular Biology & Genetics", "Physics", "Social Sciences, general",
        "Multidisciplinary", "Neuroscience & Behavior", "Space Science",
        "Immunology", "Psychiatry/Psychology", "Economics & Business"]

fail = []


def a1(ci):
    s = ""
    while ci:
        ci, r = divmod(ci - 1, 26)
        s = chr(65 + r) + s
    return s


def chk(ok, msg):
    print(("  [OK]   " if ok else "  [FAIL] ") + msg)
    if not ok:
        fail.append(msg)


print("=" * 74)
print("1) zip 成员与逐字节对比")
print("=" * 74)
za, zb = zipfile.ZipFile(NEW), zipfile.ZipFile(ORI)
chk(set(za.namelist()) == set(zb.namelist()), f"成员集合一致（{len(za.namelist())}）")
diff = [n for n in za.namelist() if n in zb.namelist() and za.read(n) != zb.read(n)]
chk(diff == ["xl/worksheets/sheet15.xml"], f"仅目标 sheet 变化：{diff}")

print("=" * 74)
print("2) XML 合法性 / zip 完整性")
print("=" * 74)
chk(za.testzip() is None, "zip.testzip() 通过")
bad = []
for n in za.namelist():
    if n.endswith((".xml", ".rels")):
        try:
            ET.fromstring(za.read(n))
        except Exception as e:
            bad.append((n, str(e)))
chk(not bad, f"全部 XML 成员合法（{len(bad)} 处非法）{bad[:3]}")
za.close()
zb.close()

print("=" * 74)
print("3) openpyxl 结构")
print("=" * 74)
wb = openpyxl.load_workbook(NEW)
chk(len(wb.sheetnames) == 19, f"19 张表（实得 {len(wb.sheetnames)}）")
ws = wb["Q1-Q4"]
exp_dim = f"A1:{a1(133 + SHIFT)}25"
chk(ws.dimensions == exp_dim, f"Q1-Q4 dimension={ws.dimensions}（期望 {exp_dim}）")
n_merge = len(ws.merged_cells.ranges)
chk(n_merge == 44 + len(NEWP) * 4, f"合并区 {n_merge}（期望 {44 + len(NEWP) * 4}）")
chk(ws.freeze_panes == f"{a1(122 + SHIFT)}4", f"冻结窗格 {ws.freeze_panes}（期望 {a1(122 + SHIFT)}4）")

print("=" * 74)
print("4) 表头期次")
print("=" * 74)
hdr = []
for c in range(2, ws.max_column + 1, 12):
    v = ws.cell(row=1, column=c).value
    if v is not None:
        hdr.append(str(v))
        if len(hdr) >= len(NEWP):
            break
exp_new = [f"{p[:4]}.{p[4:]}" for p in NEWP]
try:
    got_new = [f"{float(h):.2f}" for h in hdr]
    exp_fmt = [f"{float(h):.2f}" for h in exp_new]
except ValueError:
    got_new, exp_fmt = hdr, exp_new
chk(got_new == exp_fmt, f"前 {len(NEWP)} 期期次 = {exp_new[0]}..{exp_new[-1]}（实得 {hdr[:3]}..{hdr[-1:] if hdr else []}）")

print("=" * 74)
print("5) 存量数值原样保留（右移后逐格比对）")
print("=" * 74)
wb2 = openpyxl.load_workbook(ORI)
ws2 = wb2["Q1-Q4"]
# 5a) A 列（学科名 + 表头行）**不动**——列位移只该影响 >=B 的列
a_mism = []
for r in range(1, 26):
    o, g = ws2.cell(row=r, column=1).value, ws.cell(row=r, column=1).value
    if o != g:
        a_mism.append((r, o, g))
chk(not a_mism, f"A 列（学科名）未被列位移波及：差异 {len(a_mism)} {a_mism[:4]}")
chk([ws.cell(row=r, column=1).value for r in range(4, 26)] == DISC,
    "A4:A25 = 22 学科名齐备")
n_old = 0
mism = []
for r in range(1, 26):
    for c in range(2, 134):
        old = ws2.cell(row=r, column=c).value
        got = ws.cell(row=r, column=c + SHIFT).value
        if isinstance(old, float) and isinstance(got, float):
            n_old += 1
            if abs(old - got) > 1e-9:
                mism.append((r, c, old, got))
        elif (old is None) != (got is None):
            mism.append((r, c, old, got))
chk(not mism, f"存量值（含数值与空）右移后一致：比了 {n_old} 个数值，差异 {len(mism)} {mism[:4]}")
wb2.close()

print("=" * 74)
print("6) 新增列：有源=数值、无源=空；实体覆盖统计")
print("=" * 74)
DISC_ROW = {r: ws.cell(row=r, column=1).value for r in range(4, 26)}
cov = {"云南大学": set(), "中国大陆": set(), "全球": set()}
bad_cells = []
for i, p in enumerate(NEWP):
    base = 2 + i * 12
    for ei, e in enumerate(("云南大学", "中国大陆", "全球")):
        b = base + ei * 4
        any_val = False
        for ri in range(4, 26):
            rowvals = [ws.cell(row=ri, column=b + q).value for q in range(4)]
            nv = [v for v in rowvals if v is not None]
            if nv:
                any_val = True
                if not all(isinstance(v, (int, float)) and 0 <= v <= 100 for v in nv):
                    bad_cells.append((p, e, DISC_ROW[ri], rowvals))
                # 四分位应大致和为 100
                if len(nv) == 4 and abs(sum(nv) - 100) > 1.5:
                    bad_cells.append((p, e, DISC_ROW[ri], f"和={sum(nv):.2f}"))
        if any_val:
            cov[e].add(p)
chk(not bad_cells, f"新增数值合法（0-100 且四分位和≈100）：异常 {len(bad_cells)} {bad_cells[:4]}")
print(f"   云南大学 有值期次 {len(cov['云南大学'])}：{sorted(cov['云南大学'])}")
print(f"   中国大陆 有值期次 {len(cov['中国大陆'])}：{sorted(cov['中国大陆'])}")
print(f"   全球     有值期次 {len(cov['全球'])}：{sorted(cov['全球'])}")

print("=" * 74)
print("7) 空列确实是空（未误填）")
print("=" * 74)
empt = []
for i, p in enumerate(NEWP):
    base = 2 + i * 12
    for ei, e in enumerate(("云南大学", "中国大陆", "全球")):
        if p in cov[e]:
            continue
        b = base + ei * 4
        for ri in range(4, 26):
            for q in range(4):
                if ws.cell(row=ri, column=b + q).value is not None:
                    empt.append((p, e, DISC_ROW[ri], q))
chk(not empt, f"应空的位置确实为空：误填 {len(empt)} {empt[:4]}")
wb.close()

print()
print("全部通过 ✅" if not fail else f"有 {len(fail)} 项失败 ❌")
sys.exit(1 if fail else 0)
