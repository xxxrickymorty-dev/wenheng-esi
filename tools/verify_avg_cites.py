# -*- coding: utf-8 -*-
"""任务4 副本验证：逐项核对插入结果，任何一项不过就报错。

检查项
 1) zip 成员集合不变；除目标 sheet 外**逐字节相同**
 2) 每个成员都是合法 XML；zip.testzip() 通过
 3) openpyxl 能打开；19 张表；篇均被引 dimension=A1:EI26、合并区 46 个
 4) 表头期次 = 201903..202407 + 原有 202409..202607（+ 原有重复槽）
 5) 存量数值原样保留（按 (行, 期次) 对齐）
 6) 新增「云大」值 == 总被引 ÷ 论文数（用**原台账**重算）
 7) 新增「中国大陆 / 国际」两列为空
"""
import pathlib
import re
import sys
import warnings
import xml.etree.ElementTree as ET
import zipfile
from xml.sax.saxutils import unescape

import openpyxl

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = pathlib.Path(r"C:\Users\damn\Desktop\首届高校图书馆")
BK = ROOT / "归档" / "台账备份"
# 对「回填前」与「回填后（仅篇均被引）」两份中间态，验证只有 sheet14 变化。
ORI = BK / "云大全学科排名、阈值动态检测表.回填篇均被引前.bak.xlsx"
NEW = BK / "云大全学科排名、阈值动态检测表.回填中间态-仅篇均被引.xlsx"

DISC = ["Chemistry", "Plant & Animal Science", "Microbiology", "Engineering",
        "Materials Science", "Mathematics", "Environment/Ecology",
        "Biology & Biochemistry", "Pharmacology & Toxicology", "Geosciences",
        "Agricultural Sciences", "Computer Science", "Clinical Medicine",
        "Molecular Biology & Genetics", "Physics", "Social Sciences, general",
        "Multidisciplinary", "Neuroscience & Behavior", "Space Science",
        "Immunology", "Psychiatry/Psychology", "Economics & Business"]
NEWP = [f"{y}{m:02d}" for y in range(2019, 2025) for m in (1, 3, 5, 7, 9, 11)
        if not (y == 2019 and m < 3) and not (y == 2024 and m > 7)]

fail = []


def chk(ok, msg):
    print(("  [OK]   " if ok else "  [FAIL] ") + msg)
    if not ok:
        fail.append(msg)


def col_letter(i):
    s = ""
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def col_index(s):
    n = 0
    for ch in s:
        n = n * 26 + ord(ch) - 64
    return n


def sheet_path(zf, name):
    wb = zf.read("xl/workbook.xml").decode("utf-8")
    rels = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]*)"',
                           zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")))
    for nm, rid in re.findall(r'<sheet[^>]*name="([^"]*)"[^>]*r:id="(rId\d+)"', wb):
        if nm == name:
            return "xl/" + rels[rid].lstrip("/")
    raise SystemExit(f"找不到 {name}")


print("=" * 74)
print("1) zip 成员与逐字节对比")
print("=" * 74)
zo, zn = zipfile.ZipFile(ORI), zipfile.ZipFile(NEW)
mo = {i.filename: zo.read(i.filename) for i in zo.infolist()}
mn = {i.filename: zn.read(i.filename) for i in zn.infolist()}
chk(set(mo) == set(mn), f"成员集合一致（{len(mo)} 个）")
tgt = sheet_path(zo, "篇均被引")
diff = [k for k in mo if mo[k] != mn[k]]
chk(diff == [tgt], f"仅目标 sheet 变化：{diff}")

print()
print("=" * 74)
print("2) XML 合法性 / zip 完整性")
print("=" * 74)
chk(zn.testzip() is None, "zip.testzip() 通过")
bad = []
for k, v in mn.items():
    if k.endswith(".xml") or k.endswith(".rels"):
        try:
            ET.fromstring(v)
        except ET.ParseError as e:
            bad.append((k, str(e)))
chk(not bad, f"全部 XML 成员合法（检查 {sum(1 for k in mn if k.endswith(('.xml','.rels')))} 个）{bad[:3]}")

print()
print("=" * 74)
print("3) openpyxl 结构")
print("=" * 74)
wb = openpyxl.load_workbook(NEW)
chk(len(wb.sheetnames) == 19, f"19 张表（实得 {len(wb.sheetnames)}）")
ws = wb["篇均被引"]
dim = re.search(r'<dimension ref="([^"]*)"',
                zn.read(tgt).decode("utf-8")).group(1)
chk(dim == "A1:EI26", f"dimension = {dim}（期望 A1:EI26）")
mc = re.findall(r'<mergeCell ref="([^"]*)"', zn.read(tgt).decode("utf-8"))
chk(len(mc) == 46, f"合并区 46 个（实得 {len(mc)}）")
chk(all(m.endswith("2") for m in mc), "合并区都在第 2 行")

print()
print("=" * 74)
print("4) 表头期次")
print("=" * 74)
hdr = {}
for c in ws[2]:
    if c.value is not None and c.column >= 2:
        hdr[col_letter(c.column)] = str(c.value)
got = list(hdr.values())
chk(got[:33] == NEWP, f"前 33 期 = 201903..202407（实得 {got[:33]}）")
chk(got[33:] == [str(p) for p in
                 [202409, 202411, 202501, 202503, 202505, 202507, 202509,
                  202511, 202601, 202603, 202605, 202607, 202607]],
    f"后 13 期 = 原样保留（实得 {got[33:]}）")

print()
print("=" * 74)
print("5/6/7) 数值：存量保留 / 新增=总被引÷论文数 / 另两列空")
print("=" * 74)


def read_num(zf, name, hdr_row=1):
    """返回 ({(行名, 期次): 值}, {期次: 列号})；hdr_row = 期次表头所在行。"""
    sp = sheet_path(zf, name)
    x = zf.read(sp).decode("utf-8")
    ss = [unescape("".join(re.findall(r"<t[^>]*>(.*?)</t>", s, re.S)))
          for s in re.findall(r"<si>(.*?)</si>",
                              zf.read("xl/sharedStrings.xml").decode("utf-8"), re.S)]
    hd, names, out = {}, {}, {}
    ROW = r'<row r="(\d+)"[^>]*>(.*?)</row>'
    for m in re.finditer(ROW, x, re.S):
        rn, body = int(m.group(1)), m.group(2)
        if rn == hdr_row:
            for cm in re.finditer(r'<c r="([A-Z]+)%d"([^>]*?)(?:/>|>(.*?)</c>)' % rn, body, re.S):
                v = re.search(r"<v>(.*?)</v>", cm.group(3) or "")
                if v and re.fullmatch(r"20\d{4}", v.group(1)) and 't="s"' not in cm.group(2):
                    hd[cm.group(1)] = v.group(1)
            continue
        ca = re.search(r'<c r="A%d"([^>]*?)(?:/>|>(.*?)</c>)' % rn, body, re.S)
        if not (ca and 't="s"' in ca.group(1)):
            continue
        va = re.search(r"<v>(\d+)</v>", ca.group(2) or "")
        if not va:
            continue
        nm = ss[int(va.group(1))]
        for cn, p in hd.items():
            cm = re.search(r'<c r="%s%d"([^>]*?)(?:/>|>(.*?)</c>)' % (cn, rn), body, re.S)
            if cm:
                v = re.search(r"<v>(.*?)</v>", cm.group(2) or "")
                if v:
                    out[(nm, p)] = float(v.group(1))
    ocol = {}
    for c, p in hd.items():          # 首见优先：202607 有重复槽，取有数据的那个
        ocol.setdefault(p, col_index(c))
    return out, ocol


orig, ocol = read_num(zo, "篇均被引", hdr_row=2)
cit, _ = read_num(zo, "总被引")
pap, _ = read_num(zo, "论文数")
SHIFT = 99

# 5) 存量期次原样（按 期次 -> 原列号 + 99 定位新列）
old_periods = ["202409", "202411", "202501", "202503", "202505", "202507",
               "202509", "202511", "202601", "202603", "202605", "202607"]
row_of = {c.value: c.row for c in ws["A"] if c.value}
chg, n5 = [], 0
for (nm, p), v in orig.items():
    if p not in old_periods or nm not in DISC or nm not in row_of:
        continue
    n5 += 1
    nv = ws.cell(row=row_of[nm], column=ocol[p] + SHIFT).value
    if nv is None or abs(float(nv) - v) > 1e-9:
        chg.append((nm, p, v, nv))
chk(n5 > 0 and not chg, f"存量数值原样保留：比了 {n5} 格，差异 {len(chg)} {chg[:4]}")

# 6) 新增值
bad2 = []
for r in range(4, 26):
    nm = ws.cell(row=r, column=1).value
    if nm not in DISC:
        continue
    for i, p in enumerate(NEWP):
        c = ws.cell(row=r, column=2 + i * 3)
        want = cit.get((nm, p))
        pw = pap.get((nm, p))
        exp = round(want / pw, 10)
        if c.value is None or abs(float(c.value) - exp) > 1e-9:
            bad2.append((nm, p, c.value, exp))
chk(not bad2, f"新增 726 格 = 总被引÷论文数（不符 {len(bad2)}：{bad2[:4]}）")

# 7) 另两列空
nonempty = []
for r in range(4, 26):
    for i in range(len(NEWP)):
        for k in (1, 2):
            c = ws.cell(row=r, column=2 + i * 3 + k)
            if c.value is not None:
                nonempty.append((r, i, k, c.value))
chk(not nonempty, f"新增「中国大陆/国际」两列留空（非空 {len(nonempty)}）")

wb.close()
zo.close()
zn.close()
print()
print("=" * 74)
print(("全部通过 ✅" if not fail else f"有 {len(fail)} 项未通过 ❌"))
for f in fail:
    print("   -", f)
