# -*- coding: utf-8 -*-
"""管理员书架只从部署产物按请求生成完整数据。"""
import json
from pathlib import Path

from server import shelf_data


def _write_period(root, period="202607", *, valid=True):
    target = root / period
    target.mkdir()
    (target / "overview_facts.json").write_text(json.dumps({
        "coverage_window": "2026年5月—2026年7月", "prev_label": "2026年5月",
        "rank_intl_now": 118, "rank_intl_delta": 2, "rank_cn_uni_now": 12,
        "top_papers_now": 24, "top_papers_delta": 1, "top3_by_top_papers": ["Chemistry"],
        "best_1permille": ["Chemistry"], "top2_not_entered": [], "n_disc_1pct": 3,
    }), encoding="utf-8")
    (target / "meta.json").write_text(json.dumps({
        "assertions": [["PASS"], ["FAIL"]], "top_total_by_period": {"202607": 24},
    }), encoding="utf-8")
    if valid:
        table = ("discipline,papers,total_cites,cnci,top_papers,thr_1pct,thr_1pct_gap,prox_change,rank_cn_change\n"
                 "Chemistry,10,20,2.5,3,1,0,+1,-\n")
    else:
        table = "discipline,papers\nwrong,1\n"
    (target / "table2.csv").write_text(table, encoding="utf-8")


def test_管理员数据只纳入完整的已部署期次(tmp_path):
    _write_period(tmp_path, "202607")
    _write_period(tmp_path, "202609", valid=False)
    payload = shelf_data.build_data([{"period": "202609"}, {"period": "202607"}], out_dir=tmp_path)
    assert [book["period"] for book in payload["books"]] == ["202607"]
    book = payload["books"][0]
    assert (book["no"], book["label"], book["n_assert_pass"], book["n_assert_total"]) == (38, "2026年7月", 1, 2)
    assert book["table"] == [{"en": "Chemistry", "cn": "化学", "papers": 10, "cites": 20,
                               "cnci": 2.5, "top": 3, "thr": 1, "gap": 0,
                               "prox": "+1", "rank": "-"}]
    assert payload["trend"] == {"202607": 24}


def test_所有期次异常仍返回有效且无指标来源的_js(tmp_path, monkeypatch):
    monkeypatch.setattr(shelf_data.config, "OUT_DIR", tmp_path)
    response = shelf_data.javascript_response([{"period": "202609"}])
    assert response.media_type == "application/javascript"
    assert response.body == b'window.WH_DATA = {"books":[],"trend":{}};\n'


def test_入库的静态空壳不再携带完整指标():
    root = Path(__file__).resolve().parents[1]
    prohibited = ("rank_intl_now", "cnci", "top_papers_now", "table2.csv")
    for path in (root / "web" / "public" / "shelf" / "data.js",
                 root / "web" / "dist" / "shelf" / "data.js"):
        content = path.read_text(encoding="utf-8-sig")
        assert "window.WH_DATA" in content
        assert not any(token in content for token in prohibited)


def test_预测摘要同时带两条路线(tmp_path):
    """`rows`=被引维度、`rows2`=排名维度；展示层据此出「双路线预测」页。"""
    target = tmp_path / "202607"
    target.mkdir()
    header = ("| 学科 | 模型 | 实现 | 滚动 MAPE | 拟合 MAPE | 指标A | 指标B "
              "| 是否达标 | 接近度 | 最早过线 | 较可能 | 最晚 |\n")
    (target / "预测选模与过线.md").write_text(
        "## 被引维度\n" + header
        + "| Physics | SARIMA | V1 | 1.50% | 2.00% | 1.4 | 1.6 | 是 | 75.13% | 2027-05 | 2028-01 | 2030-03 |\n"
        + "## 排名维度\n" + header
        + "| Physics | 阻尼Holt | V1 | 2.32% | 3.00% | 2.0 | 2.6 | 是 | 79.14% | 2027-05 | 2028-09 | — |\n"
        + "**实现口径**：a；b；c\n**序列窗口**：201909～202607 共 42 期。\n",
        encoding="utf-8")
    payload = shelf_data._read_forecast(target)
    assert payload["rows"][0]["en"] == "Physics" and payload["rows"][0]["model"] == "SARIMA"
    assert payload["rows"][0]["t1"] == "2027-05" and payload["rows"][0]["t3"] == "2030-03"
    assert payload["rows2"][0]["model"] == "阻尼Holt" and payload["rows2"][0]["prox"] == "79.14%"
    assert payload["window"] == "201909～202607 共 42 期"


def test_报告正文提取_从一概述起_章标题与子标题分型(tmp_path):
    """2026-10-11（票 05）：正文块＝h1 章标题 / h2 子标题 / p 段落；封面与名词解释跳过。"""
    from docx import Document
    doc = Document()
    doc.add_paragraph("学科信息参考")
    doc.add_paragraph("2026年第5期（总第39期）")
    doc.add_paragraph("云南大学发展规划处  图书馆编      2026年9月")
    doc.add_paragraph("名词解释：")
    doc.add_paragraph("阈值")
    doc.add_paragraph("学科进入ESI前1%的最低被引频次。")
    doc.add_paragraph("一 概述")
    doc.add_paragraph("本期全球共有10674个机构登上ESI前1%榜单。")
    doc.add_paragraph("学校发展趋势").style = doc.styles["List Paragraph"]
    doc.add_paragraph("与上期数据相比，我校国内机构排名上升1位。")
    doc.add_paragraph("二 我校学科发展情况")
    doc.add_paragraph("3.顶尖论文情况")
    doc.save(tmp_path / "报告初稿_202609（试排）.docx")

    blocks = shelf_data.report_text(tmp_path)
    assert [b["t"] for b in blocks] == ["h1", "p", "h2", "p", "h1", "h2"]
    assert blocks[0]["text"] == "一 概述"
    assert blocks[1]["text"].startswith("本期全球共有")
    assert blocks[2]["text"] == "学校发展趋势"
    assert blocks[4]["text"] == "二 我校学科发展情况"
    assert blocks[5]["text"] == "3.顶尖论文情况"
    # 缓存命中：同文件再取一致（docx 未变）
    assert shelf_data.report_text(tmp_path) == blocks


def test_报告正文缺docx给空表(tmp_path):
    assert shelf_data.report_text(tmp_path) == []


def test_报告正文含插图与表格_按口径删减(tmp_path):
    """2026-10-11（rick 定）：插图与表格进正文；顶尖论文清单/目录/大表不要；
    「近期XX发展参考指标变化情况」16 模块只留化学。"""
    import base64
    from docx import Document
    png = tmp_path / "dot.png"
    png.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="))

    def img_table(rows=3, cols=2):
        t = doc.add_table(rows=rows, cols=cols)
        for r in range(rows):
            for c in range(cols):
                t.cell(r, c).paragraphs[0].add_run().add_picture(str(png))
        return t

    doc = Document()
    doc.add_paragraph("一 概述")
    doc.add_paragraph("本期全球共有10674个机构。")
    doc.add_picture(str(png))                      # 插图段（无文本）
    tbl_toc = doc.add_table(rows=2, cols=2)
    tbl_toc.cell(0, 0).text = "名词解释"            # 目录表 → 不要
    doc.add_paragraph("——近期化学发展参考指标变化情况")
    img_table()                                    # 化学指标变化模块 → 保留
    doc.add_paragraph("——近期物理发展参考指标变化情况")
    img_table()                                    # 非化学模块 → 不要
    tbl_big = doc.add_table(rows=61, cols=2)        # 顶尖论文清单（>60 行）→ 不要
    tbl_big.cell(0, 0).text = "序号"
    tbl_big2 = doc.add_table(rows=2, cols=2)        # 大表（首格 ESI学科）→ 不要
    tbl_big2.cell(0, 0).text = "ESI学科"
    tbl_prov = doc.add_table(rows=2, cols=3)        # 省表类 → 保留
    tbl_prov.cell(0, 0).text = "机构名称"
    doc.save(tmp_path / "报告初稿_202609（试排）.docx")

    blocks = shelf_data.report_text(tmp_path)
    kinds = [b["t"] for b in blocks]
    assert kinds == ["h1", "p", "img", "p", "tbl", "tbl"]
    assert blocks[2]["rid"]                          # 插图带 rid（供 report-image 路由）
    assert blocks[3]["text"].startswith("——近期化学")  # 化学图题保留
    # 留下的第一个表＝化学指标变化（单元格带图），第二个＝省表类
    assert sum(len(c.get("imgs") or []) for r in blocks[4]["rows"] for c in r) == 6
    assert blocks[5]["rows"][0][0]["text"] == "机构名称"
