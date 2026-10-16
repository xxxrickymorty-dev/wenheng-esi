# -*- coding: utf-8 -*-
"""预测章图题与导语的学科名随实际出图学科改写（2026-10-16：改选学科优先后）。

母版里的 4 条图题和章首导语都把学科名写死（第 5 期原件写法）；报告目标学科改选优先后，
必须按 forecast_charts.json 的实际学科改写，否则「免疫学的图配着分子生物学的题」。
"""
import json
import pathlib
import sys

from docx import Document
from PIL import Image

REPO = pathlib.Path(__file__).resolve().parent.parent
for p in (str(REPO), str(REPO / "src")):
    if p not in sys.path:
        sys.path.insert(0, p)

import fill_docx                                                # noqa: E402


def _png(path):
    Image.new("RGB", (4, 4), (10, 40, 30)).save(path)


def test_图题与导语随实际学科改写(tmp_path):
    ov = tmp_path / "charts" / "_overview"
    ov.mkdir(parents=True)
    charts = []
    for i, subject in ((1, "免疫学"), (2, "分子生物学与遗传学")):
        for task, kind in (("citation", "cites"), ("rank", "rank")):
            png = ov / f"forecast_{i}_{kind}.png"
            _png(png)
            charts.append({"index": i, "subject": subject, "task": task, "kind": kind,
                           "model": "SARIMA模型", "first_crossing": None,
                           "png": png.name})
    (tmp_path / "forecast_charts.json").write_text(
        json.dumps({"charts": charts}, ensure_ascii=False), encoding="utf-8")

    doc = Document()
    doc.add_paragraph("本期针对我校的分子生物与遗传学和物理学2个学科，我们利用指数平滑模型，"
                      "从被引频次和国内排名两个指标对其进行预测分析，提供以下参考信息：")
    for cap in ("分子生物与遗传学学科被引频次与ESI阈值的模型预测图",
                "分子生物与遗传学学科国内排名与学科国内ESI前1%机构数的模型预测图",
                "物理学科被引频次与ESI阈值的模型预测图",
                "物理学学科国内排名与学科国内ESI前1%机构数的模型预测图"):
        doc.add_paragraph(cap)
        doc.add_paragraph("")

    slots, warns = [], []
    fill_docx.fill_forecast_charts(doc, tmp_path, slots, warns)

    texts = [p.text.strip() for p in doc.paragraphs]
    assert "免疫学学科被引频次与ESI阈值的模型预测图" in texts
    assert "免疫学学科国内排名与学科国内ESI前1%机构数的模型预测图" in texts
    assert "分子生物学与遗传学学科被引频次与ESI阈值的模型预测图" in texts
    assert any("本期针对我校的免疫学和分子生物学与遗传学2个学科" in t for t in texts)
    assert not any(t.endswith("模型预测图") and t.startswith("物理") for t in texts)
    assert warns == []
    assert any(s["location"].startswith("预测章/导语") for s in slots)
