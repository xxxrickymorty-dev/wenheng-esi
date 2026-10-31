# -*- coding: utf-8 -*-
"""监测版（高校 ESI 报告）测试：名单完整性 + 缓存完整性 + 最简版渲染。

数据依赖 = `monitor_data/_cache_<期次>.json`（缓存入库）与 `数据与复现/ESI数据/`
（只读源）。缺缓存时按登记簿规则跳过并留探针。
"""
from __future__ import annotations

import sys
import pathlib

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "src"))

from _guard_registry import skip_unless            # noqa: E402
from monitor import compare as mcmp                # noqa: E402
from monitor import insight as mi                  # noqa: E402
from monitor import report as mr                   # noqa: E402
from monitor import schools as sch                 # noqa: E402

_PERIOD = mr.latest_period()
_cache_probe = lambda: bool(_PERIOD) and mr.cache_path(_PERIOD).exists()   # noqa: E731


def test_schools_list_integrity():
    """名单自检：中文名唯一、985/211 都有量、每条带省份。"""
    cns = [s["cn"] for s in sch.SCHOOLS]
    assert len(cns) == len(set(cns)), "中文名有重复"
    n985 = sum(1 for s in sch.SCHOOLS if s["is985"])
    assert n985 == 39 and len(sch.SCHOOLS) - n985 >= 70
    assert all(s["prov"] and s["en"] for s in sch.SCHOOLS)
    assert sch.find("海南大学") is not None
    assert sch.find("不存在大学") is None


def test_cache_covers_built_in_schools():
    """缓存完整性：109 校在机构总表/学科榜里能对上号（未入册校白名单除外）。"""
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建（monitor_data/_cache_*.json 缺）", _cache_probe)
    cache = mr.load_cache(_PERIOD)
    known = {row[0] for row in cache["cn_inst"]}
    known |= {en for v in cache["fields"].values() for en in v["schools"]}
    missing = [s["cn"] for s in sch.SCHOOLS if s["en"] not in known]
    assert missing == ["中国政法大学"], f"名单里有缓存对不上的校：{missing}"


def test_render_minimal_report_hainan():
    """基础渲染（海大）：12 学科行、2 个前 1‰、站位/俱乐部/预警/对标章俱全、无论文层。"""
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建", _cache_probe)
    cache = mr.load_cache(_PERIOD)
    html = mr.render("海南大学", _PERIOD, cache)
    assert html.count('<span class="en">') == 12, "海大应上榜 12 学科"
    assert "前 1‰ 俱乐部（2 学科）" in html
    for sec in ("学科站位", "前 1‰ 俱乐部", "进位预警", "对标"):
        assert sec in html
    assert "省内对比" not in html, "对标不限省内，省内章已撤（rick 2026-10-31 裁）"
    assert "期刊集中度" not in html and "高被引 / 热点论文" not in html, "最简版不含论文层章"
    assert "data:image" not in html, "图表只许内联 svg，不许 base64 位图"
    assert "红件清单" not in html, "监测版不借云大产线的红件清单语言"


def test_render_unregistered_school_states():
    """未入册校（中国政法大学）也出报告：如实呈现，不编数字。"""
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建", _cache_probe)
    cache = mr.load_cache(_PERIOD)
    html = mr.render("中国政法大学", _PERIOD, cache)
    assert "未收录该校" in html
    assert "本期不出学科行" in html


def test_doctor_checks_pass_on_202609():
    """doctor D1/D2/D3/D6 对当期全绿（D3 的未入册校=政法大学，白名单内不算写法错）。"""
    from monitor import doctor as mdoc
    if not _PERIOD or not (mr.ESI_ROOT / _PERIOD).is_dir() or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次数据或缓存缺席", lambda: _PERIOD is not None
                    and (mr.ESI_ROOT / _PERIOD).is_dir())
    rep = mdoc.check(_PERIOD)
    assert rep["fails"] == 0, f"doctor 有 fail：{[i for i in rep['items'] if not i['ok']]}"
    d3 = next(i for i in rep["items"] if i["id"] == "D3")
    assert "未入册 1" in d3["detail"], "应恰好一所未入册校（中国政法大学）"


def test_insight_facts_deterministic():
    """解读事实行：海大＝12 学科全入围、2 个前 1‰、总排名 871——全部带数字的关系句。"""
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建", _cache_probe)
    facts = mi.facts_of("海南大学", mr.load_cache(_PERIOD))
    joined = "；".join(facts)
    assert "上榜 12/22 学科" in joined
    assert "2 个学科进入全球前 1‰" in joined and "农业科学" in joined
    assert "全球机构总排名 871" in joined
    assert mi.number_guard("被引 68,391 领先", facts) is False
    assert mi.number_guard("被引 999,999 领先", facts) is True, "事实外数字必须拒收"


def test_render_insight_section():
    """解读段：给了就渲染并如实标注来源；不给就不出现。"""
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建", _cache_probe)
    cache = mr.load_cache(_PERIOD)
    base = mr.render("海南大学", _PERIOD, cache)
    assert "解读" not in base, "缺省不带解读段"
    with_it = mr.render("海南大学", _PERIOD, cache,
                        insight={"text": "总体格局稳中有进。", "mode": "rule"})
    assert "总体格局稳中有进" in with_it and "规则拼装" in with_it
    modelled = mr.render("海南大学", _PERIOD, cache,
                         insight={"text": "格局概括。", "mode": "model"})
    assert "模型解读（数字已过盾校验）" in modelled

def test_compare_report_render_and_generate():
    """对比件：两校矩阵、总量表、解读段；文件名确定性；未知校拒绝。"""
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建", _cache_probe)
    cache = mr.load_cache(_PERIOD)
    html = mcmp.render(["海南大学", "清华大学", "中国政法大学"], _PERIOD, cache)
    assert "总量对比" in html and "学科矩阵" in html
    assert "海南大学" in html and "清华大学" in html
    assert "未入册" in html, "政法大学应在总量表如实标注未入册"
    assert html.count('<i>/') >= 12, "矩阵格应是 排名/榜内所数"
    f1 = mcmp.generate(["海南大学", "清华大学"], _PERIOD)
    f2 = mcmp.generate(["清华大学", "海南大学"], _PERIOD)
    assert f1 == f2 and f1.exists(), "同集合不同顺序＝同一份文件（幂等）"
    import pytest
    with pytest.raises(KeyError):
        mcmp.generate(["海南大学", "不存在大学"], _PERIOD)


def test_compare_insight_facts():
    """对比事实行：总量序＋分学科龙头，全带数字（数字白名单同源）。"""
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建", _cache_probe)
    facts = mcmp.facts_of(["海南大学", "清华大学", "贵州大学"], mr.load_cache(_PERIOD))
    joined = "；".join(facts)
    assert "本次对比 3 校" in joined and "清华大学" in joined
    assert any("位次最靠前" in x for x in facts)
    assert mi.number_guard("编造 12345 名", facts) is True


# ---------------------------------------------------------------- 渲染层丰富（趋势/分位/哑铃）
_trend_probe = lambda: (mr.MONITOR_ROOT / "_trend" / "_index.json").exists()   # noqa: E731


def test_render_trend_section_hainan():
    """趋势章（海大，_trend 已入库）：三章趋势 + 折线/时间轴图 + 章号四五顺延 + 文字清单退场。"""
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建", _cache_probe)
    if not _trend_probe():
        skip_unless("趋势库未建（monitor_data/_trend/_index.json 缺）", _trend_probe)
    cache = mr.load_cache(_PERIOD)
    html = mr.render("海南大学", _PERIOD, cache)
    assert '<span class="num">三</span>趋势' in html
    assert html.count("<svg") >= 3, "总被引轨迹 + 最强学科排名 + 入榜与进位时间轴"
    assert "滚动十年窗" in html, "章首 note 应交代回填口径"
    assert "入榜与进位时间轴" in html, "甘特轴替代文字清单"
    assert html.count("起进入前 1%（") <= 3, "文字清单形态大幅退场（归并句另算）"
    assert '<span class="num">四</span>进位预警' in html
    assert '<span class="num">五</span>对标' in html


def test_render_trend_insufficient_school():
    """无趋势学校（政法不在任何一期档案里）：整章替换为不足 note，不出图。"""
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建", _cache_probe)
    cache = mr.load_cache(_PERIOD)
    html = mr.render("中国政法大学", _PERIOD, cache)
    assert "趋势数据不足，随期次积累解锁" in html
    assert "<svg" not in html


def test_render_percentile_and_rca_columns():
    """站位表新增分位/相对强度列：分位恒有（rank/m），RCA 缺全球学科时降级 —，注入后出 ×值。"""
    import re
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建", _cache_probe)
    cache = mr.load_cache(_PERIOD)
    html = mr.render("海南大学", _PERIOD, cache)
    assert "分位" in html and "相对强度" in html
    assert re.search(r"前 \d+\.\d%", html), "分位应是一位小数百分比"
    assert "部分覆盖下指数偏高" in html, "RCA 算法与局限必须写在 note 里"
    if not cache.get("global_papers"):
        assert '<td class="n">—</td>' in html, "缓存无全球学科时 RCA 降级为 —"
    gp_file = mr.ESI_ROOT / _PERIOD / "全球学科.xlsx"
    if not gp_file.exists():
        skip_unless("全球学科.xlsx 缺", lambda: gp_file.exists())
    gp = mr._global_papers(mr.ESI_ROOT / _PERIOD)
    assert len(gp) == 22, "全球学科应解析出 22 行"
    assert mr._global_papers(mr.ESI_ROOT / "190001") == {}, "缺文件 → 空字典不报错"
    cache2 = dict(cache)
    cache2["global_papers"] = gp
    html2 = mr.render("海南大学", _PERIOD, cache2)
    en = sch.find("海南大学")["en"]
    mine = sum(r["schools"][en][2] for r in cache["fields"].values() if r["schools"].get(en))
    g_sum = sum(gp.values())
    chem = cache["fields"]["Chemistry"]["schools"][en]
    rca = (chem[2] / mine) / (gp["CHEMISTRY"] / g_sum)
    assert f"×{rca:.2f}" in html2, "化学行相对强度应与手算一致"


def test_charts_line_chart():
    """line_chart：polyline + 图例 + 缺值断段 + ref 虚线 + y_invert，纯函数输出 svg。"""
    from monitor import charts as mc
    svg = mc.line_chart([("甲", [("202001", 100.0), ("202003", 150.0), ("202005", None),
                                 ("202007", 220.0)]),
                         ("乙", [("202001", 80.0), ("202003", 150.0), ("202005", 190.0),
                                 ("202007", 210.0)])])
    assert svg.startswith("<svg") and svg.endswith("</svg>")
    assert svg.count("<polyline") == 2, "两条线各一段连续 polyline"
    assert "<rect" in svg, "甲的 202005 缺值后成孤点，应以点标记承接断线"
    assert "甲" in svg and "乙" in svg, "图例放图下方"
    svg2 = mc.line_chart([("排名", [("202001", 390), ("202005", 120)])],
                         y_invert=True, ref_lines=[(30, "1‰ 线")])
    assert "stroke-dasharray" in svg2 and "1‰ 线" in svg2


def test_charts_dumbbell():
    """dumbbell：每行两圆点（a/b 各一）+ 值标注 + 图例，纯函数输出 svg。"""
    from monitor import charts as mc
    svg = mc.dumbbell([("化学", 390, 45), ("数学", 120, 96)],
                      name_a="海南大学", name_b="清华大学")
    assert svg.startswith("<svg") and svg.endswith("</svg>")
    assert svg.count("<circle") == 4, "两行 × 两端圆点"
    assert "海南大学" in svg and "清华大学" in svg
    assert "390" in svg and "45" in svg, "值标注在点旁"


def test_compare_dumbbell_section():
    """恰两校 → 哑铃小节（章二），矩阵顺延章三；三校不出现哑铃（矩阵已足够）。"""
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建", _cache_probe)
    cache = mr.load_cache(_PERIOD)
    html2 = mcmp.render(["海南大学", "清华大学"], _PERIOD, cache)
    assert "逐学科位次哑铃" in html2
    assert "<svg" in html2 and "排名越小越好" in html2
    assert '<span class="num">三</span>学科矩阵' in html2
    html3 = mcmp.render(["海南大学", "清华大学", "贵州大学"], _PERIOD, cache)
    assert "逐学科位次哑铃" not in html3
    assert '<span class="num">二</span>学科矩阵' in html3


# ---------------------------------------------------------------- 数据装配层 + 打印 + Word 导出（P1 收尾）


def test_report_data_structure():
    """_report_data：五段齐＋解读段、nums 4 项、站位表 11 列头、显示值全部字符串化。"""
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建", _cache_probe)
    cache = mr.load_cache(_PERIOD)
    d = mr._report_data("海南大学", _PERIOD, cache,
                        insight={"text": "总体格局稳中有进。", "mode": "rule"})
    assert [s["num"] for s in d["sections"]] == ["一", "二", "三", "四", "五"]
    assert d["insight"] == {"text": "总体格局稳中有进。", "mode": "rule"}
    assert len(d["nums"]) == 4
    assert all(isinstance(n["value"], str) and isinstance(n["note"], str) for n in d["nums"])
    t1 = d["sections"][0]["table"]
    assert len(t1["head"]) == 11 and len(t1["colclasses"]) == 11
    assert len(t1["subs"]) == len(t1["rows"]) == 12, "海大 12 学科行，英文名逐行对齐"
    assert all(isinstance(c, str) for row in t1["rows"] for c in row), "显示值一律字符串化"
    assert d["sections"][1]["title"] == "前 1‰ 俱乐部（2 学科）"
    assert len(d["algorithms"]) == 3
    d2 = mr._report_data("中国政法大学", _PERIOD, cache)
    assert d2["nums"] == [] and d2["insight"] is None and d2["absent_note"]
    assert d2["sections"][0]["table"]["placeholder"], "未入册校站位表带 placeholder 合并行"
    html = mr.render("海南大学", _PERIOD, cache)
    assert '<span class="pill pm">前 1‰</span>' in html, "档位徽标经 pill_cols 保留"


def test_docx_export_hainan():
    """docx 导出：落 monitor_data/<校>/<期次>.docx；读回标题/表头与 _report_data 同源；幂等。"""
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建", _cache_probe)
    import docx as pydocx
    from monitor import docx as mdx
    cache = mr.load_cache(_PERIOD)
    d = mr._report_data("海南大学", _PERIOD, cache)
    p0 = mr.MONITOR_ROOT / "海南大学" / f"{_PERIOD}.docx"
    if p0.exists():
        p0.unlink()          # 装配层措辞升级后重建一次，保证读回的是当前 _report_data
    p = mdx.export("海南大学", _PERIOD, cache)
    assert p == mr.MONITOR_ROOT / "海南大学" / f"{_PERIOD}.docx" and p.exists()
    mtime = p.stat().st_mtime
    p2 = mdx.export("海南大学", _PERIOD, cache)
    assert p2 == p and p2.stat().st_mtime == mtime, "已存在直接返回，不重写"
    doc = pydocx.Document(str(p))
    texts = [q.text for q in doc.paragraphs]
    assert texts[0] == f"海南大学 · 学科监测报告 · {_PERIOD}"
    joined = "\n".join(texts)
    assert d["nums"][2]["value"] in joined, "总被引与 HTML 同源（同一 _report_data）"
    assert "三、趋势" in joined and "入榜与进位时间轴" in joined, "趋势章文字保留（图不放 docx）"
    t = doc.tables[0]
    assert [t.rows[0].cells[i].text for i in range(3)] == ["学科", "全球排名", "分位"]
    assert t.rows[1].cells[0].text == d["sections"][0]["table"]["rows"][0][0]


def test_print_css_present():
    """打印样式：@media print 进 _CSS 并随之进入渲染产物。"""
    assert "@media print" in mr._CSS and "@page" in mr._CSS
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建", _cache_probe)
    html = mr.render("海南大学", _PERIOD, mr.load_cache(_PERIOD))
    assert "@media print" in html and "break-inside:avoid" in html


# ---------------------------------------------------------------- 呈现层二段（环图/甘特轴/条形）

def test_charts_donut():
    """donut：段＝stroke circle、图例带占比、中心 title；空输入/全零返回空壳。"""
    from monitor import charts as mc
    svg = mc.donut([("农业科学", 5000), ("化学", 3424), ("工程学", 2000)],
                   title="在榜论文合计\n10,424")
    assert svg.startswith("<svg") and svg.endswith("</svg>")
    assert svg.count("<circle") == 3, "三段环形"
    assert "48.0%" in svg and "32.8%" in svg, "图例自动算占比"
    assert "10,424" in svg and "农业科学" in svg
    assert mc.donut([]).count("<circle") == 0, "空输入 → 空 svg 壳"
    assert mc.donut([("甲", 0), ("乙", 0)]).count("<circle") == 0, "合计为 0 → 空壳"


def test_charts_entry_timeline():
    """entry_timeline：事件行有菱形/五角星标记，pre 行有实心圆＋浅色延伸；空输入空壳。"""
    from monitor import charts as mc
    svg = mc.entry_timeline(
        [{"label": "植物与动物学", "first_seen": "202001", "entry1": None,
          "entrypm": "202007", "pre": True},
         {"label": "农业科学", "first_seen": "202003", "entry1": "202005", "entrypm": None}],
        periods=["202001", "202003", "202005", "202007"])
    assert svg.startswith("<svg") and svg.endswith("</svg>")
    assert svg.count("<polygon") >= 4, "菱形＋五角星事件标记＋图例"
    assert svg.count("<circle") == 1, "pre 行轨道起点实心圆"
    assert ">202005<" in svg and ">202007<" in svg, "x 轴期次刻度"
    assert "进入前 1%" in svg and "进入前 1‰" in svg, "底部图例"
    assert mc.entry_timeline([], periods=["202001"]).count("<polygon") == 0, "空输入 → 空壳"


def test_render_presentation_charts_hainan():
    """海大呈现面：站位环图、1‰ 条带、预警接近度条形、入榜时间轴俱全；文字清单退场。"""
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建", _cache_probe)
    if not _trend_probe():
        skip_unless("趋势库未建", _trend_probe)
    cache = mr.load_cache(_PERIOD)
    html = mr.render("海南大学", _PERIOD, cache)
    assert "在榜学科论文结构" in html, "站位章 donut"
    assert "1‰线 第" in html, "俱乐部章 permille_strip 行内标注"
    assert "1× 阈值" in html, "预警章接近度条形（ref 线）"
    assert "入榜与进位时间轴" in html, "趋势章甘特轴"
    assert html.count("<svg") >= 5, "两折线 + 时间轴 + 环图 + 条带 + 条形"
    assert html.count("起进入前 1%（") <= 3, "文字清单形态大幅退场"


def test_render_timeline_all_pre_tsinghua():
    """清华（全 pre 极端）：无 1% 进位事件行，归并成一句 note；页面文字清单残留 ≤3。"""
    if not _PERIOD or not mr.cache_path(_PERIOD).exists():
        skip_unless("监测版期次缓存未建", _cache_probe)
    if not _trend_probe():
        skip_unless("趋势库未建", _trend_probe)
    cache = mr.load_cache(_PERIOD)
    d = mr._report_data("清华大学", _PERIOD, cache)
    tl = d["sections"][2].get("timeline") or {}
    assert tl, "清华趋势点充足，时间轴规格应在"
    assert tl["pre_note"].startswith("档案起点前已入围"), "归并 note 承载 pre 学科"
    html = mr.render("清华大学", _PERIOD, cache)
    assert html.count("起进入前 1%（") <= 3, "全 pre 场景无进位文字复读"
    assert "档案起点前已入围" in html
