# -*- coding: utf-8 -*-
"""ledger_patch（台账 zip 级安全改写）回归：只动目标成员 / 公式拒改 / 只写数字 / 备份不覆盖。

2026-09-18 实事故：正则拼串少个空格 → 整份台账（唯一事实源）变砖。此后写路径加了
良构闸（xlsx_io.rewrite）与 audit 四查。本文件用 openpyxl 合成的小台账钉住
ledger_patch 自身的口子；zip 原语的细节归 tests/test_xlsx_io.py，不重复。

铁律背景：台账禁止 openpyxl 存盘（会丢 6 张图表/条件格式/打印设置），
write_cells 只重写目标 sheet 的 XML、其余成员逐字节原样复制——这是 diff_members
断言的核心不变量。
"""
import pathlib
import sys
import zipfile

import openpyxl
import pytest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "src"))

import ledger_patch  # noqa: E402


def _ensure_shared_strings(p):
    """补一个空的 sharedStrings.xml 成员：openpyxl 合成件用 inlineStr、没有该成员，
    而 write_cells 硬性读它（真台账必有共享串表）。空表即可——字符串仍走 inline 路径。"""
    with zipfile.ZipFile(p) as zf:
        items = {i.filename: zf.read(i.filename) for i in zf.infolist()}
    if "xl/sharedStrings.xml" in items:
        return
    items["xl/sharedStrings.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>'
    ).encode("utf-8")
    ct = items["[Content_Types].xml"].decode("utf-8")
    ct = ct.replace(
        "</Types>",
        '<Override PartName="/xl/sharedStrings.xml" ContentType='
        '"application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
    items["[Content_Types].xml"] = ct.encode("utf-8")
    tmp = p.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in items.items():
            zf.writestr(name, data)
    tmp.replace(p)


@pytest.fixture()
def synth_ledger(tmp_path):
    """合成小台账：两个 sheet；总被引 A 列学科、B 列数字、C 列公式。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "总被引"
    ws["A1"] = "学科"
    ws["B1"] = "202003"
    ws["A2"] = "Economics & Business"
    ws["B2"] = 30
    ws["C2"] = "=B2*2"
    ws2 = wb.create_sheet("其他")
    ws2["A1"] = "别动我"
    p = tmp_path / "synth.xlsx"
    wb.save(p)
    _ensure_shared_strings(p)
    return p


# ---------------------------------------------------------------- 读

def test_sheet_names(synth_ledger):
    assert ledger_patch.sheet_names(synth_ledger) == ["总被引", "其他"]


def test_read_grid_行内字符串与数字(synth_ledger):
    # openpyxl 合成件是 inlineStr（真台账是共享串表）——两条路 cell_text 都要走得通
    grid, shared = ledger_patch.read_grid(synth_ledger, "总被引")
    assert shared == []
    assert ledger_patch.cell_text(*grid[2]["A"], shared) == "Economics & Business"
    assert ledger_patch.cell_text(*grid[2]["B"], shared) == "30"


def test_cell_text_三类单元格():
    # 行内字符串（工具新写的期次标签用这种）
    assert ledger_patch.cell_text('t="inlineStr"', "<is><t>期次</t><t>标签</t></is>", []) == "期次标签"
    # 共享字符串下标
    assert ledger_patch.cell_text(' t="s"', "<v>1</v>", ["甲", "乙"]) == "乙"
    # 数字 / 普通文本；空单元格
    assert ledger_patch.cell_text("", "<v>42</v>", []) == "42"
    assert ledger_patch.cell_text("", None, []) is None
    # HTML 实体还原
    assert ledger_patch.cell_text(' t="s"', "<v>0</v>", ["A&amp;B"]) == "A&B"


def test_locate_按列找行_按行找列_宽松匹配(synth_ledger):
    grid, shared = ledger_patch.read_grid(synth_ledger, "总被引")
    assert ledger_patch.locate(grid, shared, "A", "Economics & Business") == 2
    assert ledger_patch.locate(grid, shared, "A", " economics & business ") == 2   # 空格大小写无关
    assert ledger_patch.locate(grid, shared, 1, "202003", by="row") == "B"
    with pytest.raises(KeyError):
        ledger_patch.locate(grid, shared, "A", "不存在的学科")
    with pytest.raises(KeyError):
        ledger_patch.locate(grid, shared, 1, "不存在的表头", by="row")


# ---------------------------------------------------------------- 写

def test_write_cells_只动目标成员_其余逐字节原样(synth_ledger, tmp_path):
    before = tmp_path / "before.xlsx"
    before.write_bytes(synth_ledger.read_bytes())

    applied = ledger_patch.write_cells(synth_ledger, "总被引", {(2, "B"): 31})
    assert applied == [((2, "B"), "30", 31)]

    grid, shared = ledger_patch.read_grid(synth_ledger, "总被引")
    assert ledger_patch.cell_text(*grid[2]["B"], shared) == "31"

    with zipfile.ZipFile(synth_ledger) as zf:
        tgt = ledger_patch.sheet_member(zf, "总被引")
    d = ledger_patch.diff_members(before, synth_ledger, "总被引")
    assert d["changed"] == [tgt]                 # 只有目标 sheet 的 XML 变了
    assert d["added"] == [] and d["removed"] == []
    assert d["unexpected_changed"] == []         # 非目标成员一字节不动（2026-09-18 事故防线）


def test_write_cells_同值不写_返回空(synth_ledger, tmp_path):
    before = tmp_path / "before.xlsx"
    before.write_bytes(synth_ledger.read_bytes())
    assert ledger_patch.write_cells(synth_ledger, "总被引", {(2, "B"): 30}) == []
    assert ledger_patch.diff_members(before, synth_ledger)["changed"] == []


def test_write_cells_公式格直接拒绝(synth_ledger):
    with pytest.raises(SystemExit, match="含公式"):
        ledger_patch.write_cells(synth_ledger, "总被引", {(2, "C"): 5})


def test_write_cells_只收数字(synth_ledger):
    with pytest.raises(SystemExit, match="只写数字"):
        ledger_patch.write_cells(synth_ledger, "总被引", {(2, "B"): "31"})


def test_backup_永不覆盖_同名自动加序号(synth_ledger, tmp_path, monkeypatch):
    monkeypatch.setattr(ledger_patch, "BACKUP_DIR", tmp_path / "台账备份")
    b1 = ledger_patch.backup(synth_ledger, "改前")
    b2 = ledger_patch.backup(synth_ledger, "改前")
    assert b1.exists() and b2.exists()
    assert b1 != b2 and b2.name.endswith(".1.bak.xlsx")
    assert b1.read_bytes() == synth_ledger.read_bytes()      # copy2 全量备份


def test_audit_写入后四查全过(synth_ledger, tmp_path):
    before = tmp_path / "before.xlsx"
    before.write_bytes(synth_ledger.read_bytes())
    ledger_patch.write_cells(synth_ledger, "总被引", {(2, "B"): 31})
    res = ledger_patch.audit(before, synth_ledger, "总被引", {(2, "B"): 31})
    assert res["ok"] is True
    assert res["readback"][(2, "B")] == 31
    assert res["sheets"] == 2
    assert res["bad_xml"] == []
