# -*- coding: utf-8 -*-
"""陈列馆游客数据只能来自报告的三段白名单。"""
import csv
import json

from docx import Document

from server import shelf_preview


def _report(path, *, valid=True, cover=True, delayed_body=False, issue="2026年第5期（总第39期）",
            month="2026年9月", toc_extra=None):
    document = Document()
    cover_lines = ("   ", "学科信息参考", issue, f"云南大学发展规划处  图书馆编      {month}")
    for line in (cover_lines if cover else cover_lines[1:]):
        document.add_paragraph(line)
    toc = list(zip(shelf_preview._TOC_TITLES, ("", "1", "2", "26", "28", "30")))
    if toc_extra:
        toc.append(toc_extra)
    table = document.add_table(rows=len(toc), cols=2)
    for index, (title, page) in enumerate(toc):
        table.cell(index, 0).text = title if valid else "正文数据"
        table.cell(index, 1).text = page
    document.add_paragraph("名词解释：")
    for term in shelf_preview._GLOSSARY_TERMS:
        document.add_paragraph(term)
        document.add_paragraph(f"{term}的说明。")
    if delayed_body:
        document.add_paragraph("正文秘密：291 篇及运行路径 C:/private/out")
    document.add_paragraph("一 概述")
    document.add_paragraph("正文秘密：291 篇及运行路径 C:/private/out")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "表二秘密"
    document.save(path)


def _products(path, *, disciplines=22, n_disc=14, top=291, slots=None, assertions=None, pending=None):
    """给某期写上首页事实条的产物出处；值原样写入以便试坏形状。
    slots / assertions / pending 分别写 red_slots.json、meta.json 断言、audit_record.json
    的「手工槽清点」；None＝不写该文件（缺产物 → 对应项省略）。"""
    with (path / "table2.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["discipline", "papers", "cnci", "total_cites"])
        for index in range(disciplines):
            writer.writerow([f"学科{index}", "100", "1.0", "2000"])
    facts = {}
    if n_disc is not None:
        facts["n_disc_1pct"] = n_disc
    if top is not None:
        facts["top_papers_now"] = top
    (path / "overview_facts.json").write_text(json.dumps(facts, ensure_ascii=False), encoding="utf-8")
    if slots is not None:
        rows = [{"location": f"正文/p{i}", "old": "旧", "new": "新", "source": "sentences.py"}
                for i in range(slots)]
        (path / "red_slots.json").write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
    if assertions is not None:
        meta = {"assertions": [["PASS", f"断言{i}"] for i in range(assertions)]}
        (path / "meta.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    if pending is not None:
        audit = {"verdict": "review", "items": [{"id": "manual_slots", "n_pending": pending}]}
        (path / "audit_record.json").write_text(json.dumps(audit, ensure_ascii=False), encoding="utf-8")


def test_只提取封面目录和名词解释_正文哨兵不泄漏(tmp_path):
    period = tmp_path / "202609"
    period.mkdir()
    _report(period / "报告初稿_202609.docx")

    payload = shelf_preview.build_preview([{"period": "202609"}], out_dir=tmp_path)
    book = payload["books"][0]
    assert payload["preview"] is True and payload["trend"] == {}
    assert set(book) == {"period", "cover", "toc", "glossary", "label", "no"}
    assert book["cover"] == ["   ", "学科信息参考", "2026年第5期（总第39期）", "云南大学发展规划处  图书馆编      2026年9月"]
    assert book["toc"] == [
        {"title": title, "page": page}
        for title, page in zip(shelf_preview._TOC_TITLES, ("", "1", "2", "26", "28", "30"))
    ]
    assert book["label"] == "2026年9月" and book["no"] == 39
    assert "正文秘密" not in json.dumps(payload, ensure_ascii=False)
    assert "表二秘密" not in json.dumps(payload, ensure_ascii=False)


def test_缺文件或结构异常逐期关闭(tmp_path):
    malformed = tmp_path / "202607"
    malformed.mkdir()
    _report(malformed / "报告初稿_202607.docx", valid=False)

    payload = shelf_preview.build_preview(
        [{"period": "202609"}, {"period": "202607"}], out_dir=tmp_path)
    assert payload == {"preview": True, "books": [], "trend": {}}


def test_缺封面标志或延迟正文标题都关闭(tmp_path):
    missing_cover = tmp_path / "202601"
    delayed_heading = tmp_path / "202603"
    missing_cover.mkdir()
    delayed_heading.mkdir()
    _report(missing_cover / "报告初稿_202601.docx", cover=False)
    _report(delayed_heading / "报告初稿_202603.docx", delayed_body=True)

    payload = shelf_preview.build_preview(
        [{"period": "202601"}, {"period": "202603"}], out_dir=tmp_path)
    assert payload == {"preview": True, "books": [], "trend": {}}


def test_按期次升序并补齐书架封面字段(tmp_path):
    old = tmp_path / "202601"
    latest = tmp_path / "202609"
    old.mkdir()
    latest.mkdir()
    _report(old / "报告初稿_202601.docx", issue="2026年第1期（总第35期）", month="2026年1月")
    _report(latest / "报告初稿_202609.docx")

    payload = shelf_preview.build_preview(
        [{"period": "202609"}, {"period": "202601"}], out_dir=tmp_path)
    assert [(book["period"], book["label"], book["no"]) for book in payload["books"]] == [
        ("202601", "2026年1月", 35), ("202609", "2026年9月", 39),
    ]


def test_事实条盖章_数取自最新一期产物(tmp_path):
    old = tmp_path / "202601"
    latest = tmp_path / "202609"
    old.mkdir()
    latest.mkdir()
    _report(old / "报告初稿_202601.docx", issue="2026年第1期（总第35期）", month="2026年1月")
    _report(latest / "报告初稿_202609.docx")
    _products(old, disciplines=20, n_disc=12, top=245)
    _products(latest, slots=3, assertions=2, pending=1)

    payload = shelf_preview.build_preview(
        [{"period": "202601"}, {"period": "202609"}], out_dir=tmp_path)
    assert payload["stats"] == {
        "monitored": 22, "topOnePercent": 14, "highlyCited": 291,
        "docSlots": 3, "selfChecks": 2, "pendingManual": 1,
        "source": "202609 期流水线产物（table2、overview_facts、red_slots、meta、audit_record）",
    }


def test_事实条缺项就省项_最新一期全缺就整个不给(tmp_path):
    partial = tmp_path / "202601"
    bare = tmp_path / "202603"
    partial.mkdir()
    bare.mkdir()
    _report(partial / "报告初稿_202601.docx", issue="2026年第1期（总第35期）", month="2026年1月")
    _report(bare / "报告初稿_202603.docx", issue="2026年第2期（总第36期）", month="2026年3月")

    # 202601 的产物缺了表、数字形状也不对（真假值 / 字符串不算数）→ 一项都不给。
    _products(partial, disciplines=0, n_disc=True, top="很多")
    payload = shelf_preview.build_preview([{"period": "202601"}], out_dir=tmp_path)
    assert "stats" not in payload

    # 只剩一个形状对的数；新三样给了坏形状（该是清单的给成标量）→ 新项按缺处理，
    # 也不拖累已有数——宁可缺，不可给没有出处的数。
    _products(partial, disciplines=0, n_disc=14, top=None)
    (partial / "red_slots.json").write_text(json.dumps({"坏": True}), encoding="utf-8")
    (partial / "meta.json").write_text(json.dumps({"assertions": "不是清单"}), encoding="utf-8")
    (partial / "audit_record.json").write_text(json.dumps({"items": "不是清单"}), encoding="utf-8")
    payload = shelf_preview.build_preview([{"period": "202601"}], out_dir=tmp_path)
    assert payload["stats"] == {
        "topOnePercent": 14, "source": "202601 期流水线产物（table2、overview_facts、red_slots、meta、audit_record）",
    }

    # 新三样只有形状对的：槽位数照给、待人工 0 是真值照给（不省），其余缺 → 省。
    _products(partial, disciplines=0, n_disc=None, top=None, slots=5, pending=0)
    payload = shelf_preview.build_preview([{"period": "202601"}], out_dir=tmp_path)
    assert payload["stats"] == {
        "docSlots": 5, "pendingManual": 0,
        "source": "202601 期流水线产物（table2、overview_facts、red_slots、meta、audit_record）",
    }

    # 最新一期（202603）什么产物都没有 → 整条不给，也不回退拿旧期凑数。
    payload = shelf_preview.build_preview(
        [{"period": "202601"}, {"period": "202603"}], out_dir=tmp_path)
    assert [book["period"] for book in payload["books"]] == ["202601", "202603"]
    assert "stats" not in payload


def test_目录混入正文或封面期次错标都关闭(tmp_path):
    toc_leak = tmp_path / "202601"
    mismatched_cover = tmp_path / "202603"
    toc_leak.mkdir()
    mismatched_cover.mkdir()
    _report(toc_leak / "报告初稿_202601.docx", toc_extra=("正文秘密", "31"))
    _report(mismatched_cover / "报告初稿_202603.docx")

    payload = shelf_preview.build_preview(
        [{"period": "202601"}, {"period": "202603"}], out_dir=tmp_path)
    assert payload == {"preview": True, "books": [], "trend": {}}
