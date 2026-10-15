# -*- coding: utf-8 -*-
"""灵活报告（第二条产线）的测试。

覆盖三类：

1. **红线**（纯函数，任何环境都跑）—— 纵轴范围越权必拒、白名单必拒、观点分级必守；
2. **不可分叉**（需数据）—— 数字逐点可回读、图与文同源、复刻产物不被污染；
3. **无模型也能跑**（需数据）—— 规则后端下产出完整分析稿（本设计的核心主张）。

数据驱动的用例缺数据时走 `_guard_registry.skip_unless`（登记理由 + 探针），
与其余测试同一套约定。
"""
import hashlib
import json
import pathlib
import re

import pytest

from _guard_registry import skip_unless

from flexible import mdview, opinion, planner, polish, queries, report, spec
from flexible.datasource import DataSource

# ---------------------------------------------------------------- 夹具


def _period_dir():
    """挑一个有产物的期次目录（优先 202607，它是金标期次）。"""
    from conftest import REPO
    outs = [p for p in (REPO / "out").glob("20*") if p.is_dir()]
    if not outs:
        return None
    gold = [p for p in outs if p.name == "202607" and (p / "ledger_long.parquet").is_file()]
    if gold:
        return gold[0]
    ok = [p for p in outs if (p / "ledger_long.parquet").is_file()]
    return ok[0] if ok else None


@pytest.fixture(scope="session")
def src():
    d = _period_dir()
    if d is None:
        skip_unless("out/ 下无含 ledger_long.parquet 的期次产物",
                    lambda: bool(_period_dir()))
    return DataSource(d)


@pytest.fixture(scope="session")
def flex_report_out(src, tmp_path_factory):
    """生成一次灵活报告（规则后端），供多条用例复用。

    ⚠ 显式注入 `RuleBrain`：session 级夹具的创建早于函数级 env 隔离，
    不注入就会带着本机 .env 的真 key 出网（2026-09-29 踩——结果随网络成败漂移）。
    """
    from conftest import REPO
    import ai
    from ai._brain import RuleBrain
    d = REPO / src.dir
    with ai.using(brain=RuleBrain()):
        r = report.generate(d, "single_trend")
    return r


# ================================================================ 1. 红线
class TestSpec红线_纵轴范围越权必拒:
    """AI 不得指定图表纵轴范围 —— 不是"不推荐"，是**结构上不存在这个字段**。"""

    @pytest.mark.parametrize("key", [
        "ylim", "ymin", "ymax", "vmin", "vmax", "range", "limits",
        "axis_range", "y_range", "ylimits", "yticks", "y0", "y1", "yscale",
        # 大小写与下划线变体也必须拦住（归一化后再比）
        "YLim", "y_min", "Y-MAX", "axisRange", "YLimits",
    ])
    def test_纵轴类键一律拒绝(self, key):
        with pytest.raises(spec.SpecError, match="纵轴范围"):
            spec.validate({"kind": "cnci", "disc": "Physics", key: 1})

    def test_语义参数可以过(self):
        """反向轴 / 参考值 / 零基准 —— 只给语义，不给界。"""
        assert spec.validate({"kind": "rank_pct", "disc": "Physics", "invert": False})
        assert spec.validate({"kind": "cnci", "disc": "Physics", "include": [1.0]})
        assert spec.validate({"kind": "top_papers", "disc": "Physics", "floor_zero": True})

    def test_未知图型被拒(self):
        with pytest.raises(spec.SpecError, match="未知图型"):
            spec.validate({"kind": "scatter3d", "disc": "Physics"})

    def test_未登记的语义参数被拒(self):
        # cnci 只接受 include；给它 zero_base 必须拒（不然图型契约形同虚设）
        with pytest.raises(spec.SpecError, match="不接受参数"):
            spec.validate({"kind": "cnci", "disc": "Physics", "zero_base": True})

    def test_单学科图必须给学科(self):
        with pytest.raises(spec.SpecError, match="需要 disc"):
            spec.validate({"kind": "top_papers"})

    def test_概览图不许带学科(self):
        with pytest.raises(spec.SpecError, match="概览图"):
            spec.validate({"kind": "overview_1pct", "disc": "Physics"})

    def test_窗口长度越界被拒(self):
        with pytest.raises(spec.SpecError, match="n_periods"):
            spec.validate({"kind": "cnci", "disc": "Physics", "n_periods": 99})

    def test_图型目录不含轴范围参数(self):
        """给模型看的目录里也不能出现轴范围键（否则等于提示它去要）。"""
        for c in spec.catalog():
            assert all(k not in str(c).lower() for k in ("ylim", "vmin", "vmax"))
            assert set(c["semantic"]) <= {"invert", "include", "zero_base", "floor_zero"}


class TestQueries白名单:
    """LLM 拿数的唯一入口：未登记的一律拒。"""

    def test_未知查询类型被拒(self, src):
        with pytest.raises(queries.QueryError, match="未登记的查询类型"):
            queries.execute(src, "drop_table")

    def test_未声明参数被拒(self, src):
        with pytest.raises(queries.QueryError, match="未声明参数"):
            queries.execute(src, "series", metric="papers", disc="Chemistry", limit=1)

    def test_未知指标被拒(self, src):
        with pytest.raises(queries.QueryError, match="不在白名单"):
            queries.series(src, "secret_metric", "Chemistry")

    def test_未知学科被拒(self, src):
        with pytest.raises(queries.QueryError, match="不在名单内"):
            queries.series(src, "papers", "Not A Discipline")

    def test_未知期次被拒(self, src):
        with pytest.raises(queries.QueryError, match="不在数据范围内"):
            queries.series(src, "papers", "Chemistry", ["199901"])

    def test_发给模型的schema不含任何数值(self):
        """字段字典里只能有指标名与说明，不能夹带数据值。"""
        s = json.dumps(queries.SCHEMA, ensure_ascii=False)
        for bad in ("total_cites\": 1", "0.845", "11558"):
            assert bad not in s


class TestOpinion观点分级:
    """三级判断（2026-10-14 判断草稿口径）：描述性＝规则生成；解释性＝带标注的假设；
    专业判断＝可写但需馆员拍板——**判断句未带 `==…==` 待审线即判漏标**。"""

    def test_未带标记的判断句被报漏标(self):
        hits = opinion.guard("建议加大该学科投入，应当优先补齐短板。")
        assert hits, "未带待审线的判断句必须被拦下（新口径：漏标，而非「出现即违规」）"

    def test_带了待审线就不算漏标(self):
        assert not opinion.guard("==建议加大该学科投入==。")

    def test_手工空位里的措辞不算违规(self):
        txt = f"待人工填写：{opinion.manual('资源投入建议')}"
        assert not opinion.guard(txt)

    def test_假设缺可验证方式即报错(self):
        with pytest.raises(ValueError, match="可验证方式"):
            opinion.hypothesis("上升因为发文量增长", "")

    def test_假设来源必须是规定值(self):
        with pytest.raises(ValueError, match="来源标记"):
            opinion.hypothesis("x", "y", source="某模型")

    def test_假设渲染必带三件套(self):
        h = opinion.hypothesis("上升可能与发文量同向", "查《论文数》202601–202607 各期",
                               opinion.SRC_AI)
        out = opinion.render_hypothesis(h)
        assert "待验证假设" in out and opinion.SRC_AI in out and "可验证方式" in out

    def test_描述性判断是规则生成的(self):
        """同一个入参必须给出同一句话 —— 不经模型，确定性。"""
        pts = [{"period": "202601", "value": 1.0}, {"period": "202603", "value": 1.5}]
        a = opinion.descriptive_trend(pts, "接近度", "化学")
        b = opinion.descriptive_trend(pts, "接近度", "化学")
        # 🔴 别断言某个具体词（句式现在按数据形态分档，"上升"只是其中一种说法）
        assert a == b and any(w in a for w in ("上升", "走高", "升到"))


class TestPlanner无模型也能跑:
    """**本设计的核心主张**：模型位可插拔，无模型时流程照跑（D6 后唯一出口＝ai 内部口）。"""

    def test_未配云不碰网络(self, src, monkeypatch):
        """把 urlopen 打成炸弹：没配云时规划必须照样成立、零网络。"""
        import urllib.request

        def boom(*a, **k):
            raise AssertionError("未配云时不应发起任何网络调用")

        monkeypatch.setattr(urllib.request, "urlopen", boom)
        plan = planner.build(src, "single_trend")
        assert plan["backend"] == "rule"
        assert plan["charts"] and plan["disc"] and plan["periods"]

    def test_模型不可用即回退规则(self, src):
        """调用失败不是错误，是降级 —— 演示站永远走这条。"""
        from _ai_helpers import FailingBrain
        import ai
        with ai.using(brain=FailingBrain()):
            plan = planner.build(src, "single_trend")
        assert plan["backend"] == "rule"
        assert plan["fallback_reason"]
        assert plan["charts"], "回退后仍须给出完整规划"

    def test_规划出网前数字已占位(self, src):
        """🔴 唯一出口的硬纪律：发给模型的规划依据里没有任何数字原值。"""
        from _ai_helpers import ScriptedBrain, Turn, raw_digits
        import ai
        brain = ScriptedBrain([lambda b: Turn(text='{"disc": "Chemistry", "charts": '
                                                   '[{"kind": "cnci", "disc": "Chemistry"}]}')])
        with ai.using(brain=brain):
            plan = planner.build(src, "single_trend")
        assert plan["backend"] == "scripted"
        assert brain.briefs, "应有一次外发"
        for t in brain.briefs[0].egress_texts():
            assert not raw_digits(t), f"外发含数字原值：{t[:80]!r}"

    def test_模型给出越权键则整份回退(self, src):
        """模型想给纵轴范围 → 不接受、不崩溃，回退规则规划。"""
        from _ai_helpers import ScriptedBrain, Turn
        import ai
        with ai.using(brain=ScriptedBrain(
                [Turn(text='{"disc": "Chemistry", "ylim": "auto"}')])):
            plan = planner.build(src, "single_trend")
        assert plan["backend"] == "rule"
        assert "越权键" in plan["fallback_reason"]

    def test_模型带入新数字整份回退(self, src):
        """模型自造/回填数字 → 盾拒，整份回退规则（规划输出不许有数字）。"""
        from _ai_helpers import ScriptedBrain, Turn
        import ai
        with ai.using(brain=ScriptedBrain(
                [Turn(text='{"disc": "Chemistry", "n_periods": 7}')])):
            plan = planner.build(src, "single_trend")
        assert plan["backend"] == "rule"
        assert plan["fallback_reason"]

    def test_模型给出非法学科被过滤(self, src):
        from _ai_helpers import ScriptedBrain, Turn
        import ai
        with ai.using(brain=ScriptedBrain(
                [Turn(text='{"disc": "Chemistry", "charts": '
                           '[{"kind": "cnci", "disc": "NoSuchDisc"}, '
                           '{"kind": "cnci", "disc": "Physics"}]}')])):
            plan = planner.build(src, "single_trend")
        # 非法图型项被丢掉；剩下那条也是合法学科
        for c in plan["charts"]:
            assert c["kind"] in spec.KINDS
            if c.get("disc"):
                assert c["disc"] in src.disciplines

    def test_模型返回非JSON不崩(self, src):
        from _ai_helpers import ScriptedBrain, Turn
        import ai
        with ai.using(brain=ScriptedBrain([Turn(text="我建议关注工程学。")])):
            plan = planner.build(src, "single_trend")
        assert plan["backend"] == "rule"


class Test计划含校验点:
    """票 01：规划带 `checks`（规则生成、白名单键）；**模型塞不进来**。"""

    def test_规则计划含校验点且键为白名单(self, src):
        plan = planner.build(src, "single_trend")
        checks = plan.get("checks")
        assert isinstance(checks, list) and 3 <= len(checks) <= 5
        for c in checks:
            assert set(c) == {"kind", "target", "expect"}
            assert c["kind"] in ("source_trace", "spec_query_match", "outline_match")

    def test_横版计划也含校验点(self, src):
        plan = planner.build(src, "cross_compare")
        assert isinstance(plan.get("checks"), list)
        assert 3 <= len(plan["checks"]) <= 5

    def test_校验点确定性(self, src):
        """同一意图两次规划 → 同一套校验点（可复现）。"""
        assert planner.build(src, "single_trend")["checks"] == \
            planner.build(src, "single_trend")["checks"]

    def test_模型给的校验点被剥离(self, src):
        """模型自己塞 checks → 落盘的是**规则**那套，不是它的。"""
        from _ai_helpers import ScriptedBrain, Turn
        import ai
        text = ('{"disc": "Chemistry", '
                '"charts": [{"kind": "cnci", "disc": "Chemistry"}], '
                '"checks": [{"kind": "hack", "target": "x", "expect": "y"}]}')
        with ai.using(brain=ScriptedBrain([Turn(text=text)])):
            plan = planner.build(src, "single_trend")
        assert plan["backend"] == "scripted", "本用例要走模型轮，不然断言是空的"
        assert plan["checks"], "模型轮也必须带规则校验点"
        assert all(c["kind"] != "hack" for c in plan["checks"])

    def test_产物flex_plan落盘含校验点(self, flex_report_out):
        """缝 S1：跑完一条灵活报告，flex_plan.json 里就有 checks。"""
        d = pathlib.Path(flex_report_out["dir"])
        plan = json.loads((d / "flex_plan.json").read_text(encoding="utf-8"))
        checks = plan.get("checks")
        assert isinstance(checks, list) and 3 <= len(checks) <= 5


# ================================================================ 2. 不可分叉
class Test数字可回溯:
    """验收第 6 条：每个数字能指到查询记录与台账来源。"""

    def test_每个数据点都带来源链(self, src):
        r = queries.series(src, "total_cites", "Chemistry", src.filled_periods[-3:])
        assert r["points"]
        for p in r["points"]:
            ch = p["source"]
            assert ch["origin_sheet"] == "总被引"
            assert ch["discipline"] == "Chemistry"
            assert ch["period"] == p["period"]
            assert ch["locator"]

    def test_来源链能逐点回读(self, src):
        r = queries.series(src, "proximity_1pct", "Engineering", None)
        t = queries.trace_check(src, r)
        assert t["checked"] > 0 and t["bad"] == 0, t

    def test_横向对比的每行都带回读链(self, src):
        r = queries.compare(src, "papers", None, src.current_period())
        t = queries.trace_check(src, r)
        assert t["checked"] == r["n"] and t["bad"] == 0

    def test_异常项的链路指向两期(self, src):
        r = queries.anomaly(src, src.current_period())
        for it in r["items"]:
            assert it["source"]["period"] == r["period"]
            assert it["source_prev"]["period"] == r["prev"]


class Test排序方向:
    """「好」的在前：排名越小越好（升序），接近度越大越好（降序）。"""

    def test_接近度降序(self, src):
        rows = queries.compare(src, "proximity_1pct", None, src.current_period())["rows"]
        vs = [x["value"] for x in rows]
        assert vs == sorted(vs, reverse=True)

    def test_国内排名升序(self, src):
        rows = queries.compare(src, "rank", None, src.current_period())["rows"]
        vs = [x["value"] for x in rows]
        assert vs == sorted(vs)

    def test_摘要标了规划依据(self, src):
        sc = queries.scope(src, src.current_period())
        assert sc["trust"] == "planning_input"
        assert "不是可引用数值" in sc["note"]


# ================================================================ 3. 引擎扩 op（票 01）
class Test引擎扩op:
    """票 01：at / delta / events / compare.top —— tmp parquet 视图上确定性验证。"""

    @pytest.fixture()
    def view(self, tmp_path):
        """合成 parquet 视图：2 学科 × 3 期（Chemistry 中间期缺论文数）。"""
        import pandas as pd
        d = tmp_path / "202601"
        (d / "wide").mkdir(parents=True)
        rows = []
        papers = {"Physics": {"202601": 10.0, "202603": 12.0, "202605": 15.0},
                  "Chemistry": {"202601": 20.0, "202603": None, "202605": 25.0}}
        for disc, vals in papers.items():
            for p, v in vals.items():
                rows.append({"discipline": disc, "period": p, "value": v,
                             "metric": "papers", "source": None})
        pd.DataFrame(rows).to_parquet(d / "ledger_long.parquet")
        pd.DataFrame({"202601": [1.0, 1.0], "202603": [1.0, 1.0], "202605": [1.0, 1.0]},
                     index=["Physics", "Chemistry"]).to_parquet(d / "wide" / "total_cites.parquet")
        pd.DataFrame({"202601": [1.0, 0.0], "202603": [0.0, 1.0], "202605": [1.0, 0.0]},
                     index=["Physics", "Chemistry"]).to_parquet(d / "wide" / "entered_1pct.parquet")
        pd.DataFrame({"202601": [90.0, 95.0], "202603": [91.0, 96.0], "202605": [92.0, 97.0]},
                     index=["Physics", "Chemistry"]).to_parquet(
                         d / "wide" / "threshold_1permille.parquet")
        pd.DataFrame({"period": ["202601", "202605"], "rank_intl": [560.0, 552.0],
                      "rank_cn": [153.0, 154.0], "rank_cn_uni": [130.0, 131.0]}).to_parquet(
                          d / "wide" / "ynd_rank.parquet")
        return DataSource(d)

    def test_at_取到格(self, view):
        r = queries.at(view, "papers", "Physics", "202603")
        assert r["value"] == 12.0
        assert r["source"]["origin_sheet"] == "论文数"
        assert r["source"]["period"] == "202603" and r["source"]["discipline"] == "Physics"

    def test_at_缺数如实说(self, view):
        r = queries.at(view, "papers", "Chemistry", "202603")
        assert r["value"] is None and r["source"] is None and r["note"]

    def test_at_期次越界拒(self, view):
        with pytest.raises(queries.QueryError, match="不在数据范围内"):
            queries.at(view, "papers", "Physics", "202607")

    def test_delta_差与比(self, view):
        r = queries.delta(view, "papers", "Physics", "202605")
        assert r["prev"] == "202603" and r["value"] == 15.0 and r["prev_value"] == 12.0
        assert r["diff"] == 3.0 and r["ratio_pct"] == 25.0
        assert r["source"]["period"] == "202605" and r["source_prev"]["period"] == "202603"

    def test_delta_指定对比期(self, view):
        r = queries.delta(view, "papers", "Physics", "202605", prev="202601")
        assert r["diff"] == 5.0 and r["ratio_pct"] == 50.0

    def test_delta_缺数如实说(self, view):
        r = queries.delta(view, "papers", "Chemistry", "202605")
        assert r["value"] == 25.0 and r["prev_value"] is None
        assert r["diff"] is None and r["ratio_pct"] is None and "202603" in r["note"]

    def test_delta_首期无上期拒(self, view):
        with pytest.raises(queries.QueryError, match="没有可比期次"):
            queries.delta(view, "papers", "Physics", "202601")

    def test_events_单科进出(self, view):
        r = queries.events(view, "Physics")
        assert [(e["period"], e["type"]) for e in r["events"]] == [("202603", "out"), ("202605", "in")]
        assert r["in_since"] == {"Physics": "202601"}
        assert all(e["source"]["origin_sheet"] == "学科进入时间" for e in r["events"])

    def test_events_全部按时间排(self, view):
        r = queries.events(view)
        assert r["n"] == 4
        assert [(e["period"], e["discipline"], e["type"]) for e in r["events"]] == [
            ("202603", "Chemistry", "in"), ("202603", "Physics", "out"),
            ("202605", "Chemistry", "out"), ("202605", "Physics", "in")]
        assert r["in_since"] == {"Physics": "202601"}

    def test_events_未知学科拒(self, view):
        with pytest.raises(queries.QueryError, match="不在名单内"):
            queries.events(view, "NoSuchDisc")

    def test_compare_top_取前k(self, view):
        r = queries.compare(view, "papers", None, "202605", top=1)
        assert r["n"] == 1 and r["rows"][0]["discipline"] == "Chemistry"
        with pytest.raises(queries.QueryError, match="top"):
            queries.compare(view, "papers", None, "202605", top=0)

    def test_千分位阈值走宽表(self, view):
        r = queries.at(view, "threshold_1permille", "Physics", "202603")
        assert r["value"] == 91.0 and r["source"]["origin_sheet"] == "千分之一被引阈值"

    def test_overall_rank_云大序列(self, view):
        r = queries.execute(view, "overall_rank")
        assert [p["period"] for p in r["points"]] == ["202601", "202605"]
        assert r["points"][-1]["rank_intl"] == 552.0
        assert r["points"][-1]["source"]["origin_sheet"] == "云大ESI排名"

    def test_execute_新op可派发(self, view):
        r = queries.execute(view, "delta", metric="papers", disc="Physics", period="202605")
        assert r["kind"] == "delta"


class Test引擎真实产物:
    """真产物上的数据面补齐：转置排名序列、千分位阈值、两适配器同值。"""

    def test_转置排名序列可查(self, src):
        hits = [d for d in src.disciplines if queries.series(src, "rank_cn", d)["points"]]
        assert hits, "至少一个学科应有国内排名（转置表）序列"
        r = queries.series(src, "rank_cn", hits[0])
        ps = [p["period"] for p in r["points"]]
        assert ps == sorted(ps)
        assert all(p["source"]["origin_sheet"] == "国内排名" for p in r["points"])

    def test_千分位阈值可查(self, src):
        hits = [d for d in src.disciplines
                if queries.series(src, "threshold_1permille", d)["points"]]
        assert hits, "至少一个学科应有千分位阈值序列"

    def test_两适配器同值(self):
        """WorkbookView（当期核对副本直读）与 DataSource（parquet）对引擎同协议同值。"""
        from conftest import REPO
        xlsx = REPO / "out" / "_acquisition" / "202609" / "监测表.xlsx"
        pq = REPO / "out" / "202609" / "ledger_long.parquet"
        if not (xlsx.is_file() and pq.is_file()):
            skip_unless("out/_acquisition/202609/监测表.xlsx 或 out/202609 产物缺失",
                        lambda: xlsx.is_file() and pq.is_file())
        from flexible.workbook import WorkbookView
        wb, ds = WorkbookView(xlsx), DataSource(pq.parent)
        for metric in ("total_cites", "papers", "proximity_1pct", "threshold_1pct",
                       "threshold_1permille", "rank_cn", "rank_intl"):
            for d in ("Chemistry", "Physics"):
                a, b = wb.series(metric, d), ds.series(metric, d)
                common = a.index.intersection(b.index)
                assert len(common) >= 3, f"{metric}/{d} 无重叠期次"
                assert (a[common] - b[common]).abs().max() < 1e-9, metric
        assert wb.entered_1pct().equals(ds.entered_1pct())
        ra = queries.at(wb, "papers", "Chemistry", "202609")
        rb = queries.at(ds, "papers", "Chemistry", "202609")
        assert ra["value"] == rb["value"]
        assert ra["source"]["origin_sheet"] == rb["source"]["origin_sheet"]


class Test不污染复刻产物:
    """灵活报告只写 out/<期次>/flexible/，复刻报告的三件产物一个字节都不许动。"""

    def test_灵活产物落在flexible子目录(self, src):
        from conftest import REPO
        d = REPO / src.dir
        reps = report.list_reports(d)
        if not reps:
            skip_unless("尚未生成过灵活报告", lambda: bool(report.list_reports(d)))
        for r in reps:
            assert pathlib.Path(r["path"]).parent.parent.name == "flexible"

    def test_灵活图不在复刻图表目录里(self, src):
        from conftest import REPO
        charts = REPO / src.dir / "charts"
        if not charts.is_dir():
            skip_unless("无复刻图表目录", lambda: charts.is_dir())
        stray = [p.name for p in charts.rglob("*.png") if "__" in p.name]
        assert not stray, f"灵活报告的图混进了复刻图表目录：{stray[:3]}"

    def test_spec默认参数不改变出图(self, src, tmp_path):
        """spec 层只做「复用」：不给语义参数时，出图必须与直接调 charts 逐字节相同。"""
        import charts
        from conftest import REPO
        ld = REPO / src.dir
        a = tmp_path / "via_spec"
        b = tmp_path / "direct"
        a.mkdir()
        b.mkdir()
        spec.make(src, {"kind": "top_papers", "disc": "Chemistry"}, a)
        # 直接调（窗口与 spec 默认一致：有数据的最后 DEFAULT_WINDOW 期）
        ps = src.filled_periods[-spec.DEFAULT_WINDOW:]
        charts.chart_top_papers(_read(ld, "top_papers"), "Chemistry", b / "t.png", ps)
        h1 = hashlib.sha256((a / "top_papers__Chemistry.png").read_bytes()).hexdigest()
        h2 = hashlib.sha256((b / "t.png").read_bytes()).hexdigest()
        assert h1 == h2, "spec 层改变了出图结果 —— 它应当只复用 charts.py 的配方"


def _read(ld, name):
    import pandas as pd
    return pd.read_parquet(pathlib.Path(ld) / "wide" / f"{name}.parquet")


class Test图片路由安全:
    """灵活图的取图路由必须与复刻图同一套越界闸。"""

    @pytest.mark.parametrize("name", [
        "../charts/x.png", "..\\..\\server\\.env", "/etc/passwd", "C:/windows/win.ini",
        "a/../../b.png", "", "x.png.exe",
    ])
    def test_越界与非法后缀一律拒绝(self, name):
        from server import jobs
        assert jobs.resolve_flex_chart("out/202607", "single_trend", name) is None

    @pytest.mark.parametrize("intent", ["../..", "a/b", "..\\x", ""])
    def test_intent段也必须干净(self, intent):
        from server import jobs
        assert jobs.resolve_flex_chart("out/202607", intent, "x.png") is None


# ================================================================ 3. 端到端
class Test规则后端产出完整分析稿:
    """验收第 5 条（最重要）：无 LLM 时也能产出完整分析稿。"""

    def test_报告文件齐全(self, flex_report_out):
        d = pathlib.Path(flex_report_out["dir"])
        for f in ("report.md", "report.html", "flex_plan.json",
                  "flex_specs.json", "flex_sources.json", "flex_meta.json"):
            assert (d / f).is_file(), f"缺产物 {f}"

    def test_落盘html单文件带图(self, flex_report_out):
        """🔴 「下载报告」给的是单文件：落盘的 report.html 图必须已内联——
        下载后离开 `charts/` 目录也要能看到图（别把图留在原地）。"""
        d = pathlib.Path(flex_report_out["dir"])
        html_text = (d / "report.html").read_text(encoding="utf-8")
        assert 'src="charts/' not in html_text, "还有相对路径图——单文件下载会全断"
        assert "data:image/png;base64," in html_text

    def test_规划后端是rule且有图(self, flex_report_out):
        p = flex_report_out["plan"]
        assert p["backend"] == "rule"
        assert flex_report_out["n_charts"] >= 1

    def test_五个章节齐全(self, flex_report_out):
        md = pathlib.Path(flex_report_out["report"]).read_text(encoding="utf-8")
        for sec in ("当期概况", "关键指标走势", "相对位次",
                    "成因研判（判断草稿）", "建议事项（判断草稿）", "本稿如何产生", "数据来源链"):
            assert sec in md, f"缺章节：{sec}"

    def test_报告同时产出docx_可直接改(self, flex_report_out):
        """2026-10-15（rick）：下载报告＝docx（方便馆员修改）——出稿即产 `report.docx`。"""
        dx = pathlib.Path(flex_report_out["report"]).parent / "report.docx"
        assert dx.is_file(), "出稿应同时产 docx"
        from docx import Document
        doc = Document(str(dx))
        texts = [p.text for p in doc.paragraphs]
        assert any("当期概况" in t for t in texts), "正文应在"
        assert texts[-1].startswith("内部参考 · 非报送件"), "页脚标注一次"
        assert len(doc.tables) >= 2, "概况表＋来源链表"
        assert len(doc.inline_shapes) >= 1, "图内嵌（离开 charts/ 也带图）"

    def test_自检通过(self, flex_report_out):
        c = flex_report_out["check"]
        assert c["pass"], c
        assert c["trace"]["checked"] > 0 and c["trace"]["bad"] == 0
        assert c["has_charts"] and c["n_charts"] >= 1

    def test_一张图都没有时自检必须判不过(self, src):
        """🔴 `0/0 全对`是老毛病：真实模型选了非默认图型时旧代码出 0 张图，
        而自检照样判过。没有图的分析稿不算完整稿。"""
        from conftest import REPO
        r = report.generate(REPO / src.dir, "single_trend")
        built = {"blocks": [], "sources": [], "charts": [], "hyps": [
            {"text": "t", "verify": "v", "source": "s"}]}
        c = report.self_check("# 标题\n\n正文", src, built)
        assert not c["pass"], c
        assert not c["has_charts"]
        _ = r

    def test_模型选了非默认图型也要出图(self, src):
        """图型清单以 `spec.KINDS` 为准，报告层不许再列白名单（D6 后经 ai 内部口）。"""
        from _ai_helpers import ScriptedBrain, Turn
        from conftest import REPO
        import ai
        brain = ScriptedBrain([Turn(text='{"disc": "Agricultural Sciences", "charts": '
                                         '[{"kind": "overview_1permille"}, '
                                         '{"kind": "cnci", "disc": "Agricultural Sciences"}]}')])
        with ai.using(brain=brain):
            r = report.generate(REPO / src.dir, "single_trend")
        assert r["plan"]["backend"] == "scripted"
        assert r["n_charts"] == 2, r["plan"]["charts"]
        assert r["check"]["pass"]

    def test_正文无漏标的判断句(self, flex_report_out):
        md = pathlib.Path(flex_report_out["report"]).read_text(encoding="utf-8")
        assert not opinion.guard(md), opinion.guard(md)[:3]

    def test_专业判断留了空位(self, flex_report_out):
        md = pathlib.Path(flex_report_out["report"]).read_text(encoding="utf-8")
        assert opinion.manual("").split("：")[0] in md

    def test_每份报告都标注非报送件(self, flex_report_out):
        md = pathlib.Path(flex_report_out["report"]).read_text(encoding="utf-8")
        assert report.BANNER in md
        assert "不带盖章位" in md

    @pytest.mark.parametrize("intent", ["single_trend", "cross_compare"])
    def test_正文里的图片相对路径真能打开(self, src, intent):
        """🔴 断链不抛异常 —— 只有按相对路径真的去开一次才能发现。

        产物结构是 `<intent>/report.md` + `<intent>/charts/*.png`，
        所以正文里必须是 `charts/xxx.png`；写成 `../charts/` 会指到
        `flexible/charts/`（不存在的目录），报告看起来一切正常、图全空。
        """
        from conftest import REPO
        r = report.generate(REPO / src.dir, intent)
        rp = pathlib.Path(r["report"])
        links = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", rp.read_text(encoding="utf-8"))
        assert links, "正文里一张图都没有"
        for lk in links:
            assert not lk.startswith(("http://", "https://", "/")), f"相对路径被写成 {lk}"
            assert (rp.parent / lk).is_file(), f"图片打不开：{rp.parent / lk}"

    def test_两条意图都能出稿(self, src):
        from conftest import REPO
        r = report.generate(REPO / src.dir, "cross_compare")
        assert r["check"]["pass"]
        md = pathlib.Path(r["report"]).read_text(encoding="utf-8")
        for sec in ("本期概览", "横向对比", "成因研判（判断草稿）", "建议事项（判断草稿）"):
            assert sec in md


class Test语言不死板:
    """rick 2026-09-21：「灵活报告并不灵活，语言还是很死板」。

    死板的根因＝三条指标一律「X月至X月（7 期），X 的 Y 自 A 上升至 B，…」同一个句式连写三遍。
    判据不是"读起来顺"，是**可测的三条**：句式随数据形态变、同报告内不撞句、结果可复现。
    """

    def _pts(self, vs, start=0):
        return [{"period": f"2026{1 + i:02d}", "value": v} for i, v in enumerate(vs)]

    @pytest.mark.parametrize("vs,kw", [
        ([1, 2, 3, 4], ("走高", "上升", "升到")),          # 单调上行
        ([4, 3, 2, 1], ("走低", "下滑", "降到")),          # 单调下行
        ([1, 3, 2, 4], ("波动", "并非单边", "反复")),       # 震荡
        ([1, 1, 1, 1], ("走平",)),                        # 走平
        ([2, 2, 8, 7], ("集中在一期",)),                   # 台阶（跳一下再小幅回落）
        ([1, 4, 2, 2], ("先升后降",)),                     # 反转
    ])
    def test_句式随数据形态变(self, vs, kw):
        s = opinion.descriptive_trend(self._pts(vs), "某指标", "某学科")
        assert any(k in s for k in kw), f"{vs} → {s}"

    def test_同报告内三条指标不撞同一句(self, flex_report_out):
        md = pathlib.Path(flex_report_out["report"]).read_text(encoding="utf-8")
        seg = md.split("## 二、")[1].split("## 三、")[0]
        lines = [l for l in seg.splitlines() if l.startswith("- ")]
        assert len(lines) >= 3
        heads = [l[:14] for l in lines]
        assert len(set(heads)) == len(heads), f"三条开头撞了：{heads}"
        # 旧句式「自 A 上升至 B」不该再出现三遍
        assert seg.count("上升至") <= 1, "又退回模板句式了"

    def test_规则后端可复现(self, src):
        """句式变体必须**确定性** —— 同一份数据两次跑出同一句。"""
        from flexible import queries
        r = queries.series(src, "total_cites", "Chemistry", src.filled_periods[-5:])
        a = opinion.descriptive_trend(r["points"], "总被引频次", "化学")
        b = opinion.descriptive_trend(r["points"], "总被引频次", "化学")
        assert a == b

    def test_第二条起不再重复区间(self, src):
        from flexible import queries
        r = queries.series(src, "papers", "Chemistry", src.filled_periods[-5:])
        s = opinion.descriptive_trend(r["points"], "论文数", "化学", with_span=False)
        assert "至" not in s[:8], s

    def test_概况引导句不再是同一句(self):
        a = opinion.gap_lead("2026年7月", "化学", -120)
        b = opinion.gap_lead("2026年7月", "化学", 120)
        c = opinion.gap_lead("2026年7月", "化学", None)
        assert "还差" in a and "高出" in b and "核心指标如下" in c
        assert a != b

    def test_相对位次句不再把学科数说两遍(self):
        s = opinion.rank_sentence(22, 15, 14, "化学")
        assert s.count("22") == 1, s
        assert "后 1/3" in s


class Test润色位是可选的:
    """润色只改措辞、不改数字；不可用/失败一律原样返回（演示站永远走这条）。

    D6（2026-09-29）后经 `ai._attempt` 模型边界出网：测试用 `ai.using(brain=…)` 注入替身。
    """

    def test_默认关闭(self, monkeypatch):
        monkeypatch.delenv("WENHENG_FLEX_POLISH", raising=False)
        assert polish.enabled() is False
        assert polish.rewrite("接近度自 0.6 升至 0.8。") == "接近度自 0.6 升至 0.8。"

    def test_模型不可用时原样返回(self, monkeypatch):
        from _ai_helpers import FailingBrain
        import ai
        monkeypatch.setenv("WENHENG_FLEX_POLISH", "1")
        txt = "总被引自 9,356 升至 11,558。"
        with ai.using(brain=FailingBrain()):
            assert polish.rewrite(txt) == txt

    def test_校验通过才还原_数字一个不少(self, monkeypatch):
        from _ai_helpers import ScriptedBrain, Turn
        import ai
        monkeypatch.setenv("WENHENG_FLEX_POLISH", "1")
        txt = "接近度自 0.6 升至 0.8。"
        with ai.using(brain=ScriptedBrain(
                [Turn(text="接近度由 [[1]] 升至 [[2]]。")])):
            assert polish.rewrite(txt) == "接近度由 0.6 升至 0.8。"

    def test_校验不通过时原样返回(self, monkeypatch):
        """模型把占位符弄丢 / 冒出新数字 → 必须退回原文，不许端上桌。"""
        from _ai_helpers import ScriptedBrain, Turn
        import ai
        monkeypatch.setenv("WENHENG_FLEX_POLISH", "1")
        txt = "接近度自 0.6716 升至 0.845。"
        with ai.using(brain=ScriptedBrain([Turn(text="接近度上升了 30% 左右。")])):
            assert polish.rewrite(txt) == txt

    def test_不许动的标记不会被润掉(self, monkeypatch):
        from _ai_helpers import ScriptedBrain, Turn
        import ai
        monkeypatch.setenv("WENHENG_FLEX_POLISH", "1")
        txt = "【手工：化学资源投入建议】"
        with ai.using(brain=ScriptedBrain([Turn(text="这条被改写了。")])):
            assert polish.rewrite(txt) == txt

    def test_未配云不调模型(self, src, monkeypatch):
        import urllib.request

        def boom(*a, **k):
            raise AssertionError("未配云时绝不能碰模型（演示站无模型）")

        monkeypatch.setattr(urllib.request, "urlopen", boom)
        from conftest import REPO
        report.generate(REPO / src.dir, "single_trend")


class TestHTML渲染:
    """网页与导出件共用同一份渲染实现（避免"网页对、导出件错"）。"""

    def test_基本语法都渲染(self):
        md = ("# 标题\n\n> 引用\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n"
              "- 项一\n\n![图](../charts/x.png)\n\n**粗体** 与 `代码`\n\n---\n")
        h = mdview.render(md)
        for tag in ("<h2>", "<blockquote>", "<table", "<ul>", "<img", "<strong>",
                    "<code>", "<hr>"):
            assert tag in h, tag

    def test_图片路径可映射到服务端路由(self):
        h = mdview.render("![x](../charts/a.png)",
                          img_map={"a.png": "/p/202607/flexible/chart/t/a.png"})
        assert "/p/202607/flexible/chart/t/a.png" in h

    def test_转义HTML防止注入(self):
        h = mdview.render("普通文本 <script>alert(1)</script>")
        assert "<script>" not in h

    def test_行内图片属性引号也会转义(self):
        h = mdview.render('![x" onerror="alt](chart.png" onerror="src)')
        assert '<img src="chart.png&quot; onerror=&quot;src"' in h
        assert 'alt="x&quot; onerror=&quot;alt"' in h
        assert ' onerror="' not in h

    def test_独立件是自包含的(self):
        h = mdview.standalone("# 标题", "t")
        assert h.startswith("<!doctype html>") and "<style>" in h
        assert "http://" not in h and "https://" not in h

    def test_单文件_相对图内联成data_uri_缺的不编(self, tmp_path):
        """「下载报告」拿到的是单文件：相对路径的图内联成 data URI；缺文件/网络图原样留着。"""
        (tmp_path / "charts").mkdir()
        (tmp_path / "charts" / "a.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 16)
        h = mdview.standalone("![图](charts/a.png)\n\n![没的](charts/none.png)", "t",
                              base_dir=tmp_path)
        assert "data:image/png;base64," in h
        assert 'src="charts/a.png"' not in h
        assert 'src="charts/none.png"' in h, "缺文件不编：原样留着"
        assert "data:" not in mdview.inline_images('<img src="https://x/a.png">', tmp_path)

    def test_内联不越界(self, tmp_path):
        (tmp_path / "out.png").write_bytes(b"x")
        sub = tmp_path / "d"
        sub.mkdir()
        h = mdview.inline_images('<img src="../out.png">', sub)
        assert "data:" not in h and 'src="../out.png"' in h


class TestCLI与清单:
    def test_cli入口可跑(self, src, capsys):
        import run_flexible
        rc = run_flexible.main([str(src.dir), "cross_compare"])
        assert rc == 0
        out = capsys.readouterr().out
        assert "自检 通过" in out and "数字可回溯" in out

    def test列表能列出刚生成的(self, src):
        reps = report.list_reports(src.dir)
        assert any(r["intent"] == "cross_compare" for r in reps)
        for r in reps:
            assert pathlib.Path(r["path"]).is_file()
            assert r["meta"].get("backend")


class Test指定学科全链路:
    """rick 2026-09-29：「指定学科填了化学，报告和化学无关」——三层都不许再忽略。"""

    def test_中文名可用且横版围着它写(self, src):
        from conftest import REPO
        r = report.generate(REPO / src.dir, "cross_compare", disc="化学")
        assert r["check"]["pass"], r["check"]
        md = pathlib.Path(r["report"]).read_text(encoding="utf-8")
        assert "化学 · 当期落点" in md and "化学 · 近 7 期走势" in md
        assert "**化学**" in md, "全池对比表里指定学科要加粗"
        assert "化学" in r["plan"]["title"] and "横向对比" in r["plan"]["title"]
        assert r["n_charts"] == 4, r["plan"]["charts"]
        assert [c["disc"] for c in r["plan"]["charts"][:2]] == ["Chemistry", "Chemistry"]

    def test_未指定学科时横版保持全池口径(self, src):
        from conftest import REPO
        r = report.generate(REPO / src.dir, "cross_compare")
        md = pathlib.Path(r["report"]).read_text(encoding="utf-8")
        assert "当期落点" not in md
        assert not r["plan"].get("disc")

    def test_未知学科给人话报错(self, src):
        with pytest.raises(ValueError, match="不在监测名单"):
            planner.build(src, "single_trend", disc="量子计算")

    def test_模型在横版乱塞单科图被丢弃(self, src):
        """旧行为：模型挑「工程学」加一堆图、标题还写「工程学学科分析报告」——无指定学科时不许。"""
        from _ai_helpers import ScriptedBrain, Turn
        import ai
        brain = ScriptedBrain([Turn(text='{"title": "工程学 学科分析报告", "charts": '
                                          '[{"kind": "cnci", "disc": "Engineering"}, '
                                          '{"kind": "overview_1pct"}]}')])
        with ai.using(brain=brain):
            plan = planner.build(src, "cross_compare")
        assert [c["kind"] for c in plan["charts"]] == ["overview_1pct"]
        assert "学科横向对比" in plan["title"], plan["title"]

    def test_指定学科时模型乱选也被扳回(self, src):
        from _ai_helpers import ScriptedBrain, Turn
        import ai
        brain = ScriptedBrain([Turn(text='{"title": "工程学 学科分析报告", "charts": '
                                          '[{"kind": "cnci", "disc": "Physics"}]}')])
        with ai.using(brain=brain):
            plan = planner.build(src, "cross_compare", disc="Chemistry")
        assert "化学" in plan["title"], plan["title"]
        assert all(c["disc"] in (None, "Chemistry") for c in plan["charts"])

    def test_标题括号位统一英文名(self, src):
        """2026-10-15 rick：模型实测写「经济与商业（经济与商业）学科趋势诊断」——
        括号位统一成英文名（学科名「中文名（英文名）」口径）；散写英文名仍换中文。"""
        from _ai_helpers import ScriptedBrain, Turn
        import ai
        brain = ScriptedBrain([
            Turn(text='{"title": "农业科学（农业科学）学科趋势诊断", "charts": '
                      '[{"kind": "cnci", "disc": "Agricultural Sciences"}]}'),
            Turn(text='{"title": "Agricultural Sciences学科趋势分析", "charts": '
                      '[{"kind": "cnci", "disc": "Agricultural Sciences"}]}'),
        ])
        with ai.using(brain=brain):
            p1 = planner.build(src, "single_trend", disc="Agricultural Sciences")
            p2 = planner.build(src, "single_trend", disc="Agricultural Sciences")
        assert "（Agricultural Sciences）" in p1["title"], p1["title"]
        assert "（农业科学）" not in p1["title"], p1["title"]
        assert p2["title"] == "农业科学学科趋势分析", p2["title"]


class Test计划带候选排序:
    """considered：挑选时**本来就算过**的候选前三（规划那刻记，展示层只读）。

    语出 agent-zizhu 票 01 追加（rick 要「独白」）：`[{disc, disc_cn, proximity}]`，
    按接近度降序，首项＝最终选中的学科；规则与模型两条路径都带，模型给的一律剥。
    """

    def test_规则挑科带候选前三(self, src):
        from ai._brain import RuleBrain
        import ai
        with ai.using(brain=RuleBrain()):
            plan = planner.build(src, "single_trend")
        cons = plan["considered"]
        assert len(cons) == 3, "前 3 名候选"
        assert cons[0]["disc"] == plan["disc"], "首项＝最终选中"
        prox = [c["proximity"] for c in cons]
        assert prox == sorted(prox, reverse=True), "按接近度降序"
        assert all(c["disc_cn"] and c["disc"] for c in cons)

    def test_模型给的considered一律剥(self, src):
        from _ai_helpers import ScriptedBrain, Turn
        import ai
        brain = ScriptedBrain([Turn(text='{"disc": "Physics", "title": "物理学发展态势诊断", '
                                          '"considered": [{"disc": "Bogus", "proximity": 9}]}')])
        with ai.using(brain=brain):
            plan = planner.build(src, "single_trend")
        assert all(c["disc"] != "Bogus" for c in plan["considered"]), "白名单外的新栏一律剥"
        assert plan["considered"] and plan["considered"][0]["disc"] == plan["disc"]

    def test_considered随计划落盘(self, flex_report_out, src):
        from conftest import REPO
        p = REPO / src.dir / "flexible" / "single_trend" / "flex_plan.json"
        plan = json.loads(p.read_text(encoding="utf-8"))
        assert len(plan["considered"]) == 3
        assert plan["considered"][0]["disc"] == plan["disc"]


class Test流式生成:
    def test_事件序列与大纲同源(self, src):
        """status → plan(带大纲) → section×N（逐节标题==大纲）→ done；无假进度。"""
        from conftest import REPO
        events = list(report.generate_stream(REPO / src.dir, "cross_compare"))
        assert [e["event"] for e in events[:2]] == ["status", "plan"]
        assert events[-1]["event"] == "done"
        sec = [e for e in events if e["event"] == "section"]
        assert sec, "必须逐节发事件"
        assert [e["title"] for e in sec] == events[1]["outline"], \
            "逐节标题必须与 plan 事件的同源大纲一致"
        assert events[-1]["n_charts"] >= 1
        assert events[-1]["check_pass"] is True

    def test_流式含指定学科(self, src):
        from conftest import REPO
        events = list(report.generate_stream(REPO / src.dir, "single_trend", disc="化学"))
        plan_ev = events[1]
        assert "化学" in plan_ev["title"]
        sec = [e for e in events if e["event"] == "section"]
        assert sec and any("化学" in e.get("body", "") for e in sec)

    def test_plan事件带候选与配图_done带数据点数(self, src):
        """轨迹独白的数据都在事件里：挑学科（considered）、配图（charts）、点数（done）。"""
        from conftest import REPO
        events = list(report.generate_stream(REPO / src.dir, "single_trend"))
        plan_ev = events[1]
        assert plan_ev["considered"] and plan_ev["considered"][0]["disc"] == plan_ev["disc"]
        assert plan_ev["charts"], "配图步的数据"
        assert all(c.get("kind") and c.get("why") is not None for c in plan_ev["charts"])
        done = events[-1]
        assert done["event"] == "done"
        assert done["n_sources"] > 0, "数字能指回台账的条数"
        assert done["n_sections"] >= 1

    def test_done事件带自检明细_没过时说得清哪条(self, src):
        """票 05：校验没过要**明说哪条没过**——done 事件带自检明细，不只一个布尔。
        （票 03 字段面翻转：判断句标记完备性取代「专业判断绝不出现」。）"""
        from conftest import REPO
        events = list(report.generate_stream(REPO / src.dir, "single_trend"))
        done = events[-1]
        c = done.get("check")
        assert isinstance(c, dict), f"done 应带 check 明细：{sorted(done)}"
        for key in ("trace_ok", "trace_bad", "has_charts", "judgement_marked",
                    "judgement_leaks", "backstop_n"):
            assert key in c, f"缺自检项 {key}"
        assert c["trace_ok"] is True and c["has_charts"] is True
        assert c["trace_bad"] == 0
        assert c["judgement_marked"] is True and c["judgement_leaks"] == []
        assert c["backstop_n"] == 0

    def test_停止_终态stopped_不落产物(self, tmp_path):
        """点停 → 下一次事件点收手 → 终态 stopped；产物在收尾才写，停在中途不留稿。"""
        import shutil
        d = _period_dir()
        if d is None:
            skip_unless("out/ 下无含 ledger_long.parquet 的期次产物",
                        lambda: bool(_period_dir()))
        tmp = tmp_path / "202607"
        tmp.mkdir()
        shutil.copyfile(d / "ledger_long.parquet", tmp / "ledger_long.parquet")
        shutil.copytree(d / "wide", tmp / "wide")
        state = {"sections": 0}
        events = []
        for ev in report.generate_stream(tmp, "single_trend",
                                         should_stop=lambda: state["sections"] >= 1):
            events.append(ev)
            if ev["event"] == "section":
                state["sections"] += 1
        assert [e["event"] for e in events[:2]] == ["status", "plan"]
        assert events[-1]["event"] == "stopped", f"终态必须是 stopped：{events[-1]}"
        assert not (tmp / "flexible" / "single_trend" / "report.md").exists(), \
            "停在中途不落产物（生成在收尾才写盘）"

    @pytest.mark.parametrize("intent,kw", [
        ("single_trend", {}),
        ("cross_compare", {}),
        ("cross_compare", {"disc": "化学"}),
        ("cross_compare", {"metric": "papers"}),
    ])
    def test_大纲与正文标题一一对应(self, src, intent, kw):
        """防漂移（C8）：正文 `## 一、…` 的标题序列必须与大纲（同一份节表）逐字相同。

        节表在 `report._single_sections/_cross_sections`；动态节（指定学科、梯队分布）
        也在这份表里长出来——大纲与正文不可能再各排各的。
        """
        from conftest import REPO
        events = list(report.generate_stream(REPO / src.dir, intent, **kw))
        outline = events[1]["outline"]
        md = (REPO / src.dir / "flexible" / intent / "report.md").read_text(encoding="utf-8")
        heads = [ln[3:].strip() for ln in md.splitlines() if ln.startswith("## ")]
        # 正文内容节在前；「本稿如何产生」「数据来源链」是固定附录，不在大纲里。
        assert heads[:len(outline)] == outline, \
            f"大纲与正文标题漂移：{heads[:len(outline)]} != {outline}"
        assert heads[len(outline):] == ["本稿如何产生（可验证的架构主张）", "附：数据来源链"]


def _copy_period(tmp_path):
    """把金标期次的关键数据复制到 tmp —— 写模式的用例不许污染 out/202607 的演示产物。"""
    import shutil
    d = _period_dir()
    if d is None:
        skip_unless("out/ 下无含 ledger_long.parquet 的期次产物",
                    lambda: bool(_period_dir()))
    tmp = tmp_path / "202607"
    tmp.mkdir()
    shutil.copyfile(d / "ledger_long.parquet", tmp / "ledger_long.parquet")
    shutil.copytree(d / "wide", tmp / "wide")
    return tmp


def _wr_reply(b):
    """替身：描述节改写（规则稿＋事实照抄换头）；判断节创写（自标待审线，合法候选）。"""
    from _ai_helpers import Turn
    task = b.task
    if "【节主题】" in task:
        topic = task.split("【节主题】\n", 1)[1].split("\n\n【事实摘要】")[0]
        facts = task.split("【事实摘要】\n", 1)[1] if "【事实摘要】" in task else ""
        body = topic + "。" + facts.replace("\n- ", "；")
        return Turn(text=f"=={body}==。以上为待审判断草稿。")
    rule = task.split("【规则稿】\n", 1)[1].split("\n\n【事实摘要】")[0]
    facts = (task.split("【事实摘要】\n", 1)[1].replace("\n", "；")
             if "【事实摘要】" in task else "")
    return Turn(text="改写：" + rule + ("；" + facts if facts else ""))


class Test写作位接线:
    def test_写作位开时全篇走写作位并留痕(self, src, tmp_path, monkeypatch):
        """票 04：描述节改写 + 判断节创写，共 7 段调用；每段独立留痕。"""
        from _ai_helpers import ScriptedBrain
        import ai
        monkeypatch.setenv("WENHENG_FLEX_WRITE", "1")
        tmp = _copy_period(tmp_path)
        brain = ScriptedBrain([_wr_reply] * 10)
        with ai.using(brain=brain, data_root=tmp_path):
            r = report.generate(tmp, "single_trend")
        assert r["mode"] == "write"
        polish_briefs = [b for b in brain.briefs if getattr(b, "ability", "") == "polish"]
        assert len(polish_briefs) == 7, "概况 1 + 趋势 3 + 位次 1 + 判断节 2"
        rec_file = tmp_path / "202607" / "ai_calls.jsonl"
        rows = [json.loads(x) for x in
                rec_file.read_text(encoding="utf-8").splitlines() if x.strip()]
        assert len(rows) == 7 and all(x["capability"] == "polish" for x in rows)
        assert all(x["mode"] == "model" for x in rows)
        assert r["write"] and r["write"]["segments"] == 7
        md = pathlib.Path(r["report"]).read_text(encoding="utf-8")
        assert "改写：" in md
        cause = md.split("成因研判（判断草稿）")[1].split("建议事项（判断草稿）")[0]
        assert "==" in cause, "判断节应带待审线（自标保留）"
        assert "【手工：" not in cause, "模型成功时判断节不该是空位"
        assert r["check"]["pass"], r["check"]
        assert not opinion.guard(md)
        from docx import Document
        doc = Document(str(pathlib.Path(r["report"]).parent / "report.docx"))
        assert any(run.underline for p in doc.paragraphs for run in p.runs), \
            "待审线在 docx 里应是下划线（方便馆员直接改）"

    def test_默认关时mode是off(self, src):
        from conftest import REPO
        r = report.generate(REPO / src.dir, "single_trend")
        assert r["mode"] == "off" and r["write"] is None

    def test_配好即默认启用_不再要按钮(self, src, tmp_path, monkeypatch):
        """票 04：开关开＋模型可达 → 不给 write 参数也默认走创写。"""
        from _ai_helpers import ScriptedBrain
        import ai
        monkeypatch.setenv("WENHENG_FLEX_WRITE", "1")
        tmp = _copy_period(tmp_path)
        brain = ScriptedBrain([_wr_reply] * 10)
        with ai.using(brain=brain, data_root=tmp_path):
            r = report.generate(tmp, "single_trend")
        assert r["mode"] == "write", "配好写作位后生成默认走创写（按钮不再是门槛）"
        assert r["write"] and r["write"]["segments"] == 7


class Test产线贯通:
    """票 04：末两节创写（成因研判/建议事项）＋回退安全值＋声明一处＋默认启用。

    一律经产线口 `report.generate()` 以替身注入（缝 A），在 tmp 期次副本上跑。
    """

    def test_判断节创写_兜底补标计数(self, tmp_path, monkeypatch):
        """模型漏标 → 兜底补标并计数（透明、不判失败）；成稿无未标记判断句。"""
        from _ai_helpers import ScriptedBrain, Turn
        import ai
        monkeypatch.setenv("WENHENG_FLEX_WRITE", "1")
        tmp = _copy_period(tmp_path)

        def reply(b):
            if "【节主题】" in b.task:
                topic = b.task.split("【节主题】\n", 1)[1].split("\n\n【事实摘要】")[0]
                facts = (b.task.split("【事实摘要】\n", 1)[1].replace("\n- ", "；")
                         if "【事实摘要】" in b.task else "")
                # 题面+事实自带标记（含「建议事项」题面），只留尾句有意漏标
                return Turn(text=f"=={topic}。{facts}==。此外，建议优先关注长期积累。")
            return _wr_reply(b)

        brain = ScriptedBrain([reply] * 10)
        with ai.using(brain=brain, data_root=tmp_path):
            r = report.generate(tmp, "single_trend")
        assert r["mode"] == "write" and r["check"]["pass"], r["check"]
        assert r["write"]["backstop"] == 2, "两节各补标 1 句"
        assert r["check"]["backstop_n"] == 2
        md = pathlib.Path(r["report"]).read_text(encoding="utf-8")
        assert "建议优先关注长期积累" in md
        assert not opinion.guard(md)

    def test_写作位失败_判断节回退空位_明说原因(self, tmp_path, monkeypatch):
        from _ai_helpers import BrainError, ScriptedBrain
        import ai
        monkeypatch.setenv("WENHENG_FLEX_WRITE", "1")
        tmp = _copy_period(tmp_path)

        def reply(b):
            if "【节主题】" in b.task:
                raise BrainError("超时", code="timeout")
            return _wr_reply(b)

        brain = ScriptedBrain([reply] * 10)
        with ai.using(brain=brain, data_root=tmp_path):
            r = report.generate(tmp, "single_trend")
        assert r["mode"] == "write"
        md = pathlib.Path(r["report"]).read_text(encoding="utf-8")
        cause = md.split("成因研判（判断草稿）")[1].split("建议事项（判断草稿）")[0]
        assert "【手工：" in cause, "写作位失败 → 判断节回退空位"
        assert r["write"]["fallbacks"] == 2
        assert "brain_error:timeout" in r["write"]["reasons"], "明说哪条没过"
        assert r["check"]["pass"], r["check"]

    def test_标记语法损坏_该节回退空位(self, tmp_path, monkeypatch):
        """补标后仍无法闭环（标记语法损坏）→ 该节判不过 → 回退安全值。"""
        from _ai_helpers import ScriptedBrain, Turn
        import ai
        monkeypatch.setenv("WENHENG_FLEX_WRITE", "1")
        tmp = _copy_period(tmp_path)

        def reply(b):
            if "【节主题】" in b.task:
                topic = b.task.split("【节主题】\n", 1)[1].split("\n\n【事实摘要】")[0]
                facts = (b.task.split("【事实摘要】\n", 1)[1].replace("\n- ", "；")
                         if "【事实摘要】" in b.task else "")
                return Turn(text=f"{topic}。{facts}。==建议加大投入= 与应当优先补齐。")
            return _wr_reply(b)

        brain = ScriptedBrain([reply] * 10)
        with ai.using(brain=brain, data_root=tmp_path):
            r = report.generate(tmp, "single_trend")
        md = pathlib.Path(r["report"]).read_text(encoding="utf-8")
        cause = md.split("成因研判（判断草稿）")[1].split("建议事项（判断草稿）")[0]
        assert "【手工：" in cause, "标记语法损坏 → 该节回退空位"
        assert "guard_reject" in r["write"]["reasons"], r["write"]
        assert r["check"]["pass"], r["check"]

    def test_头部唯一声明含图例_正文无重复声明(self, tmp_path, monkeypatch):
        from _ai_helpers import ScriptedBrain
        import ai
        monkeypatch.setenv("WENHENG_FLEX_WRITE", "1")
        tmp = _copy_period(tmp_path)
        brain = ScriptedBrain([_wr_reply] * 10)
        with ai.using(brain=brain, data_root=tmp_path):
            r = report.generate(tmp, "single_trend")
        md = pathlib.Path(r["report"]).read_text(encoding="utf-8")
        assert "细下划线的句子" in md and "待审句" in md, "头部应有待审线图例"
        assert md.count("细下划线的句子") == 1, "图例只在头部说明一次"
        for gone in ("不由本产线生成", "不是结论", "非结论", "待验证假设"):
            assert gone not in md, f"重复声明未清除：{gone}"
        assert "本件为内部参考" not in md, "报告里不再重复声明（页面导语承担）"
        html = mdview.standalone(md, "t")
        assert "export-mark" in html and "内部参考 · 非报送件" in html, "下载件页脚标注一次"

    def test_规则后端_空位版_确定性(self, tmp_path):
        """demo 档（规则后端）行为与升级前一致：全篇规则稿＋空位；同输入同输出。"""
        from ai._brain import RuleBrain
        import ai
        tmp = _copy_period(tmp_path)
        with ai.using(brain=RuleBrain()):
            r1 = report.generate(tmp, "single_trend")
            md1 = pathlib.Path(r1["report"]).read_text(encoding="utf-8")
            r2 = report.generate(tmp, "single_trend")
            md2 = pathlib.Path(r2["report"]).read_text(encoding="utf-8")
        assert r1["mode"] == "off" and r1["write"] is None
        assert r1["check"]["pass"], r1["check"]
        norm = lambda s: re.sub(r"> 生成时间：.*", "> 生成时间：X", s)   # noqa: E731
        assert norm(md1) == norm(md2), "规则版必须确定性可复现"
        assert "==" not in md1, "规则版没有待审线"
        cause = md1.split("成因研判（判断草稿）")[1].split("建议事项（判断草稿）")[0]
        assert "【手工：" in cause, "规则版判断节＝空位"
        assert "细下划线的句子" not in md1, "规则版不摆空图例"

    def test_两种意图_末两节同为判断草稿(self, tmp_path, monkeypatch):
        from _ai_helpers import ScriptedBrain
        import ai
        monkeypatch.setenv("WENHENG_FLEX_WRITE", "1")
        tmp = _copy_period(tmp_path)
        brain = ScriptedBrain([_wr_reply] * 12)
        with ai.using(brain=brain, data_root=tmp_path):
            r = report.generate(tmp, "cross_compare", disc="化学")
        assert r["mode"] == "write" and r["check"]["pass"], r["check"]
        md = pathlib.Path(r["report"]).read_text(encoding="utf-8")
        for sec in ("成因研判（判断草稿）", "建议事项（判断草稿）"):
            assert sec in md
        cause = md.split("成因研判（判断草稿）")[1].split("建议事项（判断草稿）")[0]
        assert "==" in cause
        assert not opinion.guard(md)


class Test自检翻转与兜底补标:
    """票 03：判断句**未带标记即失败**（旧口径「出现判断词即失败」作废）。

    兜底词表 `建议|应当|必须|值得|需要|优先`（env `WENHENG_FLEX_MARK_WORDS` 覆盖）；
    漏标＝命中且不在成对 `==…==` 内、也不在 `【手工：…】`空位里；补标确定性、计数可读。
    """

    def test_兜底词表缺省六词(self):
        assert opinion.mark_words() == ("建议", "应当", "必须", "值得", "需要", "优先")

    def test_词表可被环境变量覆盖(self, monkeypatch):
        monkeypatch.setenv("WENHENG_FLEX_MARK_WORDS", "务必,千万")
        assert opinion.mark_words() == ("务必", "千万")
        assert opinion.guard("务必加大该学科投入。")
        assert not opinion.guard("建议加大该学科投入。"), "被覆盖掉的词不再检查"

    def test_成对标记内不算漏标(self):
        assert not opinion.guard("==建议加大该学科投入==，该判断数据来自台账。")

    def test_单侧标记不构成豁免(self):
        assert opinion.guard("==建议加大该学科投入，该判断数据来自台账。")

    def test_手工空位内不误报(self):
        assert not opinion.guard(f"- {opinion.manual('资源投入建议')}")

    def test_节标题不误报不补标(self):
        """节名「建议事项（判断草稿）」自带词表词——标题是版式结构，不是判断句。"""
        txt = "## 五、建议事项（判断草稿）\n\n正文没有判断句。"
        assert not opinion.guard(txt)
        assert opinion.backfill_marks(txt) == (txt, 0)

    def test_兜底补标确定性与计数(self):
        txt = "建议加大该学科投入。应当优先补齐短板。\n\n- 需要复核阈值口径。"
        a, na = opinion.backfill_marks(txt)
        b, nb = opinion.backfill_marks(txt)
        assert a == b and na == nb == 3, (a, na, nb)
        assert "==建议加大该学科投入==" in a
        assert not opinion.guard(a), opinion.guard(a)

    def test_补标不动已标记与空位(self):
        txt = "==建议加大投入==。\n\n- " + opinion.manual("资源投入建议")
        assert opinion.backfill_marks(txt) == (txt, 0)

    def test_自检_未补标判不过_补标后过(self):
        """`judgement_marked` 判的是**传入文本**（期望为补标后成稿）——
        无 sources 时 self_check 不触及 src，传 None 即可（纯文本语义用例）。"""
        built = {"blocks": [], "sources": [], "charts": [{"path": "x"}]}
        raw = "# 稿\n\n建议加大该学科投入。"
        c1 = report.self_check(raw, None, built)
        assert not c1["judgement_marked"] and not c1["pass"], c1
        assert c1["judgement_leaks"] and len(c1["judgement_leaks"]) <= 3
        assert c1["backstop_n"] == 0, "缺省 0"
        fixed, n = opinion.backfill_marks(raw)
        c2 = report.self_check(fixed, None, {**built, "backstop_n": n})
        assert n == 1 and c2["backstop_n"] == 1
        assert c2["judgement_marked"] and c2["judgement_leaks"] == [] and c2["pass"], c2

    def test_done流式字段贯通_旧字段清零(self, src):
        """done.check 与 check 字典同面：旧字段（no_professional_judgement /
        hypotheses_complete）清零，新字段（judgement_marked / judgement_leaks /
        backstop_n）有值。"""
        from conftest import REPO
        events = list(report.generate_stream(REPO / src.dir, "cross_compare"))
        c = events[-1]["check"]
        assert "no_professional_judgement" not in c and "hypotheses_complete" not in c
        assert c["judgement_marked"] is True and c["judgement_leaks"] == []
        assert c["backstop_n"] == 0


class Test待审线渲染:
    """票 01：`==句子==` → 细下划线 span（待审线）；网页与下载件同一渲染器、同一 class。

    只认成对 `==…==`（单行内）；单侧 `==` 原样保留、不吞正文；标记内容走既有转义链。
    """

    def test_成对标记渲染为细下划线span(self):
        h = mdview.render("总被引升至 ==11,558== 次。")
        assert '<span class="fx-mark">11,558</span>' in h
        assert "==" not in h, "字面等号不该残留"

    def test_同一行多组标记各自成span(self):
        h = mdview.render("==第一句==与==第二句==都需馆员过目")
        assert '<span class="fx-mark">第一句</span>' in h
        assert '<span class="fx-mark">第二句</span>' in h

    def test_单侧等号原样保留不吞正文(self):
        h = mdview.render("写作位降级 == 时回退空位。")
        assert "fx-mark" not in h
        assert "==" in h and "时回退空位" in h

    def test_标记不跨行配对(self):
        h = mdview.render("==跨行开始\n此处才结束==")
        assert "fx-mark" not in h

    def test_标记内容走既有转义链(self):
        h = mdview.render("==A & B <b>粗</b>==")
        assert '<span class="fx-mark">A &amp; B &lt;b&gt;粗&lt;/b&gt;</span>' in h
        assert "<b>粗</b>" not in h, "HTML 特殊字符不得逃逸"

    def test_下载件同一class与细下划线样式(self):
        h = mdview.standalone("==待审句==的说明", "t")
        assert '<span class="fx-mark">待审句</span>' in h
        assert ".fx-mark{" in h
        assert "text-decoration:underline" in h
        assert "text-decoration-thickness:1px" in h
        assert "text-underline-offset:2px" in h

    def test_旧稿声明渲染时剥离_图例保留(self):
        """2026-10-15（rick：报告里不再重复声明）：历史稿顶部的「本件为内部参考…」渲染时去掉，
        同行的待审线图例保留；整行只剩声明时直接丢弃；下载件页脚标注一次。"""
        old = ("> **本件为内部参考，非报送件** —— 由灵活报告产线（第二条产线）生成，"
               "路径由 AI 自主规划；**不带盖章位、不作为对外报送材料**。"
               "正式对外件请用复刻报告产线（母版 704 槽位、三道闸、可签发）。"
               "　细下划线的句子＝**待审句**，需馆员拍板；拍板线下完成。\n\n正文。")
        h = mdview.render(old)
        assert "本件为内部参考" not in h and "复刻报告产线" not in h
        assert "细下划线的句子" in h, "图例要留下（旧稿里与声明同一行）"
        h2 = mdview.render(old.split("　")[0] + "\n\n正文。")
        assert "本件为内部参考" not in h2, "整行只剩声明 → 整行丢弃"

    def test_下载件页脚标注一次(self):
        h = mdview.standalone("正文。", "t")
        assert "export-mark" in h
        assert "内部参考 · 非报送件" in h
        assert h.count("内部参考 · 非报送件") == 1, "声明在下载件只出现一次（页脚）"
