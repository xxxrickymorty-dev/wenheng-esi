# -*- coding: utf-8 -*-
"""槽位表宽度在 DOCX 保存和重载后仍保持一致。"""
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

import fill_docx


def test_槽位表几何保存重载后仍为固定双栏(repo):
    master = repo / "master" / "学科信息参考母版_2026第5期.docx"
    doc = Document(master)
    table = next(
        tbl for tbl in doc.tables
        if tbl.rows and tbl.rows[0].cells
        and tbl.rows[0].cells[0].text.strip().startswith("国内外排名百分位趋势")
    )

    fill_docx._apply_slot_geometry(table)
    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)

    saved = Document(stream)
    table = next(
        tbl for tbl in saved.tables
        if tbl.rows and tbl.rows[0].cells
        and tbl.rows[0].cells[0].text.strip().startswith("国内外排名百分位趋势")
    )
    grid = table._tbl.find(qn("w:tblGrid"))
    widths = [col.get(qn("w:w")) for col in grid.findall(qn("w:gridCol"))[:2]]
    table_width = table._tbl.tblPr.find(qn("w:tblW"))
    layout = table._tbl.tblPr.find(qn("w:tblLayout"))

    assert widths == ["4898", "4898"]
    assert table_width.get(qn("w:w")) == "9796"
    assert table_width.get(qn("w:type")) == "dxa"
    assert layout is not None and layout.get(qn("w:type")) == "fixed"
    for row in table.rows:
        for cell in row.cells[:2]:
            tc_width = cell._tc.tcPr.find(qn("w:tcW"))
            assert tc_width.get(qn("w:w")) == "4898"
            assert tc_width.get(qn("w:type")) == "dxa"
