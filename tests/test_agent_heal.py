# -*- coding: utf-8 -*-
"""Agent 层 M3 · 自愈闭环测试（SRS R-23 / R-24 / R-26；WH-DOC-04 §6.5）。

四层，由内向外：

  1. **规则库本身** —— 结构与口径（别名表必须覆盖全部 22 个学科的中文名）
  2. **执行闸** —— 未确认的规则类型必须被拒绝（R-24 的可核验判据）
  3. **解析口径** —— overlay 真的改变"怎么读"，且默认口径逐字不变
  4. **端到端脏例** —— 真实台账 + 一处人为脏化：拦停 → 自愈 → 断言全过，且**原件零改动**

第 4 层依赖馆员台账（`input/`，不入库），故缺数据时 `skip` 而非 fail。
脏化方式为 **zip 级 XML 替换**（不是 openpyxl 存盘）—— 后者会抹掉公式缓存值，
把"一处脏"变成"满表脏"，测出来的就不是自愈能力了。
"""
import hashlib
import json
import pathlib
import re
import zipfile

import openpyxl
import pytest

from agent import flow, heal, states, tools
from parse_ledger import DISCIPLINES_CN, DISCIPLINES_EN, _parse_scope, parse_wide


# ============================================================ 工具：zip 级脏化
def _sheet_xml(zf, title):
    """工作表名 → 它在 zip 里的 XML 路径（走 workbook.xml + rels，不猜 sheetN）。"""
    wb = zf.read("xl/workbook.xml").decode("utf-8")
    rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid = None
    for tag in re.finditer(r"<sheet\b[^>]*>", wb):
        t = re.search(r'name="([^"]*)"', tag.group(0))
        r = re.search(r'r:id="([^"]*)"', tag.group(0))
        if t and r and t.group(1) == title:
            rid = r.group(1)
            break
    if rid is None:
        raise LookupError(f"台账无工作表 {title}")
    m = re.search(r'Id="%s"[^>]*Target="([^"]*)"' % re.escape(rid), rels)
    target = m.group(1).lstrip("/")
    return target if target.startswith("xl/") else "xl/" + target.lstrip("./")


def _si_texts(ss):
    return ["".join(re.findall(r"<t[^>]*>(.*?)</t>", blk, re.S))
            for blk in re.findall(r"<si>(.*?)</si>", ss, re.S)]


def corrupt_cell_text(src, dst, sheet, old, new):
    """把某工作表**首列**里第一个等于 `old` 的单元格改成 `new`。

    只动两个 XML：目标 sheet + sharedStrings。其余条目**逐字节原样搬运**，
    所以公式、图表、条件格式、缓存值全部保住 —— 这就是"一处脏"。
    """
    with zipfile.ZipFile(src) as zf:
        sheet_path = _sheet_xml(zf, sheet)
        xml = zf.read(sheet_path).decode("utf-8")
        ss = zf.read("xl/sharedStrings.xml").decode("utf-8")

    # 定位要改的行（用 openpyxl 只读找，不改）
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    try:
        row = next(r[0].row for r in wb[sheet].iter_rows(min_col=1, max_col=1)
                   if r[0].value == old)
    finally:
        wb.close()
    coord = f"A{row}"

    texts = _si_texts(ss)
    if old not in texts:
        raise LookupError(f"{sheet}!{coord} 的值 {old!r} 不在共享字符串表里")
    if new in texts:
        new_idx = texts.index(new)
    else:
        new_idx = len(texts)
        ss = re.sub(r"</sst>\s*$", f"<si><t>{new}</t></si></sst>", ss)
        assert new_idx in range(len(_si_texts(ss))), "追加共享字符串失败"

    m = re.search(r'<c r="%s"([^>]*?)(?:/>|>.*?</c>)' % coord, xml, re.S)
    assert m, f"工作表 XML 里找不到单元格 {coord}"
    style = re.search(r'\bs="(\d+)"', m.group(1))
    cell = (f'<c r="{coord}"' + (f' s="{style.group(1)}"' if style else "")
            + f' t="s"><v>{new_idx}</v></c>')
    xml = xml[:m.start()] + cell + xml[m.end():]

    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == sheet_path:
                data = xml.encode("utf-8")
            elif item.filename == "xl/sharedStrings.xml":
                data = ss.encode("utf-8")
            zout.writestr(item, data)
    return coord


def sha256(p):
    return hashlib.sha256(pathlib.Path(p).read_bytes()).hexdigest()


# ============================================================ 工具：假工作表
class FakeWS:
    """只实现 parse_* 用到的那点接口（title + iter_rows）。"""

    def __init__(self, title, rows):
        self.title = title
        self._rows = rows

    def iter_rows(self, values_only=True, **kw):
        assert values_only
        return iter(self._rows)


def wide_rows(header, first_col_values):
    return [header] + [[v] + [1.0] * (len(header) - 1) for v in first_col_values]


def parse_diag(rows, title="总被引", **kw):
    """跑一次 parse_wide 并取回 (df, diag)——诊断通道就是断言与报告的眼睛。"""
    diag = {}
    df = parse_wide(FakeWS(title, rows), [], diag=diag, **kw)
    return df, diag


def known(diag, title="总被引"):
    return set((diag.get("known_names") or {}).get(title, []))


def unknown(diag, title="总被引"):
    return (diag.get("unknown_names") or {}).get(title, [])


# ============================================================ 1. 规则库
def test_规则库六条且错误码在文档全集内():
    rules = heal.load_rules()
    assert len(rules) == 6, "WH-DOC-04 §6.5 首批 6 类"
    assert len({r["id"] for r in rules}) == 6
    for r in rules:
        assert r["code"] in tools.ERROR_CODES, r["id"]


def test_每条规则都说清了依据与可执行档位():
    for r in heal.load_rules():
        assert r.get("doc"), f"{r['id']} 没写文档依据"
        assert r["executable"] in (heal.AUTO, heal.DECLARE, heal.NEVER)


def test_别名表覆盖全部22个学科的中文名_防漂移():
    """学科名单改了而规则库没跟上 → 中文名会再次静默丢行。这条断言就是防这个。"""
    alias = {}
    for r in heal.load_rules():
        alias.update((r.get("overlay") or {}).get("discipline_alias") or {})
    # DISCIPLINES_CN 是 {英文规范名: 中文名}；别名表是 {别名: 英文规范名}
    missing = [cn for cn in DISCIPLINES_CN.values() if cn not in alias]
    assert not missing, f"别名表缺中文名：{missing}"
    wrong = {cn: alias[cn] for en, cn in DISCIPLINES_CN.items() if alias[cn] != en}
    assert not wrong, f"别名映射与 DISCIPLINES_CN 不一致：{wrong}"
    assert len(DISCIPLINES_EN) == 22


def test_规则库缺文件时抛错而非静默空库(tmp_path):
    with pytest.raises(heal.RuleError):
        heal.load_rules(tmp_path / "不存在.yaml")


# ============================================================ 2. 执行闸（R-24）
def test_未确认的规则被拒绝执行():
    """R-24 的可核验判据：未确认规则被拒绝执行。"""
    with pytest.raises(tools.ToolError) as e:
        tools.call("apply_repair_rule", rule="header_offset")
    assert e.value.code == "RULE_UNAPPROVED"
    assert e.value.detail.get("advice"), "拒绝时必须给出人工处置建议"


def test_不可修复的规则不给过():
    with pytest.raises(tools.ToolError) as e:
        tools.call("apply_repair_rule", rule="append_period_column")
    assert e.value.code == "RULE_UNAPPROVED"


def test_不存在的规则被拒绝():
    with pytest.raises(tools.ToolError) as e:
        tools.call("apply_repair_rule", rule="我自己编的规则")
    assert e.value.code == "RULE_UNAPPROVED"


def test_已确认规则能取出overlay():
    out = tools.call("apply_repair_rule", rule="normalize_period")
    assert out["overlay"] == {"normalize_period": True}


# ============================================================ 3. 匹配与落盘
def _assert_fail(text, **extra):
    return {"code": "ASSERT_FAIL", "message": "1 项断言未过",
            "fails": [text], "step": "parse", **extra}


def test_规则库无匹配时返回None():
    r = heal.attempt({"code": "MISSING_PERIOD", "message": "缺期次"}, {})
    assert r is None, "错误码都不沾边就不该假装有方案"


def test_命中学科数断言时给出别名overlay(tmp_path):
    ctx = {"run_dir": str(tmp_path)}
    r = heal.attempt(_assert_fail("宽表 total_cites 学科数=22 (实际 21，缺 ['Chemistry'])"), ctx)
    assert "discipline_alias_map" in r["rules"]
    assert "化学" in r["overlay"]["discipline_alias"]
    assert ctx["overlay"] is r["overlay"], "overlay 必须挂到 ctx 上，否则重跑还是旧口径"
    assert (tmp_path / "repair.yaml").exists()
    y = (tmp_path / "repair.yaml").read_text(encoding="utf-8")
    for field in ("rule:", "params:", "reason:", "approved_by:"):
        assert field in y, f"repair.yaml 缺文档规定的字段 {field}"


def test_一次自愈把能命中的规则全用上():
    """自愈只有一次机会（MAX_HEAL_ATTEMPTS=1），所以不能只挑第一条规则。"""
    ctx = {}
    r = heal.attempt(_assert_fail("宽表 cnci 学科数=22 (实际 21，缺 ['Physics']) · "
                                  "宽表 cnci 202609 列为空（下期未贴数）"), ctx)
    assert set(r["rules"]) >= {"discipline_alias_map", "normalize_period"}
    assert set(r["overlay"]) == {"discipline_alias", "normalize_period"}


def test_命中但未确认的规则进建议而非overlay(tmp_path):
    ctx = {"run_dir": str(tmp_path)}
    r = heal.attempt(_assert_fail("新国际排名 期组数≥5"), ctx)
    assert not r["overlay"], "未确认规则不得进入 overlay"
    assert r["outcome"] == "refused"
    assert any(a["rule"] == "header_offset" for a in r["advice"])
    assert r["advice"][0]["advice"], "建议必须写出来，否则人工不知道该干什么"


def test_自愈钩子永不抛异常():
    """它在主干上：抛出去会把"台账脏"变成"程序崩"，两者必须分得很清楚。"""
    for bad in (None, {}, {"code": "ASSERT_FAIL"}, {"code": "ASSERT_FAIL", "fails": None},
                {"code": None, "message": None}):
        heal.attempt(bad, {})       # 不抛即通过


def test_编排把overlay传给run_parse(monkeypatch, tmp_path):
    seen = {}
    real_call = tools.call

    def spy(name, **kw):
        if name == "run_parse":
            seen.update(kw)
            return {"ok": True, "period": "202607", "out_dir": "x"}
        return real_call(name, **kw)

    monkeypatch.setattr(tools, "call", spy)
    monkeypatch.setattr(flow, "_heal_hook",
                        lambda failure, ctx: (ctx.update({"overlay": {"normalize_period": True}})
                                              or {"overlay": {"normalize_period": True},
                                                  "rules": ["normalize_period"],
                                                  "outcome": "repaired", "refused": [],
                                                  "advice": []}))
    n = {"i": 0}

    def behaviour(sid):
        if sid == "parse":
            n["i"] += 1
            if n["i"] == 1:
                raise tools.ToolError("ASSERT_FAIL", "1 项断言未过",
                                      fails=["宽表 cnci 202609 列为空（下期未贴数）"])
        return {"ok": True, "period": "202607", "out_dir": "x"}

    monkeypatch.setattr(flow, "_invoke", lambda step, ctx, result: behaviour(step["id"]))
    r = flow.run("fake.xlsx", out_root=tmp_path / "o", run_dir=tmp_path / "r")
    assert r.status == states.COMPLETED
    assert r.repairs and r.repairs[0]["rules"] == ["normalize_period"]
    assert (tmp_path / "r" / "run_result.json").exists()
    assert "repairs" in json.loads((tmp_path / "r" / "run_result.json").read_text(encoding="utf-8"))


def test_自愈逾限的诊断报告写清了自愈过程(monkeypatch, tmp_path):
    def behaviour(step, ctx, result):
        if step["id"] == "parse":
            raise tools.ToolError("ASSERT_FAIL", "1 项断言未过",
                                  fails=["宽表 total_cites 学科数=22 (实际 21，缺 ['Chemistry'])"],
                                  diagnostics={})
        return {"ok": True, "period": "202607", "out_dir": "x"}

    monkeypatch.setattr(flow, "_invoke", behaviour)
    r = flow.run("fake.xlsx", out_root=tmp_path / "o", run_dir=tmp_path / "r")
    assert r.status == states.ESCALATED
    doc = (tmp_path / "r" / "诊断报告.md").read_text(encoding="utf-8")
    assert "## 自愈过程" in doc and "已应用规则" in doc


# ============================================================ 4. 解析口径（纯函数）
def test_默认口径逐字不变_不归一不映射():
    rows = wide_rows(["学科", "2023.3"], ["Chemistry", "Plant and Animal Science"])
    df, diag = parse_diag(rows)
    assert list(df.columns) == ["2023.3"], "默认不该动期次写法"
    assert known(diag) == {"Chemistry"}, "默认不该认别名"
    assert "Plant and Animal Science" in unknown(diag)


def test_别名overlay能把中英混排救回来():
    rows = wide_rows(["学科", "202607"], ["化学", "Physics"])
    _df, diag = parse_diag(rows, alias={"化学": "Chemistry"})
    assert known(diag) == {"Chemistry", "Physics"}
    assert unknown(diag) == []


def test_脏期次overlay归一到YYYYMM():
    rows = wide_rows(["学科", "2023.3", "2023-05", "2023年7月", "202311"],
                     ["Chemistry"])
    df, _ = parse_diag(rows, normalize=True)
    assert list(df.columns) == ["202303", "202305", "202307", "202311"]


def test_表头行偏移overlay生效():
    """表头偏移的症状是**期次列整列丢失**（学科列照样能认出来），
    所以机器不猜行号 —— 猜错会整体错位且一声不响。"""
    rows = [["某期 ESI 学科数据（内部）", None],
            ["学科", "202607"],
            ["Chemistry", 1.0]]
    df0, diag0 = parse_diag(rows)
    assert list(df0.columns) == [], "默认口径下期次表头读不到（这正是故障现场）"
    assert known(diag0) == {"Chemistry"}, "学科列不受表头偏移影响"

    df, diag = parse_diag(rows, header_row=1)
    assert list(df.columns) == ["202607"]
    assert known(diag) == {"Chemistry"}


def test_诊断不写进产物_parquet字节可复现(tmp_path):
    """pandas 3 会把 df.attrs 写进 parquet 元数据。诊断必须走独立通道，
    否则每次改动诊断内容都会改动全部宽表产物的字节，"金样回测按哈希比对"就废了。"""
    import pandas as pd
    rows = wide_rows(["学科", "202607"], ["Chemistry"])
    a, _ = parse_diag(rows)
    b, _ = parse_diag(rows)
    fa, fb = tmp_path / "a.parquet", tmp_path / "b.parquet"
    a.to_parquet(fa)
    b.to_parquet(fb)
    assert not a.attrs, "parse_wide 不该往 DataFrame 上挂 attrs"
    assert fa.read_bytes() == fb.read_bytes()


# ============================================================ 5. 端到端脏例（需真实台账）
@pytest.fixture(scope="module")
def ledger_path(all_ledgers):
    """复用 conftest 的台账夹具：跳过的登记与复核只在那一处做，别在这里另起一套。"""
    return all_ledgers[-1]


@pytest.fixture(scope="module")
def dirty_alias_ledger(ledger_path, tmp_path_factory):
    """真实台账 + 一处脏：首列 Chemistry → 化学。别名规则可救。"""
    d = tmp_path_factory.mktemp("dirty")
    dst = d / "监测表.xlsx"
    coord = corrupt_cell_text(ledger_path, dst, "总被引", "Chemistry", "化学")
    return {"path": dst, "coord": coord, "src": ledger_path}


@pytest.fixture(scope="module")
def dirty_unknown_ledger(ledger_path, tmp_path_factory):
    """真实台账 + 规则库救不了的脏：改成谁都不认识的名字。"""
    d = tmp_path_factory.mktemp("dirty2")
    dst = d / "监测表.xlsx"
    corrupt_cell_text(ledger_path, dst, "总被引", "Chemistry", "某某新学科")
    return {"path": dst, "src": ledger_path}


def test_脏化只改一个格_其余逐字不变(dirty_alias_ledger):
    """先自证夹具干净，否则后面测的是别的东西。"""
    a = openpyxl.load_workbook(dirty_alias_ledger["src"], read_only=True, data_only=True)
    b = openpyxl.load_workbook(dirty_alias_ledger["path"], read_only=True, data_only=True)
    try:
        diff = [(c, ra[0], rb[0])
                for c, (ra, rb) in enumerate(zip(a["总被引"].iter_rows(values_only=True),
                                                 b["总被引"].iter_rows(values_only=True)))
                if ra != rb]
        assert len(diff) == 1, f"应只差一行，实际 {len(diff)} 行：{diff[:3]}"
    finally:
        a.close()
        b.close()


def test_学科名写错时断言点名到具体学科(dirty_alias_ledger, tmp_path):
    """⚠ 回归护栏：宽表 reindex 后 df.index 恒为 22 个规范名，
    若断言用 df.index 判名单，这两条断言会**恒真**——学科名写错也放行。"""
    from parse_ledger import load_ledger, run_assertions
    ok, report, _missing = run_assertions(load_ledger(dirty_alias_ledger["path"]))
    fails = [m for s, m in report if s == "FAIL"]
    assert not ok, "拼错的学科名必须被拦住"
    assert any("Chemistry" in m for m in fails), f"应点名缺的是谁，实际：{fails[:3]}"


@pytest.mark.parametrize("step_only", [True], ids=["parse步"])
def test_脏例端到端_拦停到自愈成功(dirty_alias_ledger, monkeypatch, tmp_path, step_only):
    """验收主用例：真实台账脏一处 → 拦停 → 一次自愈 → 断言全过 → 原件零改动。"""
    led = dirty_alias_ledger["path"]
    before = sha256(led)

    # 只跑 parse 步：自愈发生在这一步，后续四步的正确性由既有端到端用例覆盖，
    # 没必要为了一次自愈再等十几秒出 99 张图。
    monkeypatch.setattr(flow, "STEPS", tuple(s for s in flow.STEPS if s["id"] == "parse"))
    r = flow.run(str(led), out_root=tmp_path / "out", run_dir=tmp_path / "run")

    assert r.status == states.COMPLETED, f"自愈后应走完：{r.failure}"
    assert [s["attempt"] for s in r.steps] == [1, 2], "应恰好重跑一次（首轮 1 + 自愈 2）"
    assert r.repairs and "discipline_alias_map" in r.repairs[0]["rules"]

    run = tmp_path / "run"
    assert (run / "repair.yaml").exists(), "自愈决定必须落盘（LLD §5.2）"
    events = (run / "agent_events.jsonl").read_text(encoding="utf-8")
    assert "apply_repair_rule" in events, "轨迹要记下用了哪条规则"
    hist = (run / "state_history.jsonl").read_text(encoding="utf-8")
    assert "retrying" in hist and "running" in hist, "拦停→重试→继续 必须有状态留痕"

    assert sha256(led) == before, "🔴 自愈绝不能改台账原件（R-24）"


def test_规则库救不了时升级并出诊断报告(dirty_unknown_ledger, monkeypatch, tmp_path):
    led = dirty_unknown_ledger["path"]
    before = sha256(led)
    monkeypatch.setattr(flow, "STEPS", tuple(s for s in flow.STEPS if s["id"] == "parse"))
    r = flow.run(str(led), out_root=tmp_path / "out", run_dir=tmp_path / "run")

    assert r.status == states.ESCALATED
    assert [s["attempt"] for s in r.steps] == [1, 2], "仍应走满一次自愈机会再升级"
    doc = (tmp_path / "run" / "诊断报告.md").read_text(encoding="utf-8")
    assert "某某新学科" in doc, "诊断报告要点名那个认不出的名字"
    assert "自愈过程" in doc
    assert sha256(led) == before


# ============================================================ 6. declare 档 overlay 形状
# 钉住的是一处**静默失效**（2026-09-16 实测发现并修正）：`declared` 里曾多套一层参数名
# （`header_row: {...}` / `group_size: {...}`），overlay 于是变成
#     {header_offset: {header_row: {工作表: 行号}}}
# 而 `_parse_scope` 期待 {header_offset: {工作表: 行号}} → 永远读到默认值。
# 症状：馆员填了参数、规则显示已批准、自愈报 repaired，实际口径一个字节没改。
# 它不抛异常、不影响任何断言，只有形状断言能钉住 —— 下面几条就是那个钉子。
def test_declare档overlay形状与解析口径对齐():
    r = heal.rule_by_id("header_offset")
    ov = heal.overlay_of(r, params={"declared": {"总被引": 2}})
    assert ov == {"header_offset": {"总被引": 2}}, "多包一层会让声明静默失效"
    assert _parse_scope(ov, "总被引")["header_row"] == 2

    r2 = heal.rule_by_id("split_column_group")
    ov2 = heal.overlay_of(r2, params={"declared": {"新国际排名": 4}})
    assert _parse_scope(ov2, "新国际排名")["group_size"] == 4


def test_declare档声明表头在第0行不算未声明():
    """0 是合法行号。曾用 `if v` 过滤，会把 0 当成「馆员没填」丢掉并误报 RULE_UNAPPROVED。"""
    r = heal.rule_by_id("header_offset")
    ov = heal.overlay_of(r, params={"declared": {"总被引": 0}})
    assert _parse_scope(ov, "总被引")["header_row"] == 0


def test_规则库declare档的declared不得再套参数名():
    """结构性防回归：declared 的键应是工作表名，不是 header_row / group_size 这类参数名。"""
    banned = {"header_row", "group_size"}
    for r in heal.load_rules():
        if r["executable"] != heal.DECLARE:
            continue
        assert not (banned & set(r.get("declared") or {})), \
            f"{r['id']} 的 declared 又套了参数名，会与 _parse_scope 错位、静默失效"


def test_overlay只含解析口径键_不碰断言():
    """自愈只让本该被读到的数据被读到，不让断言放行 —— 改断言＝改口径，属评审事项。"""
    allowed = {"header_offset", "group_size", "normalize_period", "discipline_alias"}
    for r in heal.load_rules():
        ov = heal.overlay_of(r, params={"declared": {"某表": 1}})
        assert set(ov) <= allowed, f"{r['id']} 产出了断言口径之外的键：{set(ov) - allowed}"


@pytest.mark.parametrize("body,why", [
    ("version: 1\n", "缺 rules 列表"),
    ("version: 1\nrules: []\n", "rules 为空"),
    ("version: 1\nrules:\n  - id: x\n", "规则缺必需字段"),
    ("version: 1\nrules:\n  - {id: x, code: ASSERT_FAIL, match: {message_re: a},"
     " executable: 乱写}\n", "executable 非法"),
    ("version: 1\nrules:\n  - {id: x, code: 不存在的码, match: {message_re: a},"
     " executable: auto}\n", "错误码不在文档全集内"),
])
def test_规则库结构非法时报错(tmp_path, body, why):
    """坏规则库必须响亮地失败 —— 静默降级成「无规则」比报错难查得多。"""
    p = tmp_path / "bad.yaml"
    p.write_text(body, encoding="utf-8")
    with pytest.raises(heal.RuleError):
        heal.load_rules(p)


# ============================================================ 7. 验收判据：报告数字对
# 里程碑的验收条件从「断言全过」改成「报告数字对」（rick 2026-09-16 定）。
# 理由：断言全过只说明没触发我们**已经想到**的那几条不变量；一处没被想到的解析
# 偏斜照样能让断言全绿、而报告里的数字全是错的。逐格对齐才是"数字对"的证明。
def _table2_of(led):
    """就地从内存里算出报告大表。

    与 `compute.load_wide` 走的是同一套变换（差别只是不落 parquet 再读回来），
    所以 `table2.csv` 的每一格都在这儿了 —— 报告正文的数字全从这张表出。
    """
    import compute
    P, Q = compute.detect_periods(led["wide"])

    def pivot(key, cols):
        # write_outputs 里对 ledger[key] 做 melt → 这里照搬同一步，别自己另捏一套
        long = led[key].melt(id_vars=["discipline", "period"], value_vars=cols,
                             var_name="metric", value_name="value")
        return long.pivot_table(index="discipline", columns=["metric", "period"],
                                values="value")

    ni = pivot("new_intl", ["n_inst_1pct", "rank", "rank_pct"])
    nc = pivot("new_cn", ["n_inst_1pct", "rank", "rank_pct"])
    return compute.build_table2(led["wide"], ni, nc, led["threshold_1permille"], P, Q)


def test_自愈后的读数与干净原件逐格一致(dirty_alias_ledger):
    """**M3 的验收判据**：自愈之后，报告数字与拿干净台账跑出来的一模一样。

    这条同时钉住了自愈的两种走偏：修得不够（还有格子差着）和修过了头
    （把不该改的格子也改了）。断言的粒度是逐格，不是"没报错"。
    """
    import pandas as pd
    from parse_ledger import load_ledger

    ov = heal.attempt(
        _assert_fail("宽表 total_cites 学科数=22 (实际 21，缺 ['Chemistry'])"),
        {"run_dir": None})["overlay"]
    clean = load_ledger(dirty_alias_ledger["src"])
    healed = load_ledger(dirty_alias_ledger["path"], overlay=ov)

    for m in clean["wide"]:
        pd.testing.assert_frame_equal(clean["wide"][m], healed["wide"][m],
                                      check_exact=True, obj=f"宽表 {m}")

    t_clean, t_healed = _table2_of(clean), _table2_of(healed)
    pd.testing.assert_frame_equal(t_clean, t_healed, check_exact=True, obj="报告大表")
    # 正数自证：别拿两张空表比出"一致"
    assert len(t_healed) == 22 and t_healed["total_cites"].notna().all()
    assert t_healed.equals(t_clean)


def test_没自愈的脏台账算出来的读数确实是错的(dirty_alias_ledger):
    """"报告数字对"这条判据自身的阳性对照：不修的话数字真的不一样。

    少了这条，上面那条测试可能只是在证明"两坨一样的东西相等"。
    """
    import pandas as pd
    from parse_ledger import load_ledger
    clean = load_ledger(dirty_alias_ledger["src"])["wide"]["total_cites"]
    dirty = load_ledger(dirty_alias_ledger["path"])["wide"]["total_cites"]
    with pytest.raises(AssertionError):
        pd.testing.assert_frame_equal(clean, dirty)
    # 具体错成什么样也钉住：认不出的学科被 reindex 成 NaN —— 报告里那一行就是空的。
    # 只钉当期，未来列（下期还没贴数）本来就该空。
    P = [c for c in clean.columns if clean[c].notna().any()][-1]
    assert dirty.loc["Chemistry", P] != dirty.loc["Chemistry", P], "脏读数应缺当期值"
    assert clean.loc["Chemistry", P] == clean.loc["Chemistry", P], "干净读数当期有值"
    # 连大表都算不出来（int(NaN) 直接炸）。所以"数字不对"不是程度问题，是根本出不了稿。
    with pytest.raises(ValueError, match="NaN"):
        _table2_of(load_ledger(dirty_alias_ledger["path"]))


# ============================================================ 8. 失败运行的产物边界
# 取证复现的故障（2026-09-16）：`write_outputs` 曾排在断言闸**之前**，于是
#     一次失败运行 = out/<期次>/ 里「本次的新宽表/新大表」+「上一次的报告 docx」
# 对勘本照旧渲染那份旧报告，页面没有任何警报 —— 典型静默产错报告。
def test_断言失败时解析一个字节都不落盘(dirty_unknown_ledger, monkeypatch, tmp_path):
    """规则库救不了的脏例：两轮都失败，产物目录必须干净得像没跑过。

    ⚠ D1 例外：AI 层诊断留痕 `ai_calls.jsonl` 是**审计记录**不是产物
    （SPEC §2.2「每次调用尝试恰一条」），它就该跟着期次目录走。这里只拦产物。
    """
    monkeypatch.setattr(flow, "STEPS", tuple(s for s in flow.STEPS if s["id"] == "parse"))
    out = tmp_path / "out"
    r = flow.run(str(dirty_unknown_ledger["path"]), out_root=out, run_dir=tmp_path / "run")
    assert r.status == states.ESCALATED
    leftovers = [p for p in out.rglob("*")
                 if p.name != "ai_calls.jsonl" and p.is_file()] if out.exists() else []
    assert not leftovers, f"失败运行不该在产物目录留东西，实际：{[str(p) for p in leftovers]}"


def test_解析成功时照常落盘(dirty_unknown_ledger, monkeypatch, tmp_path):
    """上面那条的阳性对照：不是我为了让它过把落盘整个删了。"""
    monkeypatch.setattr(flow, "STEPS", tuple(s for s in flow.STEPS if s["id"] == "parse"))
    out = tmp_path / "out"
    r = flow.run(str(dirty_unknown_ledger["src"]), out_root=out, run_dir=tmp_path / "run")
    assert r.status == states.COMPLETED
    assert (out / r.period / "wide" / "total_cites.parquet").exists()
    assert (out / r.period / "meta.json").exists()


def _seed_prev_products(period_dir):
    """埋一份"上一版成功运行"的成品。"""
    period_dir.mkdir(parents=True, exist_ok=True)
    (period_dir / "报告初稿_202607（试排）.docx").write_bytes(b"OLD-REPORT")
    (period_dir / "red_slots.json").write_text('{"old": 1}', encoding="utf-8")
    (period_dir / "红色件清单.md").write_text("旧清单", encoding="utf-8")
    (period_dir / "table2.csv").write_text("discipline,ranking\n", encoding="utf-8")


def test_失败运行不会留下新表格配旧报告(monkeypatch, tmp_path):
    """复现路径：先有一版成功产物 → 再跑一次 → 在 charts 步失败。

    要求：旧成品被撤进 `_superseded/`（不是删掉，要能追溯），期内不再有 docx。
    这样看板显示的是"本期未出稿"，而不是一份跟新表格对不上的旧报告。
    """
    out = tmp_path / "out"
    d = out / "202607"
    _seed_prev_products(d)

    def behaviour(step, ctx, result):
        if step["id"] == "parse":
            d.joinpath("table2.csv").write_text("discipline,ranking\n新数据\n", encoding="utf-8")
            return {"ok": True, "period": "202607", "out_dir": str(d)}
        if step["id"] == "charts":
            raise tools.ToolError("FONT_MISSING", "模拟字体缺失")
        return {"ok": True, "period": "202607", "out_dir": str(d)}

    monkeypatch.setattr(flow, "_invoke", behaviour)
    r = flow.run("fake.xlsx", out_root=out, run_dir=tmp_path / "run")

    assert r.status == states.ESCALATED
    assert set(r.superseded) == {"报告初稿_202607（试排）.docx", "red_slots.json", "红色件清单.md"}
    assert not list(d.glob("报告初稿_*.docx")), "旧报告必须撤下，否则看板会拿它配新表格"
    assert not (d / "red_slots.json").exists(), "对勘本按它判「本期可看」，也要撤下"
    arch = list((d / flow.SUPERSEDED_DIR).rglob("报告初稿_*.docx"))
    assert len(arch) == 1 and arch[0].read_bytes() == b"OLD-REPORT", "撤下的旧稿要留着可追溯"
    assert (d / "table2.csv").read_text(encoding="utf-8").endswith("新数据\n"), \
        "中间产物不必撤 —— 它每步都被覆盖，且本次的这份是断言验过的"


def test_自愈成功后仍要撤下上版成品(dirty_alias_ledger, monkeypatch, tmp_path):
    """自愈成功≠上版没被顶掉。撤下的动作只跟"目录开始被改写"有关，与成败无关。"""
    out = tmp_path / "out"
    monkeypatch.setattr(flow, "STEPS", tuple(s for s in flow.STEPS if s["id"] == "parse"))
    first = flow.run(str(dirty_alias_ledger["src"]), out_root=out, run_dir=tmp_path / "r1")
    assert first.status == states.COMPLETED
    _seed_prev_products(out / first.period)

    second = flow.run(str(dirty_alias_ledger["path"]), out_root=out, run_dir=tmp_path / "r2")
    assert second.status == states.COMPLETED, f"自愈应成功：{second.failure}"
    assert second.superseded, "撤下清单要回填到 run_result 里，否则审计看不出稿去哪了"
    assert not list((out / first.period).glob("red_slots.json"))
    sup = list((out / first.period / flow.SUPERSEDED_DIR).rglob("red_slots.json"))
    assert len(sup) == 1 and sup[0].read_text(encoding="utf-8") == '{"old": 1}'
    assert (sup[0].parent / "_来源.txt").exists(), "要知道这批旧稿是被哪次运行撤下的"


def test_解析没过时不撤上版成品(monkeypatch, tmp_path):
    """边界：解析两轮都没过 → 期内压根没被改写 → 上一版继续有效，不该被藏起来。"""
    out = tmp_path / "out"
    d = out / "202607"
    _seed_prev_products(d)

    monkeypatch.setattr(flow, "_invoke",
                        lambda step, ctx, result: (_ for _ in ()).throw(
                            tools.ToolError("ASSERT_FAIL", "1 项断言未过", fails=["x"])))
    r = flow.run("fake.xlsx", out_root=out, run_dir=tmp_path / "run")

    assert r.status == states.ESCALATED
    assert not r.superseded
    assert (d / "报告初稿_202607（试排）.docx").exists(), "没被覆盖过的稿件不该被撤下"
    assert not (d / flow.SUPERSEDED_DIR).exists()


# —— 两条路径必须用同一份口径：演示站走 `run_pipeline.py` 直跑，不走 agent/flow.py ——
def test_成品撤下与解析落盘的顺序被钉住():
    """结构性防回归。这两处的顺序是这次修的 bug 本体，翻回去不抛异常、只静默产错报告：

      · `run_parse` 必须先过断言闸再 `write_outputs`；
      · `run_pipeline` 必须在 `write_outputs` 之前撤成品。
    """
    repo = pathlib.Path(__file__).resolve().parents[1]
    steps_src = (repo / "src" / "steps.py").read_text(encoding="utf-8")
    body = steps_src[steps_src.index("def parse("):]
    assert body.index("if ctx.probe:") < body.index("write_outputs("), \
        "probe 又开始落盘了 —— 编排预判会留下垃圾目录"
    assert body.index("if not ok and ctx.strict_assertions:") < body.index("write_outputs("), \
        "parse 又开始先落盘后判断言了 —— 失败运行会在产物目录留半成品"

    pipe_src = (repo / "src" / "run_pipeline.py").read_text(encoding="utf-8")
    assert "quarantine_final_products" in pipe_src, \
        "直跑路径（演示站主路径）漏了撤成品这一步"
    assert pipe_src.index("quarantine_final_products") < pipe_src.index("write_outputs("), \
        "直跑路径必须先撤成品再写盘"


def test_成品定义只有一处_两路径不各写一套():
    import parse_ledger
    assert flow.SUPERSEDED_DIR == parse_ledger.SUPERSEDED_DIR
    assert tuple(flow.FINAL_PRODUCTS) == tuple(parse_ledger.FINAL_PRODUCTS)


# ============================================================ 9. 原件只读护栏（运行期）
# 测试里的 sha256 断言只证明"我测的那条路径没改台账"。运行期这条护栏覆盖所有路径。
def test_台账被改写时运行期护栏拦停(monkeypatch, tmp_path):
    led = tmp_path / "监测表.xlsx"
    led.write_bytes(b"v1")

    def behaviour(step, ctx, result):
        pathlib.Path(ctx["ledger"]).write_bytes(b"v2")      # 模拟原件被改写
        return {"ok": True, "period": "202607", "out_dir": str(tmp_path / "o")}

    monkeypatch.setattr(flow, "_invoke", behaviour)
    r = flow.run(str(led), out_root=tmp_path / "out", run_dir=tmp_path / "run")

    assert r.status == states.ESCALATED, "原件被改写，本次结论不可采信"
    assert r.failure and r.failure["step"] == "guard"
    doc = (tmp_path / "run" / "诊断报告.md").read_text(encoding="utf-8")
    assert "原件只读护栏" in doc and "台账指纹" in doc
    assert b"v1" != led.read_bytes()      # 自证：护栏确实是在"真被改了"之后触发的


def test_台账没动时护栏不误报(monkeypatch, tmp_path):
    """阳性对照。护栏要是见谁都拦，那就不是护栏，是停机键。"""
    led = tmp_path / "监测表.xlsx"
    led.write_bytes(b"v1")
    monkeypatch.setattr(flow, "_invoke",
                        lambda step, ctx, result: {"ok": True, "period": "202607",
                                                   "out_dir": str(tmp_path / "o")})
    r = flow.run(str(led), out_root=tmp_path / "out", run_dir=tmp_path / "run")
    assert r.status == states.COMPLETED
    assert r.ledger_sha256 and len(r.ledger_sha256) == 64


# ============================================================ 10. 规则库故障 vs 无匹配
def test_规则库坏掉时报库故障而非无匹配(tmp_path):
    """两者都会让自愈"什么也没做"，但人工要去的方向相反：修规则库 vs 看台账。"""
    r = heal.attempt(_assert_fail("宽表 total_cites 学科数=22 (实际 21，缺 ['Chemistry'])"),
                     {"run_dir": str(tmp_path)},
                     rules_path=tmp_path / "并没有这个文件.yaml")
    assert r is not None, "库坏了不能等同于「没有匹配规则」"
    assert "规则库不可用" in r["library_error"]
    assert not r["overlay"], "库都读不出来，不能假装有方案"
    assert "规则库不可用" in (tmp_path / "repair.yaml").read_text(encoding="utf-8")


def test_规则库坏掉时轨迹不写成规则命中(monkeypatch, tmp_path):
    why = "规则库不可用：规则库 YAML 解析失败"
    monkeypatch.setattr(flow, "_heal_hook", lambda failure, ctx: {
        "rules": [], "overlay": {}, "advice": [], "outcome": "refused",
        "library_error": why, "refused": [{"rule": None, "why": why}]})

    def behaviour(step, ctx, result):
        if step["id"] == "parse":
            raise tools.ToolError("ASSERT_FAIL", "1 项断言未过", fails=["x"])
        return {"ok": True, "period": "202607"}

    monkeypatch.setattr(flow, "_invoke", behaviour)
    r = flow.run("fake.xlsx", out_root=tmp_path / "out", run_dir=tmp_path / "run")

    assert r.status == states.ESCALATED
    events = (tmp_path / "run" / "agent_events.jsonl").read_text(encoding="utf-8")
    assert "规则库不可用" in events
    assert "规则命中但不能执行" not in events, \
        "库坏了却被写成「规则命中」，会让人拿着这条去台账里找一个并不存在的脏例"
    doc = (tmp_path / "run" / "诊断报告.md").read_text(encoding="utf-8")
    assert "## 规则库不可用" in doc
