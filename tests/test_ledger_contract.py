# -*- coding: utf-8 -*-
"""`src/ledger_contract.py`：动态监测表契约面的单一入口。

钉住五件事：
1. 定位顺序（`input/` 运行时母本优先，回落 `../数据与复现/` 馆员原件）；
2. `LEDGER_NAME` 只有一份字面量（ledger / inventory / contract 是同一个对象）；
3. 三种表清单是三个真实不同的作用域（宽表 6 / 读表 17 / 写链 17）；
4. 沙盘副本与文件指纹；
5. 母版、数据源目录，以及两个沙盘工具确实走这一层、不再各写契约。
"""
from __future__ import annotations

import hashlib
import json
import os
import pathlib
import time
import zipfile

import openpyxl
import pytest

import inventory
import ledger
import ledger_contract as contract
import ledger_merge
import parse_ledger
import xlsx_io


REPO = pathlib.Path(__file__).resolve().parents[1]


# ------------------------------------------------------------------ 定位与命令
def _touch(path: pathlib.Path, data: bytes = b"PK\x03\x04") -> pathlib.Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return path


def test_locate_prefers_runtime_then_source(tmp_path):
    repo = tmp_path / "wenheng"
    repo.mkdir()
    runtime = contract.runtime_path(repo)
    source = contract.source_path(repo)

    assert contract.locate(repo) is None, "两边都没有时返回 None，不猜路径"
    _touch(source)
    assert contract.locate(repo) == source, "只有馆员原件时回落原件"
    _touch(runtime)
    assert contract.locate(repo) == runtime, "两边都在时运行时母本优先"


def test_require_exits_with_the_old_cli_message(tmp_path):
    repo = tmp_path / "wenheng"
    repo.mkdir()
    with pytest.raises(SystemExit) as exc:
        contract.require(repo)
    assert str(exc.value) == f"找不到台账：{contract.LEDGER_NAME}"


# ------------------------------------------------------------------ 常量只有一份
def test_ledger_name_is_one_object_across_modules():
    assert contract.LEDGER_NAME is ledger.LEDGER_NAME
    assert inventory.LEDGER_NAME is ledger.LEDGER_NAME, \
        "inventory 不能再写第二份文件名；它只做 ledger 的再导出"


def test_discipline_table_is_reused_not_copied():
    assert contract.DISCIPLINES_CN is parse_ledger.DISCIPLINES_CN
    assert contract.DISCIPLINES_EN == tuple(parse_ledger.DISCIPLINES_EN)
    assert len(contract.DISCIPLINES_CN) == 22


# ------------------------------------------------------------------ 表清单作用域
def test_readiness_sheets_cover_every_parser_wide_table():
    assert contract.WIDE_SHEETS == tuple(parse_ledger.WIDE_SHEETS)
    assert len(contract.WIDE_SHEETS) == 6
    assert len(contract.READINESS_SHEETS) == 17
    assert len(set(contract.READINESS_SHEETS)) == 17
    assert set(contract.WIDE_SHEETS) <= set(contract.READINESS_SHEETS)


def test_chain_scope_is_a_write_list_not_the_read_contract():
    assert contract.CHAIN_SHEETS is ledger_merge.CHAIN_SHEETS
    assert contract.SHEET_SOURCE is ledger_merge.SHEET_SOURCE
    # 两个 17 张表不是同一份：读表要「高校学科发展」，写链多一张「InCites快照」。
    assert set(contract.READINESS_SHEETS) - set(contract.CHAIN_SHEETS) == {"高校学科发展"}
    assert set(contract.CHAIN_SHEETS) - set(contract.READINESS_SHEETS) == {"InCites快照"}


# ------------------------------------------------------------------ 副本与指纹
def test_sandbox_copy_and_fingerprint(tmp_path):
    src = _touch(tmp_path / "母本.xlsx", b"PK\x03\x04ledger")
    dst = tmp_path / "sandbox" / "副本.xlsx"

    out = contract.sandbox_copy(src, dst)
    assert out == dst and dst.read_bytes() == src.read_bytes()
    expected = hashlib.sha256(src.read_bytes()).hexdigest()
    assert contract.sha256(src) == expected
    assert contract.sha256(dst) == expected


def test_sandbox_copy_skips_fresh_target_and_refresh_forces(tmp_path):
    src = _touch(tmp_path / "母本.xlsx", b"v1")
    dst = contract.sandbox_copy(src, tmp_path / "副本.xlsx")

    now = time.time()
    src.write_bytes(b"v2")
    dst.write_bytes("旧副本".encode())
    os.utime(src, (now - 100, now - 100))
    os.utime(dst, (now, now))

    contract.sandbox_copy(src, dst)
    assert dst.read_bytes() == "旧副本".encode(), "目标不比源旧时不重复复制"

    contract.sandbox_copy(src, dst, refresh=True)
    assert dst.read_bytes() == b"v2", "refresh=True 必须强制覆盖"


def test_fingerprint_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        contract.sha256(tmp_path / "没有这个文件.xlsx")


# ------------------------------------------------------------------ 母版与数据根
def test_master_and_data_roots(tmp_path):
    repo = tmp_path / "wenheng"
    repo.mkdir()

    assert contract.master_path(repo) == repo / "master" / contract.MASTER_NAME
    assert contract.data_root(repo) == tmp_path / "数据与复现"
    assert contract.esi_root(repo) == tmp_path / "数据与复现" / "ESI数据"
    assert contract.incites_root(repo) == tmp_path / "数据与复现" / "incites"


# ------------------------------------------------------------------ 两个工具都走契约
def test_both_sandbox_tools_go_through_the_contract():
    for name in ("heal_sandbox.py", "migration_check.py"):
        text = (REPO / "tools" / name).read_text(encoding="utf-8")
        assert "ledger_contract" in text, f"{name} 没走契约模块"
        assert "find_ledger" not in text, f"{name} 又写了一份定位器"
        assert "LEDGER_NAME =" not in text, f"{name} 又写了一份文件名"
        assert "REQUIRED_SHEETS" not in text, f"{name} 又写了一份表清单"
        assert "数据与复现" not in text, f"{name} 又硬编码了数据根"
        assert "学科信息参考母版" not in text, f"{name} 又硬编码了母版名"
        assert "copy2" not in text, f"{name} 没走 sandbox_copy"
        assert "hashlib" not in text, f"{name} 没走 contract.sha256"


def _tiny_ledger(path: pathlib.Path) -> pathlib.Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "总被引"
    ws.append(["学科", "202607"])
    ws.append(["Chemistry", 1.0])
    ws.append(["Physics", 2.0])
    wb.save(path)
    return path


def test_heal_sandbox_edit_uses_the_workbook_gate(tmp_path, monkeypatch):
    """脏化不再自己拼 zip：sharedStrings / sheet 的落盘统一过良构闸。"""
    import heal_sandbox as hs

    path = _tiny_ledger(tmp_path / "监测表.xlsx")
    with zipfile.ZipFile(path) as zf:
        sheet_path = xlsx_io.member_of(zf, "总被引")
        before = {i.filename: hashlib.sha256(zf.read(i.filename)).hexdigest()
                  for i in zf.infolist()}

    coord = hs.corrupt_cell_text(path, "总被引", "Chemistry", "化学")
    assert coord == "A2"

    with zipfile.ZipFile(path) as zf:
        after = {i.filename: hashlib.sha256(zf.read(i.filename)).hexdigest()
                 for i in zf.infolist()}
    changed = {name for name in before if before[name] != after.get(name)}
    assert changed <= {sheet_path, "xl/sharedStrings.xml"}
    assert sheet_path in changed

    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        assert wb["总被引"]["A2"].value == "化学"
        assert wb["总被引"]["B2"].value == 1
        assert wb["总被引"]["A3"].value == "Physics"
    finally:
        wb.close()

    # 坏 XML 被闸拦住，文件一字节不动（这是 2026-09-18 台账变砖事故的同类入口）。
    # 正常路径会转义；这里临时关掉转义，逼出「非法 XML」以验证闸门确实在链路上。
    monkeypatch.setattr(hs.html, "escape", lambda s: s)
    before_bytes = path.read_bytes()
    with pytest.raises(SystemExit):
        hs.corrupt_cell_text(path, "总被引", "化学", "a<b")
    assert path.read_bytes() == before_bytes


def _seed_migration_repo(repo: pathlib.Path) -> pathlib.Path:
    """造一份「17 张表 + 22 学科 + 最新期次」的最小换馆样本。"""
    src = repo / "input" / contract.LEDGER_NAME
    src.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = contract.READINESS_SHEETS[0]
    ws.append(["学科", "202607"])
    for cn in contract.DISCIPLINES_CN.values():
        ws.append([cn, 1.0])
    for name in contract.READINESS_SHEETS[1:]:
        wb.create_sheet(name)
    wb.save(src)

    master = contract.master_path(repo)
    master.parent.mkdir(parents=True, exist_ok=True)
    master.write_bytes(b"docx")
    contract.esi_root(repo).mkdir(parents=True, exist_ok=True)
    contract.incites_root(repo).mkdir(parents=True, exist_ok=True)
    return src


def test_migration_check_uses_contract_and_keeps_story_shape(tmp_path, monkeypatch):
    import migration_check as mc

    repo = tmp_path / "wenheng"
    src = _seed_migration_repo(repo)
    work = repo / "tmp" / "_migration_sandbox"
    before = contract.sha256(src)

    monkeypatch.setattr(mc, "REPO", repo)
    monkeypatch.setattr(mc, "WORK", work)
    assert mc.main() == 0

    story = json.loads((work / "readiness.json").read_text(encoding="utf-8"))
    assert set(story) >= {"generated_at", "school", "ledger", "checks", "ready", "note"}
    assert story["ledger"] == "兄弟馆_监测表.xlsx"
    assert story["ready"] is True
    assert [c["id"] for c in story["checks"]] == [
        "sheets", "disciplines", "period", "master", "sources"]
    assert all(c["ok"] for c in story["checks"])
    assert story["checks"][0]["label"] == f"必需的 {len(contract.READINESS_SHEETS)} 张工作表"
    assert (work / "兄弟馆_监测表.xlsx").is_file()
    assert contract.sha256(src) == before, "换馆沙盘只碰副本，原件一字节不动"
